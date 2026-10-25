from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill


IDENTITY_FIELDS = {
    "employee_name": "employee_name",
    "employee_id": "employee_id",
    "pan": "pan",
    "pay_scale": "pay_scale",
    "final_designation": "final_designation",
    "bank_account": "bank_account",
    "gpf_pran": "gpf_pran",
    "belt_no": "belt_no",
}

CORE_COMPONENTS = (
    "basic_pay",
    "da",
    "hra",
    "income_tax",
    "itax",
    "lic",
    "gpf",
    "gpf_2004",
    "ais_gpf",
    "rghs",
    "sip",
    "si_premium",
    "sil",
    "si_loan",
    "sii",
    "rop",
    "advance_salary",
    "roadways_deduction",
    "ppss_jdr",
    "hfrh",
    "eacpenf",
    "giais",
    "car_rent",
)

MATCH_STATUSES = {
    "MATCHED",
    "NEW_EMPLOYEE",
    "REMOVED_EMPLOYEE",
    "POSSIBLE_MATCH_NEEDS_REVIEW",
    "IDENTITY_CONFLICT",
    "UNMATCHED",
}


def _row_get(row: Any, key: str, default: Any = None) -> Any:
    try:
        return row[key]
    except (KeyError, IndexError, TypeError):
        return default


def _as_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return default


def _period_parts(month_year: str) -> tuple[str, str]:
    text = str(month_year or "").replace(" ", "")
    if "/" in text:
        month, year = text.split("/", 1)
        return month.strip(), year.strip()
    parts = str(month_year or "").split()
    if len(parts) >= 2:
        return parts[0].strip(), parts[-1].strip()
    return str(month_year or "").strip(), ""


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


def _pct_diff(base: float, compare: float) -> float:
    denom = abs(base) if abs(base) > 0.0001 else 1.0
    return abs(compare - base) * 100.0 / denom


def init_monthly_comparison_schema(db) -> None:
    db.executescript("""
    CREATE TABLE IF NOT EXISTS monthly_comparison_runs_v3 (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        run_name TEXT,
        base_month TEXT,
        base_year TEXT,
        base_group TEXT,
        base_source_file TEXT,
        compare_month TEXT,
        compare_year TEXT,
        compare_group TEXT,
        compare_source_file TEXT,
        match_strategy TEXT,
        amount_threshold REAL DEFAULT 1000,
        percent_threshold REAL DEFAULT 10,
        base_employee_count INTEGER DEFAULT 0,
        compare_employee_count INTEGER DEFAULT 0,
        matched_count INTEGER DEFAULT 0,
        new_employee_count INTEGER DEFAULT 0,
        removed_employee_count INTEGER DEFAULT 0,
        changed_employee_count INTEGER DEFAULT 0,
        high_variance_count INTEGER DEFAULT 0,
        needs_review_count INTEGER DEFAULT 0,
        total_gross_difference REAL DEFAULT 0,
        total_deduction_difference REAL DEFAULT 0,
        total_net_difference REAL DEFAULT 0,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS monthly_comparison_results_v3 (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        run_id INTEGER,
        match_status TEXT,
        change_type TEXT,
        employee_name TEXT,
        employee_id TEXT,
        pan TEXT,
        base_employee_row_id INTEGER,
        compare_employee_row_id INTEGER,
        base_group TEXT,
        compare_group TEXT,
        base_designation TEXT,
        compare_designation TEXT,
        base_pay_scale TEXT,
        compare_pay_scale TEXT,
        base_basic REAL DEFAULT 0,
        compare_basic REAL DEFAULT 0,
        basic_difference REAL DEFAULT 0,
        base_gross REAL DEFAULT 0,
        compare_gross REAL DEFAULT 0,
        gross_difference REAL DEFAULT 0,
        base_deduction REAL DEFAULT 0,
        compare_deduction REAL DEFAULT 0,
        deduction_difference REAL DEFAULT 0,
        base_net REAL DEFAULT 0,
        compare_net REAL DEFAULT 0,
        net_difference REAL DEFAULT 0,
        component_changes_json TEXT,
        identity_changes_json TEXT,
        issue_message TEXT,
        review_status TEXT DEFAULT 'OPEN',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    );
    CREATE INDEX IF NOT EXISTS idx_monthly_runs_created ON monthly_comparison_runs_v3(created_at);
    CREATE INDEX IF NOT EXISTS idx_monthly_results_run ON monthly_comparison_results_v3(run_id);
    CREATE INDEX IF NOT EXISTS idx_monthly_results_status ON monthly_comparison_results_v3(match_status, change_type);
    """)


def options(db) -> dict[str, Any]:
    rows = db.execute("""
        SELECT e.month_year, e.group_name, e.source_filename, b.bill_no, COUNT(*) AS employee_count
        FROM employees_v2 e
        LEFT JOIN bills_v2 b ON b.id = e.bill_id
        WHERE COALESCE(e.is_deleted,0) = 0
        GROUP BY e.month_year, e.group_name, e.source_filename, b.bill_no
        ORDER BY e.month_year, e.group_name, e.source_filename
    """).fetchall()
    months, years, groups, source_files, bill_numbers = set(), set(), set(), set(), set()
    periods = []
    for row in rows:
        month, year = _period_parts(_row_get(row, "month_year", ""))
        if month:
            months.add(month)
        if year:
            years.add(year)
        if _row_get(row, "group_name", ""):
            groups.add(_row_get(row, "group_name", ""))
        if _row_get(row, "source_filename", ""):
            source_files.add(_row_get(row, "source_filename", ""))
        if _row_get(row, "bill_no", ""):
            bill_numbers.add(_row_get(row, "bill_no", ""))
        periods.append({
            "month_year": _row_get(row, "month_year", ""),
            "month": month,
            "year": year,
            "group_name": _row_get(row, "group_name", ""),
            "source_file": _row_get(row, "source_filename", ""),
            "bill_no": _row_get(row, "bill_no", ""),
            "employee_count": _row_get(row, "employee_count", 0),
        })
    return {
        "months": sorted(months),
        "years": sorted(years),
        "groups": sorted(groups),
        "source_files": sorted(source_files),
        "bill_numbers": sorted(bill_numbers),
        "periods": periods,
    }


def _period_where(prefix: str, filters: dict[str, Any]) -> tuple[str, list[Any]]:
    conditions = ["COALESCE(e.is_deleted,0) = 0"]
    params: list[Any] = []
    month = str(filters.get(f"{prefix}_month") or "").strip()
    year = str(filters.get(f"{prefix}_year") or "").strip()
    group = str(filters.get(f"{prefix}_group") or "").strip()
    source_file = str(filters.get(f"{prefix}_source_file") or "").strip()
    if month and year:
        conditions.append("REPLACE(e.month_year, ' ', '') IN (?, ?)")
        params.extend([f"{month}/{year}", f"{month}/{year}".replace("/", " / ")])
    elif month:
        conditions.append("e.month_year LIKE ?")
        params.append(f"%{month}%")
    elif year:
        conditions.append("e.month_year LIKE ?")
        params.append(f"%{year}%")
    if group:
        conditions.append("e.group_name = ?")
        params.append(group)
    if source_file:
        conditions.append("e.source_filename = ?")
        params.append(source_file)
    return " AND ".join(conditions), params


def load_employee_snapshot(db, prefix: str, filters: dict[str, Any]) -> list[dict[str, Any]]:
    where, params = _period_where(prefix, filters)
    rows = db.execute(f"""
        SELECT
          e.id AS employee_row_id,
          e.bill_id,
          e.source_filename AS source_file,
          e.group_name,
          e.month_year,
          e.raw_name AS employee_name,
          e.employee_id,
          e.pan_raw AS pan,
          e.gpf_pran_raw AS gpf_pran,
          e.bank_acc AS bank_account,
          e.belt_no,
          COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation, '') AS final_designation,
          e.pay_scale_raw AS pay_scale,
          e.gross_amount,
          e.sum_deduction,
          e.net_total
        FROM employees_v2 e
        WHERE {where}
        ORDER BY e.source_filename, e.sr_no, e.id
    """, params).fetchall()
    ids = [_row_get(row, "employee_row_id") for row in rows]
    components_by_employee: dict[int, dict[str, float]] = defaultdict(dict)
    if ids:
        ph = ",".join("?" for _ in ids)
        comp_rows = db.execute(f"""
            SELECT employee_id, normalized_head, SUM(amount) AS amount
            FROM employee_components_v2
            WHERE employee_id IN ({ph})
              AND COALESCE(normalized_head, '') != ''
            GROUP BY employee_id, normalized_head
        """, ids).fetchall()
        for row in comp_rows:
            components_by_employee[_row_get(row, "employee_id")][_row_get(row, "normalized_head")] = _as_float(_row_get(row, "amount"))
    snapshot = []
    for row in rows:
        item = dict(row)
        components = components_by_employee.get(item["employee_row_id"], {})
        item["components"] = components
        item["basic"] = components.get("basic_pay", 0.0)
        snapshot.append(item)
    return snapshot


def _match_key(row: dict[str, Any], strategy: str) -> tuple[str, str] | None:
    strategy = str(strategy or "smart").lower()
    fields = {
        "employee_id": "employee_id",
        "pan": "pan",
        "gpf_pran": "gpf_pran",
        "bank_account": "bank_account",
    }
    if strategy in fields:
        value = _norm(row.get(fields[strategy]))
        return (strategy, value) if value else None
    for key in ("employee_id", "pan", "gpf_pran", "bank_account"):
        value = _norm(row.get(key))
        if value:
            return (key, value)
    return None


def _index_snapshot(rows: list[dict[str, Any]], strategy: str) -> tuple[dict[tuple[str, str], dict[str, Any]], list[dict[str, Any]]]:
    index = {}
    weak = []
    for row in rows:
        key = _match_key(row, strategy)
        if key and key not in index:
            index[key] = row
        else:
            weak.append(row)
    return index, weak


def _identity_conflict(base: dict[str, Any], compare: dict[str, Any]) -> bool:
    if _norm(base.get("employee_id")) and _norm(compare.get("employee_id")) and _norm(base.get("employee_id")) != _norm(compare.get("employee_id")):
        return True
    if _norm(base.get("pan")) and _norm(compare.get("pan")) and _norm(base.get("pan")) != _norm(compare.get("pan")):
        return True
    return False


def _identity_changes(base: dict[str, Any] | None, compare: dict[str, Any] | None) -> dict[str, dict[str, Any]]:
    changes = {}
    if not base or not compare:
        return changes
    for out_key, field in IDENTITY_FIELDS.items():
        left = base.get(field, "")
        right = compare.get(field, "")
        if _norm(left) != _norm(right):
            changes[out_key] = {"base": left, "compare": right}
    return changes


def _component_changes(base: dict[str, Any] | None, compare: dict[str, Any] | None, amount_threshold: float) -> dict[str, dict[str, Any]]:
    base_components = (base or {}).get("components", {})
    compare_components = (compare or {}).get("components", {})
    heads = set(base_components) | set(compare_components) | set(CORE_COMPONENTS)
    changes = {}
    for head in sorted(heads):
        left = _as_float(base_components.get(head))
        right = _as_float(compare_components.get(head))
        diff = right - left
        if abs(diff) >= 0.01:
            status = "CHANGED"
            if abs(left) < 0.01 and abs(right) >= amount_threshold:
                status = "STARTED"
            elif abs(right) < 0.01 and abs(left) >= amount_threshold:
                status = "STOPPED"
            changes[head] = {"base": left, "compare": right, "difference": diff, "status": status}
    return changes


def _change_type(base: dict[str, Any] | None, compare: dict[str, Any] | None, match_status: str, component_changes: dict[str, Any], identity_changes: dict[str, Any], amount_threshold: float, percent_threshold: float) -> str:
    if match_status == "NEW_EMPLOYEE":
        return "NEW_EMPLOYEE"
    if match_status == "REMOVED_EMPLOYEE":
        return "REMOVED_EMPLOYEE"
    if match_status in {"POSSIBLE_MATCH_NEEDS_REVIEW", "IDENTITY_CONFLICT"}:
        return "NEEDS_REVIEW"
    if not base or not compare:
        return "NEEDS_REVIEW"
    if identity_changes:
        if "pay_scale" in identity_changes:
            return "PAY_SCALE_CHANGED"
        if "final_designation" in identity_changes:
            return "DESIGNATION_CHANGED"
        return "IDENTITY_CHANGED"
    if abs(_as_float(compare.get("basic")) - _as_float(base.get("basic"))) >= 0.01:
        return "BASIC_CHANGED"
    for item in component_changes.values():
        if item.get("status") == "STARTED":
            return "COMPONENT_STARTED"
        if item.get("status") == "STOPPED":
            return "COMPONENT_STOPPED"
    if _is_high_variance(base, compare, component_changes, amount_threshold, percent_threshold):
        return "HIGH_VARIANCE"
    if abs(_as_float(compare.get("gross_amount")) - _as_float(base.get("gross_amount"))) >= amount_threshold:
        return "GROSS_CHANGED"
    if abs(_as_float(compare.get("sum_deduction")) - _as_float(base.get("sum_deduction"))) >= amount_threshold:
        return "DEDUCTION_CHANGED"
    if abs(_as_float(compare.get("net_total")) - _as_float(base.get("net_total"))) >= amount_threshold:
        return "NET_CHANGED"
    if component_changes:
        return "DEDUCTION_CHANGED"
    return "NO_CHANGE"


def _is_high_variance(base: dict[str, Any], compare: dict[str, Any], component_changes: dict[str, Any], amount_threshold: float, percent_threshold: float) -> bool:
    for field in ("gross_amount", "net_total"):
        left = _as_float(base.get(field))
        right = _as_float(compare.get(field))
        if abs(right - left) >= amount_threshold and _pct_diff(left, right) >= percent_threshold:
            return True
    deduction_diff = _as_float(compare.get("sum_deduction")) - _as_float(base.get("sum_deduction"))
    if deduction_diff >= amount_threshold and _pct_diff(_as_float(base.get("sum_deduction")), _as_float(compare.get("sum_deduction"))) >= percent_threshold:
        return True
    for item in component_changes.values():
        if item.get("status") in {"STARTED", "STOPPED"} and abs(_as_float(item.get("difference"))) >= amount_threshold:
            return True
    return False


def _result_row(run_id: int, base: dict[str, Any] | None, compare: dict[str, Any] | None, match_status: str, amount_threshold: float, percent_threshold: float, issue: str = "") -> dict[str, Any]:
    component_changes = _component_changes(base, compare, amount_threshold)
    identity_changes = _identity_changes(base, compare)
    change_type = _change_type(base, compare, match_status, component_changes, identity_changes, amount_threshold, percent_threshold)
    review_status = "OPEN" if change_type in {"HIGH_VARIANCE", "NEEDS_REVIEW"} or match_status in {"POSSIBLE_MATCH_NEEDS_REVIEW", "IDENTITY_CONFLICT"} else ""
    base = base or {}
    compare = compare or {}
    return {
        "run_id": run_id,
        "match_status": match_status,
        "change_type": change_type,
        "employee_name": compare.get("employee_name") or base.get("employee_name") or "",
        "employee_id": compare.get("employee_id") or base.get("employee_id") or "",
        "pan": compare.get("pan") or base.get("pan") or "",
        "base_employee_row_id": base.get("employee_row_id"),
        "compare_employee_row_id": compare.get("employee_row_id"),
        "base_group": base.get("group_name", ""),
        "compare_group": compare.get("group_name", ""),
        "base_designation": base.get("final_designation", ""),
        "compare_designation": compare.get("final_designation", ""),
        "base_pay_scale": base.get("pay_scale", ""),
        "compare_pay_scale": compare.get("pay_scale", ""),
        "base_basic": _as_float(base.get("basic")),
        "compare_basic": _as_float(compare.get("basic")),
        "basic_difference": _as_float(compare.get("basic")) - _as_float(base.get("basic")),
        "base_gross": _as_float(base.get("gross_amount")),
        "compare_gross": _as_float(compare.get("gross_amount")),
        "gross_difference": _as_float(compare.get("gross_amount")) - _as_float(base.get("gross_amount")),
        "base_deduction": _as_float(base.get("sum_deduction")),
        "compare_deduction": _as_float(compare.get("sum_deduction")),
        "deduction_difference": _as_float(compare.get("sum_deduction")) - _as_float(base.get("sum_deduction")),
        "base_net": _as_float(base.get("net_total")),
        "compare_net": _as_float(compare.get("net_total")),
        "net_difference": _as_float(compare.get("net_total")) - _as_float(base.get("net_total")),
        "component_changes_json": json.dumps(component_changes, ensure_ascii=False, sort_keys=True),
        "identity_changes_json": json.dumps(identity_changes, ensure_ascii=False, sort_keys=True),
        "issue_message": issue,
        "review_status": review_status,
    }


def run_comparison(db, filters: dict[str, Any]) -> dict[str, Any]:
    init_monthly_comparison_schema(db)
    amount_threshold = _as_float(filters.get("amount_threshold"), 1000.0) or 1000.0
    percent_threshold = _as_float(filters.get("percent_threshold"), 10.0) or 10.0
    match_strategy = str(filters.get("match_strategy") or "smart").lower()
    base_rows = load_employee_snapshot(db, "base", filters)
    compare_rows = load_employee_snapshot(db, "compare", filters)
    run_name = str(filters.get("run_name") or "").strip() or f"Monthly Comparison {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    cur = db.execute("""
        INSERT INTO monthly_comparison_runs_v3 (
            run_name, base_month, base_year, base_group, base_source_file,
            compare_month, compare_year, compare_group, compare_source_file,
            match_strategy, amount_threshold, percent_threshold,
            base_employee_count, compare_employee_count
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        run_name,
        filters.get("base_month", ""), filters.get("base_year", ""), filters.get("base_group", ""), filters.get("base_source_file", ""),
        filters.get("compare_month", ""), filters.get("compare_year", ""), filters.get("compare_group", ""), filters.get("compare_source_file", ""),
        match_strategy, amount_threshold, percent_threshold, len(base_rows), len(compare_rows),
    ))
    run_id = cur.lastrowid
    base_index, base_weak = _index_snapshot(base_rows, match_strategy)
    compare_index, compare_weak = _index_snapshot(compare_rows, match_strategy)
    matched_base_keys = set()
    result_rows = []
    for key, compare in compare_index.items():
        base = base_index.get(key)
        if base:
            matched_base_keys.add(key)
            status = "IDENTITY_CONFLICT" if _identity_conflict(base, compare) else "MATCHED"
            issue = "Identity fields conflict for matched employee" if status == "IDENTITY_CONFLICT" else ""
            result_rows.append(_result_row(run_id, base, compare, status, amount_threshold, percent_threshold, issue))
        else:
            result_rows.append(_result_row(run_id, None, compare, "NEW_EMPLOYEE", amount_threshold, percent_threshold, "Employee appears only in comparison period"))
    for key, base in base_index.items():
        if key not in matched_base_keys and key not in compare_index:
            result_rows.append(_result_row(run_id, base, None, "REMOVED_EMPLOYEE", amount_threshold, percent_threshold, "Employee appears only in base period"))
    for row in compare_weak:
        candidate = _name_candidate(row, base_weak + base_rows)
        result_rows.append(_result_row(run_id, candidate, row, "POSSIBLE_MATCH_NEEDS_REVIEW", amount_threshold, percent_threshold, "Weak identity; review before matching"))
    for result in result_rows:
        _insert_result(db, result)
    summary = summarize_results(result_rows, len(base_rows), len(compare_rows))
    db.execute("""
        UPDATE monthly_comparison_runs_v3
        SET matched_count = ?, new_employee_count = ?, removed_employee_count = ?,
            changed_employee_count = ?, high_variance_count = ?, needs_review_count = ?,
            total_gross_difference = ?, total_deduction_difference = ?, total_net_difference = ?
        WHERE id = ?
    """, (
        summary["matched_count"], summary["new_employee_count"], summary["removed_employee_count"],
        summary["changed_employee_count"], summary["high_variance_count"], summary["needs_review_count"],
        summary["total_gross_difference"], summary["total_deduction_difference"], summary["total_net_difference"],
        run_id,
    ))
    summary.update({
        "run_id": run_id,
        "base_period": f"{filters.get('base_month','')}/{filters.get('base_year','')}".strip("/"),
        "comparison_period": f"{filters.get('compare_month','')}/{filters.get('compare_year','')}".strip("/"),
    })
    return summary


def _name_candidate(row: dict[str, Any], candidates: list[dict[str, Any]]) -> dict[str, Any] | None:
    row_key = (_norm(row.get("employee_name")), _norm(row.get("pay_scale")), _norm(row.get("final_designation")))
    for candidate in candidates:
        candidate_key = (_norm(candidate.get("employee_name")), _norm(candidate.get("pay_scale")), _norm(candidate.get("final_designation")))
        if row_key == candidate_key and row_key[0]:
            return candidate
    return None


def _insert_result(db, result: dict[str, Any]) -> None:
    columns = [
        "run_id", "match_status", "change_type", "employee_name", "employee_id", "pan",
        "base_employee_row_id", "compare_employee_row_id", "base_group", "compare_group",
        "base_designation", "compare_designation", "base_pay_scale", "compare_pay_scale",
        "base_basic", "compare_basic", "basic_difference", "base_gross", "compare_gross",
        "gross_difference", "base_deduction", "compare_deduction", "deduction_difference",
        "base_net", "compare_net", "net_difference", "component_changes_json",
        "identity_changes_json", "issue_message", "review_status",
    ]
    db.execute(
        "INSERT INTO monthly_comparison_results_v3 (" + ",".join(columns) + ") VALUES (" + ",".join("?" for _ in columns) + ")",
        [result.get(col) for col in columns],
    )


def summarize_results(rows: list[dict[str, Any]], base_count: int, compare_count: int) -> dict[str, Any]:
    return {
        "base_employee_count": base_count,
        "comparison_employee_count": compare_count,
        "matched_count": sum(1 for row in rows if row["match_status"] == "MATCHED"),
        "new_employee_count": sum(1 for row in rows if row["match_status"] == "NEW_EMPLOYEE"),
        "removed_employee_count": sum(1 for row in rows if row["match_status"] == "REMOVED_EMPLOYEE"),
        "changed_employee_count": sum(1 for row in rows if row["change_type"] != "NO_CHANGE"),
        "high_variance_count": sum(1 for row in rows if row["change_type"] == "HIGH_VARIANCE"),
        "needs_review_count": sum(1 for row in rows if row["review_status"] == "OPEN" or row["match_status"] in {"POSSIBLE_MATCH_NEEDS_REVIEW", "IDENTITY_CONFLICT"}),
        "total_gross_difference": sum(_as_float(row.get("gross_difference")) for row in rows if row["match_status"] == "MATCHED"),
        "total_deduction_difference": sum(_as_float(row.get("deduction_difference")) for row in rows if row["match_status"] == "MATCHED"),
        "total_net_difference": sum(_as_float(row.get("net_difference")) for row in rows if row["match_status"] == "MATCHED"),
    }


def list_runs(db) -> list[dict[str, Any]]:
    init_monthly_comparison_schema(db)
    return [dict(row) for row in db.execute("SELECT * FROM monthly_comparison_runs_v3 ORDER BY id DESC").fetchall()]


def results(db, args: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    init_monthly_comparison_schema(db)
    args = dict(args or {})
    where, params = [], []
    for key, column in [("run_id", "run_id"), ("status", "match_status"), ("change_type", "change_type"), ("review_status", "review_status")]:
        if args.get(key):
            where.append(f"{column} = ?")
            params.append(args[key])
    search = str(args.get("search") or args.get("employee_search") or "").strip()
    if search:
        where.append("(employee_name LIKE ? OR employee_id LIKE ? OR pan LIKE ?)")
        params.extend([f"%{search}%"] * 3)
    min_diff = args.get("amount_difference_gt") or args.get("min_difference")
    if min_diff not in (None, ""):
        where.append("(ABS(gross_difference) >= ? OR ABS(deduction_difference) >= ? OR ABS(net_difference) >= ?)")
        amount = _as_float(min_diff)
        params.extend([amount, amount, amount])
    sql = "SELECT * FROM monthly_comparison_results_v3"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY run_id DESC, CASE match_status WHEN 'IDENTITY_CONFLICT' THEN 1 WHEN 'POSSIBLE_MATCH_NEEDS_REVIEW' THEN 2 WHEN 'NEW_EMPLOYEE' THEN 3 WHEN 'REMOVED_EMPLOYEE' THEN 4 ELSE 9 END, employee_name"
    return [dict(row) for row in db.execute(sql, params).fetchall()]


def mark_reviewed(db, result_id: int) -> int:
    cur = db.execute("UPDATE monthly_comparison_results_v3 SET review_status = 'REVIEWED' WHERE id = ?", (result_id,))
    return cur.rowcount


def delete_run(db, run_id: int) -> None:
    db.execute("DELETE FROM monthly_comparison_results_v3 WHERE run_id = ?", (run_id,))
    db.execute("DELETE FROM monthly_comparison_runs_v3 WHERE id = ?", (run_id,))


def export_workbook(db, export_dir: str | Path, run_id: int, metadata: dict[str, Any]) -> Path:
    init_monthly_comparison_schema(db)
    run = db.execute("SELECT * FROM monthly_comparison_runs_v3 WHERE id = ?", (run_id,)).fetchone()
    if not run:
        raise ValueError("Monthly comparison run not found")
    rows = results(db, {"run_id": run_id})
    base_period = f"{_row_get(run, 'base_month', '')}_{_row_get(run, 'base_year', '')}".strip("_") or "base"
    compare_period = f"{_row_get(run, 'compare_month', '')}_{_row_get(run, 'compare_year', '')}".strip("_") or "comparison"
    file_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", f"Monthly_Comparison_{base_period}_vs_{compare_period}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    path = Path(export_dir) / file_name
    meta_rows = [
        {"Field": "Application", "Value": metadata.get("app_name", "IFMS Report Analyzer")},
        {"Field": "Designed and Developed by", "Value": "Neeraj Purohit"},
        {"Field": "Role", "Value": "Senior Assistant, Rajasthan Police"},
        {"Field": "Email", "Value": metadata.get("developer_email", "")},
        {"Field": "Base Period", "Value": f"{_row_get(run, 'base_month', '')}/{_row_get(run, 'base_year', '')}"},
        {"Field": "Comparison Period", "Value": f"{_row_get(run, 'compare_month', '')}/{_row_get(run, 'compare_year', '')}"},
        {"Field": "Thresholds", "Value": f"amount={_row_get(run, 'amount_threshold', '')}, percent={_row_get(run, 'percent_threshold', '')}"},
        {"Field": "Export Timestamp", "Value": datetime.now().isoformat(timespec="seconds")},
        {"Field": "App Build Label", "Value": metadata.get("build_label", "")},
        {"Field": "Disclaimer", "Value": metadata.get("disclaimer", "")},
    ]
    summary_rows = [dict(run)]
    sheets = {
        "Export Metadata": pd.DataFrame(meta_rows),
        "Comparison Summary": pd.DataFrame(summary_rows),
        "Matched Employees": pd.DataFrame([r for r in rows if r["match_status"] == "MATCHED"]),
        "New Employees": pd.DataFrame([r for r in rows if r["match_status"] == "NEW_EMPLOYEE"]),
        "Removed Employees": pd.DataFrame([r for r in rows if r["match_status"] == "REMOVED_EMPLOYEE"]),
        "Pay Changes": pd.DataFrame([r for r in rows if r["change_type"] in {"BASIC_CHANGED", "PAY_SCALE_CHANGED", "GROSS_CHANGED"}]),
        "Allowance Changes": pd.DataFrame([r for r in rows if "ALLOWANCE" in r["change_type"]]),
        "Deduction Changes": pd.DataFrame([r for r in rows if r["change_type"] in {"DEDUCTION_CHANGED", "COMPONENT_STARTED", "COMPONENT_STOPPED"}]),
        "Net Salary Changes": pd.DataFrame([r for r in rows if r["change_type"] == "NET_CHANGED"]),
        "Identity Changes": pd.DataFrame([r for r in rows if r["change_type"] in {"IDENTITY_CHANGED", "DESIGNATION_CHANGED", "PAY_SCALE_CHANGED", "NEEDS_REVIEW"}]),
        "High Variance": pd.DataFrame([r for r in rows if r["change_type"] == "HIGH_VARIANCE"]),
        "Needs Review": pd.DataFrame([r for r in rows if r["review_status"] == "OPEN"]),
        "Raw Comparison Rows": pd.DataFrame(rows),
        "Data Quality Notes": pd.DataFrame([
            {"Note": "Monthly comparison is read-only and does not modify Employee Data or component ledger tables."},
            {"Note": "Name-based matches are review candidates only."},
            {"Note": "Human review of the data is compulsory before using the data."},
        ]),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, index=False, sheet_name=safe_name)
            ws = writer.sheets[safe_name]
            ws.freeze_panes = "A2"
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 42)
    return path

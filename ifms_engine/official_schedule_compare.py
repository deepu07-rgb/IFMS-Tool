from __future__ import annotations

import hashlib
import os
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import fitz
import pandas as pd
from openpyxl.styles import Font, PatternFill

from ifms_engine.dictionary import extract_english_designation
from ifms_engine.schedule_engine_v3 import query_schedule_entries


SUPPORTED_SCHEDULES = {
    "income tax": "income_tax",
    "lic": "lic",
    "gpf contribution": "gpf",
    "gpf 2004": "gpf_2004",
    "si premium": "si_premium",
    "si loan": "si_loan",
    "sii": "sii",
    "rghs": "rghs",
    "advance salary": "advance_salary",
}

UNSUPPORTED_HINTS = {
    "road ways": "roadways_deduction",
    "roadways": "roadways_deduction",
    "ppss jdr": "ppss_jdr",
    "rop": "rop",
}

SUPPORTED_LAYOUT_TYPES = {
    "income_tax",
    "lic",
    "gpf",
    "gpf_2004",
    "si_premium",
    "si_loan",
    "sii",
    "rghs",
    "roadways_deduction",
    "advance_salary",
    "ppss_jdr",
}

NOISE_PREFIXES = (
    "RAJ IFMS",
    "Disclaimer:",
    "Print Date:",
    "Government of Rajasthan",
    "Page ",
)

FOOTER_PREFIXES = (
    "Grand Total:",
    "Rupees in Words:",
    "Sign ",
    "Signer:",
    "Reason:",
    "Location:",
)


@dataclass
class Line:
    text: str
    page_no: int


def _row_get(row: Any, key: str, default: Any = None) -> Any:
    try:
        return row[key]
    except (KeyError, IndexError, TypeError):
        return default


def _amount(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    text = str(value).replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def _amount_text(value: str) -> bool:
    return bool(re.fullmatch(r"-?\d{1,3}(?:,\d{2,3})*\.\d{2}|-?\d+\.\d{2}", value.strip()))


def _int_text(value: str) -> bool:
    return bool(re.fullmatch(r"\d{1,4}", value.strip()))


def _employee_id(value: str) -> bool:
    return bool(re.fullmatch(r"RJ[A-Z]{2}\d{8,12}", value.strip(), re.I))


def _pan(value: str) -> bool:
    return bool(re.fullmatch(r"[A-Z]{5}\d{4}[A-Z]", value.strip(), re.I))


def _clean_lines(lines: list[Line]) -> list[Line]:
    cleaned = []
    for line in lines:
        text = line.text.strip()
        if not text:
            continue
        if any(text.startswith(prefix) for prefix in NOISE_PREFIXES):
            continue
        cleaned.append(Line(text, line.page_no))
    return cleaned


def _heading_name(heading: str) -> str:
    name = re.sub(r"^Schedule\s+of\s+", "", heading, flags=re.I).strip()
    name = re.sub(r"\s*\(Budget Head.*$", "", name, flags=re.I).strip()
    return name


def schedule_type_from_heading(heading: str) -> tuple[str, str]:
    lower = _heading_name(heading).lower()
    if lower in SUPPORTED_SCHEDULES:
        return SUPPORTED_SCHEDULES[lower], "SUPPORTED"
    if "road ways" in lower or "roadways" in lower:
        return "roadways_deduction", "SUPPORTED"
    if "ppss jdr" in lower:
        return "ppss_jdr", "SUPPORTED"
    for hint, key in UNSUPPORTED_HINTS.items():
        if hint in lower:
            return key, "NEEDS_LAYOUT_MAPPING"
    return re.sub(r"[^a-z0-9]+", "_", lower).strip("_") or "unknown", "NEEDS_LAYOUT_MAPPING"


def classify_pdf(path: str | Path) -> dict[str, Any]:
    doc = fitz.open(path)
    first_pages = "\n".join(doc[i].get_text("text") for i in range(min(doc.page_count, 2)))
    headings = re.findall(r"^\s*Schedule\s+of\s+.+$", first_pages, flags=re.I | re.M)
    if headings:
        kind = "official_schedule"
    elif "Salary Bill (Regular (Inner Sheet))" in first_pages:
        kind = "inner_sheet"
    else:
        kind = "unknown"
    return {"classification": kind, "pages_total": doc.page_count, "headings": [h.strip() for h in headings]}


def extract_pdf_lines(path: str | Path) -> tuple[list[Line], int]:
    doc = fitz.open(path)
    lines: list[Line] = []
    for page_index, page in enumerate(doc, start=1):
        for raw in page.get_text("text").splitlines():
            text = raw.strip()
            if text:
                lines.append(Line(text, page_index))
    return lines, doc.page_count


def detect_sections(path: str | Path) -> tuple[list[dict[str, Any]], int]:
    lines, pages_total = extract_pdf_lines(path)
    heading_indexes = [
        idx for idx, line in enumerate(lines)
        if re.match(r"^Schedule\s+of\s+.+$", line.text, flags=re.I)
    ]
    sections: list[dict[str, Any]] = []
    for pos, start_idx in enumerate(heading_indexes):
        end_idx = heading_indexes[pos + 1] if pos + 1 < len(heading_indexes) else len(lines)
        section_lines = lines[start_idx:end_idx]
        heading = section_lines[0].text
        schedule_type, parse_status = schedule_type_from_heading(heading)
        section = {
            "schedule_type": schedule_type,
            "schedule_heading": heading,
            "budget_head": _search_section(section_lines, r"Budget Head\s+([^)]+)"),
            "office_name": _search_section(section_lines, r"Office Name:\s*(.+)"),
            "ddo_code": _search_section(section_lines, r"DDO Code:\s*(.+)"),
            "service_head": _search_section(section_lines, r"Service Head:\s*(.+)"),
            "bill_no": "",
            "bill_date": "",
            "month_year": _search_section(section_lines, r"Month/Year:\s*(.+)"),
            "start_page": section_lines[0].page_no,
            "end_page": section_lines[-1].page_no,
            "grand_total": _parse_grand_total(section_lines),
            "parse_status": parse_status,
            "raw_section_text": "\n".join(line.text for line in section_lines),
            "_lines": section_lines,
        }
        bill = _search_section(section_lines, r"Bill No\.\s*&\s*Date:\s*(.+)")
        if "&" in bill:
            left, right = bill.split("&", 1)
            section["bill_no"] = left.strip()
            section["bill_date"] = right.strip()
        else:
            section["bill_no"] = bill.strip()
        sections.append(section)
    return sections, pages_total


def _search_section(lines: list[Line], pattern: str) -> str:
    text = "\n".join(line.text for line in lines[:80])
    match = re.search(pattern, text, flags=re.I)
    return match.group(1).strip() if match else ""


def _parse_grand_total(lines: list[Line]) -> float:
    for idx, line in enumerate(lines):
        if line.text.strip().lower().startswith("grand total"):
            amounts = [_amount(candidate.text) for candidate in lines[idx + 1:idx + 8] if _amount_text(candidate.text)]
            return amounts[-1] if amounts else 0.0
    return 0.0


def parse_official_pdf(path: str | Path, source_file: str | None = None) -> dict[str, Any]:
    classification = classify_pdf(path)
    sections, pages_total = detect_sections(path)
    parsed_sections = []
    rows: list[dict[str, Any]] = []
    for section in sections:
        section_rows: list[dict[str, Any]] = []
        if section["parse_status"] == "SUPPORTED":
            section_rows = parse_section_rows(section, source_file or Path(path).name)
        section["rows_detected"] = len(section_rows)
        parsed_sections.append({k: v for k, v in section.items() if k != "_lines"})
        rows.extend(section_rows)
    return {
        "classification": classification["classification"],
        "pages_total": pages_total,
        "sections": parsed_sections,
        "rows": rows,
        "sections_detected": len(sections),
        "rows_detected": len(rows),
    }


def parse_section_rows(section: dict[str, Any], source_file: str) -> list[dict[str, Any]]:
    lines = _clean_lines(section.get("_lines", []))
    starts = []
    for idx, line in enumerate(lines[:-1]):
        if _int_text(line.text) and _employee_id(lines[idx + 1].text):
            starts.append(idx)
    rows = []
    for pos, start in enumerate(starts):
        end = starts[pos + 1] if pos + 1 < len(starts) else len(lines)
        block = lines[start:end]
        row = _parse_row_block(section, block, source_file)
        if row:
            rows.append(row)
    return rows


def _parse_row_block(section: dict[str, Any], block: list[Line], source_file: str) -> dict[str, Any] | None:
    trimmed_block = []
    for line in block:
        if any(line.text.strip().startswith(prefix) for prefix in FOOTER_PREFIXES):
            break
        trimmed_block.append(line)
    block = trimmed_block
    texts = [line.text.strip() for line in block]
    if len(texts) < 4 or not _int_text(texts[0]) or not _employee_id(texts[1]):
        return None
    schedule_type = section["schedule_type"]
    try:
        if schedule_type == "income_tax":
            row = _parse_income_tax(texts)
        elif schedule_type == "lic":
            row = _parse_lic(texts)
        elif schedule_type == "gpf":
            row = _parse_gpf(texts)
        elif schedule_type == "gpf_2004":
            row = _parse_gpf_2004(texts)
        elif schedule_type == "si_premium":
            row = _parse_si_premium(texts)
        elif schedule_type == "si_loan":
            row = _parse_simple_amount_schedule(texts, schedule_type)
        elif schedule_type == "sii":
            row = _parse_simple_amount_schedule(texts, schedule_type)
        elif schedule_type == "rghs":
            row = _parse_basic_designation_amount(texts, schedule_type)
        elif schedule_type == "roadways_deduction":
            row = _parse_cooperative_amount(texts, schedule_type)
        elif schedule_type == "advance_salary":
            row = _parse_advance_salary(texts)
        elif schedule_type == "ppss_jdr":
            row = _parse_cooperative_amount(texts, schedule_type)
        else:
            return None
    except ValueError:
        return _review_row(section, block, source_file)
    row.update({
        "schedule_type": schedule_type,
        "source_file": source_file,
        "page_no": block[0].page_no,
        "designation_en": extract_english_designation(row.get("designation_raw", "")),
        "raw_text_block": "\n".join(texts),
        "confidence": row.get("confidence", 1.0),
        "row_status": row.get("row_status", "PARSED"),
    })
    if not row.get("employee_id") or not row.get("total_amount"):
        row["confidence"] = min(float(row.get("confidence") or 0.6), 0.6)
        row["row_status"] = "NEEDS_REVIEW"
    return row


def _parse_income_tax(texts: list[str]) -> dict[str, Any]:
    pan_idx = next((i for i, item in enumerate(texts) if _pan(item)), -1)
    if pan_idx < 4 or len(texts) < pan_idx + 7:
        raise ValueError("income tax row layout not recognized")
    nums = [item for item in texts[pan_idx + 1:] if _amount_text(item)]
    if len(nums) < 6:
        raise ValueError("income tax amount columns missing")
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": " ".join(texts[3:pan_idx]),
        "pan": texts[pan_idx],
        "basic_pay": _amount(nums[0]),
        "component_amount": _amount(nums[1]),
        "component_amount_2": _amount(nums[2]) + _amount(nums[3]) + _amount(nums[4]),
        "total_amount": _amount(nums[5]),
    }


def _parse_lic(texts: list[str]) -> dict[str, Any]:
    amount_indexes = [i for i, item in enumerate(texts) if _amount_text(item)]
    if len(amount_indexes) < 3:
        raise ValueError("lic amount columns missing")
    basic_idx, premium_idx, total_idx = amount_indexes[0], amount_indexes[1], amount_indexes[2]
    policy_idx = basic_idx + 1
    if policy_idx >= len(texts):
        raise ValueError("lic policy column missing")
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": " ".join(texts[3:basic_idx]),
        "basic_pay": _amount(texts[basic_idx]),
        "policy_no": texts[policy_idx],
        "component_amount": _amount(texts[premium_idx]),
        "total_amount": _amount(texts[total_idx]),
    }


def _parse_gpf(texts: list[str]) -> dict[str, Any]:
    amount_indexes = [i for i, item in enumerate(texts) if _amount_text(item)]
    if len(amount_indexes) < 3:
        raise ValueError("gpf amount columns missing")
    basic_idx, contrib_idx, loan_idx = amount_indexes[0], amount_indexes[1], amount_indexes[2]
    total_idx = amount_indexes[3] if len(amount_indexes) > 3 else loan_idx
    gpf_idx = basic_idx - 1
    if gpf_idx < 3:
        raise ValueError("gpf number column missing")
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": " ".join(texts[3:gpf_idx]),
        "gpf_no": texts[gpf_idx],
        "basic_pay": _amount(texts[basic_idx]),
        "component_amount": _amount(texts[contrib_idx]),
        "component_amount_2": _amount(texts[loan_idx]),
        "total_amount": _amount(texts[total_idx]),
    }


def _parse_gpf_2004(texts: list[str]) -> dict[str, Any]:
    amount_indexes = [i for i, item in enumerate(texts) if _amount_text(item)]
    if len(amount_indexes) < 3:
        raise ValueError("gpf 2004 amount columns missing")
    basic_idx, contribution_idx, total_idx = amount_indexes[0], amount_indexes[1], amount_indexes[-1]
    loan_text = texts[contribution_idx + 1] if contribution_idx + 1 < total_idx else ""
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": " ".join(texts[3:basic_idx]),
        "basic_pay": _amount(texts[basic_idx]),
        "component_amount": _amount(texts[contribution_idx]),
        "component_amount_2": _amount(loan_text),
        "total_amount": _amount(texts[total_idx]),
    }


def _parse_si_premium(texts: list[str]) -> dict[str, Any]:
    amount_indexes = [i for i, item in enumerate(texts) if _amount_text(item)]
    if len(amount_indexes) < 3:
        raise ValueError("si premium amount columns missing")
    basic_idx = amount_indexes[0]
    amounts = [_amount(texts[i]) for i in amount_indexes[1:]]
    total_amount = amounts[-1]
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": "",
        "policy_no": texts[3] if basic_idx > 3 else "",
        "basic_pay": _amount(texts[basic_idx]),
        "component_amount": amounts[0] if amounts else 0,
        "component_amount_2": sum(amounts[1:-1]) if len(amounts) > 2 else 0,
        "total_amount": total_amount,
    }


def _parse_simple_amount_schedule(texts: list[str], schedule_type: str) -> dict[str, Any]:
    amount_indexes = [i for i, item in enumerate(texts) if _amount_text(item)]
    if len(amount_indexes) < 2:
        raise ValueError(f"{schedule_type} amount columns missing")
    basic_idx = amount_indexes[0]
    total_idx = amount_indexes[-1]
    policy_no = ""
    if basic_idx > 3 and not any(ch in texts[basic_idx - 1] for ch in "()"):
        policy_no = texts[basic_idx - 1]
        designation = " ".join(texts[3:basic_idx - 1])
    else:
        designation = " ".join(texts[3:basic_idx])
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": designation,
        "policy_no": policy_no,
        "basic_pay": _amount(texts[basic_idx]),
        "component_amount": _amount(texts[amount_indexes[1]]),
        "component_amount_2": sum(_amount(texts[i]) for i in amount_indexes[2:-1]),
        "total_amount": _amount(texts[total_idx]),
    }


def _parse_basic_designation_amount(texts: list[str], schedule_type: str) -> dict[str, Any]:
    amount_indexes = [i for i, item in enumerate(texts) if _amount_text(item)]
    if len(amount_indexes) < 2:
        raise ValueError(f"{schedule_type} amount columns missing")
    basic_idx = amount_indexes[0]
    total_idx = amount_indexes[-1]
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": " ".join(texts[3:basic_idx]),
        "basic_pay": _amount(texts[basic_idx]),
        "component_amount": _amount(texts[amount_indexes[1]]),
        "component_amount_2": sum(_amount(texts[i]) for i in amount_indexes[2:-1]),
        "total_amount": _amount(texts[total_idx]),
    }


def _parse_cooperative_amount(texts: list[str], schedule_type: str) -> dict[str, Any]:
    amount_indexes = [i for i, item in enumerate(texts) if _amount_text(item)]
    if not amount_indexes:
        raise ValueError(f"{schedule_type} deduction amount missing")
    total_idx = amount_indexes[-1]
    pre_amount = texts[3:total_idx]
    gpf_no = pre_amount[-1] if pre_amount else ""
    identity_columns = 2 if schedule_type == "ppss_jdr" else 1
    designation_start = min(3 + identity_columns, len(texts))
    designation_end = max(designation_start, total_idx - 1)
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": " ".join(texts[designation_start:designation_end]),
        "gpf_no": gpf_no,
        "basic_pay": 0,
        "component_amount": _amount(texts[total_idx]),
        "total_amount": _amount(texts[total_idx]),
    }


def _parse_advance_salary(texts: list[str]) -> dict[str, Any]:
    amount_indexes = [i for i, item in enumerate(texts) if _amount_text(item)]
    if not amount_indexes:
        raise ValueError("advance salary amount missing")
    total_idx = amount_indexes[-1]
    return {
        "sr_no": int(texts[0]),
        "employee_id": texts[1],
        "employee_name": texts[2],
        "designation_raw": " ".join(texts[4:total_idx]) if len(texts) > 4 else "",
        "policy_no": texts[3] if len(texts) > 3 else "",
        "basic_pay": 0,
        "component_amount": _amount(texts[total_idx]),
        "total_amount": _amount(texts[total_idx]),
    }


def _review_row(section: dict[str, Any], block: list[Line], source_file: str) -> dict[str, Any]:
    texts = [line.text.strip() for line in block]
    return {
        "sr_no": int(texts[0]) if texts and _int_text(texts[0]) else None,
        "employee_id": texts[1] if len(texts) > 1 else "",
        "employee_name": texts[2] if len(texts) > 2 else "",
        "designation_raw": "",
        "designation_en": "",
        "schedule_type": section["schedule_type"],
        "source_file": source_file,
        "page_no": block[0].page_no if block else section.get("start_page", 0),
        "raw_text_block": "\n".join(texts),
        "confidence": 0.4,
        "row_status": "NEEDS_REVIEW",
    }


def init_official_schedule_schema(db) -> None:
    db.executescript("""
    CREATE TABLE IF NOT EXISTS official_schedule_uploads_v3 (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        source_file TEXT,
        original_file_name TEXT,
        file_hash TEXT,
        pages_total INTEGER DEFAULT 0,
        sections_detected INTEGER DEFAULT 0,
        rows_detected INTEGER DEFAULT 0,
        parsed_at TEXT,
        status TEXT,
        error_message TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS official_schedule_sections_v3 (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        upload_id INTEGER,
        schedule_type TEXT,
        schedule_heading TEXT,
        budget_head TEXT,
        office_name TEXT,
        ddo_code TEXT,
        service_head TEXT,
        bill_no TEXT,
        bill_date TEXT,
        month_year TEXT,
        start_page INTEGER,
        end_page INTEGER,
        rows_detected INTEGER DEFAULT 0,
        grand_total REAL DEFAULT 0,
        parse_status TEXT,
        raw_section_text TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS official_schedule_rows_v3 (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        upload_id INTEGER,
        section_id INTEGER,
        schedule_type TEXT,
        source_file TEXT,
        page_no INTEGER,
        sr_no INTEGER,
        employee_name TEXT,
        employee_id TEXT,
        designation_raw TEXT,
        designation_en TEXT,
        pan TEXT,
        gpf_no TEXT,
        policy_no TEXT,
        basic_pay REAL DEFAULT 0,
        component_amount REAL DEFAULT 0,
        component_amount_2 REAL DEFAULT 0,
        total_amount REAL DEFAULT 0,
        raw_text_block TEXT,
        confidence REAL DEFAULT 0,
        row_status TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS official_schedule_comparison_v3 (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        upload_id INTEGER,
        section_id INTEGER,
        schedule_type TEXT,
        comparison_scope TEXT,
        derived_employee_row_id INTEGER,
        official_row_id INTEGER,
        employee_id TEXT,
        employee_name TEXT,
        pan TEXT,
        official_amount REAL DEFAULT 0,
        derived_amount REAL DEFAULT 0,
        difference REAL DEFAULT 0,
        match_status TEXT,
        match_confidence REAL DEFAULT 0,
        issue_message TEXT,
        reviewed_status TEXT DEFAULT 'OPEN',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        reviewed_at TEXT
    );
    CREATE INDEX IF NOT EXISTS idx_official_upload_hash ON official_schedule_uploads_v3(file_hash);
    CREATE INDEX IF NOT EXISTS idx_official_sections_upload ON official_schedule_sections_v3(upload_id);
    CREATE INDEX IF NOT EXISTS idx_official_rows_upload ON official_schedule_rows_v3(upload_id, section_id);
    CREATE INDEX IF NOT EXISTS idx_official_rows_employee ON official_schedule_rows_v3(employee_id);
    CREATE INDEX IF NOT EXISTS idx_official_cmp_upload ON official_schedule_comparison_v3(upload_id, section_id, schedule_type);
    """)


def store_official_upload(db, pdf_path: str | Path, original_file_name: str, source_file: str) -> dict[str, Any]:
    init_official_schedule_schema(db)
    file_hash = hashlib.sha256(Path(pdf_path).read_bytes()).hexdigest()
    parsed = parse_official_pdf(pdf_path, source_file)
    status = "PARSED" if parsed["classification"] == "official_schedule" else "REJECTED"
    if status == "REJECTED":
        raise ValueError("Uploaded PDF is not classified as an official schedule PDF")
    cur = db.execute("""
        INSERT INTO official_schedule_uploads_v3 (
            source_file, original_file_name, file_hash, pages_total, sections_detected,
            rows_detected, parsed_at, status, error_message
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        source_file, original_file_name, file_hash, parsed["pages_total"],
        parsed["sections_detected"], parsed["rows_detected"],
        datetime.now().isoformat(timespec="seconds"), status, "",
    ))
    upload_id = cur.lastrowid
    section_id_by_position: dict[int, int] = {}
    for pos, section in enumerate(parsed["sections"]):
        cur = db.execute("""
            INSERT INTO official_schedule_sections_v3 (
                upload_id, schedule_type, schedule_heading, budget_head, office_name,
                ddo_code, service_head, bill_no, bill_date, month_year, start_page,
                end_page, rows_detected, grand_total, parse_status, raw_section_text
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            upload_id, section.get("schedule_type", ""), section.get("schedule_heading", ""),
            section.get("budget_head", ""), section.get("office_name", ""),
            section.get("ddo_code", ""), section.get("service_head", ""),
            section.get("bill_no", ""), section.get("bill_date", ""), section.get("month_year", ""),
            section.get("start_page", 0), section.get("end_page", 0), section.get("rows_detected", 0),
            section.get("grand_total", 0), section.get("parse_status", ""), section.get("raw_section_text", ""),
        ))
        section_id_by_position[pos] = cur.lastrowid
    section_positions = [s for s in parsed["sections"] if s.get("parse_status") == "SUPPORTED"]
    section_id_by_type: dict[str, list[int]] = defaultdict(list)
    for pos, section in enumerate(parsed["sections"]):
        section_id_by_type[section.get("schedule_type", "")].append(section_id_by_position[pos])
    used_by_type: dict[str, int] = defaultdict(int)
    for row in parsed["rows"]:
        stype = row.get("schedule_type", "")
        section_ids = section_id_by_type.get(stype) or []
        section_id = section_ids[min(used_by_type[stype], len(section_ids) - 1)] if section_ids else None
        db.execute("""
            INSERT INTO official_schedule_rows_v3 (
                upload_id, section_id, schedule_type, source_file, page_no, sr_no,
                employee_name, employee_id, designation_raw, designation_en, pan,
                gpf_no, policy_no, basic_pay, component_amount, component_amount_2,
                total_amount, raw_text_block, confidence, row_status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            upload_id, section_id, stype, row.get("source_file", source_file),
            row.get("page_no", 0), row.get("sr_no"), row.get("employee_name", ""),
            row.get("employee_id", ""), row.get("designation_raw", ""), row.get("designation_en", ""),
            row.get("pan", ""), row.get("gpf_no", ""), row.get("policy_no", ""),
            row.get("basic_pay", 0), row.get("component_amount", 0), row.get("component_amount_2", 0),
            row.get("total_amount", 0), row.get("raw_text_block", ""), row.get("confidence", 0),
            row.get("row_status", "PARSED"),
        ))
    return upload_summary(db, upload_id)


def upload_summary(db, upload_id: int) -> dict[str, Any]:
    upload = db.execute("SELECT * FROM official_schedule_uploads_v3 WHERE id = ?", (upload_id,)).fetchone()
    if not upload:
        raise ValueError("Official schedule upload not found")
    sections = [dict(row) for row in db.execute(
        "SELECT * FROM official_schedule_sections_v3 WHERE upload_id = ? ORDER BY id", (upload_id,)
    ).fetchall()]
    supported = [row for row in sections if row.get("parse_status") == "SUPPORTED"]
    unsupported = [row for row in sections if row.get("parse_status") != "SUPPORTED"]
    return {
        "upload_id": upload_id,
        "source_file": _row_get(upload, "source_file", ""),
        "pages_total": _row_get(upload, "pages_total", 0),
        "sections_detected": _row_get(upload, "sections_detected", 0),
        "supported_sections": supported,
        "unsupported_sections": unsupported,
        "rows_detected": _row_get(upload, "rows_detected", 0),
        "status": _row_get(upload, "status", ""),
        "warnings": [f"{row['schedule_heading']} needs layout mapping" for row in unsupported],
    }


def official_aggregates(db, upload_id: int, section_id: int | None = None, schedule_type: str | None = None) -> list[dict[str, Any]]:
    params: list[Any] = [upload_id]
    where = ["upload_id = ?", "COALESCE(row_status, '') != 'NEEDS_REVIEW'"]
    if section_id:
        where.append("section_id = ?")
        params.append(section_id)
    if schedule_type:
        where.append("schedule_type = ?")
        params.append(schedule_type)
    rows = [dict(row) for row in db.execute(
        f"SELECT * FROM official_schedule_rows_v3 WHERE {' AND '.join(where)} ORDER BY employee_id, id",
        params,
    ).fetchall()]
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for row in rows:
        key = (row.get("schedule_type") or "", row.get("employee_id") or row.get("pan") or row.get("employee_name") or "")
        current = grouped.setdefault(key, {
            "schedule_type": row.get("schedule_type", ""),
            "official_row_id": row.get("id"),
            "employee_id": row.get("employee_id", ""),
            "employee_name": row.get("employee_name", ""),
            "pan": row.get("pan", ""),
            "gpf_no": row.get("gpf_no", ""),
            "basic_pay": row.get("basic_pay", 0),
            "official_amount": 0.0,
            "row_count": 0,
        })
        current["official_amount"] += _amount(row.get("total_amount"))
        current["row_count"] += 1
    return list(grouped.values())


def run_comparison(db, upload_id: int, filters: dict[str, Any] | None = None) -> dict[str, Any]:
    filters = dict(filters or {})
    init_official_schedule_schema(db)
    section_id = int(filters.get("section_id") or 0) or None
    schedule_type = filters.get("schedule_type") or None
    if section_id and not schedule_type:
        section = db.execute("SELECT schedule_type FROM official_schedule_sections_v3 WHERE id = ?", (section_id,)).fetchone()
        schedule_type = _row_get(section, "schedule_type", None)
    db.execute(
        "DELETE FROM official_schedule_comparison_v3 WHERE upload_id = ? AND (? IS NULL OR section_id = ?) AND (? IS NULL OR schedule_type = ?)",
        (upload_id, section_id, section_id, schedule_type, schedule_type),
    )
    official = official_aggregates(db, upload_id, section_id, schedule_type)
    official_by_id = {(row.get("schedule_type"), row.get("employee_id")): row for row in official if row.get("employee_id")}
    schedule_types = sorted({row["schedule_type"] for row in official} | ({schedule_type} if schedule_type else set()))
    derived: list[dict[str, Any]] = []
    for stype in schedule_types:
        if stype not in SUPPORTED_LAYOUT_TYPES:
            continue
        args = {k: v for k, v in filters.items() if k in {"group_name", "source_file", "bill_no", "month_year", "show_deleted"} and v not in ("", None)}
        args["schedule_type"] = stype
        derived.extend(query_schedule_entries(db, args, paginate=False).get("entries", []))
    derived_by_id = {(row.get("schedule_type"), row.get("employee_id")): row for row in derived if row.get("employee_id")}
    results: list[dict[str, Any]] = []
    for key, off in official_by_id.items():
        der = derived_by_id.get(key)
        if der:
            diff = round(_amount(off["official_amount"]) - _amount(der.get("amount")), 2)
            status = "MATCHED" if abs(diff) < 0.01 else "AMOUNT_MISMATCH"
            issue = "" if status == "MATCHED" else "Official amount differs from derived schedule amount"
            confidence = 1.0
        else:
            diff = round(_amount(off["official_amount"]), 2)
            status = "EXTRA_IN_OFFICIAL"
            issue = "Official row has no derived schedule employee match"
            confidence = 0.9
        results.append(_comparison_row(upload_id, section_id, off, der, diff, status, confidence, issue))
    for key, der in derived_by_id.items():
        if key in official_by_id:
            continue
        off = {"schedule_type": der.get("schedule_type", ""), "employee_id": der.get("employee_id", ""), "employee_name": der.get("employee_name", ""), "pan": der.get("pan", ""), "official_amount": 0}
        diff = round(0 - _amount(der.get("amount")), 2)
        results.append(_comparison_row(upload_id, section_id, off, der, diff, "MISSING_IN_OFFICIAL", 0.9, "Derived schedule row is missing from official schedule"))
    for result in results:
        db.execute("""
            INSERT INTO official_schedule_comparison_v3 (
                upload_id, section_id, schedule_type, comparison_scope, derived_employee_row_id,
                official_row_id, employee_id, employee_name, pan, official_amount, derived_amount,
                difference, match_status, match_confidence, issue_message, reviewed_status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            result["upload_id"], result["section_id"], result["schedule_type"], result["comparison_scope"],
            result.get("derived_employee_row_id"), result.get("official_row_id"), result.get("employee_id", ""),
            result.get("employee_name", ""), result.get("pan", ""), result.get("official_amount", 0),
            result.get("derived_amount", 0), result.get("difference", 0), result.get("match_status", ""),
            result.get("match_confidence", 0), result.get("issue_message", ""), "OPEN",
        ))
    return comparison_summary(db, upload_id, section_id, schedule_type)


def _comparison_row(upload_id: int, section_id: int | None, off: dict[str, Any], der: dict[str, Any] | None, diff: float, status: str, confidence: float, issue: str) -> dict[str, Any]:
    return {
        "upload_id": upload_id,
        "section_id": section_id,
        "schedule_type": off.get("schedule_type") or (der or {}).get("schedule_type", ""),
        "comparison_scope": "employee_aggregate",
        "derived_employee_row_id": (der or {}).get("employee_db_id"),
        "official_row_id": off.get("official_row_id"),
        "employee_id": off.get("employee_id") or (der or {}).get("employee_id", ""),
        "employee_name": off.get("employee_name") or (der or {}).get("employee_name", ""),
        "pan": off.get("pan") or (der or {}).get("pan", ""),
        "official_amount": _amount(off.get("official_amount")),
        "derived_amount": _amount((der or {}).get("amount")),
        "difference": diff,
        "match_status": status,
        "match_confidence": confidence,
        "issue_message": issue,
    }


def comparison_rows(db, args: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    args = dict(args or {})
    params = []
    where = []
    for key, column in [("upload_id", "upload_id"), ("section_id", "section_id"), ("schedule_type", "schedule_type"), ("status", "match_status")]:
        value = args.get(key)
        if value not in (None, ""):
            where.append(f"{column} = ?")
            params.append(value)
    search = str(args.get("search") or args.get("employee_search") or "").strip()
    if search:
        where.append("(employee_id LIKE ? OR employee_name LIKE ? OR pan LIKE ?)")
        params.extend([f"%{search}%"] * 3)
    sql = "SELECT * FROM official_schedule_comparison_v3"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY CASE match_status WHEN 'AMOUNT_MISMATCH' THEN 1 WHEN 'MISSING_IN_OFFICIAL' THEN 2 WHEN 'EXTRA_IN_OFFICIAL' THEN 3 WHEN 'NEEDS_REVIEW' THEN 4 ELSE 9 END, employee_name"
    return [dict(row) for row in db.execute(sql, params).fetchall()]


def comparison_summary(db, upload_id: int, section_id: int | None = None, schedule_type: str | None = None) -> dict[str, Any]:
    rows = comparison_rows(db, {"upload_id": upload_id, "section_id": section_id or "", "schedule_type": schedule_type or ""})
    official_total = sum(_amount(row.get("official_amount")) for row in rows if row.get("match_status") != "MISSING_IN_OFFICIAL")
    derived_total = sum(_amount(row.get("derived_amount")) for row in rows if row.get("match_status") != "EXTRA_IN_OFFICIAL")
    counts = defaultdict(int)
    for row in rows:
        counts[row.get("match_status") or ""] += 1
    return {
        "upload_id": upload_id,
        "schedule_type": schedule_type or "all",
        "official_rows": len([row for row in rows if row.get("match_status") != "MISSING_IN_OFFICIAL"]),
        "official_aggregate_rows": len([row for row in rows if row.get("official_row_id")]),
        "derived_rows": len([row for row in rows if row.get("derived_employee_row_id")]),
        "matched_count": counts["MATCHED"],
        "missing_in_official_count": counts["MISSING_IN_OFFICIAL"],
        "extra_in_official_count": counts["EXTRA_IN_OFFICIAL"],
        "amount_mismatch_count": counts["AMOUNT_MISMATCH"],
        "identity_mismatch_count": counts["IDENTITY_MISMATCH"],
        "needs_review_count": counts["NEEDS_REVIEW"] + counts["LOW_CONFIDENCE_MATCH"] + counts["OFFICIAL_LAYOUT_NEEDS_REVIEW"],
        "official_total": official_total,
        "derived_total": derived_total,
        "difference_total": round(official_total - derived_total, 2),
        "status": "MATCHED" if rows and abs(official_total - derived_total) < 0.01 and counts["MATCHED"] == len(rows) else "NEEDS_REVIEW",
    }


def export_comparison_workbook(db, export_dir: str | Path, upload_id: int, filters: dict[str, Any], metadata: dict[str, Any]) -> Path:
    filters = dict(filters or {})
    section_id = int(filters.get("section_id") or 0) or None
    schedule_type = filters.get("schedule_type") or None
    rows = comparison_rows(db, {"upload_id": upload_id, "section_id": section_id or "", "schedule_type": schedule_type or ""})
    if not rows:
        run_comparison(db, upload_id, filters)
        rows = comparison_rows(db, {"upload_id": upload_id, "section_id": section_id or "", "schedule_type": schedule_type or ""})
    upload = db.execute("SELECT * FROM official_schedule_uploads_v3 WHERE id = ?", (upload_id,)).fetchone()
    sections = [dict(row) for row in db.execute("SELECT * FROM official_schedule_sections_v3 WHERE upload_id = ? ORDER BY id", (upload_id,)).fetchall()]
    official_rows = [dict(row) for row in db.execute("SELECT * FROM official_schedule_rows_v3 WHERE upload_id = ? ORDER BY schedule_type, employee_name, id", (upload_id,)).fetchall()]
    aggregates = official_aggregates(db, upload_id, section_id, schedule_type)
    derived = []
    for stype in sorted({row.get("schedule_type") for row in rows if row.get("schedule_type")}):
        derived.extend(query_schedule_entries(db, {"schedule_type": stype}, paginate=False).get("entries", []))
    summary = comparison_summary(db, upload_id, section_id, schedule_type)
    bill_no = next((s.get("bill_no") for s in sections if s.get("bill_no")), "all")
    prefix = f"Official_Schedule_Comparison_{bill_no}_{schedule_type or 'all'}"
    file_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", prefix).strip("_") + "_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".xlsx"
    path = Path(export_dir) / file_name
    meta_rows = [
        {"Field": "Application", "Value": metadata.get("app_name", "IFMS Report Analyzer")},
        {"Field": "Designed and Developed by", "Value": "Neeraj Purohit"},
        {"Field": "Role", "Value": "Senior Assistant, Rajasthan Police"},
        {"Field": "Email", "Value": metadata.get("developer_email", "")},
        {"Field": "Export Timestamp", "Value": datetime.now().isoformat(timespec="seconds")},
        {"Field": "Upload File", "Value": _row_get(upload, "original_file_name", "")},
        {"Field": "Bill No", "Value": bill_no},
        {"Field": "Month/Year", "Value": next((s.get("month_year") for s in sections if s.get("month_year")), "")},
        {"Field": "Schedule Type", "Value": schedule_type or "all"},
        {"Field": "App Build Label", "Value": metadata.get("build_label", "")},
        {"Field": "Disclaimer", "Value": metadata.get("disclaimer", "")},
    ]
    sheets = {
        "Export Metadata": pd.DataFrame(meta_rows),
        "Upload Summary": pd.DataFrame([dict(upload)] if upload else []),
        "Section Summary": pd.DataFrame(sections),
        "Comparison Summary": pd.DataFrame([summary]),
        "Matched Rows": pd.DataFrame([r for r in rows if r.get("match_status") == "MATCHED"]),
        "Missing in Official": pd.DataFrame([r for r in rows if r.get("match_status") == "MISSING_IN_OFFICIAL"]),
        "Extra in Official": pd.DataFrame([r for r in rows if r.get("match_status") == "EXTRA_IN_OFFICIAL"]),
        "Amount Mismatches": pd.DataFrame([r for r in rows if r.get("match_status") == "AMOUNT_MISMATCH"]),
        "Identity Mismatches": pd.DataFrame([r for r in rows if r.get("match_status") == "IDENTITY_MISMATCH"]),
        "Needs Review": pd.DataFrame([r for r in rows if "REVIEW" in str(r.get("match_status", ""))]),
        "Official Parsed Rows": pd.DataFrame(official_rows),
        "Official Employee Aggregates": pd.DataFrame(aggregates),
        "Derived Schedule Rows": pd.DataFrame(derived),
        "Unsupported Sections": pd.DataFrame([s for s in sections if s.get("parse_status") != "SUPPORTED"]),
        "Data Quality Notes": pd.DataFrame([
            {"Note": "Official schedule comparison is read-only and does not modify Employee Data or derived schedules."},
            {"Note": "Unsupported layouts are stored as NEEDS_LAYOUT_MAPPING for manual review."},
            {"Note": "LIC is compared by employee-level aggregate because multiple policy rows may exist."},
            {"Note": "Human review of the data is compulsory before using the data."},
        ]),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, index=False, sheet_name=safe_name)
            ws = writer.sheets[safe_name]
            ws.freeze_panes = "A2"
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 42)
    return path


def file_sha256(path: str | Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

import pytest
from flask import Flask
from ifms_suite.web import register_suite

@pytest.fixture
def client(tmp_path):
    app=Flask(__name__);app.config.update(TESTING=True,DATABASE=str(tmp_path/'suite.db'),UPLOAD_FOLDER=str(tmp_path/'uploads'))
    register_suite(app)
    return app.test_client()

def test_persist_calculation_and_reject_finalization_without_review(client):
    payload=dict(employee_code='RJ1',allowance='DA',start='2025-01-01',end='2025-01-31',segments=[dict(start='2025-01-01',end='2025-01-31',basic='50000',due_rate='55',drawn_rate='53')])
    response=client.post('/suite/api/claims',json=payload)
    assert response.status_code==201
    result=response.json
    assert result['calculation']['total']=='1000.00'
    assert client.get('/suite/api/claims').json[0]['id']==result['id']
    assert client.post(f"/suite/api/claims/{result['id']}/finalize",json={}).status_code==400
    assert client.post(f"/suite/api/claims/{result['id']}/finalize",json={'reviewed':True,'rule_reference':'Office verified rate order','drawn_reference':'Reviewed salary record'}).status_code==200
    assert client.get(f"/suite/api/claims/{result['id']}/excel").data[:2]==b'PK'
    assert client.get(f"/suite/api/claims/{result['id']}/pdf").data[:4]==b'%PDF'

def test_conflicting_master_not_silently_overwritten(client):
    assert client.post('/suite/api/employees',json={'employee_code':'RJ1','name':'A'}).status_code==201
    assert client.post('/suite/api/employees',json={'employee_code':'RJ1','name':'B'}).status_code==409
    assert client.get('/suite/api/employees').json[0]['name']=='A'

def test_invalid_calculation_rejected(client):
    assert client.post('/suite/api/claims',json={'start':'2025-01-01','end':'2025-01-31','segments':[]}).status_code==400

def test_finalization_normalizes_dates_and_allowance(client):
    p=dict(employee_code='RJ1',allowance='DA',start='2025-01-01',end='2025-01-31',segments=[dict(start='2025-01-01',end='2025-01-31',basic='50000',due_rate='55',drawn_rate='53')])
    review=dict(reviewed=True,rule_reference='Order123',drawn_reference='Bill1')
    first=client.post('/suite/api/claims',json=p).json['id']
    assert client.post(f'/suite/api/claims/{first}/finalize',json=review).status_code==200
    p.update(allowance=' da ',start='20250101',end='20250131')
    second=client.post('/suite/api/claims',json=p).json['id']
    assert client.post(f'/suite/api/claims/{second}/finalize',json=review).status_code==400

def test_exports_preserve_review_and_force_user_text(client):
    import io
    from openpyxl import load_workbook
    import pymupdf
    p=dict(employee_code='=1+1',allowance='DA',start='2025-01-01',end='2025-01-31',segments=[dict(start='2025-01-01',end='2025-01-31',basic='50000',due_rate='55',drawn_rate='53')])
    id=client.post('/suite/api/claims',json=p).json['id']
    pdf=client.get(f'/suite/api/claims/{id}/pdf').data
    assert 'DRAFT' in ''.join(page.get_text() for page in pymupdf.open(stream=pdf,filetype='pdf'))
    review=dict(reviewed=True,rule_reference='Order123',drawn_reference='Bill1')
    client.post(f'/suite/api/claims/{id}/finalize',json=review)
    w=load_workbook(io.BytesIO(client.get(f'/suite/api/claims/{id}/excel').data))
    assert w.active['B1'].data_type=='s'
    assert any('Order123' in str(c.value) for row in w['Calculation basis'] for c in row)

def test_review_metadata_cannot_relabel_claim(client):
    import io
    from openpyxl import load_workbook
    p=dict(employee_code='RJ1',allowance='DA',start='2025-01-01',end='2025-01-31',segments=[dict(start='2025-01-01',end='2025-01-31',basic='50000',due_rate='55',drawn_rate='53')])
    id=client.post('/suite/api/claims',json=p).json['id']
    client.post(f'/suite/api/claims/{id}/finalize',json=dict(reviewed=True,rule_reference='Order',drawn_reference='Bill',employee_code='OTHER'))
    w=load_workbook(io.BytesIO(client.get(f'/suite/api/claims/{id}/excel').data))
    assert w.active['B1'].value=='RJ1'

def test_concurrent_overlapping_finalizations_allow_only_one(client):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    p=dict(employee_code='RJ1',allowance='DA',start='2025-01-01',end='2025-01-31',segments=[dict(start='2025-01-01',end='2025-01-31',basic='50000',due_rate='55',drawn_rate='53')])
    ids=[client.post('/suite/api/claims',json=p).json['id'] for _ in range(8)]
    barrier=Barrier(8)
    def finalize(id):
        with client.application.test_client() as c:
            barrier.wait()
            return c.post(f'/suite/api/claims/{id}/finalize',json=dict(reviewed=True,rule_reference='Order',drawn_reference='Bill')).status_code
    with ThreadPoolExecutor(max_workers=8) as ex:statuses=list(ex.map(finalize,ids))
    assert statuses.count(200)==1
    assert statuses.count(400)==7

def test_leave_encashment_and_salary_claims_have_distinct_entitlements(client):
    p=dict(employee_code='RJ1',allowance='DA',start='2025-01-01',end='2025-01-31',segments=[dict(start='2025-01-01',end='2025-01-31',basic='50000',due_rate='55',drawn_rate='53')])
    review=dict(reviewed=True,rule_reference='Order',drawn_reference='Bill')
    salary=client.post('/suite/api/claims',json=p).json['id']
    assert client.post(f'/suite/api/claims/{salary}/finalize',json=review).status_code==200
    leave=client.post('/suite/api/claims',json=dict(p,entitlement_type='leave_encashment')).json['id']
    assert client.post(f'/suite/api/claims/{leave}/finalize',json=review).status_code==200

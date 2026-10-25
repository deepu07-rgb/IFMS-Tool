from io import BytesIO

import openpyxl
import pytest

from app import app
from scripts.qa_employee_master import build_workbook
from scripts.verify_v2_all_samples import run_verification


pytestmark = pytest.mark.v2


def _employees(client, **params):
    query = {"page": 1, "per_page": 200, **params}
    resp = client.get("/api/employees", query_string=query)
    assert resp.status_code == 200
    return resp.get_json()


def _advanced(filters):
    import json
    return json.dumps(filters)


def test_v3_employee_filters_filtered_export_and_qa(tmp_path):
    db_path = tmp_path / "v3_employee_filters.db"
    export_dir = tmp_path / "exports"
    export_dir.mkdir()

    report = run_verification(db_path=str(db_path), verbose=False)
    assert report.passed, "\n".join(report.failure_messages())

    app.config["DATABASE"] = str(db_path)
    app.config["EXPORT_FOLDER"] = str(export_dir)
    app.config["USE_V2_PARSER"] = True
    client = app.test_client()

    all_rows = _employees(client, per_page=50)
    assert all_rows["total"] == 1118
    assert all_rows["per_page"] == 50
    assert len(all_rows["employees"]) == 50
    assert all_rows["quality"]["total_employees"] == 1118

    add = client.post(
        "/api/v2/dictionary",
        json={
            "category": "designation",
            "raw_text": "IPS, JUNIOR SCALE",
            "final_text": "IPS",
            "match_type": "contains",
            "active": True,
            "apply_existing": False,
        },
    )
    assert add.status_code == 200
    apply = client.post("/api/v3/dictionary/apply", json={"mapping_id": add.get_json()["id"]})
    assert apply.status_code == 200

    ips = _employees(
        client,
        advanced_filters=_advanced([{"column": "final_designation", "operator": "contains", "value": "IPS"}]),
    )
    assert ips["total"] > 0
    assert all("IPS" in r["final_designation"] for r in ips["employees"])

    itax_gt = _employees(
        client,
        advanced_filters=_advanced([{"column": "ITax", "operator": "greater than", "value": "5000"}]),
    )
    assert itax_gt["total"] > 0
    assert all(float(r["itax"] or 0) > 5000 for r in itax_gt["employees"])

    itax_eq = _employees(
        client,
        advanced_filters=_advanced([{"column": "ITax", "operator": "equals", "value": "5000"}]),
    )
    assert all(float(r["itax"] or 0) == 5000 for r in itax_eq["employees"])

    basic_between = _employees(
        client,
        advanced_filters=_advanced([{"column": "Basic", "operator": "between", "value": "50000", "value2": "60000"}]),
    )
    assert basic_between["total"] > 0
    assert all(50000 <= float(r["basic"] or 0) <= 60000 for r in basic_between["employees"])

    gpf2004_exists = _employees(
        client,
        advanced_filters=_advanced([{"column": "GPF 2004", "operator": "exists"}]),
    )
    assert gpf2004_exists["total"] > 0
    assert all(float(r["gpf_2004"] or 0) != 0 for r in gpf2004_exists["employees"])

    needs_review = _employees(client, employee_data_status="NEEDS_REVIEW")
    assert needs_review["total"] > 0
    assert all(r["employee_data_status"] == "NEEDS_REVIEW" for r in needs_review["employees"])

    combined = _employees(
        client,
        advanced_filters=_advanced([
            {"column": "final_designation", "operator": "contains", "value": "IPS"},
            {"column": "ITax", "operator": "greater than", "value": "5000"},
        ]),
    )
    assert all("IPS" in r["final_designation"] and float(r["itax"] or 0) > 5000 for r in combined["employees"])

    pay_scale = _employees(
        client,
        advanced_filters=_advanced([{"column": "pay_scale", "operator": "between", "value": "L8", "value2": "L12"}]),
    )
    assert pay_scale["total"] > 0
    for row in pay_scale["employees"]:
        import re
        m = re.search(r"\d+", row["pay_scale"])
        assert m and 8 <= int(m.group(0)) <= 12

    page_2 = _employees(
        client,
        page=2,
        per_page=10,
        advanced_filters=_advanced([{"column": "ITax", "operator": "greater than", "value": "5000"}]),
    )
    assert page_2["page"] == 2
    assert len(page_2["employees"]) <= 10

    no_result = _employees(
        client,
        advanced_filters=_advanced([{"column": "employee_name", "operator": "equals", "value": "NO SUCH EMPLOYEE"}]),
    )
    assert no_result["total"] == 0
    assert no_result["employees"] == []

    export_resp = client.get(
        "/api/export_filtered",
        query_string={
            "advanced_filters": _advanced([{"column": "ITax", "operator": "greater than", "value": "5000"}])
        },
    )
    assert export_resp.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(export_resp.data), data_only=True)
    assert {"Export Metadata", "Filtered Employee Data", "Applied Filters", "Data Quality Summary"}.issubset(wb.sheetnames)
    ws = wb["Filtered Employee Data"]
    headers = [c.value for c in ws[1]]
    itax_col = headers.index("ITax") + 1
    exported_values = [row[itax_col - 1].value for row in ws.iter_rows(min_row=2)]
    assert exported_values
    assert all(float(v or 0) > 5000 for v in exported_values)

    qa_path = build_workbook(str(db_path), str(export_dir))
    assert qa_path.endswith(".xlsx")
    qa_wb = openpyxl.load_workbook(qa_path, read_only=True)
    assert {
        "QA Summary",
        "Needs Review Rows",
        "Missing Critical Fields",
        "Manual Edits",
        "Dictionary Applied Rows",
        "Component Review Rows",
        "Final Designation Summary",
        "Missing Pay Scale",
        "Missing Employee ID",
        "Suspicious PAN",
    }.issubset(set(qa_wb.sheetnames))

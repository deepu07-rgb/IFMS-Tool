"""Synthetic supplied-form export contract; no private source data."""

from io import BytesIO
from copy import deepcopy
import pytest
import pymupdf
from openpyxl import load_workbook


def test_partial_month_dates_remain_visible_inside_supplied_month_column():
    from ifms_suite.supplied_form import export_supplied_excel

    bill, metadata = case()
    bill["rows"][0].update(start="2026-01-16", end="2026-01-31")
    sheet = load_workbook(export_supplied_excel(bill, metadata)).active
    assert "16/01/2026" in sheet["A6"].value
    assert "31/01/2026" in sheet["A6"].value


def case():
    rows = []
    for month in ("2026-01", "2026-02"):
        rows.append(
            {
                "month": month,
                "components": [
                    {
                        "head": "BASIC",
                        "due": "1200",
                        "drawn": "1000",
                        "prior_settlement": "0",
                        "difference": "200",
                    },
                    {
                        "head": "DA",
                        "due": "600",
                        "drawn": "500",
                        "prior_settlement": "0",
                        "difference": "100",
                    },
                ],
            }
        )
    return {
        "status": "ready",
        "rows": rows,
        "deductions": [],
        "totals": {
            "due": "3600",
            "drawn": "3000",
            "difference": "600",
            "prior_settlement": "0",
            "total_deductions": "0",
            "net_payable": "600",
        },
    }, {
        "supplied_template_id": "jodhpur_classic",
        "employee_name": "Synthetic employee",
        "employee_code": "SYN-01",
    }


def test_printed_grid_groups_order_and_excel_editable_values():
    from ifms_suite.supplied_form import export_supplied_excel

    bill, metadata = case()
    before = deepcopy(bill)
    book = load_workbook(export_supplied_excel(bill, metadata))
    assert book.sheetnames == ["Due drawn statement"]
    sheet = book.active
    assert sheet["D4"].value == "PAY TO BE DRAWN"
    assert sheet["J4"].value == "PAY ALREADY DRAWN"
    assert sheet["P4"].value == "PAY DIFFERENCE DRAWN"
    assert [sheet.cell(5, c).value for c in range(4, 10)] == [
        "BASIC\nPAY",
        "D.A.",
        "H.R.A.",
        "H.D.A.",
        "CCA",
        "TOTAL",
    ]
    assert sheet["D6"].value == 1200
    assert sheet["J6"].value == 1000
    assert sheet["P6"].value == 200
    assert sheet["AA6"].value == 300
    assert sheet.protection.sheet is False
    assert sheet.page_setup.orientation == "landscape"
    assert not any(cell.data_type == "f" for row in sheet for cell in row)
    assert bill == before


@pytest.mark.parametrize(
    "template,expected",
    [
        ("jodhpur_classic", "CM\nCORONA"),
        ("jodhpur_gpf_arrear", "GPF ARREAR\nALREADY\nDEDUCTED"),
    ],
)
def test_pdf_uses_clean_grid_and_explicit_variant(template, expected):
    from ifms_suite.supplied_form import export_supplied_pdf, export_supplied_excel

    bill, metadata = case()
    metadata["supplied_template_id"] = template
    xls = load_workbook(export_supplied_excel(bill, metadata)).active
    assert expected in [c.value for row in xls for c in row]
    stream = export_supplied_pdf(bill, metadata)
    assert isinstance(stream, BytesIO) and stream.tell() == 0
    with pymupdf.open(stream=stream.getvalue(), filetype="pdf") as doc:
        text = "".join(p.get_text() for p in doc)
        assert len(doc) == 1
        assert not doc[0].get_images()
        for value in (
            "OFFICE OF THE POLICE COMMISSIONER JODHPUR",
            "PAY TO BE DRAWN",
            "PAY ALREADY DRAWN",
            "PAY DIFFERENCE DRAWN",
            "Synthetic employee",
            "TOTAL",
            "600",
            "Drawing and Disbursing Officer",
        ):
            assert value in text


@pytest.mark.parametrize(
    "change,reason",
    [
        ("template", "template"),
        ("head", "head"),
        ("credit", "reconcile"),
        ("blocked", "ready"),
        ("unresolved", "amount"),
        ("deductions", "deduction"),
    ],
)
def test_unrepresentable_output_is_blocked(change, reason):
    from ifms_suite.supplied_form import export_supplied_pdf, export_supplied_excel

    bill, metadata = case()
    if change == "template":
        metadata["supplied_template_id"] = "unmatched"
    if change == "head":
        bill["rows"][0]["components"][0]["head"] = "SPECIAL"
    if change == "credit":
        bill["rows"][0]["components"][0]["prior_settlement"] = "10"
    if change == "blocked":
        bill["status"] = "blocked"
    if change == "unresolved":
        bill["rows"][0]["components"][0]["due"] = None
    if change == "deductions":
        bill["totals"]["total_deductions"] = "20"
    for export in (export_supplied_excel, export_supplied_pdf):
        with pytest.raises(ValueError, match=reason):
            export(bill, metadata)


def test_reviewed_monthly_deductions_and_bill_references_use_existing_columns():
    from ifms_suite.supplied_form import export_supplied_excel

    bill, metadata = case()
    bill["totals"].update(total_deductions="20", net_payable="580")
    metadata["supplied_form_rows"] = {
        "2026-01": {
            "bill_no": "=SYNTHETIC",
            "bill_date": "2026-01-31",
            "tv_no": "TV-SYN",
            "tv_date": "2026-02-01",
            "deductions": {"GPF": "20"},
            "total_deductions": "20",
        }
    }
    sheet = load_workbook(export_supplied_excel(bill, metadata)).active
    assert sheet["B6"].value == "=SYNTHETIC\n2026-01-31"
    assert sheet["B6"].data_type == "s"
    assert sheet["V6"].value == 20
    assert sheet["Z6"].value == 20
    assert sheet["AA6"].value == 280


def test_unchanged_unprinted_heads_are_omitted_from_selected_head_totals():
    from ifms_suite.supplied_form import export_supplied_excel

    bill, metadata = case()
    for row in bill["rows"]:
        row["components"].append(
            {
                "head": "RECOVERY",
                "due": "-25",
                "drawn": "-25",
                "prior_settlement": "0",
                "difference": "0",
            }
        )
    bill["totals"].update(due="3550", drawn="2950")
    sheet = load_workbook(export_supplied_excel(bill, metadata)).active
    assert sheet["I6"].value == 1800
    assert sheet["O6"].value == 1500
    assert sheet["AA6"].value == 300


def test_source_bill_and_treasury_details_are_carried_when_unambiguous():
    from ifms_suite.supplied_form import export_supplied_excel

    bill, metadata = case()
    bill["rows"][0]["evidence"] = [
        {
            "bill_no": "SYN-BILL",
            "bill_date": "2026-01-30",
            "treasury_no": "SYN-TV",
            "treasury_date": "2026-02-01",
        }
    ]
    sheet = load_workbook(export_supplied_excel(bill, metadata)).active
    assert sheet["B6"].value == "SYN-BILL\n2026-01-30"
    assert sheet["C6"].value == "SYN-TV\n2026-02-01"


def test_pdf_continuation_pages_keep_signature_space_below_months():
    from ifms_suite.supplied_form import export_supplied_pdf

    bill, metadata = case()
    bill["rows"] = [
        dict(bill["rows"][0], month=f"2026-{month:02}") for month in range(1, 13)
    ]
    bill["totals"].update(
        due="21600", drawn="18000", difference="3600", net_payable="3600"
    )
    with pymupdf.open(
        stream=export_supplied_pdf(bill, metadata).getvalue(), filetype="pdf"
    ) as doc:
        assert len(doc) == 2
        for page in doc:
            words = page.get_text("words")
            month_bottom = max(w[3] for w in words if w[4].startswith("2026-"))
            signature_top = min(w[1] for w in words if w[4] == "Assistant")
            assert month_bottom < signature_top


def test_pdf_never_silently_replaces_unsupported_name_characters():
    from ifms_suite.supplied_form import export_supplied_pdf

    bill, metadata = case()
    metadata["employee_name"] = (
        "Synthetic \u0915\u0930\u094d\u092e\u091a\u093e\u0930\u0940"
    )
    with pytest.raises(ValueError, match="font"):
        export_supplied_pdf(bill, metadata)


def test_explicit_engine_gpf_adjustment_is_shown_only_in_total_row():
    from ifms_suite.supplied_form import export_supplied_excel, export_supplied_pdf

    bill, metadata = case()
    bill["deductions"] = [
        {"name": "GPF adjustment", "mode": "fixed", "value": "20", "amount": "20.00"}
    ]
    bill["totals"].update(total_deductions="20", net_payable="580")
    sheet = load_workbook(export_supplied_excel(bill, metadata)).active
    for r in (6, 7):
        assert sheet.cell(r, 22).value is None
        assert sheet.cell(r, 26).value is None
        assert sheet.cell(r, 27).value is None
    assert sheet["V8"].value == 20
    assert sheet["Z8"].value == 20
    assert sheet["AA8"].value == 580
    with pymupdf.open(
        stream=export_supplied_pdf(bill, metadata).getvalue(), filetype="pdf"
    ) as doc:
        assert "580" in doc[0].get_text()
        assert "None" not in doc[0].get_text()
        assert "DRAFT" in doc[0].get_text()


def test_unmapped_engine_adjustment_cannot_disappear_from_total_row():
    from ifms_suite.supplied_form import export_supplied_excel

    bill, metadata = case()
    bill["deductions"] = [{"name": "Unmapped fee", "amount": "20"}]
    bill["totals"].update(total_deductions="20", net_payable="580")
    with pytest.raises(ValueError, match="deduction"):
        export_supplied_excel(bill, metadata)


def credited_case():
    bill, metadata = case()
    bill["rows"][0]["components"][0].update(prior_settlement="30", difference="170")
    bill["rows"][1]["components"][1].update(prior_settlement="20", difference="80")
    bill["totals"].update(prior_settlement="50", difference="550", net_payable="550")
    bill["evidence"] = [
        {
            "kind": "prior_credit",
            "bill_no": "SYN-PRIOR-7",
            "pay_month": "2026-02",
            "reference": "Synthetic credit order",
        }
    ]
    return bill, metadata


def test_verified_prior_settlements_use_already_drawn_and_disclose_source_without_mutation():
    from ifms_suite.supplied_form import export_supplied_excel, export_supplied_pdf

    bill, metadata = credited_case()
    original = deepcopy(bill)
    sheet = load_workbook(export_supplied_excel(bill, metadata)).active
    assert sheet["J6"].value == 1030
    assert sheet["K7"].value == 520
    assert sheet["P6"].value == 170
    assert sheet["Q7"].value == 80
    assert sheet["I8"].value == 3600
    assert sheet["O8"].value == 3050
    assert sheet["U8"].value == 550
    assert sheet["AA8"].value == 550
    values = " ".join(str(cell.value) for row in sheet for cell in row)
    note = "Pay already drawn includes verified prior arrear settlements"
    assert note in values
    assert "SYN-PRIOR-7" in values and "Synthetic credit order" in values
    with pymupdf.open(
        stream=export_supplied_pdf(bill, metadata).getvalue(), filetype="pdf"
    ) as doc:
        text = "".join(page.get_text() for page in doc)
        assert note in text
        assert "SYN-PRIOR-7" in text and "Synthetic credit order" in text
        assert "3050" in text and "550" in text
    assert bill == original


@pytest.mark.parametrize("credit", ["-1", None, "NaN"])
def test_negative_or_unresolved_prior_settlement_cannot_be_printed(credit):
    from ifms_suite.supplied_form import export_supplied_excel, export_supplied_pdf

    bill, metadata = credited_case()
    bill["rows"][0]["components"][0]["prior_settlement"] = credit
    for export in (export_supplied_excel, export_supplied_pdf):
        with pytest.raises(ValueError):
            export(bill, metadata)


def test_unprinted_head_with_credit_cannot_be_omitted_even_with_zero_remaining_difference():
    from ifms_suite.supplied_form import export_supplied_excel

    bill, metadata = credited_case()
    bill["rows"][0]["components"].append(
        {
            "head": "SPECIAL",
            "due": "10",
            "drawn": "0",
            "prior_settlement": "10",
            "difference": "0",
        }
    )
    bill["totals"]["prior_settlement"] = "60"
    with pytest.raises(ValueError, match="head"):
        export_supplied_excel(bill, metadata)

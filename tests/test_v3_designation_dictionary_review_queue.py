import json
import os
import sqlite3
from io import BytesIO
from urllib.parse import quote

import openpyxl
import pytest

from app import app, init_db, process_file_v2
from ifms_engine.db import init_v2_schema


pytestmark = pytest.mark.v2


PDF10 = os.path.join("test_fixtures", "ifms_inner_sheets", "10.pdf")


def _setup(tmp_path):
    db_path = tmp_path / "v3_designation.db"
    export_dir = tmp_path / "exports"
    export_dir.mkdir()
    app.config["DATABASE"] = str(db_path)
    app.config["EXPORT_FOLDER"] = str(export_dir)
    app.config["USE_V2_PARSER"] = True
    init_db()
    init_v2_schema(app.config["DATABASE"])
    return app.test_client(), db_path


def _process_pdf10():
    assert os.path.exists(PDF10)
    process_file_v2(PDF10, "10.pdf", dict(app.config))


def _db_rows(db_path, sql, params=()):
    db = sqlite3.connect(db_path)
    db.row_factory = sqlite3.Row
    rows = db.execute(sql, params).fetchall()
    db.close()
    return rows


def test_v3_designation_dictionary_review_queue_filter_and_export(tmp_path):
    client, db_path = _setup(tmp_path)
    _process_pdf10()

    rows = _db_rows(
        db_path,
        """
        SELECT raw_name, raw_designation_full, final_designation_en, designation_source
        FROM employees_v2
        WHERE source_filename = '10.pdf'
        ORDER BY sr_no
        """,
    )
    by_name = {r["raw_name"]: r for r in rows}
    assert "IPS, JUNIOR SCALE" in by_name["PRATEEK SINGH"]["raw_designation_full"]
    assert by_name["PRATEEK SINGH"]["final_designation_en"] != "Assistant Sub-Inspector"
    assert by_name["KAMAL SHEKHAWAT"]["final_designation_en"] != "Sub-Inspector"
    assert by_name["SHARAT KAVIRAJ IPS"]["final_designation_en"] != "Inspector"
    assert by_name["MANISH KUMAR CHOUDHARY"]["final_designation_en"] != "Sub-Inspector"

    review_count = _db_rows(
        db_path,
        "SELECT COUNT(*) AS c FROM review_queue_v3 WHERE issue_type = 'UNCONFIRMED_DESIGNATION' AND status = 'OPEN'",
    )[0]["c"]
    assert review_count >= 5

    add_resp = client.post(
        "/api/v2/dictionary",
        json={
            "category": "designation",
            "raw_text": "IPS, JUNIOR SCALE",
            "final_text": "IPS",
            "match_type": "contains",
            "priority": 100,
            "apply_existing": True,
        },
    )
    assert add_resp.status_code == 200
    mapping_id = add_resp.get_json()["id"]

    mapped = _db_rows(
        db_path,
        """
        SELECT COUNT(*) AS c
        FROM employees_v2
        WHERE raw_designation_full LIKE '%IPS, JUNIOR SCALE%'
          AND final_designation_en = 'IPS'
          AND designation_source = 'dictionary_contains'
        """,
    )[0]["c"]
    assert mapped >= 3

    clear_resp = client.post("/api/delete_all")
    assert clear_resp.status_code == 200
    assert _db_rows(db_path, "SELECT COUNT(*) AS c FROM dictionary_mappings_v3 WHERE id = ?", (mapping_id,))[0]["c"] == 1

    _process_pdf10()
    auto_mapped = _db_rows(
        db_path,
        """
        SELECT COUNT(*) AS c
        FROM employees_v2
        WHERE raw_designation_full LIKE '%IPS, JUNIOR SCALE%' AND final_designation_en = 'IPS'
        """,
    )[0]["c"]
    assert auto_mapped >= 3

    hi_resp = client.get("/api/employees?designation_language=hindi&page=1&per_page=10")
    assert hi_resp.status_code == 200
    hi_rows = [r for r in hi_resp.get_json()["employees"] if "IPS, JUNIOR SCALE" in r["raw_designation_full"]]
    assert hi_rows
    assert hi_rows[0]["designation"] == "IPS"
    assert hi_rows[0]["final_designation_hi"] == ""

    filt = json.dumps([{"column": "ITax", "operator": "greater than", "value": 5000}])
    filtered_resp = client.get("/api/employees?advanced_filters=" + quote(filt))
    assert filtered_resp.status_code == 200
    filtered = filtered_resp.get_json()["employees"]
    assert filtered
    assert all(row["itax"] > 5000 for row in filtered)

    export_resp = client.get("/api/export_filtered?advanced_filters=" + quote(filt))
    assert export_resp.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(export_resp.data), data_only=True)
    assert "Export Metadata" in wb.sheetnames
    assert "Employee Master" in wb.sheetnames
    headers = [cell.value for cell in wb["Employee Master"][1]]
    assert "raw_designation_full" in headers
    assert "final_designation_en" in headers
    itax_idx = headers.index("Income Tax") + 1
    for row in wb["Employee Master"].iter_rows(min_row=2):
        assert (row[itax_idx - 1].value or 0) > 5000

    delete_resp = client.delete(f"/api/v2/dictionary/{mapping_id}")
    assert delete_resp.status_code == 200
    client.post("/api/delete_all")
    _process_pdf10()
    back_to_review = _db_rows(
        db_path,
        """
        SELECT COUNT(*) AS c
        FROM review_queue_v3
        WHERE raw_value LIKE '%IPS, JUNIOR SCALE%'
          AND issue_type = 'UNCONFIRMED_DESIGNATION'
          AND status = 'OPEN'
        """,
    )[0]["c"]
    assert back_to_review >= 3

"""Formula exports and safe literal-input round trips, using synthetic salaries."""

from io import BytesIO
from copy import deepcopy
import pytest
from openpyxl import load_workbook
from test_worksheet import example


def test_workbook_is_formula_based_in_existing_supplied_grid():
    from ifms_suite.worksheet_excel import export_workbook

    p = example()
    w = load_workbook(export_workbook(p))
    assert w.sheetnames == ["Setup", "Monthly inputs", "Due drawn statement", "Guide"]
    s = w["Due drawn statement"]
    assert s["D4"].value == "PAY TO BE DRAWN"
    assert s["J4"].value == "PAY ALREADY DRAWN"
    assert s["P4"].value == "PAY DIFFERENCE DRAWN"
    assert s["F6"].data_type == "f"
    assert s["R6"].data_type == "f"
    assert s["AA6"].data_type == "f"
    assert s["AA7"].data_type == "f"
    assert s.page_setup.orientation == "landscape"
    assert w.calculation.fullCalcOnLoad is True


def test_blank_template_does_not_create_financial_zeros():
    from ifms_suite.worksheet_excel import export_workbook
    from ifms_suite.worksheet import COLUMNS

    p = example()
    p["rows"][0].pop("drawn_HRA")
    w = load_workbook(export_workbook(p))
    index = next(i for i, c in enumerate(COLUMNS, 1) if c[0] == "drawn_HRA")
    assert w["Monthly inputs"].cell(6, index).value is None
    checks = w["Monthly inputs"].cell(6, len(COLUMNS) + 1)
    assert checks.data_type == "f"
    assert "COUNT(" in checks.value
    assert "INCOMPLETE" in checks.value


def test_round_trip_reads_literals_and_ignores_tampered_statement_totals():
    from ifms_suite.worksheet_excel import export_workbook, import_workbook
    from ifms_suite.worksheet import calculate_worksheet

    p = example()
    w = load_workbook(export_workbook(p))
    w["Due drawn statement"]["AA6"] = "=999999"
    stream = BytesIO()
    w.save(stream)
    stream.seek(0)
    imported = import_workbook(stream)
    assert imported == p
    assert calculate_worksheet(imported)["bill"]["totals"]["net_payable"] == "4960.00"


def test_user_text_remains_text_in_excel():
    from ifms_suite.worksheet_excel import export_workbook, import_workbook

    p = example(employee_name='=HYPERLINK("https://invalid")')
    p["rows"][0]["salary_reference"] = "=1+1"
    w = load_workbook(export_workbook(p))
    assert w["Setup"]["B3"].data_type == "s"
    from ifms_suite.worksheet import COLUMNS

    c = next(i for i, v in enumerate(COLUMNS, 1) if v[0] == "salary_reference")
    assert w["Monthly inputs"].cell(6, c).data_type == "s"
    assert import_workbook(export_workbook(p))["employee_name"] == p["employee_name"]


def test_input_formula_is_not_accepted_as_salary():
    from ifms_suite.worksheet_excel import export_workbook, import_workbook

    w = load_workbook(export_workbook(example()))
    w["Monthly inputs"]["E6"] = "=24800"
    stream = BytesIO()
    w.save(stream)
    stream.seek(0)
    with pytest.raises(ValueError, match="formula|Formula"):
        import_workbook(stream)


def test_salary_template_and_reader_preserve_complete_reconcilable_components():
    from ifms_suite.worksheet_excel import (
        salary_template,
        read_salaries,
        SALARY_COLUMNS,
    )

    w = load_workbook(salary_template())
    record = dict(
        employee_code="SYN-01",
        employee_name="Synthetic Person",
        pay_month="2040-07",
        bill_no="SYN1",
        BASIC=24800,
        DA=14880,
        HRA=0,
        HDA=3307,
        CCA=0,
        OTHER_ALLOWANCES=3350,
        total_deductions=5065,
        printed_gross=46337,
        printed_net=41272,
        reference="Synthetic verified salary",
    )
    w["Salary ledger"].append([record[k] for k in SALARY_COLUMNS])
    stream = BytesIO()
    w.save(stream)
    stream.seek(0)
    rows = read_salaries(stream)
    assert rows[0]["allowances"]["OTHER_ALLOWANCES"] == "3350"
    assert rows[0]["allowances"]["HRA"] == "0"
    assert rows[0]["printed_gross"] == "46337"
    assert rows[0]["deductions"] == {"TOTAL_DEDUCTIONS": "5065"}
    assert rows[0]["reviewed"] is False


def test_monthly_input_headers_cannot_be_silently_remapped():
    from ifms_suite.worksheet_excel import export_workbook, import_workbook

    w = load_workbook(export_workbook(example()))
    w["Monthly inputs"]["E5"] = "HRA"
    stream = BytesIO()
    w.save(stream)
    stream.seek(0)
    with pytest.raises(ValueError, match="header|template"):
        import_workbook(stream)


def test_excel_can_leave_unused_trailing_rows_blank(tmp_path):
    """Recalculate real formulas; deleting/clearing a spare row must not hide totals."""
    import shutil, subprocess
    from ifms_suite.worksheet_excel import export_workbook

    executable = shutil.which("soffice")
    if not executable:
        pytest.skip("Development LibreOffice unavailable; no runtime dependency")
    p = example()
    p["rows"].append({})
    source = tmp_path / "inputs.xlsx"
    source.write_bytes(export_workbook(p).read())
    output = tmp_path / "calculated"
    output.mkdir()
    process = subprocess.run(
        [
            executable,
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert process.returncode == 0, process.stdout + process.stderr
    w = load_workbook(output / "inputs.xlsx", data_only=True)
    assert w["Due drawn statement"]["AA6"].value == 4960
    assert w["Due drawn statement"]["AA7"].value is None
    assert w["Due drawn statement"]["AA8"].value == 4960


def test_rate_exceptions_and_hra_category_survive_excel_roundtrip():
    from ifms_suite.worksheet_excel import export_workbook, import_workbook

    p = example(hra_category="Y")
    p["rows"][0].update(
        da_override_reason="Synthetic specific exception",
        hra_override_reason="Synthetic HRA exception",
        hra_exempt=False,
        hra_exemption_reference="",
    )
    result = import_workbook(export_workbook(p))
    assert result.get("hra_category") == "Y"
    assert result["rows"][0]["da_override_reason"] == "Synthetic specific exception"
    assert result["rows"][0]["hra_exempt"] is False


def test_previous_published_v1_template_is_still_importable():
    from pathlib import Path
    from ifms_suite.worksheet_excel import import_workbook

    with (
        Path(__file__).resolve().parents[1]
        / "examples/IFMS-Formula-Due-Drawn-Template.xlsx"
    ).open("rb") as stream:
        p = import_workbook(stream)
    assert len(p["rows"]) == 12
    assert p["rows"][0]["month"] == "2026-01"

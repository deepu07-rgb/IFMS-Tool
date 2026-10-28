"""The bill download uses the supplied office grid, never a generic fallback."""

import io
import pymupdf
from openpyxl import load_workbook
from test_suite_workflow import client
from test_bill_routes import ready_case


def test_bill_download_defaults_to_supplied_grid_and_selects_other_supplied_variant(
    client, monkeypatch
):
    ident, selection = ready_case(client, monkeypatch)
    bill = client.post(f"/suite/api/imports/{ident}/bill_prepare", json=selection).json
    url = f"/suite/api/bills/{bill['bill_id']}"
    pdf = client.get(url + "/pdf")
    assert pdf.status_code == 200
    with pymupdf.open(stream=pdf.data, filetype="pdf") as document:
        text = "".join(page.get_text() for page in document)
        assert "PAY ALREADY DRAWN" in text
        assert "PAY TO BE DRAWN" in text
        assert "Drawing and Disbursing Officer" in text
        assert len(document) == 1
    excel = client.get(url + "/excel?template=jodhpur_gpf_arrear")
    assert excel.status_code == 200
    sheet = load_workbook(io.BytesIO(excel.data)).active
    assert sheet["D4"].value == "PAY TO BE DRAWN"
    assert "GPF ARREAR" in sheet["V4"].value
    assert client.get(url + "/pdf?template=generic").status_code == 400


def test_blocked_bill_has_no_office_form_payable_output(client, monkeypatch):
    from ifms_suite.storage import database

    ident, selection = ready_case(client, monkeypatch)
    bill = client.post(f"/suite/api/imports/{ident}/bill_prepare", json=selection).json
    # A blocked saved case is not an issued due/drawn form, even with rows.
    with client.application.app_context(), database() as db:
        import json

        result = dict(
            bill,
            status="blocked",
            blockers=[{"code": "source", "message": "Unresolved source"}],
        )
        result.pop("bill_id")
        db.execute(
            "UPDATE suite_bills SET result=?,status=? WHERE id=?",
            (json.dumps(result), "blocked", bill["bill_id"]),
        )
    response = client.get(f"/suite/api/bills/{bill['bill_id']}/pdf")
    assert response.status_code == 400
    assert "ready" in response.json["error"].lower()


def test_supplied_form_blocks_stale_source_or_rate_dependencies(client, monkeypatch):
    import json
    from ifms_suite.storage import database

    ident, selection = ready_case(client, monkeypatch)
    bill = client.post(f"/suite/api/imports/{ident}/bill_prepare", json=selection).json
    with client.application.app_context(), database() as db:
        row = db.execute("SELECT * FROM suite_rates LIMIT 1").fetchone()
        rate = json.loads(row["payload"])
        rate["rate"] = "99"
        db.execute(
            "UPDATE suite_rates SET payload=? WHERE id=?", (json.dumps(rate), row["id"])
        )
    response = client.get(f"/suite/api/bills/{bill['bill_id']}/pdf")
    assert response.status_code == 400
    assert "changed" in response.json["error"].lower()


def test_finalized_form_cannot_outlive_its_original_evidence(client, monkeypatch):
    from pathlib import Path

    ident, selection = ready_case(client, monkeypatch)
    bill = client.post(f"/suite/api/imports/{ident}/bill_prepare", json=selection).json
    review = {
        "reviewed": True,
        "deductions_reviewed": True,
        "reference": "Synthetic originals checked",
    }
    assert (
        client.post(
            f"/suite/api/bills/{bill['bill_id']}/finalize", json=review
        ).status_code
        == 200
    )
    source = next(
        (Path(client.application.config["UPLOAD_FOLDER"]) / "suite").glob("*.pdf")
    )
    source.unlink()
    response = client.get(f"/suite/api/bills/{bill['bill_id']}/excel")
    assert response.status_code == 400
    assert "source" in response.json["error"].lower()


def test_ledger_paid_credit_flows_into_same_supplied_drawn_columns(client, monkeypatch):
    from ifms_suite.storage import database
    from ifms_suite.payroll import reviewed_payment

    ident, selection = ready_case(client, monkeypatch)
    with client.application.app_context(), database() as db:
        reviewed_payment(
            db,
            dict(
                employee_code=selection["employee_code"],
                bill_no="987654321",
                pay_month="2026-03",
                bill_type="DA arrear",
                reviewed=True,
                reference="Synthetic prior payment verified",
                allowances={"DA": "50"},
                deductions={},
                printed_gross="50",
                printed_deductions="0",
                printed_net="50",
                entitlement_period={"start": "2026-01-01", "end": "2026-01-31"},
            ),
        )
    result = client.post(
        f"/suite/api/imports/{ident}/bill_prepare", json=selection
    ).json
    assert result["status"] == "ready", result
    assert result["totals"]["prior_settlement"] == "50.00"
    assert result["totals"]["difference"] == "270.00"
    response = client.get(f"/suite/api/bills/{result['bill_id']}/excel")
    assert response.status_code == 200, response.data[:300]
    sheet = load_workbook(io.BytesIO(response.data)).active
    assert sheet["K6"].value == 550
    assert sheet["U8"].value == 270
    assert any(
        "includes verified prior arrear" in str(cell.value)
        for row in sheet
        for cell in row
    )

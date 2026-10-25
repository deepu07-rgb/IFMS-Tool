import os
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from app import app
from scripts.verify_v2_all_samples import run_verification


pytestmark = [pytest.mark.v2, pytest.mark.smoke]


REQUIRED_SHEETS = {
    "Export Metadata",
    "Bill Register",
    "Employee Master",
    "Allowance Ledger",
    "Deduction Ledger",
    "Derived Schedule Summary",
    "Derived Schedule Details",
    "Group Summary",
    "Designation Summary",
    "Pay Scale Summary",
    "Budget Head Summary",
    "Validation Report",
    "Unknown Dictionary Review",
}


def _rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:]]


def _sum(rows, column):
    return round(sum(float(row.get(column) or 0) for row in rows))


def test_v2_professional_excel_export_all_samples(tmp_path):
    db_path = tmp_path / "v2_professional_export.db"
    export_dir = tmp_path / "exports"
    export_dir.mkdir()

    report = run_verification(db_path=str(db_path), verbose=False)
    assert report.passed, "\n".join(report.failure_messages())

    app.config["DATABASE"] = str(db_path)
    app.config["EXPORT_FOLDER"] = str(export_dir)
    app.config["USE_V2_PARSER"] = True

    response = app.test_client().get("/api/export")
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.content_type

    exported = sorted(export_dir.glob("IFMS_V2_Verified_Export_*.xlsx"))
    assert exported, "Expected /api/export to create a verified V2 workbook"
    assert exported[-1].exists()

    wb = openpyxl.load_workbook(BytesIO(response.data), data_only=True)
    assert REQUIRED_SHEETS.issubset(set(wb.sheetnames))

    bill_rows = _rows(wb["Bill Register"])
    employee_rows = _rows(wb["Employee Master"])
    group_rows = _rows(wb["Group Summary"])
    allowance_rows = _rows(wb["Allowance Ledger"])
    deduction_rows = _rows(wb["Deduction Ledger"])
    schedule_rows = _rows(wb["Derived Schedule Details"])

    assert len(bill_rows) == 14
    assert len(employee_rows) == 1118

    assert _sum(group_rows, "gross") == 87156799
    assert _sum(group_rows, "deduction") == 18425996
    assert _sum(group_rows, "net") == 68730803

    assert allowance_rows
    assert deduction_rows

    schedule_heads = {row["normalized_head"] for row in schedule_rows}
    assert {"lic", "gpf", "gpf_2004", "rghs", "sip"}.issubset(schedule_heads)

    rop_rows = [row for row in schedule_rows if row["schedule_type"] == "rop"]
    assert rop_rows
    assert all(row["normalized_head"] == "rop" for row in rop_rows)
    assert any(float(row["component_amount"] or 0) < 0 for row in rop_rows)

    employee_ids = {
        (row["source_file"], row["sr_no"]): row["employee_id"]
        for row in employee_rows
        if row.get("employee_id")
    }
    for row in schedule_rows:
        expected_id = employee_ids.get((row["source_file"], row["sr_no"]))
        if expected_id:
            assert row["employee_id"], f"Blank schedule employee_id for {row['source_file']} sr={row['sr_no']}"

    component_columns = [
        "Basic", "DA", "HRA", "CCA", "HDA", "Mess Allowance", "FTA", "ROP",
        "LIC", "GPF", "GPF 2004", "RGHS", "SIP", "SIL", "Income Tax",
    ]
    present_component_columns = [col for col in component_columns if col in employee_rows[0]]
    assert present_component_columns
    assert any(
        any(float(row.get(col) or 0) != 0 for row in employee_rows)
        for col in present_component_columns
    )

    assert Path(exported[-1]).name.startswith("IFMS_V2_Verified_Export_")

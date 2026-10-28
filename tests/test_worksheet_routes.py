"""Actual HTTP assisted workflow and existing persistent salary ledger integration."""

from copy import deepcopy
from io import BytesIO
import json
from openpyxl import load_workbook
from test_suite_workflow import client
from test_worksheet import example


def test_save_reload_edit_revision_and_audit(client):
    p = example()
    response = client.post("/suite/api/worksheets", json={"payload": p})
    assert response.status_code == 201, response.data
    saved = response.json
    assert saved["revision"] == 1
    assert saved["result"]["bill"]["totals"]["net_payable"] == "4960.00"
    p["rows"][0]["ded_GPF"] = "100"
    changed = client.post(
        f"/suite/api/worksheets/{saved['id']}", json={"payload": p, "revision": 1}
    )
    assert changed.status_code == 200
    assert changed.json["revision"] == 2
    assert changed.json["result"]["bill"]["totals"]["net_payable"] == "4860.00"
    assert (
        client.post(
            f"/suite/api/worksheets/{saved['id']}", json={"payload": p, "revision": 1}
        ).status_code
        == 409
    )
    loaded = client.get(f"/suite/api/worksheets/{saved['id']}").json
    assert loaded["payload"] == p
    assert loaded["audit_count"] == 2
    assert len(client.get("/suite/api/worksheets").json) == 1
    from ifms_suite.storage import database

    with client.application.app_context(), database() as db:
        audit = json.loads(
            db.execute(
                "SELECT previous_payload FROM suite_worksheet_audit WHERE worksheet_id=? AND revision=2",
                (saved["id"],),
            ).fetchone()[0]
        )
        assert audit["rows"][0]["ded_GPF"] == "0"


def test_incomplete_draft_gets_excel_but_cannot_export_a_pdf(client):
    p = example()
    p["rows"][0]["hra_rate"] = ""
    saved = client.post("/suite/api/worksheets", json={"payload": p}).json
    assert saved["result"]["status"] == "incomplete"
    assert client.get(f"/suite/api/worksheets/{saved['id']}/excel").data[:2] == b"PK"
    response = client.get(f"/suite/api/worksheets/{saved['id']}/pdf")
    assert response.status_code == 400
    assert "incomplete" in response.json["error"].lower()


def test_excel_round_trip_recalculates_and_pdf_uses_supplied_form(client):
    from ifms_suite.worksheet_excel import export_workbook

    w = load_workbook(export_workbook(example()))
    w["Due drawn statement"]["AA6"] = "=9999999"
    stream = BytesIO()
    w.save(stream)
    stream.seek(0)
    preview = client.post(
        "/suite/api/worksheets/preview-import",
        data={"file": (stream, "synthetic.xlsx")},
    )
    assert preview.status_code == 200, preview.data
    assert preview.json["result"]["bill"]["totals"]["net_payable"] == "4960.00"
    saved = client.post(
        "/suite/api/worksheets", json={"payload": preview.json["payload"]}
    ).json
    pdf = client.get(f"/suite/api/worksheets/{saved['id']}/pdf")
    assert pdf.status_code == 200
    import pymupdf

    text = "".join(p.get_text() for p in pymupdf.open(stream=pdf.data, filetype="pdf"))
    assert "PAY TO BE DRAWN" in text
    assert "4960" in text
    assert "DRAFT" in text


def salaries():
    return [
        dict(
            employee_code="SYN-01",
            employee_name="Synthetic Person",
            pay_month="2040-07",
            bill_no="SYN1",
            allowances=dict(
                BASIC="24800",
                DA="14880",
                HRA="0",
                HDA="3307",
                CCA="0",
                OTHER_ALLOWANCES="3350",
            ),
            deductions={"TOTAL_DEDUCTIONS": "5065"},
            printed_gross="46337",
            printed_deductions="5065",
            printed_net="41272",
            reference="Synthetic monthly original",
        )
    ]


def test_structured_salary_save_requires_review_and_reconciles_before_ledger_reuse(
    client,
):
    records = salaries()
    assert (
        client.post(
            "/suite/api/worksheets/salary-save", json={"records": records}
        ).status_code
        == 400
    )
    bad = deepcopy(records)
    bad[0]["printed_net"] = "1"
    assert (
        client.post(
            "/suite/api/worksheets/salary-save", json={"records": bad, "reviewed": True}
        ).status_code
        == 400
    )
    good = client.post(
        "/suite/api/worksheets/salary-save", json={"records": records, "reviewed": True}
    )
    assert good.status_code == 201, good.data
    ledger = client.get("/suite/api/ledger?employee_code=SYN-01").json
    assert ledger["payment_count"] == 1
    assert ledger["months"][0]["regular_salary"][0]["printed_net"] == "41272"
    loaded = client.get(
        "/suite/api/worksheets/ledger-load?employee_code=SYN-01&start=2040-07-01&end=2040-07-31"
    )
    assert loaded.status_code == 200, loaded.data
    row = loaded.json["rows"][0]
    assert row["drawn_HRA"] == "0.00"
    assert row["ledger_payment_id"] == good.json["ids"][0]
    p = example()
    p["rows"][0].update(row)
    calc = client.post("/suite/api/worksheets/calculate", json=p)
    assert calc.json["status"] == "ready", calc.json
    assert calc.json["bill"]["totals"]["net_payable"] == "4960.00"


def test_late_duplicate_salary_rolls_back_whole_batch(client):
    records = salaries()
    assert (
        client.post(
            "/suite/api/worksheets/salary-save",
            json={"records": records, "reviewed": True},
        ).status_code
        == 201
    )
    first = deepcopy(records[0])
    first.update(pay_month="2040-06", bill_no="SYN0")
    response = client.post(
        "/suite/api/worksheets/salary-save",
        json={"records": [first, *records], "reviewed": True},
    )
    assert response.status_code == 400
    assert (
        client.get("/suite/api/ledger?employee_code=SYN-01").json["payment_count"] == 1
    )


def test_saved_ledger_link_changes_block_calculation_and_export(client):
    client.post(
        "/suite/api/worksheets/salary-save",
        json={"records": salaries(), "reviewed": True},
    )
    linked = client.get(
        "/suite/api/worksheets/ledger-load?employee_code=SYN-01&start=2040-07-01&end=2040-07-31"
    ).json
    p = example()
    p["rows"][0].update(linked["rows"][0])
    saved = client.post("/suite/api/worksheets", json={"payload": p}).json
    p["rows"][0]["drawn_HRA"] = "1"
    changed = client.post("/suite/api/worksheets/calculate", json=p)
    assert changed.json["status"] == "incomplete"
    assert any(e["field"] == "ledger_snapshot" for e in changed.json["errors"])
    from ifms_suite.storage import database

    with client.application.app_context(), database() as db:
        value = json.loads(
            db.execute("SELECT payload FROM suite_payments").fetchone()[0]
        )
        value["reference"] = "Changed source review"
        db.execute("UPDATE suite_payments SET payload=?", (json.dumps(value),))
    assert (
        client.get(f"/suite/api/worksheets/{saved['id']}").json["result"]["status"]
        == "incomplete"
    )
    assert client.get(f"/suite/api/worksheets/{saved['id']}/pdf").status_code == 400
    assert client.get(f"/suite/api/worksheets/{saved['id']}/excel").status_code == 400


def test_salary_preview_does_not_save_or_approve_records(client):
    from ifms_suite.worksheet_excel import salary_template, SALARY_COLUMNS

    w = load_workbook(salary_template())
    vals = dict(
        employee_code="SYN-01",
        employee_name="Synthetic Person",
        pay_month="2040-07",
        bill_no="SYN1",
        BASIC=24800,
        DA=14880,
        HRA=0,
        HDA=3307,
        CCA=0,
        OTHER_ALLOWANCES=3350,
        total_deductions=5065,
        printed_gross=46337,
        printed_net=41272,
        reference="Synthetic original",
    )
    w["Salary ledger"].append([vals[k] for k in SALARY_COLUMNS])
    stream = BytesIO()
    w.save(stream)
    stream.seek(0)
    preview = client.post(
        "/suite/api/worksheets/salary-preview", data={"file": (stream, "salary.xlsx")}
    )
    assert preview.status_code == 200, preview.data
    assert preview.json["records"][0]["reviewed"] is False
    assert (
        client.get("/suite/api/ledger?employee_code=SYN-01").json["payment_count"] == 0
    )

"""Synthetic bill register export contracts; no employee or uploaded data."""

from copy import deepcopy

import pymupdf
import pytest
from openpyxl import load_workbook

from ifms_suite.bill_register_exports import (
    OFFICIAL_COLUMNS,
    OFFICIAL_REFERENCE,
    PAPER_SIZES,
    export_register,
    map_official,
)


def test_raw_excel_is_literal_editable_and_preserves_selected_order_and_values():
    payload = {
        "columns": ["Bill", "Gross", "Net", "Blank", "Text"],
        "rows": [
            {
                "Bill": "0000123",
                "Gross": "1,200.00",
                "Net": "0010.00",
                "Blank": None,
                "Text": '=HYPERLINK("https://example.test","x")',
                "Other": "omit",
            },
            {"Bill": "+001", "Gross": "-2.50", "Net": "0", "Text": "@literal"},
        ],
    }
    before = deepcopy(payload)
    book = load_workbook(export_register(payload))
    sheet = book.active
    assert payload == before
    assert list(next(sheet.values)) == payload["columns"]
    assert list(sheet.values)[1] == (
        "0000123",
        "1,200.00",
        "0010.00",
        None,
        '=HYPERLINK("https://example.test","x")',
    )
    assert list(sheet.values)[2] == ("+001", "-2.50", "0", None, "@literal")
    assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert sheet["A2"].data_type == "s"
    assert not sheet.protection.sheet
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.print_title_rows == "$1:$1"
    assert sheet.page_setup.fitToHeight == 0


@pytest.mark.parametrize("paper", ["A4", "FS"])
def test_pdf_exact_page_size_and_all_columns_paginate_without_clipping(paper):
    columns = [f"Column{n:02}" for n in range(19)]
    rows = [
        {column: f"marker{n:02}_{r:03}" for n, column in enumerate(columns)}
        for r in range(80)
    ]
    pdf = pymupdf.open(
        stream=export_register(
            {"columns": columns, "rows": rows}, format="pdf", paper=paper
        ).getvalue(),
        filetype="pdf",
    )
    assert len(pdf) > 3
    text = "\n".join(page.get_text() for page in pdf)
    for column in columns:
        assert column in text
    for row in rows:
        for value in row.values():
            assert value in text
    for page in pdf:
        assert (page.rect.width, page.rect.height) == pytest.approx(
            PAPER_SIZES[paper], abs=0.01
        )
        for x0, y0, x1, y1, *_ in page.get_text("blocks"):
            assert x0 >= 23 and x1 <= page.rect.width - 23
            assert y0 >= 0 and y1 < page.rect.height


def test_pdf_splits_a_single_very_tall_row_and_escapes_markup():
    value = " ".join(f"word{i:04}" for i in range(1800))
    payload = {
        "columns": ["Detail", "Markup"],
        "rows": [{"Detail": value, "Markup": "<b>literal</b> & 001"}],
    }
    pdf = pymupdf.open(
        stream=export_register(payload, format="pdf").getvalue(), filetype="pdf"
    )
    assert len(pdf) > 1
    text = " ".join(page.get_text() for page in pdf)
    assert "word0000" in text and "word1799" in text
    assert "<b>literal</b> & 001" in text


def test_fs_workbook_uses_explicit_custom_dimensions():
    sheet = load_workbook(
        export_register({"columns": ["A"], "rows": []}, paper="FS")
    ).active
    assert sheet.page_setup.paperWidth == "13in"
    assert sheet.page_setup.paperHeight == "8.5in"
    assert sheet.page_setup.orientation == "landscape"
    assert "13 x 8.5" in sheet.oddFooter.center.text


def test_excel_rejects_long_cells_instead_of_silently_truncating():
    with pytest.raises(ValueError, match="32767"):
        export_register({"columns": ["Detail"], "rows": [{"Detail": "x" * 32768}]})


def test_pdf_never_silently_substitutes_unsupported_characters():
    with pytest.raises(ValueError, match="use Excel"):
        export_register(
            {"columns": ["Detail"], "rows": [{"Detail": "\U00020000"}]}, format="pdf"
        )


def test_original_header_labels_and_draft_annotation_preserve_table_values():
    payload = {
        "columns": ["billNo"],
        "column_labels": {"billNo": "Bill No."},
        "rows": [{"billNo": "0001"}],
        "status": "draft",
    }
    book = load_workbook(export_register(payload))
    assert list(book.active.values) == [("Bill No.",), ("0001",)]
    assert "DRAFT" in book["Export status"]["A1"].value
    assert "DRAFT" in book.active.oddFooter.center.text
    pdf = pymupdf.open(
        stream=export_register(payload, format="pdf").getvalue(), filetype="pdf"
    )
    assert all("DRAFT" in page.get_text() for page in pdf)
    assert "Bill No." in pdf[0].get_text() and "0001" in pdf[0].get_text()


def test_official_is_blocked_until_authoritative_form_and_columns_are_verified():
    assert OFFICIAL_REFERENCE["verified"] is False
    assert OFFICIAL_REFERENCE["form"] is None
    assert OFFICIAL_REFERENCE["page"] is None
    assert not OFFICIAL_COLUMNS
    assert "finance.rajasthan.gov.in" in OFFICIAL_REFERENCE["url"]
    for action in [
        lambda: map_official([{"billNo": "001", "gross": "10"}]),
        lambda: export_register({"columns": ["billNo"], "rows": []}, mode="official"),
    ]:
        with pytest.raises(ValueError, match="not verified"):
            action()


@pytest.mark.parametrize(
    "kwargs", [{"mode": "fake"}, {"format": "csv"}, {"paper": "Legal"}]
)
def test_rejects_unknown_export_options(kwargs):
    with pytest.raises(ValueError):
        export_register({"columns": ["A"], "rows": []}, **kwargs)


@pytest.mark.parametrize(
    "payload",
    [
        None,
        {},
        {"columns": [], "rows": []},
        {"columns": ["A", "A"], "rows": []},
        {"columns": ["A"], "rows": ["bad"]},
    ],
)
def test_rejects_invalid_payload(payload):
    with pytest.raises(ValueError):
        export_register(payload)

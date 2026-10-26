import io
from openpyxl import load_workbook
import pymupdf
from ifms_suite.exports import excel,pdf
from ifms_suite.calculation import calculate_claim

def test_both_exports_show_deductions_and_zero_cash():
    p=dict(employee_code='RJ1',allowance='DA',start='2025-01-01',end='2025-01-31',status='draft')
    c=calculate_claim(p['start'],p['end'],[dict(start=p['start'],end=p['end'],basic='50000',due_rate='55',drawn_rate='53')],[dict(name='GPF',mode='percent',value='100')])
    w=load_workbook(excel(p,c));assert 'Deductions' in w.sheetnames
    assert w['Deductions']['B2'].value==1000
    text=''.join(page.get_text() for page in pymupdf.open(stream=pdf(p,c).getvalue(),filetype='pdf'))
    assert 'Cash payable: INR 0.00' in text
    assert 'GPF' in text

def test_fixed_allowance_does_not_require_a_basic_pay_value_for_export():
    p=dict(employee_code='RJ1',allowance='CCA',start='2025-01-01',end='2025-01-31')
    c=calculate_claim(p['start'],p['end'],[dict(start=p['start'],end=p['end'],basis='fixed',basic='',due_amount='700',drawn_amount='500')])
    assert load_workbook(excel(p,c)).active['K3'].value==200

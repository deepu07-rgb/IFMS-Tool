import io
from openpyxl import load_workbook
from test_suite_workflow import client


def test_ga55_preserves_original_pdf_and_exports_structured_excel(client):
    import pymupdf

    d = pymupdf.open()
    p = d.new_page()
    p.insert_text(
        (30, 40),
        "March\n2025\nRegular Salary\n123456789\nBasic\n50000.00\n50,000.00\nGPF\n1000.00\n1,000.00\n49,000.00",
    )
    data = d.tobytes()
    d.close()
    id = client.post(
        "/suite/api/imports",
        data={"kind": "ga55", "file": (io.BytesIO(data), "ga.pdf")},
    ).json["id"]
    assert client.get(f"/suite/api/imports/{id}/pdf").data == data
    response = client.get(f"/suite/api/imports/{id}/excel")
    assert response.status_code == 200
    w = load_workbook(io.BytesIO(response.data))
    assert "Payment records" in w.sheetnames
    assert "Basic" in [c.value for c in w["Payment records"][2]]
    assert w["Payment records"]["A1"].value.startswith("DRAFT")


def test_excel_statement_exports_readable_pdf_and_structured_rows(client):
    from openpyxl import Workbook
    import pymupdf

    w = Workbook()
    s = w.active
    s.title = "Arrear"
    s.append(
        [
            "Sr",
            "Month",
            "Pay Due",
            "",
            "",
            "",
            "Pay Drawn",
            "",
            "",
            "",
            "Pay Difference",
            "",
            "",
            "",
        ]
    )
    s.append(
        [
            "",
            "",
            "Basic",
            "DA",
            "HRA",
            "Total",
            "Basic",
            "DA",
            "HRA",
            "Total",
            "Basic",
            "DA",
            "HRA",
            "Total",
        ]
    )
    s.append(
        [
            1,
            "Jan-2025",
            50000,
            27500,
            10000,
            87500,
            50000,
            26500,
            10000,
            86500,
            0,
            1000,
            0,
            1000,
        ]
    )
    b = io.BytesIO()
    w.save(b)
    original = b.getvalue()
    b.seek(0)
    r = client.post(
        "/suite/api/imports", data={"kind": "due_drawn", "file": (b, "statement.xlsx")}
    )
    id = r.json["id"]
    response = client.get(f"/suite/api/imports/{id}/pdf")
    assert response.status_code == 200
    with pymupdf.open(stream=response.data, filetype="pdf") as d:
        text = "".join(p.get_text() for p in d)
        assert "2025-01" in text and "1000.00" in text and "DRAFT" in text
    assert client.get(f"/suite/api/imports/{id}/source").data == original
    export = client.get(f"/suite/api/imports/{id}/excel")
    w = load_workbook(io.BytesIO(export.data))
    assert "Statement review" in w.sheetnames
    assert w["Statement review"]["A3"].value == "2025-01"


def test_statement_export_preserves_extra_heads_and_unresolved_cells(client):
    import json
    from ifms_suite.storage import database
    from openpyxl import load_workbook

    payload = {
        "file_kind": "xlsx",
        "statement_rows": [
            {
                "month": "2025-02",
                "due": {
                    "Basic": "50000.00",
                    "DA": "27500.00",
                    "HDA": None,
                    "CCA": "500.00",
                    "HRA": "10000.00",
                },
                "drawn": {
                    "Basic": "50000.00",
                    "DA": "26500.00",
                    "HDA": "900.00",
                    "CCA": "400.00",
                    "HRA": "9000.00",
                },
                "printed_due": "89000.00",
                "printed_drawn": "86800.00",
                "printed_difference": "2200.00",
                "calculated_difference": "2200.00",
                "source_row": 3,
                "source_sheet": "Arrear",
                "issues": ["Due HDA is missing or unreadable"],
            }
        ],
    }
    with client.application.app_context(), database() as db:
        source_id = db.execute(
            "INSERT INTO suite_imports(sha256,filename,payload) VALUES(?,?,?)",
            ("synthetic-expanded-statement", "synthetic.xlsx", json.dumps(payload)),
        ).lastrowid
    response = client.get(f"/suite/api/imports/{source_id}/excel")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook["Statement review"]
    headers = [cell.value for cell in sheet[2]]
    assert sheet.cell(3, headers.index("Due HDA") + 1).value is None
    assert sheet.cell(3, headers.index("Drawn HDA") + 1).value == 900
    assert sheet.cell(3, headers.index("Due CCA") + 1).value == 500
    assert "HDA" in sheet.cell(3, headers.index("Issues") + 1).value
    response = client.get(f"/suite/api/imports/{source_id}/pdf")
    assert response.status_code == 200
    import pymupdf

    with pymupdf.open(stream=response.data, filetype="pdf") as report:
        text = "".join(page.get_text() for page in report)
    assert "Due HDA" in text
    assert "Drawn CCA" in text
    assert "900.00" in text
    assert "not read" in text.lower()

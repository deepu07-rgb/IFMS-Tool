import json
from io import BytesIO
from openpyxl import load_workbook
from test_suite_workflow import client
from ifms_suite.storage import database


def test_filtered_rows_and_exports_use_same_signed_outer_head(client):
    payload = dict(
        kind="outer",
        file_kind="pdf",
        pages=[],
        extracted={
            "Reference Number": "REF",
            "Bill Date Parsed": "2026-07-01",
            "Bill Name": "Salary",
            "Gross Amount": 90,
            "Net Amount": 85,
            "Allowances": {"BASIC": 100, "ROP": -10},
            "Deductions": {"GPF": 5},
        },
    )
    with client.application.app_context(), database() as db:
        db.execute(
            "INSERT INTO suite_imports(sha256,filename,payload) VALUES(?,?,?)",
            ("a" * 64, "outer.pdf", json.dumps(payload)),
        )
    query = "kind=outer&amount_field=allowance:ROP&minimum=-11&maximum=-9"
    response = client.get("/suite/api/document_records?" + query)
    assert response.status_code == 200
    assert response.json["count"] == 1
    assert response.json["records"][0]["allowances"]["ROP"] == -10
    excel = client.get("/suite/api/document_records/excel?" + query)
    assert excel.status_code == 200
    w = load_workbook(BytesIO(excel.data))
    assert w.active.max_row == 3
    assert w.active.auto_filter.ref
    pdf = client.get("/suite/api/document_records/pdf?" + query)
    assert pdf.status_code == 200
    import pymupdf

    text = "".join(p.get_text() for p in pymupdf.open(stream=pdf.data, filetype="pdf"))
    assert "ROP" in text and "-10" in text
    assert client.get("/suite/api/document_records?kind=ga55").json["count"] == 0


def test_ga55_case_extraction_is_visible_even_without_manual_bill_corrections(client):
    payload={'kind':'ga55','machine_facts':{'bills':[{'employee_code':'SYNTHETIC','pay_month':'2026-07','bill_type':'Regular Salary','bill_no':'123','allowances':{'BASIC':'100'},'deductions':{},'printed_gross':'100','printed_net':None}]},'machine_facts_version':'v6'}
    with client.application.app_context(),database() as db:
        db.execute('INSERT INTO suite_imports(sha256,filename,payload) VALUES(?,?,?)',('b'*64,'scan.pdf',json.dumps(payload)))
    result=client.get('/suite/api/document_records?kind=ga55')
    assert result.status_code==200
    assert result.json['count']==1
    assert result.json['records'][0]['gross']=='100'

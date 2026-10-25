import json
import re
import sqlite3
from io import BytesIO

import openpyxl
import pytest

from app import app, init_db
from ifms_engine.db import init_v2_schema


pytestmark = pytest.mark.v2


def _advanced(filters):
    return json.dumps(filters)


def _client_with_seeded_db(tmp_path):
    db_path = tmp_path / "advanced_filter_ui_api.db"
    export_dir = tmp_path / "exports"
    export_dir.mkdir()

    app.config["DATABASE"] = str(db_path)
    app.config["EXPORT_FOLDER"] = str(export_dir)
    app.config["USE_V2_PARSER"] = True
    init_db()
    init_v2_schema(str(db_path))

    db = sqlite3.connect(str(db_path))
    bill_id = db.execute(
        """
        INSERT INTO bills_v2 (
            source_filename, group_name, month_year, bill_no,
            parsed_gross, parsed_deduction, parsed_net,
            official_gross, official_deduction, official_net,
            verification_status
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            "filter_fixture.pdf",
            "FILTER GROUP",
            "04/2026",
            "BILL-1",
            330000,
            33000,
            297000,
            330000,
            33000,
            297000,
            "VERIFIED",
        ),
    ).lastrowid

    rows = [
        {
            "name": "ALPHA OFFICER",
            "employee_id": "EMP001",
            "pan": "AAAAA1111A",
            "pay_scale": "L10c",
            "raw_designation": "IPS, JUNIOR SCALE",
            "gross": 85000,
            "deduction": 10000,
            "net": 75000,
            "components": {
                "basic_pay": ("allowance", 55000),
                "da": ("allowance", 30000),
                "income_tax": ("deduction", 6000),
                "gpf_2004": ("deduction", 1000),
                "gpf": ("deduction", 3000),
            },
        },
        {
            "name": "BETA CONSTABLE",
            "employee_id": "EMP002",
            "pan": "BBBBB2222B",
            "pay_scale": "L7",
            "raw_designation": "CONSTABLE",
            "final_designation": "Constable",
            "gross": 60000,
            "deduction": 7000,
            "net": 53000,
            "components": {
                "basic_pay": ("allowance", 45000),
                "da": ("allowance", 15000),
                "income_tax": ("deduction", 5000),
                "gpf": ("deduction", 2000),
            },
        },
        {
            "name": "GAMMA IPS",
            "employee_id": "EMP003",
            "pan": "CCCCC3333C",
            "pay_scale": "L12c",
            "raw_designation": "IPS, SENIOR TIME SCALE",
            "gross": 125000,
            "deduction": 15000,
            "net": 110000,
            "components": {
                "basic_pay": ("allowance", 70000),
                "da": ("allowance", 55000),
                "income_tax": ("deduction", 12000),
                "gpf": ("deduction", 3000),
            },
        },
        {
            "name": "DELTA INSPECTOR",
            "employee_id": "EMP004",
            "pan": "DDDDD4444D",
            "pay_scale": "L14c",
            "raw_designation": "INSPECTOR",
            "final_designation": "Inspector",
            "gross": 60000,
            "deduction": 1000,
            "net": 59000,
            "components": {
                "basic_pay": ("allowance", 65000),
                "da": ("allowance", -5000),
                "rghs": ("deduction", 1000),
            },
        },
    ]

    for sr_no, row in enumerate(rows, start=1):
        final_designation = row.get("final_designation", "Needs Review")
        employee_id = db.execute(
            """
            INSERT INTO employees_v2 (
                bill_id, source_filename, page_no, sr_no, group_name, month_year,
                raw_name, normalized_name, raw_designation, raw_designation_full,
                raw_designation_en, normalized_designation, final_designation,
                final_designation_en, designation_source, designation_needs_review,
                employee_data_status, row_verification_status, employee_id,
                pan_raw, pay_scale_raw, gross_amount, sum_deduction, net_total
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                bill_id,
                "filter_fixture.pdf",
                1,
                sr_no,
                "FILTER GROUP",
                "04/2026",
                row["name"],
                row["name"],
                row["raw_designation"],
                row["raw_designation"],
                row["raw_designation"],
                final_designation,
                final_designation,
                final_designation,
                "seed",
                0,
                "VERIFIED",
                "VERIFIED",
                row["employee_id"],
                row["pan"],
                row["pay_scale"],
                row["gross"],
                row["deduction"],
                row["net"],
            ),
        ).lastrowid
        for head, (component_type, amount) in row["components"].items():
            db.execute(
                """
                INSERT INTO employee_components_v2 (
                    employee_id, bill_id, component_type, raw_head,
                    normalized_head, amount, included_in_total, confidence, source_page
                ) VALUES (?,?,?,?,?,?,?,?,?)
                """,
                (employee_id, bill_id, component_type, head, head, amount, 1, 1.0, 1),
            )
    db.commit()
    db.close()
    return app.test_client()


def _employees(client, **params):
    response = client.get("/api/employees", query_string={"page": 1, "per_page": 50, **params})
    assert response.status_code == 200
    return response.get_json()


def _levels(rows):
    for row in rows:
        match = re.search(r"\d+", row["pay_scale"])
        assert match
        yield int(match.group(0))


def test_v3_advanced_filter_ui_and_api(tmp_path):
    client = _client_with_seeded_db(tmp_path)

    add = client.post(
        "/api/v2/dictionary",
        json={
            "category": "designation",
            "raw_text": "IPS",
            "final_text": "IPS",
            "match_type": "contains",
            "active": True,
            "apply_existing": False,
        },
    )
    assert add.status_code == 200
    apply = client.post("/api/v3/dictionary/apply", json={"mapping_id": add.get_json()["id"]})
    assert apply.status_code == 200
    assert apply.get_json()["updated_count"] == 2

    itax_gt = _employees(
        client,
        advanced_filters=_advanced([{"column": "ITax", "operator": "greater than", "value": "5000"}]),
    )
    assert {row["employee_id"] for row in itax_gt["employees"]} == {"EMP001", "EMP003"}
    assert all(float(row["itax"] or 0) > 5000 for row in itax_gt["employees"])
    assert any("Income Tax greater than 5000" == chip["text"] for chip in itax_gt["filters"]["chips"])

    itax_eq = _employees(
        client,
        advanced_filters=_advanced([{"column": "ITax", "operator": "equals", "value": "5000"}]),
    )
    assert [row["employee_id"] for row in itax_eq["employees"]] == ["EMP002"]

    basic_between = _employees(
        client,
        advanced_filters=_advanced([{"column": "Basic Pay", "operator": "between", "value": "50000", "value2": "60000"}]),
    )
    assert [row["employee_id"] for row in basic_between["employees"]] == ["EMP001"]

    pay_scale_between = _employees(
        client,
        advanced_filters=_advanced([{"column": "Pay Scale", "operator": "between", "value": "L8", "value2": "L12"}]),
    )
    assert {row["employee_id"] for row in pay_scale_between["employees"]} == {"EMP001", "EMP003"}
    assert all(8 <= level <= 12 for level in _levels(pay_scale_between["employees"]))

    gpf_2004_exists = _employees(
        client,
        advanced_filters=_advanced([{"column": "GPF 2004", "operator": "exists"}]),
    )
    assert [row["employee_id"] for row in gpf_2004_exists["employees"]] == ["EMP001"]

    ips = _employees(
        client,
        advanced_filters=_advanced([{"column": "Final Designation", "operator": "contains", "value": "IPS"}]),
    )
    assert {row["employee_id"] for row in ips["employees"]} == {"EMP001", "EMP003"}
    assert all("IPS" in row["final_designation"] for row in ips["employees"])

    combined = _employees(
        client,
        advanced_filters=_advanced([
            {"column": "Final Designation", "operator": "contains", "value": "IPS"},
            {"column": "ITax", "operator": "greater than", "value": "5000"},
        ]),
    )
    assert {row["employee_id"] for row in combined["employees"]} == {"EMP001", "EMP003"}

    no_result = _employees(
        client,
        advanced_filters=_advanced([{"column": "Employee Name", "operator": "equals", "value": "NO SUCH EMPLOYEE"}]),
    )
    assert no_result["total"] == 0
    assert no_result["employees"] == []
    assert no_result["filters"]["no_result_message"] == "No matching records found."

    export = client.get(
        "/api/export_filtered",
        query_string={
            "advanced_filters": _advanced([{"column": "ITax", "operator": "greater than", "value": "5000"}])
        },
    )
    assert export.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(export.data), data_only=True)
    assert "Filtered Employee Data" in workbook.sheetnames
    sheet = workbook["Filtered Employee Data"]
    headers = [cell.value for cell in sheet[1]]
    id_col = headers.index("employee_id") + 1
    itax_col = headers.index("ITax") + 1
    exported_ids = [row[id_col - 1].value for row in sheet.iter_rows(min_row=2)]
    exported_itax = [row[itax_col - 1].value for row in sheet.iter_rows(min_row=2)]
    assert set(exported_ids) == {"EMP001", "EMP003"}
    assert all(float(value or 0) > 5000 for value in exported_itax)

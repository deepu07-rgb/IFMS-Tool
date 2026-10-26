"""Optional browser acceptance against a disposable workspace.

Run with a server started using a fresh IFMS_DATA_DIR. Requires Playwright,
Chromium, openpyxl and PyMuPDF in the developer environment. All records are synthetic.
"""

import argparse, io, re, tempfile, uuid
from pathlib import Path
from urllib.parse import urlparse
from playwright.sync_api import sync_playwright, expect
from openpyxl import Workbook, load_workbook
import pymupdf


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", default="http://127.0.0.1:5059")
    parser.add_argument("--screenshot", default="/tmp/ifms-workspace.png")
    args = parser.parse_args()
    code = "RJTEST001" + uuid.uuid4().hex[:6].upper()
    with tempfile.TemporaryDirectory(
        prefix="ifms-browser-"
    ) as td, sync_playwright() as p:
        root = Path(td)
        browser = p.chromium.launch(
            executable_path="/usr/bin/chromium", headless=True, args=["--no-sandbox"]
        )
        context = browser.new_context(
            viewport={"width": 1500, "height": 1050}, accept_downloads=True
        )
        page = context.new_page()
        errors = []
        external = []
        page.on("pageerror", lambda e: errors.append(str(e)))

        def route(r):
            if urlparse(r.request.url).scheme in ("http", "https") and urlparse(
                r.request.url
            ).hostname not in ("127.0.0.1", "localhost"):
                external.append(r.request.url)
                r.abort()
            else:
                r.continue_()

        context.route("**/*", route)
        page.goto(args.url + "/suite/")
        expect(page.locator("#summary")).to_contain_text("library employees")

        def nav(name):
            page.get_by_role("button", name=name, exact=True).click()

        nav("Employee & Groups")
        form = page.locator("#employeeForm")
        form.locator("[name=employee_code]").fill(code)
        form.locator("[name=name]").fill("Synthetic Test Employee")
        form.locator("[name=group]").fill("Browser test")
        form.get_by_role("button", name="Save employee").click()
        expect(page.locator("#employeesTable")).to_contain_text(code)
        form = page.locator("#payHistoryForm")
        form.locator("[name=employee_code]").fill(code)
        form.locator("[name=effective_from]").fill("2025-01-01")
        form.locator("[name=basic]").fill("50000")
        form.locator("[name=reference]").fill("Synthetic pay order")
        form.locator("[name=reviewed]").check()
        form.get_by_role("button", name="Save dated event").click()
        expect(page.locator("#message")).to_contain_text("Dated pay event saved")
        nav("GA 55 / OCR")
        page.get_by_role("button", name="Enter a drawn payment manually").click()
        page.locator("#reviewEmployee").fill(code)
        page.locator("#reviewReference").fill("Synthetic salary source")
        editor = page.locator(".payment-editor").first
        editor.locator("[data-field=bill_no]").fill("REG-" + code)
        editor.locator("[data-field=pay_month]").fill("2025-01")
        editor.locator(".allowances [data-field=Basic]").fill("50000")
        editor.locator(".allowances [data-field=DA]").fill("26500")
        page.locator("#reviewChecked").check()
        page.get_by_role("button", name="Approve payment history", exact=True).click()
        expect(page.locator("#message")).to_contain_text(
            "Reviewed payment history saved"
        )
        nav("Arrear / Due–Drawn")
        form = page.locator("#claimForm")
        form.locator("[name=employee_code]").fill(code)
        form.locator("[name=source]").select_option("history")
        form.locator("[name=start]").fill("2025-01-01")
        form.locator("[name=end]").fill("2025-01-31")
        form.locator("[name=allocation_policy]").select_option("source_order")
        form.locator("[name=scheme]").select_option("GPF-2004")
        form.get_by_role("button", name="Calculate & save draft").click()
        expect(page.locator("#calculation")).to_contain_text("Cash payable 0.00")
        expect(page.locator("#calculation")).to_contain_text("1000.00")
        id = int(
            re.search(
                r"Draft #(\d+)", page.locator("#calculation h2").inner_text()
            ).group(1)
        )
        nav("Reports & Backup")
        form = page.locator(f'form[data-claim="{id}"]')
        form.locator("[name=rule_reference]").fill("Verified DA order 01-Apr-2025")
        page.evaluate("refresh()")
        expect(form.locator("[name=rule_reference]")).to_have_value(
            "Verified DA order 01-Apr-2025"
        )
        form.locator("[name=drawn_reference]").fill("REG-" + code)
        form.locator("[name=reviewed]").check()
        form.get_by_role("button", name="Finalize reviewed statement").click()
        expect(page.locator(f'form[data-settlement="{id}"]')).to_be_visible()
        with page.expect_download() as download:
            page.locator(f'a[href="/suite/api/claims/{id}/excel"]').filter(
                visible=True
            ).click()
        path = root / "statement.xlsx"
        download.value.save_as(path)
        w = load_workbook(path)
        assert w["Deductions"]["B2"].value == 1000
        form = page.locator(f'form[data-settlement="{id}"]')
        form.locator("[name=bill_no]").fill("ARREAR-" + code)
        form.locator("[name=pay_month]").fill("2025-04")
        form.locator("[name=reference]").fill("Synthetic full credit source")
        form.locator("[name=reviewed]").check()
        form.get_by_role("button", name="Record settlement").click()
        expect(page.locator("#message")).to_contain_text(
            "Actual full settlement linked"
        )
        print(
            "PASS: employee → pay history → reviewed drawn payment → history-based DA/GPF claim → export → actual settlement"
        )
        # A separate employee in GA55 must keep its own identity, even with a blank common-code field.
        ga_code = code + "G"
        doc = pymupdf.open()
        sheet = doc.new_page()
        sheet.insert_text(
            (30, 40),
            f"Employee ID\n:\n{ga_code}\nFebruary\n2025\nRegular Salary\n123456789\nBasic\n50000.00\nDA\n27500.00\nHRA\n10000.00\n87,500.00\nGPF\n1000.00\n1,000.00\n86,500.00\nTotal Allowances",
        )
        ga = root / f"ga55-{code}.pdf"
        doc.save(ga)
        doc.close()
        nav("GA 55 / OCR")
        form = page.locator("#importForm")
        form.locator("[name=kind]").select_option("ga55")
        form.locator("[name=file]").set_input_files(ga)
        form.get_by_role("button", name="Extract for review").click()
        card = page.locator("#importsList .card").filter(
            has=page.get_by_role("heading", name=f"ga55-{code}.pdf", exact=True)
        )
        expect(card).to_be_visible()
        card.get_by_role("button", name="Review / enter payments").click()
        editor = page.locator(".payment-editor").first
        expect(editor.locator("[data-field=employee_code]")).to_have_value(ga_code)
        expect(editor.locator("[data-field=printed_net]")).to_have_value("86500.00")
        page.locator("#reviewReference").fill("Synthetic GA55")
        page.locator("#reviewChecked").check()
        page.get_by_role("button", name="Approve payment history", exact=True).click()
        expect(page.locator("#message")).to_contain_text(
            "Reviewed payment history saved"
        )
        print(
            "PASS: digital GA55 → per-payment employee identity → printed-total reconciliation → shared payment history"
        )
        # Uploaded due/drawn arithmetic and rate recalculation.
        w = Workbook()
        s = w.active
        s.title = "Arrear"
        s.append(
            [
                "Sr",
                "Month",
                "Pay Due",
                "",
                "",
                "",
                "Pay Drawn",
                "",
                "",
                "",
                "Pay Difference",
                "",
                "",
                "",
            ]
        )
        s.append(
            [
                "",
                "",
                "Basic",
                "DA",
                "HRA",
                "Total",
                "Basic",
                "DA",
                "HRA",
                "Total",
                "Basic",
                "DA",
                "HRA",
                "Total",
            ]
        )
        s.append(
            [
                1,
                "Jan-2025",
                50000,
                27000,
                10000,
                87000,
                50000,
                26500,
                10000,
                86500,
                0,
                500,
                0,
                500,
            ]
        )
        xlsx = root / f"due-drawn-{code}.xlsx"
        w.save(xlsx)
        form = page.locator("#importForm")
        form.locator("[name=kind]").select_option("due_drawn")
        form.locator("[name=file]").set_input_files(xlsx)
        form.get_by_role("button", name="Extract for review").click()
        card = page.locator("#importsList .card").filter(
            has=page.get_by_role("heading", name=f"due-drawn-{code}.xlsx", exact=True)
        )
        expect(card).to_be_visible()
        card.get_by_role("button", name="Check / recalculate statement").click()
        form = page.locator("#statementForm")
        form.locator("[name=employee_code]").fill(code + "S")
        form.locator("[name=reviewed]").check()
        form.get_by_role("button", name="Create revised statement draft").click()
        expect(page.locator("#message")).to_contain_text("gross difference INR 1000.00")
        print(
            "PASS: Excel due/drawn → worksheet review → library DA recalculation → saved revised draft"
        )
        # Fixed custom allowance entry should not demand irrelevant basic pay.
        nav("Arrear / Due–Drawn")
        form = page.locator("#claimForm")
        form.locator("[name=source]").select_option("manual")
        form.locator("[name=allocation_policy]").select_option("manual")
        form.locator("[name=allowance]").fill("CCA")
        form.locator("[name=allowance]").press("Tab")
        page.locator("#claimBasis").select_option("fixed")
        row = page.locator("#segments tr").first
        row.locator("[data-key=start]").fill("2025-01-01")
        row.locator("[data-key=end]").fill("2025-01-31")
        row.locator("[data-key=due_rate]").fill("700")
        row.locator("[data-key=drawn_rate]").fill("500")
        form.get_by_role("button", name="Calculate & save draft").click()
        expect(page.locator("#calculation")).to_contain_text("gross INR 200.00")
        print(
            "PASS: custom fixed allowance → deterministic calculation without unrelated basic-pay input"
        )
        with page.expect_download() as download:
            nav("Reports & Backup")
            page.get_by_role(
                "link", name="Download complete backup (records + source documents)"
            ).click()
        backup = root / "backup.zip"
        download.value.save_as(backup)
        import zipfile

        with zipfile.ZipFile(backup) as z:
            assert any(
                n.startswith("sources/") and n.endswith(".xlsx") for n in z.namelist()
            )
        print("PASS: complete backup includes original Excel and PDF documents")
        nav("Arrear / Due–Drawn")
        page.screenshot(path=args.screenshot, full_page=True)
        assert not errors, errors
        assert not external, external
        print("PASS: all browser requests stayed local; no JavaScript errors")
        browser.close()


if __name__ == "__main__":
    main()

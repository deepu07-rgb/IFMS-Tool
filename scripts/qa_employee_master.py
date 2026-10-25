import argparse
import os
import sqlite3
import sys
from datetime import datetime

import pandas as pd
from openpyxl.styles import Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(__file__))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from app import app, employee_status_case_sql, init_db, process_file_v2
from ifms_engine.db import init_v2_schema


FIXTURE_DIR = os.path.join(ROOT, "test_fixtures", "ifms_inner_sheets")
EXPORT_DIR = os.path.join(ROOT, "exports")


def _connect(db_path):
    db = sqlite3.connect(db_path)
    db.row_factory = sqlite3.Row
    return db


def _process_fixtures_if_needed(db_path):
    app.config["DATABASE"] = db_path
    app.config["USE_V2_PARSER"] = True
    init_db()
    init_v2_schema(db_path)
    db = _connect(db_path)
    count = db.execute("SELECT COUNT(*) FROM employees_v2").fetchone()[0]
    db.close()
    if count:
        return
    pdfs = sorted(
        [p for p in os.listdir(FIXTURE_DIR) if p.lower().endswith(".pdf")],
        key=lambda name: int(os.path.splitext(name)[0]) if os.path.splitext(name)[0].isdigit() else name,
    )
    for pdf in pdfs:
        process_file_v2(os.path.join(FIXTURE_DIR, pdf), pdf, dict(app.config))


def _df(db, sql, params=()):
    return pd.DataFrame([dict(r) for r in db.execute(sql, params).fetchall()])


def build_workbook(db_path, export_dir=EXPORT_DIR):
    os.makedirs(export_dir, exist_ok=True)
    _process_fixtures_if_needed(db_path)
    db = _connect(db_path)
    status_case = employee_status_case_sql()

    qa_summary = _df(db, f"""
        SELECT 'total employees' AS metric, COUNT(*) AS value FROM employees_v2
        UNION ALL SELECT 'unresolved issues', COUNT(*) FROM validation_issues_v2 WHERE COALESCE(status, 'OPEN') != 'RESOLVED'
        UNION ALL SELECT 'component mismatch issues', COUNT(*) FROM validation_issues_v2 WHERE issue_type = 'COMPONENT_EXTRACTION_MISMATCH' AND COALESCE(status, 'OPEN') != 'RESOLVED'
        UNION ALL SELECT 'dictionary mappings', COUNT(*) FROM dictionary_mappings_v3 WHERE category = 'designation'
        UNION ALL SELECT 'dictionary applied rows', COUNT(*) FROM employees_v2 WHERE COALESCE(designation_source,'') LIKE 'dictionary_%'
    """)
    status_counts = _df(db, f"""
        SELECT status AS metric, COUNT(*) AS value
        FROM (SELECT {status_case} AS status FROM employees_v2 e)
        GROUP BY status
        ORDER BY status
    """)
    qa_summary = pd.concat([qa_summary, status_counts], ignore_index=True)

    base_select = f"""
        SELECT e.source_filename AS source_file, e.group_name, e.bill_id, e.page_no, e.sr_no,
               e.raw_name AS employee_name, e.employee_id, e.pan_raw AS pan, e.pay_scale_raw AS pay_scale,
               COALESCE(e.raw_designation_full, e.raw_designation, '') AS raw_designation_full,
               COALESCE(e.raw_designation_en, '') AS raw_designation_en,
               COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation, '') AS final_designation,
               e.designation_source, {status_case} AS employee_data_status,
               e.gross_amount AS gross, e.sum_deduction AS total_deduction, e.net_total AS net_total
        FROM employees_v2 e
    """
    needs_review = _df(db, base_select + f" WHERE ({status_case}) IN ('NEEDS_REVIEW','MISSING_CRITICAL_FIELD','COMPONENT_REVIEW_REQUIRED') ORDER BY e.source_filename, e.sr_no")
    missing_critical = _df(db, base_select + f" WHERE ({status_case}) = 'MISSING_CRITICAL_FIELD' ORDER BY e.source_filename, e.sr_no")
    manual_edits = _df(db, base_select + f" WHERE ({status_case}) = 'MANUALLY_EDITED' ORDER BY e.source_filename, e.sr_no")
    dictionary_applied = _df(db, base_select + " WHERE COALESCE(e.designation_source,'') LIKE 'dictionary_%' ORDER BY e.source_filename, e.sr_no")
    component_review = _df(db, base_select + f" WHERE ({status_case}) = 'COMPONENT_REVIEW_REQUIRED' ORDER BY e.source_filename, e.sr_no")
    missing_pay_scale = _df(db, base_select + " WHERE COALESCE(e.pay_scale_raw,'') = '' ORDER BY e.source_filename, e.sr_no")
    missing_employee_id = _df(db, base_select + " WHERE COALESCE(e.employee_id,'') = '' ORDER BY e.source_filename, e.sr_no")
    suspicious_pan = _df(db, base_select + """
        WHERE COALESCE(e.pan_raw,'') != '' AND e.pan_raw NOT GLOB '[A-Z][A-Z][A-Z][A-Z][A-Z][0-9][0-9][0-9][0-9][A-Z]'
        ORDER BY e.source_filename, e.sr_no
    """)
    final_designation_summary = _df(db, """
        SELECT COALESCE(final_designation, final_designation_en, normalized_designation, 'Needs Review') AS final_designation,
               COUNT(*) AS employee_count,
               COALESCE(SUM(gross_amount),0) AS gross,
               COALESCE(SUM(sum_deduction),0) AS total_deduction,
               COALESCE(SUM(net_total),0) AS net_total
        FROM employees_v2
        GROUP BY final_designation
        ORDER BY employee_count DESC, final_designation
    """)
    missing_counts = _df(db, """
        SELECT 'employee_name' AS field, SUM(CASE WHEN COALESCE(raw_name,'') = '' THEN 1 ELSE 0 END) AS missing_count FROM employees_v2
        UNION ALL SELECT 'employee_id', SUM(CASE WHEN COALESCE(employee_id,'') = '' THEN 1 ELSE 0 END) FROM employees_v2
        UNION ALL SELECT 'pay_scale', SUM(CASE WHEN COALESCE(pay_scale_raw,'') = '' THEN 1 ELSE 0 END) FROM employees_v2
        UNION ALL SELECT 'pan', SUM(CASE WHEN COALESCE(pan_raw,'') = '' THEN 1 ELSE 0 END) FROM employees_v2
    """)
    qa_summary = pd.concat([qa_summary, missing_counts.rename(columns={"field": "metric", "missing_count": "value"})], ignore_index=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(export_dir, f"employee_master_qa_{timestamp}.xlsx")
    sheets = {
        "QA Summary": qa_summary,
        "Needs Review Rows": needs_review,
        "Missing Critical Fields": missing_critical,
        "Manual Edits": manual_edits,
        "Dictionary Applied Rows": dictionary_applied,
        "Component Review Rows": component_review,
        "Final Designation Summary": final_designation_summary,
        "Missing Pay Scale": missing_pay_scale,
        "Missing Employee ID": missing_employee_id,
        "Suspicious PAN": suspicious_pan,
    }
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=name[:31])
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = "A2"
            if ws.max_row and ws.max_column:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
                max_len = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 36)
    db.close()
    return path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--db", default=app.config["DATABASE"])
    parser.add_argument("--export-dir", default=EXPORT_DIR)
    args = parser.parse_args()
    path = build_workbook(args.db, args.export_dir)
    print(path)


if __name__ == "__main__":
    main()

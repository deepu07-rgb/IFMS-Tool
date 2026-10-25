import argparse
import os
import sqlite3
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

from app import EMPLOYEE_COMPONENT_FIELDS, app, init_db, process_file_v2
from ifms_engine.db import init_v2_schema


COMPONENT_LABELS = {
    "basic_pay": "Basic",
    "da": "DA",
    "hra": "HRA",
    "cca": "CCA",
    "hda": "HDA",
    "mess_allowance": "Mess Allowance",
    "fta": "FTA",
    "rop": "ROP",
    "lic": "LIC",
    "gpf": "GPF",
    "gpf_2004": "GPF 2004",
    "rghs": "RGHS",
    "sip": "SIP",
    "sil": "SIL",
    "sii": "SII",
    "income_tax": "Income Tax",
    "roadways_deduction": "Roadways Deduction",
    "advance_salary": "Advance Salary",
    "ppss_jdr": "PPSS JDR",
}


def _pdfs_from_args(pdf_arg):
    if pdf_arg:
        pdf = Path(pdf_arg)
        if not pdf.exists():
            raise FileNotFoundError(pdf)
        return [pdf]
    fixture_dir = ROOT / "test_fixtures" / "ifms_inner_sheets"
    return sorted(fixture_dir.glob("*.pdf"), key=lambda p: int(p.stem) if p.stem.isdigit() else 9999)


def _component_map(db):
    rows = db.execute("SELECT employee_id, normalized_head, raw_head, amount, component_type FROM employee_components_v2").fetchall()
    by_employee = {}
    audit = []
    for row in rows:
        by_employee.setdefault(row["employee_id"], {})
        by_employee[row["employee_id"]][row["normalized_head"]] = by_employee[row["employee_id"]].get(row["normalized_head"], 0) + (row["amount"] or 0)
        audit.append(dict(row))
    return by_employee, audit


def _employee_issue_map(db):
    rows = db.execute(
        """
        SELECT employee_id, issue_type, severity, message
        FROM validation_issues_v2
        WHERE COALESCE(status, 'OPEN') != 'RESOLVED'
        ORDER BY employee_id, severity DESC, issue_type
        """
    ).fetchall()
    issue_map = {}
    for row in rows:
        if row["employee_id"] is None:
            continue
        issue_map.setdefault(row["employee_id"], []).append(f"{row['issue_type']}:{row['severity']}")
    return issue_map


def build_audit_workbook(pdfs, output_dir):
    temp_dir = tempfile.TemporaryDirectory()
    db_path = Path(temp_dir.name) / "employee_master_audit.db"
    app.config["DATABASE"] = str(db_path)
    app.config["USE_V2_PARSER"] = True
    init_db()
    init_v2_schema(app.config["DATABASE"])

    for pdf in pdfs:
        process_file_v2(str(pdf), pdf.name, dict(app.config))

    db = sqlite3.connect(app.config["DATABASE"])
    db.row_factory = sqlite3.Row
    comp_by_employee, component_audit = _component_map(db)
    issue_map = _employee_issue_map(db)
    employees = db.execute(
        """
        SELECT e.*, b.bill_no
        FROM employees_v2 e
        LEFT JOIN bills_v2 b ON b.id = e.bill_id
        ORDER BY e.source_filename, e.sr_no
        """
    ).fetchall()

    master_rows = []
    missing_rows = []
    mismatch_rows = []
    confidence_rows = []
    for emp in employees:
        comps = comp_by_employee.get(emp["id"], {})
        issues = issue_map.get(emp["id"], [])
        status = emp["row_verification_status"] or ("NEEDS_REVIEW" if issues else "VERIFIED")
        row = {
            "source_file": emp["source_filename"],
            "group_name": emp["group_name"],
            "page_no": emp["page_no"],
            "sr_no": emp["sr_no"],
            "employee_name": emp["raw_name"],
            "employee_id": emp["employee_id"],
            "pan": emp["pan_raw"],
            "pay_scale": emp["pay_scale_raw"],
            "gpf_pran": emp["gpf_pran_raw"],
            "belt_no": emp["belt_no"],
            "raw_designation": emp["raw_designation"],
            "normalized_designation": emp["normalized_designation"],
            "printed_gross": emp["gross_amount"],
            "printed_deduction": emp["sum_deduction"],
            "printed_net": emp["net_total"],
            "employee_data_status": status,
            "issues": "; ".join(issues),
        }
        for head, label in COMPONENT_LABELS.items():
            row[label] = comps.get(head, 0)
        master_rows.append(row)

        for field in ["raw_name", "employee_id", "pay_scale_raw", "gross_amount", "sum_deduction", "net_total"]:
            if not emp[field]:
                missing_rows.append({
                    "source_file": emp["source_filename"],
                    "group_name": emp["group_name"],
                    "sr_no": emp["sr_no"],
                    "employee_name": emp["raw_name"],
                    "missing_field": field,
                    "status": "NEEDS_REVIEW",
                })

        allowance_total = sum(v for h, v in comps.items() if h in {"basic_pay", "da", "hra", "cca", "hda", "mess_allowance", "fta", "rop", "car_rent", "hfrh", "eacpenf", "giais", "kma_allowance", "mta", "special_pay", "wash_allowance"})
        deduction_total = sum(v for h, v in comps.items() if h in {"lic", "gpf", "gpf_2004", "rghs", "sip", "sil", "sii", "income_tax"})
        if round(allowance_total) != round(emp["gross_amount"] or 0) or round(deduction_total) != round(emp["sum_deduction"] or 0):
            mismatch_rows.append({
                "source_file": emp["source_filename"],
                "group_name": emp["group_name"],
                "page_no": emp["page_no"],
                "sr_no": emp["sr_no"],
                "employee_name": emp["raw_name"],
                "component_gross": allowance_total,
                "printed_gross": emp["gross_amount"],
                "component_deduction": deduction_total,
                "printed_deduction": emp["sum_deduction"],
                "status": "NEEDS_REVIEW",
            })

        confidence_rows.append({
            "source_file": emp["source_filename"],
            "sr_no": emp["sr_no"],
            "employee_name": emp["raw_name"],
            "identity_confidence": "HIGH" if emp["employee_id"] and emp["pay_scale_raw"] and emp["pan_raw"] else "REVIEW",
            "totals_confidence": "HIGH" if round((emp["gross_amount"] or 0) - (emp["sum_deduction"] or 0)) == round(emp["net_total"] or 0) else "REVIEW",
            "component_confidence": "HIGH" if not issues else "REVIEW",
        })

    component_rows = []
    for row in component_audit:
        emp = db.execute("SELECT source_filename, group_name, page_no, sr_no, raw_name, employee_id FROM employees_v2 WHERE id = ?", (row["employee_id"],)).fetchone()
        if not emp:
            continue
        component_rows.append({
            "source_file": emp["source_filename"],
            "group_name": emp["group_name"],
            "page_no": emp["page_no"],
            "sr_no": emp["sr_no"],
            "employee_name": emp["raw_name"],
            "employee_id": emp["employee_id"],
            "component_type": row["component_type"],
            "raw_head": row["raw_head"],
            "normalized_head": row["normalized_head"],
            "amount": row["amount"],
        })
    db.close()
    temp_dir.cleanup()

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"employee_master_audit_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    sheets = {
        "Employee Master Audit": pd.DataFrame(master_rows),
        "Component Audit": pd.DataFrame(component_rows),
        "Missing Field Report": pd.DataFrame(missing_rows),
        "Row-Level Mismatches": pd.DataFrame(mismatch_rows),
        "Field Extraction Confidence": pd.DataFrame(confidence_rows),
    }
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            if df.empty:
                df = pd.DataFrame(columns=["status"])
            df.to_excel(writer, index=False, sheet_name=name[:31])
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                width = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 32)
    return output_path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pdf", help="Process one PDF. If omitted, all sample PDFs are audited.")
    parser.add_argument("--output-dir", default=str(ROOT / "exports"))
    args = parser.parse_args()
    output_path = build_audit_workbook(_pdfs_from_args(args.pdf), Path(args.output_dir))
    print(f"Audit workbook: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

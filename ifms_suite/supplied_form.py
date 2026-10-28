"""Clean vector reconstruction of the two supplied Jodhpur due/drawn grids.

No scanned image is embedded. Unknown layouts or values that cannot fit the
printed columns raise ValueError; callers must not fall back to a generic bill.
"""

from decimal import Decimal, InvalidOperation
from io import BytesIO
from datetime import date
from calendar import monthrange

OFFICE = "OFFICE OF THE POLICE COMMISSIONER JODHPUR"
HEADS = ("BASIC", "DA", "HRA", "HDA", "CCA")
SALARY_LABELS = ("BASIC\nPAY", "D.A.", "H.R.A.", "H.D.A.", "CCA", "TOTAL")
TEMPLATES = {
    "jodhpur_classic": {
        "deduction_keys": ("GPF", "SI", "CM_CORONA", "I_TAX"),
        "deduction_labels": ("G.P.F.", "SI", "CM\nCORONA", "I TAX"),
    },
    "jodhpur_gpf_arrear": {
        "deduction_keys": (
            "GPF_TO_DEDUCT",
            "GPF_ALREADY_DEDUCTED",
            "SIF_ALREADY_DEDUCTED",
            "I_TAX",
        ),
        "deduction_labels": (
            "GPF ARREAR\nTO BE\nDEDUCTED",
            "GPF ARREAR\nALREADY\nDEDUCTED",
            "SIF ARREAR\nALREADY\nDEDUCTED",
            "I TAX",
        ),
    },
}


def _money(value):
    try:
        amount = Decimal(str(value))
        if not amount.is_finite() or amount != amount.quantize(Decimal(".01")):
            raise ValueError("Unresolved or invalid amount")
        return amount
    except (InvalidOperation, TypeError):
        raise ValueError("Unresolved or invalid amount") from None


def _prepared(bill, metadata):
    template_id = metadata.get("supplied_template_id")
    if template_id not in TEMPLATES:
        raise ValueError("Choose a supported supplied template; no generic fallback")
    if bill.get("status") != "ready" or bill.get("blockers"):
        raise ValueError("Supplied form requires a ready bill")
    template = TEMPLATES[template_id]
    result = []
    seen = set()
    prior_total = Decimal(0)
    for row in bill.get("rows", []):
        month = row["month"]
        if month in seen:
            raise ValueError("Duplicate supplied form month")
        seen.add(month)
        heads = {}
        for component in row.get("components", []):
            head = str(component["head"]).upper()
            if head in heads:
                raise ValueError("Duplicate supplied form head: " + head)
            prior = _money(component.get("prior_settlement", "0"))
            if prior < 0:
                raise ValueError("Prior settlement amount must be nonnegative")
            prior_total += prior
            amounts = [
                _money(component.get(key)) for key in ("due", "drawn", "difference")
            ]
            if amounts[0] - amounts[1] - prior != amounts[2]:
                raise ValueError("Supplied form amount difference does not reconcile")
            if head not in HEADS:
                if amounts[2] != 0 or prior != 0:
                    raise ValueError("Unsupported changed supplied form head: " + head)
                continue
            amounts[1] += prior
            heads[head] = amounts
        if not heads:
            raise ValueError("Missing supplied form head amounts")
        detail = {}
        for key, aliases in {
            "bill_no": ("bill_no",),
            "bill_date": ("bill_date",),
            "tv_no": ("tv_no", "treasury_no"),
            "tv_date": ("tv_date", "treasury_date"),
        }.items():
            values = {
                str(item[alias])
                for item in [row] + row.get("evidence", [])
                if isinstance(item, dict)
                for alias in aliases
                if item.get(alias)
            }
            if len(values) == 1:
                detail[key] = values.pop()
        detail.update(metadata.get("supplied_form_rows", {}).get(month, {}))
        deductions = detail.get("deductions", {})
        if any(key not in template["deduction_keys"] for key in deductions):
            raise ValueError("Unsupported supplied form deduction column")
        deducted = _money(detail.get("total_deductions", "0"))
        deduction_values = [
            _money(deductions.get(key, "0")) for key in template["deduction_keys"]
        ]
        if any(v < 0 for v in deduction_values) or deducted < 0:
            raise ValueError("Invalid supplied form deduction amount")
        # Newer GPF columns record current/prior credits, not additive deductions.
        if template_id == "jodhpur_classic" and sum(deduction_values) != deducted:
            raise ValueError("Monthly deduction columns do not reconcile")
        salary = []
        for side in range(3):
            values = [heads.get(head, (Decimal(0),) * 3)[side] for head in HEADS]
            salary.extend(values + [sum(values)])
        month_label = month
        if row.get("start") or row.get("end"):
            start, end = date.fromisoformat(row["start"]), date.fromisoformat(
                row["end"]
            )
            if start > end or str(start)[:7] != month or str(end)[:7] != month:
                raise ValueError("Supplied form dates must cover one stated month")
            if start.day != 1 or end.day != monthrange(end.year, end.month)[1]:
                month_label = (
                    start.strftime("%d/%m/%Y") + "\n" + end.strftime("%d/%m/%Y")
                )
        values = [
            month_label,
            "\n".join(
                filter(
                    None,
                    [str(detail.get("bill_no", "")), str(detail.get("bill_date", ""))],
                )
            ),
            "\n".join(
                filter(
                    None, [str(detail.get("tv_no", "")), str(detail.get("tv_date", ""))]
                )
            ),
        ]
        values += salary + deduction_values + [deducted, salary[-1] - deducted]
        result.append(values)
    if not result:
        raise ValueError("Supplied form needs monthly rows")
    if set(metadata.get("supplied_form_rows", {})) - seen:
        raise ValueError("Supplied form deductions contain unmatched months")
    totals = bill.get("totals", {})
    if "prior_settlement" in totals and prior_total != _money(
        totals["prior_settlement"]
    ):
        raise ValueError("Supplied form prior settlement totals do not reconcile")
    aggregate = None
    monthly_deducted = sum(row[25] for row in result)
    total_deducted = _money(totals.get("total_deductions"))
    if monthly_deducted != total_deducted:
        if monthly_deducted or any(
            "deductions" in detail or "total_deductions" in detail
            for detail in metadata.get("supplied_form_rows", {}).values()
        ):
            raise ValueError("Monthly deduction mapping does not reconcile")
        allocated = {key: Decimal(0) for key in template["deduction_keys"]}
        aliases = {
            "gpf": "GPF",
            "gpf adjustment": "GPF",
            "si": "SI",
            "si adjustment": "SI",
            "i tax": "I_TAX",
            "i tax adjustment": "I_TAX",
            "income tax": "I_TAX",
            "income tax adjustment": "I_TAX",
            "cm corona": "CM_CORONA",
        }
        for deduction in bill.get("deductions", []):
            name = " ".join(
                str(deduction.get("name", "")).lower().replace(".", "").split()
            )
            key = aliases.get(name)
            if template_id == "jodhpur_gpf_arrear" and key == "GPF":
                key = "GPF_TO_DEDUCT"
            if key not in allocated:
                raise ValueError(
                    "Unsupported engine deduction for supplied form: " + name
                )
            amount = _money(deduction.get("amount"))
            if amount < 0:
                raise ValueError("Invalid engine deduction amount")
            allocated[key] += amount
        if sum(allocated.values()) != total_deducted:
            raise ValueError("Explicit deduction mapping required")
        aggregate = [allocated[key] for key in template["deduction_keys"]] + [
            total_deducted,
            _money(totals.get("net_payable")),
        ]
        # An aggregate adjustment has no evidenced monthly allocation.
        for values in result:
            values[21:] = [None] * 6
    if sum(row[20] for row in result) != _money(totals.get("difference")):
        raise ValueError("Supplied form totals do not reconcile: difference")
    if _money(totals.get("difference")) - total_deducted != _money(
        totals.get("net_payable")
    ):
        raise ValueError("Supplied form totals do not reconcile: net_payable")
    return template, result, aggregate


def _grid(template, rows, total_values=None):
    """Shared cell coordinates, one based, for spreadsheet and vector PDF."""
    cells = [
        (1, 1, 1, 27, OFFICE),
        (2, 1, 2, 15, ""),
        (2, 16, 2, 27, ""),
        (4, 1, 5, 1, "MONTH\nAND\nYEAR"),
        (4, 2, 5, 2, "BILL NO AND DATE"),
        (4, 3, 5, 3, "T.V. NO\nDATE"),
    ]
    for col, title in (
        (4, "PAY TO BE DRAWN"),
        (10, "PAY ALREADY DRAWN"),
        (16, "PAY DIFFERENCE DRAWN"),
    ):
        cells.append((4, col, 4, col + 5, title))
        cells.extend(
            (5, col + offset, 5, col + offset, label)
            for offset, label in enumerate(SALARY_LABELS)
        )
    for col, label in enumerate(
        template["deduction_labels"] + ("TOTAL\nDEDUCTION", "NET\nAMOUNT"), 22
    ):
        cells.append((4, col, 5, col, label))
    total = total_values if total_values is not None else _total(rows)
    for r, values in enumerate(rows + [total], 6):
        cells.extend((r, c, r, c, value) for c, value in enumerate(values, 1))
    return cells


def _total(rows, aggregate=None):
    total = [
        (
            sum(row[c] for row in rows if row[c] is not None)
            if c >= 3
            else ("TOTAL" if c == 1 else "")
        )
        for c in range(27)
    ]
    if aggregate is not None:
        total[21:] = aggregate
    elif all(row[26] is None for row in rows):
        total[21:] = [None] * 6
    return total


def _credit_disclosure(bill):
    if not any(
        _money(component.get("prior_settlement", "0"))
        for row in bill.get("rows", [])
        for component in row.get("components", [])
    ):
        return []
    notes = ["Pay already drawn includes verified prior arrear settlements"]
    references = []
    for item in bill.get("evidence", []):
        if not isinstance(item, dict) or item.get("kind") != "prior_credit":
            continue
        parts = [
            str(item[key])
            for key in ("bill_no", "pay_month", "reference", "document_id")
            if item.get(key)
        ]
        if parts:
            reference = " / ".join(parts)
            if reference not in references:
                references.append(reference)
    if references:
        notes.append("Prior settlement sources: " + "; ".join(references))
    return notes


def export_supplied_excel(bill, metadata):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    template, rows, aggregate = _prepared(bill, metadata)
    book = Workbook()
    sheet = book.active
    sheet.title = "Due drawn statement"
    border = Border(*([Side(style="thin", color="777777")] * 4))
    bottom = max(19, len(rows) + 14)
    for row in sheet.iter_rows(min_row=4, max_row=bottom, max_col=27):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.font = Font(name="Arial", size=8)
    for r, c, er, ec, value in _grid(template, rows, _total(rows, aggregate)):
        if (r, c) != (er, ec):
            sheet.merge_cells(start_row=r, start_column=c, end_row=er, end_column=ec)
        cell = sheet.cell(r, c, float(value) if isinstance(value, Decimal) else value)
        if isinstance(cell.value, str):
            cell.data_type = "s"
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.font = Font(
            name="Arial", size=14 if r == 1 else (7 if r < 6 else 9), bold=r == 1
        )
        if isinstance(value, Decimal):
            cell.number_format = '0.##;-0.##;"-"'
    for position, value in [
        ("A2", "NAME = " + str(metadata.get("employee_name", ""))),
        ("P2", "EMP. ID = " + str(metadata.get("employee_code", ""))),
    ]:
        sheet[position] = value
        sheet[position].data_type = "s"
    # Blank signatures, never copies of source signatures or rubber stamps.
    for start, end, label in (
        (9, 14, "Assistant Accounts Officer-II"),
        (19, 26, "Forwarding Officer"),
        (15, 23, "Drawing and Disbursing Officer"),
    ):
        r = bottom - (3 if start != 15 else 0)
        sheet.merge_cells(start_row=r, start_column=start, end_row=r, end_column=end)
        sheet.cell(r, start, label)
        sheet.cell(r, start).alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.cell(r, start).font = Font(size=8)
    for col in range(1, 28):
        sheet.column_dimensions[get_column_letter(col)].width = (
            7 if col not in (2, 3) else 12
        )
    for r in range(1, bottom + 1):
        sheet.row_dimensions[r].height = 24
    sheet.row_dimensions[4].height = 20
    sheet.row_dimensions[5].height = 42
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.25, bottom=0.25, header=0, footer=0
    )
    notes = _credit_disclosure(bill)
    for offset, text in enumerate(notes, 2):
        r = bottom + offset
        sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=27)
        sheet.cell(r, 1, text)
        sheet.cell(r, 1).data_type = "s"
        sheet.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(r, 1).font = Font(size=8)
        sheet.row_dimensions[r].height = 30
    sheet.print_area = f"A1:AA{bottom + len(notes) + 1 if notes else bottom}"
    sheet.oddFooter.center.text = (
        "DRAFT - SOURCE REVIEW REQUIRED"
        if metadata.get("status") != "finalized"
        else "FINALIZED AFTER SOURCE REVIEW"
    )
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def export_supplied_pdf(bill, metadata):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A3, landscape

    template, rows, aggregate = _prepared(bill, metadata)
    stream = BytesIO()
    width, height = landscape(A3)
    pdf = canvas.Canvas(stream, pagesize=(width, height))
    pdf.setTitle("Due drawn statement")
    # Keep the source's dense, wide grid; paginate monthly rows in the same form.
    for offset in range(0, len(rows), 8):
        page_rows = rows[offset : offset + 8]
        margin, top = 28, height - 38
        widths = [42, 78, 55] + [37] * 18 + [43] * 4 + [56, 64]
        scale = (width - margin * 2) / sum(widths)
        xs = [margin]
        for value in widths:
            xs.append(xs[-1] + value * scale)
        heights = [27, 25, 10, 17, 49] + [28] * 19
        ys = [top]
        for value in heights:
            ys.append(ys[-1] - value)

        def put(r, c, er, ec, value, grid=False):
            left, right, upper, lower = xs[c - 1], xs[ec], ys[r - 1], ys[er]
            if grid:
                pdf.setLineWidth(0.35)
                pdf.rect(left, lower, right - left, upper - lower)
            text = (
                ("" if value is None else str(value))
                if not isinstance(value, Decimal)
                else (
                    format(value, "f").rstrip("0").rstrip(".")
                    if value % 1
                    else str(int(value))
                )
            )
            if isinstance(value, Decimal) and not value:
                text = "-"
            try:
                text.encode("cp1252")
            except UnicodeEncodeError:
                raise ValueError(
                    "Supplied PDF font does not support these characters; use an editable Excel form or a reviewed transliteration"
                ) from None
            lines = text.split("\n")
            size = 15 if r == 1 else (11 if r == 2 else (6 if r < 6 else 8))
            size = min(
                size,
                min(
                    (
                        (right - left - 4)
                        / max(pdf.stringWidth(line, "Helvetica", 1), 1)
                        for line in lines
                    ),
                    default=size,
                ),
            )
            pdf.setFont("Helvetica", size)
            baseline = (
                (upper + lower) / 2 + (len(lines) - 1) * size * 0.58 - size * 0.35
            )
            for line in lines:
                pdf.drawCentredString((left + right) / 2, baseline, line)
                baseline -= size * 1.16

        total = (
            _total(rows, aggregate) if offset + 8 >= len(rows) else _total(page_rows)
        )
        if offset + 8 < len(rows):
            total[1] = "SUBTOTAL"
        for cell in _grid(template, page_rows, total):
            put(*cell, grid=cell[0] >= 4)
        put(2, 1, 2, 15, "NAME = " + str(metadata.get("employee_name", "")))
        put(2, 16, 2, 27, "EMP. ID = " + str(metadata.get("employee_code", "")))
        # Extend the blank printed grid below totals, as on the supplied forms.
        for r in range(len(page_rows) + 7, 22):
            for c in range(1, 28):
                put(r, c, r, c, "", grid=True)
        put(17, 9, 17, 14, "Assistant Accounts Officer-II")
        put(17, 19, 17, 26, "Forwarding Officer")
        put(20, 15, 20, 24, "Drawing and Disbursing Officer")
        from reportlab.lib.utils import simpleSplit

        note_lines = []
        for note in _credit_disclosure(bill):
            try:
                note.encode("cp1252")
            except UnicodeEncodeError:
                raise ValueError(
                    "Supplied PDF font does not support prior settlement reference characters"
                ) from None
            note_lines.extend(simpleSplit(note, "Helvetica", 7, width - margin * 2))
        if len(note_lines) > 10:
            raise ValueError(
                "Prior settlement source references do not fit supplied form footer"
            )
        pdf.setFont("Helvetica", 7)
        for line_number, line in enumerate(note_lines):
            pdf.drawString(margin, 43 + (len(note_lines) - line_number - 1) * 9, line)
        pdf.drawString(
            margin,
            24,
            (
                "DRAFT - SOURCE REVIEW REQUIRED"
                if metadata.get("status") != "finalized"
                else "FINALIZED AFTER SOURCE REVIEW"
            ),
        )
        pdf.showPage()
    pdf.save()
    stream.seek(0)
    return stream

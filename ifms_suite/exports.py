"""Both formats read the same saved calculation; neither recomputes amounts."""

from io import BytesIO
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import A4, landscape

COLUMNS = [
    "start",
    "end",
    "days",
    "denominator",
    "basic",
    "due_rate",
    "drawn_rate",
    "due",
    "drawn",
    "prior_settlement",
    "difference",
]


def excel(payload, calculation):
    w = Workbook()
    s = w.active
    s.title = "Due Drawn"
    s.append(
        [
            "Employee",
            payload.get("employee_code", ""),
            "Allowance",
            payload.get("allowance", ""),
        ]
    )
    s["B1"].data_type = "s"
    s["D1"].data_type = "s"
    s.append(COLUMNS)
    for row in calculation["rows"]:
        s.append(
            [
                (
                    float(Decimal(row[k]))
                    if k
                    in (
                        "basic",
                        "due_rate",
                        "drawn_rate",
                        "due",
                        "drawn",
                        "prior_settlement",
                        "difference",
                    )
                    else row[k]
                )
                for k in COLUMNS
            ]
        )
    s.append(["TOTAL"] + [""] * 9 + [float(Decimal(calculation["total"]))])
    s.freeze_panes = "C3"
    s.auto_filter.ref = f"A2:K{s.max_row-1}"
    for c in s[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="174D52")
    for column in s.columns:
        s.column_dimensions[column[0].column_letter].width = 18
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.orientation = "landscape"
    s.page_setup.paperSize = s.PAPERSIZE_A4
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.print_title_rows = "1:2"
    audit = w.create_sheet("Calculation basis")
    audit.append(["Property", "Value"])
    for k, v in calculation.items():
        if k != "rows":
            audit.append([k, str(v)])
    audit.append(
        [
            "Rate reference",
            payload.get("rule_reference", "User supplied — verify before official use"),
        ]
    )
    audit.append(["Drawn reference", payload.get("drawn_reference", "Unreviewed")])
    audit.append(["Status", payload.get("status", "draft").upper()])
    audit.append(["Entitlement", payload.get("entitlement_type", "salary")])
    for row in audit:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    profile = w.create_sheet("Employee details")
    for key in (
        "employee_code",
        "employee_name",
        "designation",
        "office_name",
        "ddo_code",
        "group",
        "entitlement_type",
    ):
        profile.append(
            [key, payload.get(key, "salary" if key == "entitlement_type" else "")]
        )
        profile.cell(profile.max_row, 2).data_type = "s"
    deductions = w.create_sheet("Deductions")
    deductions.append(["Head", "Amount"])
    for d in calculation.get("deductions", []):
        deductions.append([d["name"], float(Decimal(d["amount"]))])
        deductions.cell(deductions.max_row, 1).data_type = "s"
    deductions.append(
        ["Total deductions", float(Decimal(calculation.get("total_deductions", "0")))]
    )
    deductions.append(
        [
            "Cash payable",
            float(Decimal(calculation.get("net_payable", calculation["total"]))),
        ]
    )
    sources = w.create_sheet("Sources")
    sources.append(["From", "Through", "Drawn source", "Rate / rule reference"])
    for row in calculation["rows"]:
        sources.append(
            [
                row["start"],
                row["end"],
                row.get("source", ""),
                row.get("rule_reference", ""),
            ]
        )
    for row in sources:
        for cell in row:
            cell.data_type = "s"
    b = BytesIO()
    w.save(b)
    b.seek(0)
    return b


def pdf(payload, calculation):
    b = BytesIO()
    c = Canvas(b, pagesize=landscape(A4))
    width, height = landscape(A4)

    def header():
        c.setFont("Helvetica-Bold", 13)
        c.drawString(
            25,
            height - 30,
            "IFMS Due–Drawn Statement — " + payload.get("status", "draft").upper(),
        )
        c.setFont("Helvetica", 9)
        c.drawString(
            25,
            height - 48,
            f"Employee: {payload.get('employee_name','')} ({payload.get('employee_code','')})    {payload.get('allowance','')} / {payload.get('entitlement_type','salary').replace('_',' ')}    {payload.get('start','')} to {payload.get('end','')}",
        )
        c.setFont("Helvetica", 7)
        for i, k in enumerate(COLUMNS):
            c.drawString(25 + i * 72, height - 70, k.replace("_", " "))

    header()
    y = height - 88
    for row in calculation["rows"]:
        if y < 65:
            c.showPage()
            header()
            y = height - 88
        for i, k in enumerate(COLUMNS):
            c.drawString(25 + i * 72, y, str(row[k]))
        y -= 16
    c.setFont("Helvetica-Bold", 11)
    c.drawString(25, y - 15, "Remaining gross difference: INR " + calculation["total"])
    y -= 40
    for d in calculation.get("deductions", []):
        if y < 85:
            c.showPage()
            header()
            y = height - 88
        c.setFont("Helvetica", 9)
        c.drawString(25, y, str(d["name"]) + ": INR " + d["amount"])
        y -= 15
    if y < 85:
        c.showPage()
        header()
        y = height - 88
    c.setFont("Helvetica-Bold", 11)
    c.drawString(
        25,
        y,
        "Cash payable: INR " + calculation.get("net_payable", calculation["total"]),
    )
    c.setFont("Helvetica", 8)
    c.drawString(
        25,
        40,
        "Rate reference: " + str(payload.get("rule_reference", "Unreviewed"))[:120],
    )
    c.drawString(
        25,
        25,
        "Drafts require review. Check deductions and cash allocation against applicable orders. No digital signature applied.",
    )
    c.save()
    b.seek(0)
    return b

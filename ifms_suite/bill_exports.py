"""Consolidated, auditable draft exports for prepared multi-month bills."""

from decimal import Decimal
from datetime import date
from calendar import monthrange
from html import escape
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


_DRAFT = "DRAFT - HUMAN REVIEW REQUIRED"
_REVIEW = (
    "Human review of source evidence does not certify this bill or authorize payment."
)
_HEADERS = [
    "Month",
    "Days",
    "Month days",
    "Head",
    "Due",
    "Drawn",
    "Prior settlement",
    "Difference",
]


def _title(metadata):
    return (
        "FINALIZED AFTER SOURCE REVIEW"
        if metadata.get("status") == "finalized"
        else _DRAFT
    )


def _text(value):
    if value is None:
        return "Unresolved"
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=True, default=str)
    return str(value)


def _label(key):
    return str(key).replace("_", " ").capitalize()


def _number(value):
    return None if value is None else float(Decimal(str(value)))


def _rows(bill):
    for row in bill.get("rows", []):
        days = row.get("days", row.get("covered_days"))
        month_days = row.get("month_days", row.get("days_in_month"))
        if row.get("start") and row.get("end"):
            start, end = date.fromisoformat(row["start"]), date.fromisoformat(
                row["end"]
            )
            if days is None:
                days = (end - start).days + 1
            if month_days is None:
                month_days = monthrange(start.year, start.month)[1]
        for component in row.get("components", []):
            yield row, component, [
                row.get("month", ""),
                days,
                month_days,
                component.get("head", ""),
                *[
                    component.get(key)
                    for key in ("due", "drawn", "prior_settlement", "difference")
                ],
            ]


def _audit(bill):
    """Retain detailed engine evidence, including fields added to its contract."""
    for key, value in bill.items():
        if key not in {"rows", "totals"}:
            yield _label(key), value
    for row in bill.get("rows", []):
        month = str(row.get("month", ""))
        for key, value in row.items():
            if key != "components":
                yield month + " / " + _label(key), value
        for component in row.get("components", []):
            prefix = month + " / " + str(component.get("head", ""))
            for key, value in component.items():
                if key not in {
                    "head",
                    "due",
                    "drawn",
                    "prior_settlement",
                    "difference",
                }:
                    yield prefix + " / " + _label(key), value


def _append(sheet, values):
    sheet.append(values)
    # Explicit strings prevent formula execution without changing signed numbers.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def export_bill_excel(bill, metadata):
    """Return a draft workbook with monthly heads, totals, and source audit trail."""
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    _append(summary, [_title(metadata)])
    _append(summary, ["Field", "Value"])
    _append(summary, ["Status", bill.get("status", "unknown")])
    _append(summary, ["Mode", bill.get("mode", "unknown")])
    _append(summary, ["Human review", _REVIEW])
    for key, value in metadata.items():
        _append(summary, [_label(key), _text(value)])
    for key, value in bill.get("totals", {}).items():
        _append(summary, [_label(key), _number(value)])
    monthly = book.create_sheet("Monthly components")
    _append(monthly, [_title(metadata)])
    _append(monthly, _HEADERS)
    for _, _, values in _rows(bill):
        _append(monthly, values[:4] + [_number(v) for v in values[4:]])
    audit = book.create_sheet("Evidence and review")
    _append(audit, [_title(metadata)])
    _append(audit, ["Context", "Details"])
    for context, value in _audit(bill):
        # Excel cells have a 32,767-character limit; preserve long evidence in parts.
        content = _text(value)
        for start in range(0, len(content) or 1, 30000):
            _append(audit, [context, content[start : start + 30000]])
    for sheet in book:
        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = (
            f"A2:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
        )
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:2"
        for cell in sheet[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174D52")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = (
                22 if sheet == monthly else 45
            )
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00;[Red]-#,##0.00"
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def export_bill_pdf(bill, metadata):
    """Return a readable consolidated draft, with rate/source and review details."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    styles["Normal"].fontSize = 8
    styles["Normal"].leading = 11

    def paragraph(value):
        return Paragraph(escape(_text(value)).replace("\n", "<br/>"), styles["Normal"])

    story = [
        Paragraph(_title(metadata), styles["Heading1"]),
        paragraph(
            "Status: "
            + str(bill.get("status", "unknown"))
            + " | Mode: "
            + str(bill.get("mode", "unknown"))
        ),
        paragraph(_REVIEW),
        Spacer(1, 8),
    ]
    for key, value in metadata.items():
        story.append(paragraph(_label(key) + ": " + _text(value)))
    story.append(Spacer(1, 10))
    table_rows = [[paragraph(h) for h in _HEADERS]]
    for _, _, values in _rows(bill):
        rendered = [_text(v) for v in values[:4]]
        rendered.extend(
            "Unresolved" if v is None else f"{Decimal(str(v)):.2f}" for v in values[4:]
        )
        table_rows.append([paragraph(v) for v in rendered])
    table = Table(
        table_rows, repeatRows=1, colWidths=[72, 48, 60, 105, 100, 100, 100, 100]
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E4EFEF")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.extend(
        [table, Spacer(1, 10), Paragraph("Consolidated totals", styles["Heading2"])]
    )
    for key, value in bill.get("totals", {}).items():
        story.append(
            paragraph(
                _label(key)
                + ": "
                + ("Unresolved" if value is None else f"{Decimal(str(value)):.2f}")
            )
        )
    story.append(Paragraph("Source evidence, rates and review", styles["Heading2"]))
    for context, value in _audit(bill):
        # Keep each paragraph bounded so large source trails can span pages.
        content = _text(value)
        for start in range(0, len(content) or 1, 4000):
            story.append(paragraph(context + ": " + content[start : start + 4000]))
            story.append(Spacer(1, 3))
    document.build(story)
    output.seek(0)
    return output

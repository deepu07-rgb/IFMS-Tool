"""Filtered Outer/GA55 data view with matching PDF and Excel review reports."""

from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
from html import escape
import json
from flask import request, jsonify, send_file
from .storage import database
from .record_filters import filter_records


def _rows(db):
    records = []
    for source in db.execute("SELECT * FROM suite_imports ORDER BY id DESC"):
        payload = json.loads(source["payload"])
        kind = payload.get("kind")
        common = dict(
            import_id=source["id"],
            document=source["filename"],
            document_type=kind,
            status=source["status"],
        )
        if kind == "outer" and payload.get("extracted"):
            p = payload["extracted"]
            month = ""
            value = str(p.get("Bill Date Parsed") or p.get("Bill Date") or "")
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                try:
                    month = datetime.strptime(value[:10], fmt).strftime("%Y-%m")
                    break
                except ValueError:
                    pass
            records.append(
                dict(
                    common,
                    employee_code="",
                    pay_month=month,
                    bill_type=p.get("Bill Name", ""),
                    bill_no=str(p.get("Reference Number", "")),
                    allowances=p.get("Allowances", {}),
                    deductions=p.get("Deductions", {}),
                    gross=p.get("Gross Amount"),
                    net=p.get("Net Amount"),
                )
            )
        if kind == "ga55":
            approved = [
                json.loads(r["payload"])
                for r in db.execute(
                    "SELECT payload FROM suite_payments WHERE import_id=?",
                    (source["id"],),
                )
            ]
            candidates = payload.get("ga55", {}).get("payments", [])
            if not approved and payload.get("machine_facts"):
                if any(payload.get(k) for k in ("source_bill_reviews","bill_field_reviews")):
                    from .bill_routes import case_facts
                    candidates = case_facts(payload, source["sha256"]).get("bills", [])
                else:
                    candidates = payload["machine_facts"].get("bills", [])
            for p in approved or candidates:
                records.append(
                    dict(
                        common,
                        employee_code=p.get("employee_code") or "",
                        pay_month=p.get("pay_month") or "",
                        bill_type=p.get("bill_type") or "",
                        bill_no=p.get("bill_no") or "",
                        allowances=p.get("allowances", {}),
                        deductions=p.get("deductions", {}),
                        gross=p.get("printed_gross"),
                        net=p.get("printed_net"),
                    )
                )
    return records


def register_document_records(suite):
    @suite.get("/api/document_records")
    @suite.get("/api/document_records/<format>")
    def document_records(format=None):
        with database() as db:
            all_rows = _rows(db)
        rows = filter_records(all_rows, request.args)
        amount_fields = ["gross", "net"] + sorted(
            {
                prefix + ":" + h.upper()
                for r in all_rows
                for prefix, key in [
                    ("allowance", "allowances"),
                    ("deduction", "deductions"),
                ]
                for h in r[key]
            }
        )
        if format is None:
            return jsonify(
                records=rows,
                count=len(rows),
                amount_fields=amount_fields,
                notice="Unapproved OCR records require source review. Outer month is bill date; GA55 month is pay month.",
            )
        if format not in ("pdf", "excel"):
            return jsonify(error="Choose PDF or Excel"), 404
        if not rows:
            return jsonify(error="No rows match these filters"), 404
        data = BytesIO()
        if format == "excel":
            from openpyxl import Workbook
            from .document_exports import format_sheet

            w = Workbook()
            s = w.active
            s.title = "Filtered records"
            s.append(["FILTERED DATA REVIEW — verify original source"])
            allowances = sorted({h for r in rows for h in r["allowances"]})
            deductions = sorted({h for r in rows for h in r["deductions"]})
            s.append(
                [
                    "Document",
                    "Type",
                    "Employee",
                    "Month (GA55 pay / outer bill)",
                    "Bill type",
                    "Bill number",
                    "Gross",
                    "Net",
                    *allowances,
                    *["Deduction: " + h for h in deductions],
                ]
            )
            for r in rows:
                s.append(
                    [
                        r["document"],
                        r["document_type"],
                        r["employee_code"],
                        r["pay_month"],
                        r["bill_type"],
                        r["bill_no"],
                        _number(r["gross"]),
                        _number(r["net"]),
                        *[_number(r["allowances"].get(h)) for h in allowances],
                        *[_number(r["deductions"].get(h)) for h in deductions],
                    ]
                )
            format_sheet(s)
            s.auto_filter.ref = f"A2:{s.cell(s.max_row,s.max_column).coordinate}"
            w.save(data)
        else:
            from reportlab.platypus import (
                SimpleDocTemplate,
                Paragraph,
                Spacer,
                Table,
                TableStyle,
            )
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.pagesizes import A4, landscape

            styles = getSampleStyleSheet()
            story = [
                Paragraph(
                    "Filtered extracted data — source review report", styles["Heading1"]
                ),
                Paragraph(
                    f"{len(rows)} records. Original PDFs are available separately. Missing amounts remain unresolved.",
                    styles["Normal"],
                ),
            ]
            for r in rows:
                story.extend(
                    [
                        Spacer(1, 12),
                        Paragraph(
                            escape(
                                r["document"]
                                + " · "
                                + r["pay_month"]
                                + " · "
                                + r["employee_code"]
                            ),
                            styles["Heading2"],
                        ),
                        Paragraph(
                            escape(
                                "Bill "
                                + r["bill_no"]
                                + " · "
                                + r["bill_type"]
                                + " · Gross "
                                + str(
                                    r["gross"] if r["gross"] is not None else "not read"
                                )
                                + " · Net "
                                + str(r["net"] if r["net"] is not None else "not read")
                            ),
                            styles["Normal"],
                        ),
                    ]
                )
                cells = [["Type", "Component", "Amount"]] + [
                    [
                        label,
                        Paragraph(escape(str(h)), styles["Normal"]),
                        str(v) if v is not None else "not read",
                    ]
                    for key, label in [
                        ("allowances", "Allowance"),
                        ("deductions", "Deduction"),
                    ]
                    for h, v in r[key].items()
                ]
                table = Table(cells, colWidths=[100, 350, 100], repeatRows=1)
                table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.3, "#bbbbbb"),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                story.append(table)
            SimpleDocTemplate(
                data, pagesize=landscape(A4), leftMargin=28, rightMargin=28
            ).build(story)
        data.seek(0)
        return send_file(
            data,
            as_attachment=True,
            download_name="filtered-records." + ("pdf" if format == "pdf" else "xlsx"),
            mimetype=(
                "application/pdf"
                if format == "pdf"
                else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )


def _number(value):
    if value in (None, ""):
        return None
    from .calculation import money

    return float(money(value))

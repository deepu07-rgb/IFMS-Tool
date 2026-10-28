"""Small macro-free formula workbook; imported statement results are never trusted."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import json

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

from .calculation import money
from .supplied_form import HEADS, TEMPLATES, export_supplied_excel
from .worksheet import (
    VERSION,
    COLUMNS,
    AMOUNT_FIELDS,
    DEDUCTION_KEYS,
    REFERENCES,
    KINDS,
)

SETUP = {
    "employee_name": 3,
    "employee_code": 4,
    "kind": 5,
    "rounding": 6,
    "template_id": 7,
    "include_hda": 8,
}
BOOLEAN_FIELDS = ("history_reviewed", "deductions_reviewed", "hra_exempt")
EXTRA_FIELDS = {
    "da_override_reason",
    "hra_override_reason",
    "hra_exempt",
    "hra_exemption_reference",
}
LEGACY_COLUMNS = tuple(c for c in COLUMNS if c[0] not in EXTRA_FIELDS)
NUMERIC = AMOUNT_FIELDS | {"new_basic", "da_rate", "hra_rate", "hda_divisor"}
SALARY_COLUMNS = (
    "employee_code",
    "employee_name",
    "pay_month",
    "bill_no",
    *HEADS,
    "OTHER_ALLOWANCES",
    "total_deductions",
    "printed_gross",
    "printed_net",
    "reference",
)
BLUE = PatternFill("solid", fgColor="E5F2FF")


def _literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _open(stream):
    data = stream.read(5_000_001)
    if len(data) > 5_000_000:
        raise ValueError("Use a macro-free Excel workbook smaller than 5 MB")
    try:
        with ZipFile(BytesIO(data)) as archive:
            files = archive.infolist()
            if len(files) > 500 or sum(f.file_size for f in files) > 30_000_000:
                raise ValueError("Workbook contents exceed the supported size")
            if any(
                "vbaproject" in f.filename.lower()
                or "externallinks/" in f.filename.lower()
                for f in files
            ):
                raise ValueError("Macros and external workbook links are not supported")
        book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
        if any(s.max_row > 5000 or s.max_column > 100 for s in book):
            raise ValueError("Workbook exceeds 5000 rows or 100 columns")
        return book
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ValueError("Cannot read workbook: " + str(exc)) from None


def _value(cell):
    if cell.data_type == "f":
        raise ValueError(
            f"Input formula in {cell.parent.title}!{cell.coordinate}; paste the verified value instead"
        )
    value = cell.value
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return value


def _list(sheet, column, formula, first=6, last=125):
    validation = DataValidation(
        type="list", formula1='"' + formula + '"', allow_blank=True
    )
    validation.errorTitle = "Choose a listed value"
    validation.error = "Select one of the supported choices."
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"{column}{first}:{column}{last}")


def export_workbook(payload, validation=None):
    rows = payload.get("rows") or [{}]
    if (
        not isinstance(rows, list)
        or len(rows) > 120
        or any(not isinstance(r, dict) for r in rows)
    ):
        raise ValueError("Use 1–120 monthly rows")
    template_id = payload.get("template_id", "jodhpur_classic")
    if template_id not in TEMPLATES:
        raise ValueError("Choose a supported supplied form")
    # Borrow the existing supplied grid, typography, print area and signatures.
    shell = dict(
        status="ready",
        rows=[
            dict(
                month=f"ROW {i + 1}",
                components=[
                    dict(head=h, due="0", drawn="0", difference="0") for h in HEADS
                ],
            )
            for i in range(len(rows))
        ],
        deductions=[],
        totals=dict(
            due="0",
            drawn="0",
            difference="0",
            prior_settlement="0",
            total_deductions="0",
            net_payable="0",
        ),
    )
    book = load_workbook(
        export_supplied_excel(
            shell, dict(supplied_template_id=template_id, status="draft")
        )
    )
    setup = book.create_sheet("Setup", 0)
    inputs = book.create_sheet("Monthly inputs", 1)
    guide = book.create_sheet("Guide")
    setup.append(["IFMS assisted due–drawn workbook"])
    setup.append(["Template version", VERSION])
    for key, r in SETUP.items():
        setup.cell(r, 1, key.replace("_", " ").title())
        value = payload.get(key, "")
        if key == "include_hda":
            value = "Yes" if value is True else "No"
        _literal(setup.cell(r, 2), value)
        setup.cell(r, 2).fill = BLUE
    setup["A14"] = "HRA category (Y / Z / employee-specific)"
    _literal(setup["B14"], payload.get("hra_category"))
    setup["B14"].fill = BLUE
    setup["A15"] = (
        "Current rate library is checked in the app, not by standalone formulas."
    )
    setup["A10"] = (
        "Fill blue cells. The supplied statement updates automatically in Excel."
    )
    setup["A11"] = (
        "Return this workbook to the app to recheck its literal inputs and export PDF."
    )
    setup["A12"] = (
        "DRAFT: rates, entitlement and payment history need office verification."
    )
    setup.column_dimensions["A"].width = 64
    setup.column_dimensions["B"].width = 42
    for key, options in (
        ("kind", ",".join(KINDS)),
        ("rounding", "rupee,paise"),
        ("template_id", ",".join(TEMPLATES)),
        ("include_hda", "Yes,No"),
    ):
        _list(setup, "B", options, SETUP[key], SETUP[key])
    inputs["A1"] = "MONTHLY INPUTS — blue cells are editable; blanks are unresolved"
    inputs["A2"] = (
        "Drawn and due values are full-month amounts. Dates prorate them. Earlier paid amounts and extra deductions are actual allocations; never prorated again."
    )
    inputs["A3"] = (
        "Do not repeat previous payments already included in Drawn. One effective change per month; use manual monthly amounts for blended changes."
    )
    for col, (key, label, group) in enumerate(COLUMNS, 1):
        cell = inputs.cell(5, col, key)
        cell.comment = Comment(label, "IFMS")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17626C")
        cell.alignment = Alignment(wrap_text=True)
        inputs.column_dimensions[get_column_letter(col)].width = (
            30 if group == "references" else 19
        )
        if group == "link":
            inputs.column_dimensions[get_column_letter(col)].hidden = True
        for i, row in enumerate(rows, 6):
            value = row.get(key)
            if key in NUMERIC and value not in (None, ""):
                try:
                    value = float(money(value))
                except ValueError:
                    pass  # Invalid literal stays visible and fails the row check.
            if key in BOOLEAN_FIELDS and isinstance(value, bool):
                value = "Yes" if value else "No"
            if key in ("start", "end") and value:
                try:
                    value = date.fromisoformat(value)
                except (ValueError, TypeError):
                    pass
            cell = inputs.cell(i, col)
            _literal(cell, value)
            cell.fill = BLUE
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if key in NUMERIC:
                cell.number_format = "0.00"
            elif key in ("start", "end"):
                cell.number_format = "yyyy-mm-dd"
        if key == "presentation":
            _list(inputs, get_column_letter(col), "period,whole_month")
        if key in BOOLEAN_FIELDS:
            _list(inputs, get_column_letter(col), "Yes,No")
    last_check = len(COLUMNS) + 1
    inputs.cell(5, last_check, "INPUT CHECK")
    inputs.column_dimensions[get_column_letter(last_check)].width = 22
    inputs.freeze_panes = "E6"
    inputs.auto_filter.ref = f"A5:{get_column_letter(last_check)}{len(rows) + 5}"
    inputs.row_dimensions[5].height = 48
    for i in range(6, len(rows) + 6):
        inputs.cell(i, last_check, _check_formula(i))
        inputs.row_dimensions[i].height = 36
    from openpyxl.formatting.rule import CellIsRule

    inputs.conditional_formatting.add(
        f"{get_column_letter(last_check)}6:{get_column_letter(last_check)}{len(rows)+5}",
        CellIsRule(
            operator="equal",
            formula=['"INCOMPLETE"'],
            fill=PatternFill("solid", fgColor="FFDDCC"),
        ),
    )
    statement = book["Due drawn statement"]
    statement["A2"] = '= "NAME = " & Setup!B3'
    statement["P2"] = '= "EMP. ID = " & Setup!B4'
    for i in range(6, len(rows) + 6):
        check = _ref(i, "CHECK") + '="OK"'
        month = _ref(i, "month")
        start, end = _ref(i, "start"), _ref(i, "end")
        statement.cell(
            i,
            1,
            f'=IF({check},IF(AND(DAY({start})=1,DAY({end})=DAY(DATE(YEAR({end}),MONTH({end})+1,0))),{month},TEXT({start},"dd/mm/yyyy")&CHAR(10)&TEXT({end},"dd/mm/yyyy")),"")',
        )
        for col, a, b in ((2, "bill_no", "bill_date"), (3, "tv_no", "tv_date")):
            statement.cell(
                i,
                col,
                f'=IF({check},{_ref(i,a)}&IF({_ref(i,b)}="","",CHAR(10)&{_ref(i,b)}),"")',
            )
        fraction = f"(({end}-{start}+1)/DAY(DATE(YEAR({start}),MONTH({start})+1,0)))"
        precision = 'IF(Setup!$B$6="rupee",0,2)'
        for h, head in enumerate(HEADS):
            drawn = _ref(i, "drawn_" + head)
            monthly = _monthly_formula(i, head)
            period_due = f"({monthly})*{fraction}"
            whole_due = f"{drawn}*(1-{fraction})+({monthly})*{fraction}"
            statement.cell(
                i,
                h + 4,
                f'=IF({check},ROUND(IF({_ref(i,"presentation")}="whole_month",{whole_due},{period_due}),{precision}),"")',
            )
            statement.cell(
                i,
                h + 10,
                f'=IF({check},IF({_ref(i,"presentation")}="whole_month",{drawn},ROUND({drawn}*{fraction},{precision}))+{_ref(i,"prior_"+head)},"")',
            )
            statement.cell(
                i,
                h + 16,
                f'=IF({check},{get_column_letter(h+4)}{i}-{get_column_letter(h+10)}{i},"")',
            )
        for col, first in ((9, 4), (15, 10), (21, 16)):
            statement.cell(
                i,
                col,
                f'=IF({check},SUM({get_column_letter(first)}{i}:{get_column_letter(col-1)}{i}),"")',
            )
        for offset in range(4):
            classic = TEMPLATES["jodhpur_classic"]["deduction_keys"][offset]
            newer = TEMPLATES["jodhpur_gpf_arrear"]["deduction_keys"][offset]
            statement.cell(
                i,
                22 + offset,
                f'=IF({check},IF(Setup!$B$7="jodhpur_classic",{_ref(i,"ded_"+classic)},{_ref(i,"ded_"+newer)}),"")',
            )
        statement.cell(
            i,
            26,
            f'=IF({check},IF(Setup!$B$7="jodhpur_classic",SUM(V{i}:Y{i}),V{i}+Y{i}),"")',
        )
        statement.cell(i, 27, f'=IF({check},U{i}-Z{i},"")')
    # Header variant also updates if a user changes Setup while working offline.
    for n in range(4):
        a = TEMPLATES["jodhpur_classic"]["deduction_labels"][n].replace('"', '""')
        b = TEMPLATES["jodhpur_gpf_arrear"]["deduction_labels"][n].replace('"', '""')
        statement.cell(4, 22 + n, f'=IF(Setup!$B$7="jodhpur_classic","{a}","{b}")')
    total_row = len(rows) + 6
    check_col = get_column_letter(last_check)
    checks = f"'Monthly inputs'!{check_col}6:{check_col}{total_row-1}"
    all_ready = f'AND(COUNTIF({checks},"INCOMPLETE")=0,COUNTIF({checks},"OK")>0)'
    for col in range(4, 28):
        letter = get_column_letter(col)
        statement.cell(
            total_row, col, f'=IF({all_ready},SUM({letter}6:{letter}{total_row-1}),"")'
        )
    statement.oddFooter.center.text = "DRAFT — CHECK INPUT STATUS BEFORE SIGNING"
    statement.freeze_panes = "D6"
    statement.page_setup.fitToHeight = 0
    statement.print_title_rows = "1:5"
    guide_lines = [
        "QUICK START",
        "1. In Setup enter employee, claim type, rupee/paise rounding and the supplied form variant.",
        "2. In Monthly inputs enter dates and actual full-month drawn Basic, DA, HRA, HDA and CCA. Enter 0 explicitly for a confirmed absent head.",
        "3. Enter revised basic / sanctioned DA and HRA percentages required by the claim type. HDA uses (Basic + unrounded DA) / the explicitly approved divisor only when Include HDA is Yes.",
        "4. Manual mode accepts five explicit full-month due amounts. It can represent reviewed blended rates within a month; do not put guessed values in it.",
        "5. Enter earlier payments by head and the month they cover. Enter zeros only after checking there is no additional payment. Do not enter a payment twice if already included in Drawn.",
        "6. Enter only ADDITIONAL deductions from this arrear, not the entire salary deductions. Already-deducted GPF and SIF in the newer form are disclosures, not additional subtraction.",
        "7. Add all four references, set both checks to Yes, and ensure every INPUT CHECK is OK. The statement then displays the calculated amounts and totals.",
        "8. Import this workbook back into Excel Arrear Maker for independent recalculation, saved revisions and PDF. Statement formulas/cached results are ignored on import; input cells must contain literal values.",
        "For fewer months, clear all editable inputs in unused trailing rows. Keep the formula rows and columns in place. Use the app to generate a different date range or rebuild the form.",
        "DATES AND ROUNDING",
        "Dates include both endpoints. Calendar month length is used, including leap February. Each due/drawn head is rounded separately after proration; monthly differences are then added.",
        "period: show only the selected days. whole_month: retain original drawn pay for unaffected days and show a full monthly drawn salary. Prior payments and extra deductions are exact allocated amounts and never prorated again.",
        "LIMITS",
        "One effective segment per month, up to 120 months, five printed salary heads. Unchanged Mess/FTA or other heads are excluded from this selected-head form, not treated as new arrears. Rates and entitlement are human-reviewed inputs, not automatically certified rules.",
        "The workbook does not authorize payment. Keep the source documents and have the office verify and sign its draft. Handwriting OCR is optional transcription and still requires review.",
        "Saved-ledger link columns are hidden. If you intentionally change drawn values from a saved record, use the app's manual-values action to detach the link and record the correction reference.",
    ]
    for i, text in enumerate(guide_lines, 1):
        guide.cell(i, 1, text)
        guide.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[i].height = 48 if len(text) > 100 else 28
    guide.column_dimensions["A"].width = 125
    if validation is not None:
        checks = book.create_sheet("Rate checks")
        checks.append(
            [
                "Validation snapshot only; return workbook to app for current-library checking after edits."
            ]
        )
        checks.append(["Row", "Head / field", "Status", "Details / source"])
        for item in validation.get("rate_checks", []):
            checks.append(
                [
                    item["row"] + 1,
                    item["head"],
                    item["status"],
                    json.dumps(item, ensure_ascii=False),
                ]
            )
        for item in validation.get("warnings", []) + validation.get("errors", []):
            checks.append(
                [
                    item["row"] + 1 if item.get("row") is not None else "Setup",
                    item.get("field", ""),
                    "Review required",
                    item.get("message", ""),
                ]
            )
        for cells in checks:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ("A", "B", "C", "D"):
            checks.column_dimensions[col].width = 24 if col != "D" else 90
    book.calculation = CalcProperties(
        calcId=191029, fullCalcOnLoad=True, forceFullCalc=True
    )
    book.active = 0
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def _ref(row, field):
    col = (
        len(COLUMNS) + 1
        if field == "CHECK"
        else next(i for i, c in enumerate(COLUMNS, 1) if c[0] == field)
    )
    return f"'Monthly inputs'!{get_column_letter(col)}{row}"


def _monthly_formula(i, head):
    kind, hda = "Setup!$B$5", "Setup!$B$8"
    drawn = _ref(i, "drawn_" + head)
    basic = f'IF({kind}="increment",{_ref(i,"new_basic")},{_ref(i,"drawn_BASIC")})'
    da = f'({basic})*{_ref(i,"da_rate")}/100'
    if head == "BASIC":
        auto = basic
    elif head == "DA":
        auto = f'IF(OR({kind}="increment",{kind}="da"),{da},{drawn})'
    elif head == "HRA":
        auto = f'IF(OR({kind}="increment",{kind}="hra"),({basic})*{_ref(i,"hra_rate")}/100,{drawn})'
    elif head == "HDA":
        auto = f'IF(AND(OR({kind}="increment",{kind}="da"),{hda}="Yes"),(({basic})+({da}))/{_ref(i,"hda_divisor")},{drawn})'
    else:
        auto = drawn
    return f'IF({kind}="manual",{_ref(i,"due_"+head)},{auto})'


def _numeric_checks(i, fields):
    refs = [_ref(i, f) for f in fields]
    parts = [
        f'COUNT({",".join(refs)})={len(refs)}',
        f'MIN({",".join(refs)})>=0',
        f'MAX({",".join(refs)})<=100000000',
    ]
    for f, ref in zip(fields, refs):
        if f in AMOUNT_FIELDS or f == "new_basic":
            parts.append(f"{ref}=ROUND({ref},2)")
        if f in ("da_rate", "hra_rate"):
            parts.append(f"{ref}<=1000")
        if f == "hda_divisor":
            parts.append(f"{ref}>0")
    return "AND(" + ",".join(parts) + ")"


def _check_formula(i):
    common = ["drawn_" + h for h in HEADS] + ["prior_" + h for h in HEADS]
    parts = [_numeric_checks(i, common)]
    a, b, m = _ref(i, "start"), _ref(i, "end"), _ref(i, "month")
    parts += [
        f"ISNUMBER({a})",
        f"ISNUMBER({b})",
        f"{a}<={b}",
        f'TEXT({a},"yyyy-mm")={m}',
        f'TEXT({b},"yyyy-mm")={m}',
        f'OR({_ref(i,"presentation")}="period",{_ref(i,"presentation")}="whole_month")',
    ]
    if i > 6:
        parts += [f'{a}={_ref(i-1,"end")}+1', f'{m}<>{_ref(i-1,"month")}']
    parts += [f"LEN(TRIM({_ref(i,k)}))>0" for k in REFERENCES]
    parts += [f'{_ref(i,k)}="Yes"' for k in ("history_reviewed", "deductions_reviewed")]
    parts += [
        "LEN(TRIM(Setup!B3))>0",
        "LEN(TRIM(Setup!B4))>0",
        'OR(Setup!B6="rupee",Setup!B6="paise")',
        'OR(Setup!B8="Yes",Setup!B8="No")',
    ]
    kind_checks = {
        "hra": ["hra_rate"],
        "da": ["da_rate"],
        "increment": ["new_basic", "da_rate", "hra_rate"],
        "manual": ["due_" + h for h in HEADS],
    }
    choice = "FALSE"
    for kind, fields in kind_checks.items():
        choice = f'IF(Setup!B5="{kind}",{_numeric_checks(i,fields)},{choice})'
    parts.append(choice)
    parts.append(
        f'IF(AND(OR(Setup!B5="increment",Setup!B5="da"),Setup!B8="Yes"),{_numeric_checks(i,["hda_divisor"])},TRUE)'
    )
    choice = "FALSE"
    for template, conf in TEMPLATES.items():
        checks = [_numeric_checks(i, ["ded_" + k for k in conf["deduction_keys"]])]
        checks += [
            f'OR({_ref(i,"ded_"+k)}="",{_ref(i,"ded_"+k)}=0)'
            for k in DEDUCTION_KEYS
            if k not in conf["deduction_keys"]
        ]
        choice = f'IF(Setup!B7="{template}",AND({",".join(checks)}),{choice})'
    parts.append(choice)
    used = f"COUNTA('Monthly inputs'!A{i}:{get_column_letter(len(COLUMNS))}{i})"
    return (
        f'=IF({used}=0,"UNUSED",IFERROR(IF(AND('
        + ",".join(parts)
        + '),"OK","INCOMPLETE"),"INCOMPLETE"))'
    )


def import_workbook(stream):
    book = _open(stream)
    if not {"Setup", "Monthly inputs"} <= set(book.sheetnames):
        raise ValueError("Use the downloaded IFMS formula workbook template")
    version = _value(book["Setup"]["B2"])
    if version not in (VERSION, "ifms-assisted-worksheet-v1"):
        raise ValueError("Use the downloaded IFMS formula workbook template")
    sheet = book["Monthly inputs"]
    columns = LEGACY_COLUMNS if version == "ifms-assisted-worksheet-v1" else COLUMNS
    if [_value(sheet.cell(5, i)) for i in range(1, len(columns) + 1)] != [
        c[0] for c in columns
    ]:
        raise ValueError("Monthly input headers differ from the template")
    result = {key: _value(book["Setup"].cell(r, 2)) for key, r in SETUP.items()}
    category = _value(book["Setup"]["B14"])
    if category is not None:
        result["hra_category"] = category
    if result["include_hda"] not in ("Yes", "No"):
        raise ValueError("Include HDA must be Yes or No")
    result["include_hda"] = result["include_hda"] == "Yes"
    rows = []
    for cells in sheet.iter_rows(min_row=6, max_col=len(columns)):
        if not any(c.value not in (None, "") for c in cells):
            continue
        row = {}
        for (key, _, _), cell in zip(columns, cells):
            value = _value(cell)
            if value is None:
                continue
            if (
                key in NUMERIC
                and isinstance(value, (int, float))
                and not isinstance(value, bool)
            ):
                value = format(Decimal(str(value)), "f")
            if key in BOOLEAN_FIELDS:
                if value not in ("Yes", "No"):
                    raise ValueError("Review fields must be Yes or No")
                value = value == "Yes"
            row[key] = value
        rows.append(row)
    if not 1 <= len(rows) <= 120:
        raise ValueError("Use 1–120 monthly input rows")
    result["rows"] = rows
    return result


def salary_template():
    book = Workbook()
    sheet = book.active
    sheet.title = "Salary ledger"
    sheet.append(list(SALARY_COLUMNS))
    sheet.freeze_panes = "E2"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = 22
    guide = book.create_sheet("Guide")
    for text in (
        "Enter literal full-month actual salary figures. Confirm absent heads with 0; blanks are unresolved.",
        "OTHER_ALLOWANCES is the sum of all heads outside BASIC / DA / HRA / HDA / CCA, including signed recoveries where applicable.",
        "total_deductions excludes informational deductions already outside printed net. Printed gross and printed net must reconcile exactly.",
        "Use one regular salary per employee/month. Prior arrear bills belong in worksheet allocated earlier-payment fields, not this regular-salary import.",
        "Upload to preview, check the original records, and explicitly confirm before saving to the employee ledger.",
    ):
        guide.append([text])
        guide.cell(guide.max_row, 1).alignment = Alignment(wrap_text=True)
        guide.row_dimensions[guide.max_row].height = 40
    guide.column_dimensions["A"].width = 110
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def read_salaries(stream):
    book = _open(stream)
    if "Salary ledger" not in book:
        raise ValueError("Use the Salary ledger template")
    sheet = book["Salary ledger"]
    if [_value(sheet.cell(1, i)) for i in range(1, len(SALARY_COLUMNS) + 1)] != list(
        SALARY_COLUMNS
    ):
        raise ValueError("Salary ledger headers differ from the template")
    rows = []
    for cells in sheet.iter_rows(min_row=2, max_col=len(SALARY_COLUMNS)):
        if not any(c.value not in (None, "") for c in cells):
            continue
        raw = {key: _value(cell) for key, cell in zip(SALARY_COLUMNS, cells)}
        for key, value in raw.items():
            if value in (None, "") or isinstance(value, bool):
                raise ValueError(
                    f"Salary row {cells[0].row}: enter {key}, including an explicit 0 when applicable"
                )
        rows.append(
            dict(
                employee_code=str(raw["employee_code"]).strip().upper(),
                employee_name=str(raw["employee_name"]),
                pay_month=str(raw["pay_month"]),
                bill_no=str(raw["bill_no"]),
                bill_type="Regular Salary (structured Excel)",
                allowances={k: str(raw[k]) for k in (*HEADS, "OTHER_ALLOWANCES")},
                deductions={"TOTAL_DEDUCTIONS": str(raw["total_deductions"])},
                printed_gross=str(raw["printed_gross"]),
                printed_deductions=str(raw["total_deductions"]),
                printed_net=str(raw["printed_net"]),
                reviewed=False,
                reference=str(raw["reference"]),
            )
        )
    if not 1 <= len(rows) <= 1200:
        raise ValueError("Use 1–1200 salary records")
    return rows

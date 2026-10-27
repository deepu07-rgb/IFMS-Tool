import json
from io import BytesIO
from decimal import Decimal
from pathlib import Path
from flask import jsonify, send_file, current_app
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .storage import database


def statement_pdf(filename, payload, status):
    """A review report for workbook values; never mislabel it the original form."""
    from html import escape
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors

    b = BytesIO()
    doc = SimpleDocTemplate(b, pagesize=landscape(A4), leftMargin=24, rightMargin=24)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(escape(status), styles["Heading2"]),
        Paragraph("Workbook review report: " + escape(filename), styles["Normal"]),
        Paragraph(
            "Cached source values; verify employee, worksheet and rates. Download the original workbook for its exact layout.",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]
    rows = [
        [
            "Month",
            "Sheet / row",
            "Due",
            "Drawn",
            "Printed difference",
            "Calculated difference",
        ]
    ]
    for r in payload.get("statement_rows", []):
        rows.append(
            [
                r["month"],
                Paragraph(
                    escape(
                        str(r.get("source_sheet", ""))
                        + " / "
                        + str(r.get("source_row", ""))
                    ),
                    styles["Normal"],
                ),
                r["printed_due"],
                r["printed_drawn"],
                r.get("printed_difference") or "",
                r["calculated_difference"],
            ]
        )
    table = Table(rows, repeatRows=1, colWidths=[65, 200, 105, 105, 120, 130])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#174D52")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    for r in payload.get("statement_rows", []):
        details = []
        for side in ("due", "drawn"):
            for head, amount in r.get(side, {}).items():
                details.append(
                    escape(side.title() + " " + head)
                    + ": "
                    + escape(str(amount) if amount is not None else "not read")
                )
        story.append(
            Paragraph(escape(r["month"]) + ": " + "; ".join(details), styles["Normal"])
        )
        if r.get("issues"):
            story.append(
                Paragraph(
                    escape(r["month"] + ": " + "; ".join(r["issues"])), styles["Normal"]
                )
            )
    doc.build(story)
    b.seek(0)
    return b


def format_sheet(sheet):
    for c in sheet[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="174D52")
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "D3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 18


def register_document_exports(suite):
    @suite.get("/api/imports/<int:id>/<kind>")
    def export_document(id, kind):
        with database() as db:
            row = db.execute("SELECT * FROM suite_imports WHERE id=?", (id,)).fetchone()
            if not row:
                return jsonify(error="Import not found"), 404
            approved = [
                dict(json.loads(r["payload"]), employee_code=r["employee_code"])
                for r in db.execute(
                    "SELECT * FROM suite_payments WHERE import_id=?", (id,)
                )
            ]
        payload = json.loads(row["payload"])
        status = (
            "REVIEWED"
            if row["status"] == "reviewed"
            else "DRAFT - SOURCE REVIEW REQUIRED"
        )
        if kind == "pdf":
            if payload.get("file_kind") == "xlsx":
                return send_file(
                    statement_pdf(row["filename"], payload, status),
                    mimetype="application/pdf",
                    as_attachment=True,
                    download_name=f"statement-review-{id}.pdf",
                )
            path = (
                Path(current_app.config["UPLOAD_FOLDER"])
                / "suite"
                / (row["sha256"] + ".pdf")
            )
            return send_file(
                path,
                mimetype="application/pdf",
                as_attachment=True,
                download_name="original-" + row["filename"],
            )
        if kind != "excel":
            return jsonify(error="Unknown export format"), 404
        w = Workbook()
        s = w.active
        s.title = "Payment records"
        status = (
            "REVIEWED"
            if row["status"] == "reviewed"
            else "DRAFT — SOURCE REVIEW REQUIRED"
        )
        s.append([status + " · " + row["filename"]])
        records = approved or payload.get("ga55", {}).get("payments", [])
        allowances = sorted({k for p in records for k in p.get("allowances", {})})
        deductions = sorted({k for p in records for k in p.get("deductions", {})})
        headers = [
            "Employee code",
            "Pay month",
            "Bill type",
            "Bill number",
            *allowances,
            *["Deduction: " + k for k in deductions],
            "Entitlement from",
            "Entitlement through",
        ]
        s.append(headers)
        for p in records:
            values = [
                p.get("employee_code", ""),
                p["pay_month"],
                p["bill_type"],
                p["bill_no"],
            ]
            for kind, heads in [("allowances", allowances), ("deductions", deductions)]:
                values.extend(
                    float(Decimal(p[kind][k])) if k in p.get(kind, {}) else None
                    for k in heads
                )
            period = p.get("entitlement_period") or {}
            values.extend([period.get("start", ""), period.get("end", "")])
            s.append(values)
        if payload.get("statement_rows"):
            heads = list(
                dict.fromkeys(
                    head
                    for r in payload["statement_rows"]
                    for side in ("due", "drawn")
                    for head in r.get(side, {})
                )
            )
            s = w.create_sheet("Statement review")
            s.append([status + " · " + row["filename"]])
            s.append(
                [
                    "Month",
                    "Worksheet",
                    "Source row",
                    "Page",
                    *[
                        side.title() + " " + head
                        for side in ("due", "drawn")
                        for head in heads
                    ],
                    "Printed Due",
                    "Printed Drawn",
                    "Printed difference",
                    "Calculated difference",
                    "Issues",
                ]
            )
            for r in payload["statement_rows"]:
                s.append(
                    [
                        r["month"],
                        r.get("source_sheet", ""),
                        r.get("source_row", ""),
                        r.get("source_page", ""),
                        *[
                            (
                                float(Decimal(r.get(k, {}).get(h)))
                                if r.get(k, {}).get(h) is not None
                                else None
                            )
                            for k in ("due", "drawn")
                            for h in heads
                        ],
                        r["printed_due"],
                        r["printed_drawn"],
                        r.get("printed_difference"),
                        r["calculated_difference"],
                        "; ".join(r.get("issues", [])),
                    ]
                )
        if payload.get("extracted"):
            s = w.create_sheet("Outer sheet")
            s.append([status])
            s.append(["Field / component", "Type", "Amount / value"])
            for key, v in payload["extracted"].items():
                if key in ("Allowances", "Deductions"):
                    for name, amount in v.items():
                        s.append([name, key, amount])
                elif not isinstance(v, (dict, list)):
                    s.append([key, "Bill metadata", v])
        raw = w.create_sheet("Source text")
        raw.append([status])
        raw.append(["Page", "Text"])
        for i, text in enumerate(payload.get("pages", []), 1):
            for offset in range(0, len(text) or 1, 30000):
                raw.append([i, text[offset : offset + 30000]])
        for sheet in w:
            format_sheet(sheet)
        b = BytesIO()
        w.save(b)
        b.seek(0)
        return send_file(
            b,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"extracted-{id}.xlsx",
        )

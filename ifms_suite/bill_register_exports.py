"""Literal bill register exports, with a fail-closed official-form boundary.

No GF&AR form number or column layout has been authenticated. Consequently
``official`` cannot export until an authoritative form has been reviewed and
its mapping implemented here. Raw exports never calculate or normalize values.
"""

from html import escape
from io import BytesIO
from math import floor
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Table,
    TableStyle,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import reportlab


OFFICIAL_COLUMNS = ()
OFFICIAL_REFERENCE = {
    "url": "https://finance.rajasthan.gov.in/website/Rules.aspx",
    "form": None,
    "page": None,
    "verified": False,
    "verification": (
        "Official Rajasthan GF&AR Bill Register form not verified: official "
        "Finance Rajasthan sources were unavailable (network proxy HTTP 403). "
        "No form identifier, columns, or mapping is asserted."
    ),
}

# FS is the explicit option used by this application, not a universal definition
# of foolscap. Both formats use landscape orientation.
PAPER_SIZES = {"A4": landscape(A4), "FS": (13 * inch, 8.5 * inch)}
PAPER_LABELS = {"A4": "A4 landscape", "FS": "FS landscape (13 x 8.5 inches)"}


def map_official(rows):
    """Refuse to manufacture government columns or infer undocumented fields."""
    raise ValueError(
        "Official GF&AR Bill Register export is unavailable. "
        + OFFICIAL_REFERENCE["verification"]
        + " Use raw mode to export the selected IFMS columns."
    )


def _text(value):
    return "" if value is None else str(value)


def _validated(payload):
    if not isinstance(payload, dict):
        raise ValueError("Register payload must be an object.")
    columns, rows = payload.get("columns"), payload.get("rows")
    if (
        not isinstance(columns, list)
        or not columns
        or not all(isinstance(column, str) and column for column in columns)
    ):
        raise ValueError("Select at least one named register column.")
    if len(set(columns)) != len(columns):
        raise ValueError("Register column names must be unique.")
    if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
        raise ValueError("Register rows must be a list of objects.")
    return columns, [[_text(row.get(column)) for column in columns] for row in rows]


def _excel(columns, rows, paper, draft=False):
    if any(len(value) > 32767 for values in [columns, *rows] for value in values):
        raise ValueError(
            "Excel cells support at most 32767 characters; use PDF for long register text."
        )
    book = Workbook()
    sheet = book.active
    sheet.title = "Bill register"
    for values in [columns, *rows]:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # '=...', '+...', leading zeros and numeric-looking IDs stay text.
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174D52")
    for index, column in enumerate(columns, 1):
        width = max([len(column), *[len(row[index - 1]) for row in rows]])
        sheet.column_dimensions[get_column_letter(index)].width = min(
            45, max(14, width + 2)
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.print_area = sheet.dimensions
    sheet.page_setup.orientation = "landscape"
    if paper == "A4":
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    else:
        # A custom page prevents substitution with the 8.5 x 14 legal preset.
        sheet.page_setup.paperWidth = "13in"
        sheet.page_setup.paperHeight = "8.5in"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = (
        ("DRAFT | " if draft else "") + PAPER_LABELS[paper] + " | Page &P of &N"
    )
    if draft:
        status = book.create_sheet("Export status")
        status.append(["DRAFT - source register requires review"])
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def _pdf(columns, rows, paper, draft=False):
    # Embed a font rather than silently replacing currency and other Unicode
    # characters with boxes. ReportLab ships Vera as a portable Latin fallback.
    font_path = next(
        (
            path
            for path in [
                Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                Path(reportlab.__file__).parent / "fonts" / "Vera.ttf",
            ]
            if path.is_file()
        ),
        None,
    )
    if font_path is None:
        raise ValueError("A PDF text font is unavailable; use the Excel export.")
    font = TTFont("BillRegisterText", str(font_path))
    pdfmetrics.registerFont(font)
    if any(
        ord(char) not in font.face.charToGlyph
        for values in [columns, *rows]
        for value in values
        for char in value
        if not char.isspace()
    ):
        raise ValueError(
            "PDF font cannot represent some register characters; use Excel to preserve the original text."
        )
    output = BytesIO()
    size = PAPER_SIZES[paper]
    doc = SimpleDocTemplate(
        output,
        pagesize=size,
        leftMargin=24,
        rightMargin=24,
        topMargin=38,
        bottomMargin=32,
        title="Bill register - raw IFMS columns",
    )
    style = ParagraphStyle(
        "Register cell",
        fontName="BillRegisterText",
        fontSize=8,
        leading=10,
        splitLongWords=True,
        spaceAfter=0,
        spaceBefore=0,
    )
    heading = ParagraphStyle("Register heading", parent=style, fontSize=10, leading=13)

    def paragraph(value):
        return Paragraph(escape(value).replace("\n", "<br/>"), style)

    # Wide registers get consecutive column groups. Every group retains all rows
    # and states its column range; tall cells split across pages without truncation.
    available = size[0] - 48
    group_size = max(1, floor(available / 105))
    story = []
    for start in range(0, len(columns), group_size):
        stop = min(start + group_size, len(columns))
        if story:
            story.append(PageBreak())
        story.append(
            Paragraph(
                f"Raw IFMS bill register | Columns {start + 1}-{stop} of {len(columns)}",
                heading,
            )
        )
        data = [[paragraph(value) for value in columns[start:stop]]]
        data.extend([[paragraph(value) for value in row[start:stop]] for row in rows])
        table = Table(
            data,
            colWidths=[available / (stop - start)] * (stop - start),
            repeatRows=1,
            splitByRow=1,
            splitInRow=1,
            hAlign="LEFT",
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCECEE")),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#8A999A")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(table)

    def footer(canvas, _doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(24, 18, ("DRAFT | " if draft else "") + PAPER_LABELS[paper])
        canvas.drawRightString(size[0] - 24, 18, f"Page {_doc.page}")

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    return output


def export_register(payload, mode="raw", format="excel", paper="A4"):
    """Return an editable XLSX or paginated PDF stream; never mutate payload.

    ``payload`` contains selected column keys and row dictionaries. Missing/None
    values remain blank, strings remain literal, and amounts are not adjusted.
    ``FS`` explicitly means a 13 x 8.5 inch landscape page in this application.
    """
    if mode not in {"raw", "official"}:
        raise ValueError("Register mode must be raw or official.")
    if format not in {"excel", "pdf"}:
        raise ValueError("Register format must be excel or pdf.")
    if paper not in PAPER_SIZES:
        raise ValueError("Register paper must be A4 or FS (13 x 8.5 inches landscape).")
    if mode == "official":
        map_official(payload.get("rows", []) if isinstance(payload, dict) else [])
    columns, rows = _validated(payload)
    labels = payload.get("column_labels", {})
    if not isinstance(labels, dict) or not all(
        isinstance(value, str) for value in labels.values()
    ):
        raise ValueError("Register column_labels must map column keys to text labels.")
    display_columns = [labels.get(column, column) for column in columns]
    draft = (
        bool(payload.get("draft")) or str(payload.get("status", "")).lower() == "draft"
    )
    return (_excel if format == "excel" else _pdf)(display_columns, rows, paper, draft)

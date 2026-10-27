"""Read explicitly labelled due/drawn tables without repairing source amounts."""

from datetime import datetime, date
import re
from .calculation import money, rounded


def month_value(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m")
    text = " ".join(str(v or "").replace("\n", " ").split())
    for fmt in (
        "%b-%Y",
        "%b-%y",
        "%B %Y",
        "%b %Y",
        "%B-%Y",
        "%Y-%m",
        "%d/%m/%Y",
        "%d-%m-%Y",
    ):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m")
        except ValueError:
            pass
    return None


def _label(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


_GROUPS = {
    "paydue": "due",
    "duepay": "due",
    "due": "due",
    "amountdue": "due",
    "paydrawn": "drawn",
    "drawnpay": "drawn",
    "drawn": "drawn",
    "amountdrawn": "drawn",
    "paydifference": "difference",
    "difference": "difference",
    "arrear": "difference",
    "arrears": "difference",
}
_COMPONENTS = {
    "basic": "Basic",
    "basicpay": "Basic",
    "da": "DA",
    "dearnessallowance": "DA",
    "hra": "HRA",
    "houserentallowance": "HRA",
    "hda": "HDA",
    "cca": "CCA",
    "citycompensatoryallowance": "CCA",
    "total": "Total",
    "grosstotal": "Total",
    "totalpay": "Total",
}


def _layout(rows, index, aliases):
    headings = [
        (i, _GROUPS[_label(v)])
        for i, v in enumerate(rows[index])
        if _label(v) in _GROUPS
    ]
    if not {"due", "drawn"}.issubset(group for _, group in headings):
        return None
    if len({group for _, group in headings}) != len(headings):
        return None
    # Headers may span blank/annotation rows, but never search across financial rows.
    for offset in range(1, 4):
        if index + offset >= len(rows):
            break
        header = rows[index + offset]
        if any(month_value(v) for v in header):
            break
        groups = {}
        for pos, (start, group) in enumerate(headings):
            end = (
                headings[pos + 1][0]
                if pos + 1 < len(headings)
                else max(len(header), len(rows[index]))
            )
            mapping = {}
            for column in range(start, end):
                name = (
                    aliases.get(_label(header[column]))
                    if column < len(header)
                    else None
                )
                if name is None or name in mapping:
                    break
                mapping[name] = column
            else:
                if "Total" in mapping and len(mapping) > 1:
                    groups[group] = mapping
        if len(groups) != len(headings):
            continue
        if set(groups["due"]) != set(groups["drawn"]):
            continue
        month_columns = [
            i
            for i, v in enumerate(rows[index][: headings[0][0]])
            if _label(v) in {"month", "paymonth", "monthyear", "period"}
        ]
        if len(month_columns) != 1:
            continue
        return index + offset, month_columns[0], groups
    return None


def parse_statement_table(rows, *, component_aliases=None, reference_rows=None):
    """Read explicitly labelled groups; absent amounts remain unresolved.

    Optional reference_rows are independently prepared month/due/drawn records.
    They are compared as supplied, never used to fill or repair statement cells.
    Unknown or ambiguous header layouts are deliberately left unparsed.
    """
    aliases = dict(_COMPONENTS)
    for label, component in (component_aliases or {}).items():
        if (
            not isinstance(component, str)
            or not component.strip()
            or component == "Total"
        ):
            raise ValueError("A component mapping requires a non-total component name")
        aliases[_label(label)] = component
    references = {}
    for reference in reference_rows or []:
        references.setdefault(month_value(reference.get("month")), []).append(reference)
    result = []
    layout = None
    for index, row in enumerate(rows):
        if any(_label(v) in _GROUPS for v in row):
            layout = _layout(rows, index, aliases)
            continue
        if layout is None:
            continue
        header_end, month_column, groups = layout
        if index <= header_end or month_column >= len(row):
            continue
        month = month_value(row[month_column])
        if not month:
            continue
        issues = []
        values = {}
        for group, mapping in groups.items():
            values[group] = {}
            for component, column in mapping.items():
                try:
                    amount = money(row[column]) if column < len(row) else money(None)
                except ValueError:
                    amount = None
                    issues.append(
                        f"{group.title()} {component} is missing or unreadable"
                    )
                values[group][component] = amount
        due, drawn = values["due"], values["drawn"]
        for group, amounts in [("Pay Due", due), ("Pay Drawn", drawn)]:
            if all(v is not None for v in amounts.values()) and rounded(
                sum(v for k, v in amounts.items() if k != "Total")
            ) != rounded(amounts["Total"]):
                issues.append(f"{group} components do not match printed total")
        difference = (
            rounded(due["Total"] - drawn["Total"])
            if due["Total"] is not None and drawn["Total"] is not None
            else None
        )
        printed = values.get("difference", {}).get("Total")
        if (
            printed is not None
            and difference is not None
            and rounded(printed) != difference
        ):
            issues.append("Printed difference does not match Due minus Drawn")
        for component, amount in values.get("difference", {}).items():
            if component == "Total" or amount is None:
                continue
            if component not in due or component not in drawn:
                issues.append(
                    f"Printed {component} difference has no matching Due/Drawn component"
                )
            elif (
                due[component] is not None
                and drawn[component] is not None
                and rounded(amount) != rounded(due[component] - drawn[component])
            ):
                issues.append(
                    f"Printed {component} difference does not match Due minus Drawn"
                )
        matches = references.get(month, [])
        if reference_rows is not None and len(matches) != 1:
            issues.append(
                "Independent source comparison incomplete: missing or ambiguous month"
            )
        elif matches:
            for side in ("due", "drawn"):
                source = matches[0].get(side, {})
                for component in dict.fromkeys([*values[side], *source]):
                    if component not in source:
                        issues.append(
                            f"{side.title()} {component} source comparison incomplete"
                        )
                        continue
                    expected = source[component]
                    try:
                        expected = money(expected)
                    except ValueError:
                        issues.append(
                            f"{side.title()} {component} source value is unresolved"
                        )
                        continue
                    actual = values[side].get(component)
                    if actual is None:
                        issues.append(
                            f"{side.title()} {component} source comparison incomplete"
                        )
                    elif rounded(actual) != rounded(expected):
                        issues.append(
                            f"{side.title()} {component} does not match independent source"
                        )

        def formatted(value):
            return f"{value:.2f}" if value is not None else None

        result.append(
            dict(
                month=month,
                due={k: formatted(v) for k, v in due.items() if k != "Total"},
                drawn={k: formatted(v) for k, v in drawn.items() if k != "Total"},
                printed_due=formatted(due["Total"]),
                printed_drawn=formatted(drawn["Total"]),
                printed_difference=(
                    formatted(rounded(printed)) if printed is not None else None
                ),
                calculated_difference=formatted(difference),
                issues=issues,
                source_row=index + 1,
                requires_review=True,
            )
        )
    return result


def parse_workbook(data):
    from io import BytesIO
    from zipfile import ZipFile
    from openpyxl import load_workbook

    with ZipFile(BytesIO(data)) as z:
        if sum(f.file_size for f in z.infolist()) > 200 * 1024 * 1024:
            raise ValueError("Workbook expanded size exceeds 200 MB")
    w = load_workbook(BytesIO(data), data_only=True, read_only=True)
    records = []
    source = []
    try:
        for s in w:
            if s.max_row * s.max_column > 300000:
                continue
            rows = list(s.iter_rows(values_only=True))
            parsed = parse_statement_table(rows)
            for r in parsed:
                r["source_sheet"] = s.title
            records.extend(parsed)
            if parsed:
                source.append(
                    "Worksheet: "
                    + s.title
                    + "\n"
                    + "\n".join(
                        " | ".join(str(v or "") for v in row[:40]) for row in rows[:300]
                    )
                )
    finally:
        w.close()
    return dict(
        statement_rows=records,
        pages=source,
        requires_review=True,
        notice="Cached workbook values were read; formulas were not executed. Compare sources and verify rates before creating a claim.",
    )


def parse_pdf_tables(data):
    import pymupdf

    d = pymupdf.open(stream=data, filetype="pdf")
    records = []
    try:
        for page in d:
            tables = page.find_tables().tables
            if not tables:
                tables = page.find_tables(strategy="text").tables
            for t in tables:
                for r in parse_statement_table(t.extract()):
                    r["source_page"] = page.number + 1
                    records.append(r)
    finally:
        d.close()
    return records

"""
IFMS 3.0 Salary Analytics Tool - v7
Local web application for parsing IFMS PDFs and analyzing salary data.
v7.1: Fixed multi-page schedule parsing - continuation pages without headers, blank page handling - all 10 schedule types, auto group detection.
"""
import os, json, sqlite3, io, re, unicodedata, tempfile, sys, platform
from datetime import datetime
import concurrent.futures
import threading
from datetime import datetime
from pathlib import Path
from ifms_engine.dictionary import (
    DictionaryManager,
    clean_designation_key,
    extract_designation_parts,
    extract_english_designation,
)
from ifms_engine.backup_engine import (
    create_full_backup,
    ensure_backup_folder,
    export_dictionary_payload,
    export_schedule_mapping_payload,
    import_dictionary_payload,
    import_schedule_mapping_payload,
    list_backups,
    restore_full_backup,
    safe_named_file,
)

executor = concurrent.futures.ThreadPoolExecutor(max_workers=4)

APP_NAME = "IFMS Report Analyzer"
BUILD_LABEL = "Phase3-Full-Stable"
SERVER_STARTED_AT = datetime.now().isoformat(timespec="seconds")
DEVELOPER_CREDIT = "Designed and Developed by - Neeraj Purohit, Senior Assistant, Rajasthan Police"
DEVELOPER_EMAIL = "rjjo200924001769@rajasthan.in"
APP_DISCLAIMER = "This app is designed to help increase office productivity. Human review of the data is compulsory before using the data."

def _designation_match_key(value):
    return clean_designation_key(value)

def _dict_match(raw_text, pattern, match_type):
    raw_en = extract_english_designation(raw_text)
    pat_en = extract_english_designation(pattern)
    if not raw_en or not pat_en:
        return False
    match_type = (match_type or "exact").lower()
    raw_key = _designation_match_key(raw_en)
    pat_key = _designation_match_key(pat_en)
    if match_type == "exact":
        return raw_key == pat_key
    if match_type == "contains":
        return pat_key in raw_key
    if match_type == "regex":
        try:
            return (
                re.search(pat_en, raw_en, re.IGNORECASE) is not None
                or re.search(pat_en, raw_key, re.IGNORECASE) is not None
            )
        except re.error:
            return False
    return False


def match_designation_dictionary(db, raw_designation_full):
    parts = extract_designation_parts(raw_designation_full)
    raw_designation_en = parts["raw_designation_en"]
    raw_designation_key = parts["raw_designation_key"]
    if not raw_designation_key:
        return None
    rows = db.execute("""
        SELECT id, raw_text, COALESCE(raw_text_key, '') AS raw_text_key,
               COALESCE(final_text, final_text_en, '') AS final_text,
               match_type, priority
        FROM dictionary_mappings_v3
        WHERE active = 1 AND category = 'designation'
        ORDER BY priority DESC, id DESC
    """).fetchall()
    for row in rows:
        raw_text = row["raw_text"] if isinstance(row, sqlite3.Row) else row[1]
        match_type = (row["match_type"] if isinstance(row, sqlite3.Row) else row[4]) or "exact"
        if _dict_match(raw_designation_en, raw_text, match_type):
            final_text = row["final_text"] if isinstance(row, sqlite3.Row) else row[3]
            mapping_id = row["id"] if isinstance(row, sqlite3.Row) else row[0]
            return {
                "mapping_id": mapping_id,
                "raw_designation_en": raw_designation_en,
                "raw_designation_key": raw_designation_key,
                "raw_text": raw_text,
                "final_text": extract_english_designation(final_text) or str(final_text or "").strip(),
                "match_type": match_type,
                "source": f"dictionary_{match_type.lower()}",
                "confidence": 1.0,
            }
    return None

def apply_designation_dictionary(db, raw_designation_full):
    parts = extract_designation_parts(raw_designation_full)
    raw_designation_en = parts["raw_designation_en"]
    raw_designation_key = parts["raw_designation_key"]
    if not raw_designation_key:
        return {
            "final_en": "Needs Review",
            "final_hi": "",
            "final_designation": "Needs Review",
            "raw_designation_en": raw_designation_en,
            "raw_designation_hi": parts["raw_designation_hi"],
            "raw_designation_key": raw_designation_key,
            "source": "needs_review",
            "confidence": 0.0,
            "needs_review": True,
            "suggested_value": "",
        }
    match = match_designation_dictionary(db, raw_designation_full)
    if match:
        final_text = match["final_text"]
        return {
            **match,
            "final_en": final_text,
            "final_hi": "",
            "final_designation": final_text,
            "needs_review": False,
            "suggested_value": final_text,
        }
    return None

def resolve_designation_v3(db, raw_designation):
    parts = extract_designation_parts(raw_designation)
    raw_designation_en = parts["raw_designation_en"]
    raw_designation_key = parts["raw_designation_key"]
    if not raw_designation_key:
        return {
            "final_en": "Needs Review",
            "final_hi": "",
            "final_designation": "Needs Review",
            "raw_designation_en": raw_designation_en,
            "raw_designation_hi": parts["raw_designation_hi"],
            "raw_designation_key": raw_designation_key,
            "source": "needs_review",
            "confidence": 0.0,
            "needs_review": True,
            "suggested_value": "",
        }

    dictionary_match = apply_designation_dictionary(db, raw_designation)
    if dictionary_match:
        return dictionary_match

    system = DictionaryManager.resolve_system_designation(raw_designation_en)
    if system:
        system["suggested_value"] = system["final_en"]
        system["final_designation"] = system["final_en"]
        system["raw_designation_en"] = raw_designation_en
        system["raw_designation_hi"] = parts["raw_designation_hi"]
        system["raw_designation_key"] = raw_designation_key
        return system

    return {
        "final_en": "Needs Review",
        "final_hi": "",
        "final_designation": "Needs Review",
        "raw_designation_en": raw_designation_en,
        "raw_designation_hi": parts["raw_designation_hi"],
        "raw_designation_key": raw_designation_key,
        "source": "needs_review",
        "confidence": 0.0,
        "needs_review": True,
        "suggested_value": "",
    }

def audit_log_v3(db, action_type, table_name="", record_id=None, field_name="", old_value="", new_value="",
                 source_file="", page_no=None, sr_no=None, employee_name="", employee_id="", applied_scope="", user_note=""):
    db.execute("""
        INSERT INTO edit_audit_log_v3 (
            action_type, table_name, record_id, field_name, old_value, new_value,
            source_file, page_no, sr_no, employee_name, employee_id, applied_scope, user_note
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (
        action_type, table_name, record_id, field_name, str(old_value or ""), str(new_value or ""),
        source_file, page_no, sr_no, employee_name, employee_id, applied_scope, user_note
    ))

def refresh_employee_designation_status(db, employee_id):
    open_reviews = db.execute("""
        SELECT COUNT(*) FROM review_queue_v3
        WHERE employee_id = ? AND COALESCE(status, 'OPEN') != 'RESOLVED'
    """, (employee_id,)).fetchone()[0]
    open_high = db.execute("""
        SELECT COUNT(*) FROM validation_issues_v2
        WHERE employee_id = ? AND COALESCE(status, 'OPEN') != 'RESOLVED'
          AND (issue_type = 'UNCONFIRMED_DESIGNATION' OR UPPER(COALESCE(severity,'')) IN ('HIGH','CRITICAL'))
    """, (employee_id,)).fetchone()[0]
    status = "NEEDS_REVIEW" if open_reviews or open_high else "VERIFIED"
    db.execute("""
        UPDATE employees_v2
        SET designation_needs_review = ?, employee_data_status = ?, row_verification_status = ?
        WHERE id = ?
    """, (1 if open_reviews else 0, status, status, employee_id))
    return status

def employee_manual_edit_exists_sql():
    return """EXISTS (
        SELECT 1 FROM edit_audit_log_v3 a
        WHERE a.table_name IN ('employees_v2', 'employee_components_v2')
          AND CAST(a.record_id AS TEXT) = CAST(e.id AS TEXT)
          AND a.action_type IN ('employee_edit', 'employee_component_edit')
    )"""

def employee_status_case_sql():
    missing = """(COALESCE(e.raw_name,'') = '' OR COALESCE(e.employee_id,'') = '' OR COALESCE(e.pay_scale_raw,'') = '')"""
    open_unconfirmed = """EXISTS (
        SELECT 1 FROM validation_issues_v2 v
        WHERE v.employee_id = e.id AND COALESCE(v.status, 'OPEN') != 'RESOLVED'
          AND v.issue_type = 'UNCONFIRMED_DESIGNATION'
    )"""
    open_high = """EXISTS (
        SELECT 1 FROM validation_issues_v2 v
        WHERE v.employee_id = e.id AND COALESCE(v.status, 'OPEN') != 'RESOLVED'
          AND UPPER(COALESCE(v.severity,'')) IN ('HIGH','CRITICAL')
    )"""
    open_component = """EXISTS (
        SELECT 1 FROM validation_issues_v2 v
        WHERE v.employee_id = e.id AND COALESCE(v.status, 'OPEN') != 'RESOLVED'
          AND v.issue_type = 'COMPONENT_EXTRACTION_MISMATCH'
    )"""
    open_any = """EXISTS (
        SELECT 1 FROM validation_issues_v2 v
        WHERE v.employee_id = e.id AND COALESCE(v.status, 'OPEN') != 'RESOLVED'
    )"""
    manual = employee_manual_edit_exists_sql()
    return f"""CASE
        WHEN {missing} THEN 'MISSING_CRITICAL_FIELD'
        WHEN {open_component} THEN 'COMPONENT_REVIEW_REQUIRED'
        WHEN COALESCE(e.designation_needs_review,0) = 1 OR {open_unconfirmed} OR {open_high} THEN 'NEEDS_REVIEW'
        WHEN {manual} THEN 'MANUALLY_EDITED'
        WHEN {open_any} THEN 'VERIFIED_WITH_WARNINGS'
        ELSE 'VERIFIED'
    END"""

def recalculate_employee_data_status(db, employee_id):
    row = db.execute("""
        SELECT raw_name, employee_id, pan_raw, pay_scale_raw, gross_amount, sum_deduction,
               net_total, designation_needs_review
        FROM employees_v2
        WHERE id = ?
    """, (employee_id,)).fetchone()
    if not row:
        return "NEEDS_REVIEW"
    missing_critical = not (row["raw_name"] and row["employee_id"] and row["pay_scale_raw"])
    open_high = db.execute("""
        SELECT COUNT(*)
        FROM validation_issues_v2
        WHERE employee_id = ?
          AND COALESCE(status, 'OPEN') != 'RESOLVED'
          AND UPPER(COALESCE(severity,'')) IN ('HIGH','CRITICAL')
    """, (employee_id,)).fetchone()[0]
    open_component = db.execute("""
        SELECT COUNT(*)
        FROM validation_issues_v2
        WHERE employee_id = ?
          AND COALESCE(status, 'OPEN') != 'RESOLVED'
          AND issue_type = 'COMPONENT_EXTRACTION_MISMATCH'
    """, (employee_id,)).fetchone()[0]
    open_unconfirmed = db.execute("""
        SELECT COUNT(*)
        FROM validation_issues_v2
        WHERE employee_id = ?
          AND COALESCE(status, 'OPEN') != 'RESOLVED'
          AND issue_type = 'UNCONFIRMED_DESIGNATION'
    """, (employee_id,)).fetchone()[0]
    open_any = db.execute("""
        SELECT COUNT(*)
        FROM validation_issues_v2
        WHERE employee_id = ? AND COALESCE(status, 'OPEN') != 'RESOLVED'
    """, (employee_id,)).fetchone()[0]
    manual_edit = db.execute("""
        SELECT COUNT(*)
        FROM edit_audit_log_v3
        WHERE table_name IN ('employees_v2', 'employee_components_v2')
          AND CAST(record_id AS TEXT) = CAST(? AS TEXT)
          AND action_type IN ('employee_edit', 'employee_component_edit')
    """, (employee_id,)).fetchone()[0]
    if missing_critical:
        status = "MISSING_CRITICAL_FIELD"
    elif open_component:
        status = "COMPONENT_REVIEW_REQUIRED"
    elif row["designation_needs_review"] or open_unconfirmed or open_high:
        status = "NEEDS_REVIEW"
    elif manual_edit:
        status = "MANUALLY_EDITED"
    elif open_any:
        status = "VERIFIED_WITH_WARNINGS"
    else:
        status = "VERIFIED"
    db.execute("""
        UPDATE employees_v2
        SET employee_data_status = ?, row_verification_status = ?
        WHERE id = ?
    """, (status, status, employee_id))
    return status

def _matching_employee_ids_for_designation(db, raw_text, match_type="exact", scope="global", bill_id=None, source_file=None):
    sql = "SELECT id, COALESCE(raw_designation_full, raw_designation, '') AS raw FROM employees_v2 WHERE COALESCE(raw_designation_full, raw_designation, '') != '' AND COALESCE(is_deleted, 0) = 0"
    params = []
    if scope == "pdf" and bill_id is not None:
        sql += " AND bill_id = ?"
        params.append(bill_id)
    if source_file:
        sql += " AND source_filename = ?"
        params.append(source_file)
    rows = db.execute(sql, params).fetchall()
    return [r["id"] for r in rows if _dict_match(r["raw"], raw_text, match_type)]

def apply_designation_to_employees(db, raw_text, final_en, final_hi="", match_type="exact", scope="global", bill_id=None, mapping_id=None):
    if not raw_text:
        return {"matched_count": 0, "updated_count": 0, "resolved_issue_count": 0, "mapping_id": mapping_id}
    ids = _matching_employee_ids_for_designation(db, raw_text, match_type, scope, bill_id)
    if not ids:
        return {"matched_count": 0, "updated_count": 0, "resolved_issue_count": 0, "mapping_id": mapping_id}
    placeholders = ",".join("?" * len(ids))
    source = "user_manual" if scope in ("employee", "pdf") else f"dictionary_{match_type}"
    final_text = extract_english_designation(final_en) or str(final_en or "").strip()
    for emp_id in ids:
        row = db.execute("""
            SELECT COALESCE(raw_designation_full, raw_designation, '') AS raw
            FROM employees_v2
            WHERE id = ?
        """, (emp_id,)).fetchone()
        raw_full = row["raw"] if row else raw_text
        parts = extract_designation_parts(raw_full)
        raw_en = parts["raw_designation_en"]
        db.execute("""
            UPDATE employees_v2
            SET raw_designation_en = ?,
                raw_designation_hi = COALESCE(NULLIF(raw_designation_hi, ''), ?),
                raw_designation_key = ?,
                final_designation = ?,
                final_designation_en = ?,
                final_designation_hi = '',
                final_designation_hi_display = COALESCE(NULLIF(final_designation_hi_display, ''), ?),
                normalized_designation = ?,
                designation_source = ?,
                designation_confidence = 1,
                designation_needs_review = 0
            WHERE id = ?
        """, (
            raw_en,
            parts["raw_designation_hi"],
            parts["raw_designation_key"],
            final_text,
            final_text,
            parts["raw_designation_hi"],
            final_text,
            source,
            emp_id,
        ))
        audit_log_v3(
            db, "dictionary_designation_apply", "employees_v2", emp_id,
            "final_designation", "", final_text,
            applied_scope=scope,
            user_note=f"mapping_id={mapping_id or ''}"
        )
    changed = len(ids)
    unresolved_before = db.execute(f"""
        SELECT COUNT(*) FROM validation_issues_v2
        WHERE issue_type = 'UNCONFIRMED_DESIGNATION'
          AND COALESCE(status, 'OPEN') != 'RESOLVED'
          AND employee_id IN ({placeholders})
    """, ids).fetchone()[0] + db.execute(f"""
        SELECT COUNT(*) FROM review_queue_v3
        WHERE issue_type = 'UNCONFIRMED_DESIGNATION'
          AND COALESCE(status, 'OPEN') != 'RESOLVED'
          AND employee_id IN ({placeholders})
    """, ids).fetchone()[0]
    db.execute(f"""
        UPDATE validation_issues_v2 SET status = 'RESOLVED', resolved_at = CURRENT_TIMESTAMP
        WHERE issue_type = 'UNCONFIRMED_DESIGNATION'
          AND employee_id IN ({placeholders})
    """, ids)
    db.execute(f"""
        UPDATE review_queue_v3 SET status = 'RESOLVED', resolved_at = CURRENT_TIMESTAMP, action = ?
        WHERE issue_type = 'UNCONFIRMED_DESIGNATION'
          AND employee_id IN ({placeholders})
    """, [scope] + ids)
    for emp_id in ids:
        recalculate_employee_data_status(db, emp_id)
    return {
        "matched_count": len(ids),
        "updated_count": changed,
        "resolved_issue_count": unresolved_before,
        "mapping_id": mapping_id,
    }

def default_designation_match_type(raw_text):
    raw = clean_designation_key(raw_text)
    long_officer_markers = ("IPS", "RPS", "JUNIOR SCALE", "SENIOR TIME SCALE", "SENIOR SCALE", "SELECTION SCALE", "SUPER TIME SCALE")
    return "contains" if any(marker in raw for marker in long_officer_markers) or len(raw) > 35 else "exact"

def save_dictionary_mapping_v3(db, raw_text, final_en, final_hi="", match_type="exact", priority=100,
                               category="designation", active=1, created_by="user", notes=""):
    raw_clean = extract_english_designation(raw_text) if category == "designation" else str(raw_text or "").strip()
    raw_key = clean_designation_key(raw_clean) if category == "designation" else str(raw_clean).strip().casefold()
    final_text = extract_english_designation(final_en) if category == "designation" else str(final_en or "").strip()
    match_type = (match_type or "exact").lower()
    active = int(active)
    existing = db.execute("""
        SELECT id, active FROM dictionary_mappings_v3
        WHERE category = ?
          AND COALESCE(raw_text_key, '') = ?
          AND COALESCE(final_text, final_text_en, '') = ?
          AND COALESCE(match_type, 'exact') = ?
          AND active = ?
        ORDER BY id DESC
        LIMIT 1
    """, (category, raw_key, final_text, match_type, active)).fetchone()
    if existing:
        return existing["id"]
    inactive = db.execute("""
        SELECT id FROM dictionary_mappings_v3
        WHERE category = ?
          AND COALESCE(raw_text_key, '') = ?
          AND COALESCE(final_text, final_text_en, '') = ?
          AND COALESCE(match_type, 'exact') = ?
          AND active = 0
        ORDER BY id DESC
        LIMIT 1
    """, (category, raw_key, final_text, match_type)).fetchone()
    if active and inactive:
        mapping_id = inactive["id"]
        db.execute("""
            UPDATE dictionary_mappings_v3
            SET raw_text = ?, raw_text_key = ?, final_text = ?, final_text_en = ?,
                final_text_hi = '', priority = ?, active = 1, created_by = ?,
                notes = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
        """, (raw_clean, raw_key, final_text, final_text, int(priority or 0), created_by, notes, mapping_id))
        audit_log_v3(db, "designation_dictionary_reactivate", "dictionary_mappings_v3", mapping_id, "active", 0, 1, applied_scope="global")
        return mapping_id
    db.execute("""
        INSERT INTO dictionary_mappings_v3 (
            category, raw_text, raw_text_key, final_text, final_text_en, final_text_hi,
            match_type, priority, active, created_by, notes
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?)
    """, (category, raw_clean, raw_key, final_text, final_text, "", match_type, int(priority or 0), active, created_by, notes))
    mapping_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
    audit_log_v3(db, "designation_dictionary_add", "dictionary_mappings_v3", mapping_id, "raw_text", "", raw_clean, applied_scope="global")
    return mapping_id

def process_file_v2(fpath, fname, app_config):
    import sqlite3
    # This runs in background
    try:
        db = sqlite3.connect(app_config['DATABASE'])
        db.row_factory = sqlite3.Row
        db.execute("""
            UPDATE upload_jobs_v2
            SET status = 'PROCESSING',
                current_stage = 'extracting text',
                message = 'Extracting PDF text',
                percent_complete = 2,
                started_at = COALESCE(started_at, CURRENT_TIMESTAMP),
                error_message = NULL
            WHERE source_filename = ? AND status IN ('QUEUED', 'PROCESSING')
        """, (fname,))
        db.commit()
        
        from ifms_engine.pdf_reader import PDFReader
        from ifms_engine.header_parser import HeaderParser
        from ifms_engine.employee_block_detector import EmployeeBlockDetector
        from ifms_engine.employee_parser import EmployeeParser
        from ifms_engine.validator import EngineValidator
        from ifms_engine.models import ValidationIssue
        
        reader = PDFReader(fpath)
        words = reader.extract_words()
        db.execute("""
            UPDATE upload_jobs_v2
            SET current_stage = 'detecting rows',
                message = 'Detecting employee rows',
                percent_complete = 10,
                total_pages = ?
            WHERE source_filename = ? AND status = 'PROCESSING'
        """, (reader.get_page_count(), fname))
        db.commit()
        meta = HeaderParser.extract_metadata(words)
        blocks = EmployeeBlockDetector.detect_blocks(words)
        db.execute("""
            UPDATE upload_jobs_v2
            SET current_stage = 'parsing employees',
                message = 'Parsing employee rows',
                detected_employees = ?,
                percent_complete = 18
            WHERE source_filename = ? AND status = 'PROCESSING'
        """, (len(blocks), fname))
        db.commit()
        
        db.execute('''INSERT INTO bills_v2 (
            source_filename, group_name, month_year, bill_no, reference_no, bill_date,
            office_id, ddo_name, ddo_code, demand_no, object_head, budget_head_full,
            sf_amount, ca_amount, na_amount, tan_no, page_count,
            official_gross, official_deduction, official_net
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                   (
                       fname,
                       meta.get('group_name', ''),
                       meta.get('month_year', ''),
                       meta.get('bill_no', ''),
                       meta.get('reference_no', ''),
                       meta.get('bill_date', ''),
                       meta.get('office_id', ''),
                       meta.get('ddo_name', ''),
                       meta.get('ddo_code', ''),
                       meta.get('demand_no', ''),
                       meta.get('object_head', ''),
                       meta.get('budget_head_full', ''),
                       meta.get('sf_amount', 0.0),
                       meta.get('ca_amount', 0.0),
                       meta.get('na_amount', 0.0),
                       meta.get('tan_no', ''),
                       meta.get('page_count', reader.get_page_count()),
                       meta.get('official_gross', 0.0),
                       meta.get('official_deduction', 0.0),
                       meta.get('official_net', 0.0),
                   ))
        bill_id = db.execute('SELECT last_insert_rowid()').fetchone()[0]
        db.execute(
            "UPDATE upload_jobs_v2 SET total_pages = ?, processed_pages = 0 WHERE source_filename = ? AND status = 'PROCESSING'",
            (reader.get_page_count(), fname)
        )
        
        parsed_gross, parsed_ded, parsed_net = 0, 0, 0
        review_issue_count = 0
        warning_issue_count = 0
        saved_count = 0
        total_blocks = len(blocks)
        for i, block in enumerate(blocks):
            emp = EmployeeParser.parse_employee(block)
            designation_parts = extract_designation_parts(emp.raw_designation)
            designation_resolution = resolve_designation_v3(db, emp.raw_designation)
            emp.normalized_designation = designation_resolution["final_en"] or "Needs Review"
            if designation_resolution["needs_review"]:
                review_issue_count += 1
                emp.issues.append(ValidationIssue(
                    issue_type='UNCONFIRMED_DESIGNATION',
                    severity='HIGH',
                    message='Designation is not confirmed by dictionary or high-confidence system rules.',
                    raw_value=emp.raw_designation,
                    employee_id=emp.employee_id,
                    suggested_value=designation_resolution.get("suggested_value", "")
                ))
            EngineValidator.validate_employee(emp)

            missing_critical = []
            for field_name, field_value in [
                ('employee_name', emp.raw_name),
                ('employee_id', emp.employee_id),
                ('pan', emp.pan_raw),
                ('pay_scale', emp.pay_scale_raw),
                ('printed_gross', emp.gross_amount),
                ('printed_deduction', emp.sum_deduction),
                ('printed_net', emp.net_total),
            ]:
                if not field_value:
                    missing_critical.append(field_name)
            for field_name in missing_critical:
                issue_type = {
                    'employee_id': 'MISSING_EMPLOYEE_ID',
                    'pay_scale': 'MISSING_PAY_SCALE',
                    'pan': 'MISSING_PAN',
                }.get(field_name, 'MISSING_EMPLOYEE_MASTER_FIELD')
                emp.issues.append(ValidationIssue(
                    issue_type=issue_type,
                    severity='HIGH',
                    message=f'Employee master field is missing: {field_name}',
                    raw_value=emp.raw_text_block,
                    employee_id=emp.employee_id,
                    suggested_value=f'Review {field_name} extraction for this row'
                ))
            warning_issue_count += sum(1 for issue in emp.issues if issue.severity not in ('HIGH', 'CRITICAL'))
            employee_data_status = 'NEEDS_REVIEW' if designation_resolution["needs_review"] or any(
                issue.severity in ('HIGH', 'CRITICAL') or issue.issue_type in ('MISSING_EMPLOYEE_MASTER_FIELD', 'PRINTED_TOTAL_MISMATCH')
                for issue in emp.issues
            ) else 'VERIFIED'
            
            db.execute("""INSERT INTO employees_v2 (
                bill_id, source_filename, page_no, sr_no, group_name, month_year,
                raw_name, normalized_name, raw_designation, raw_designation_full,
                raw_designation_en, raw_designation_hi, raw_designation_key,
                normalized_designation, final_designation, final_designation_en, final_designation_hi, final_designation_hi_display,
                designation_display_language, designation_confidence, designation_source,
                designation_needs_review, designation_mapping_status, employee_data_status,
                employee_id, gpf_pran_raw, dob, belt_no, pan_raw, pay_scale_raw, bank_acc, aadhar_raw,
                gross_amount, sum_deduction, net_total, row_verification_status, raw_text_block)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (bill_id, fname, emp.page_no, emp.sr_no, meta['group_name'], meta['month_year'],
                 emp.raw_name, emp.raw_name, emp.raw_designation, emp.raw_designation,
                 designation_parts["raw_designation_en"],
                 designation_parts["raw_designation_hi"],
                 designation_parts["raw_designation_key"],
                 emp.normalized_designation, designation_resolution.get("final_designation", designation_resolution["final_en"]),
                 designation_resolution["final_en"], "", designation_parts["raw_designation_hi"],
                 'English', designation_resolution["confidence"], designation_resolution["source"],
                 1 if designation_resolution["needs_review"] else 0, designation_resolution["source"], employee_data_status,
                 emp.employee_id, emp.gpf_pran_raw, emp.dob, emp.belt_no, emp.pan_raw, emp.pay_scale_raw, emp.bank_acc, emp.aadhar_raw,
                 emp.gross_amount, emp.sum_deduction, emp.net_total, employee_data_status, emp.raw_text_block)
            )
            emp_id = db.execute('SELECT last_insert_rowid()').fetchone()[0]
            saved_count += 1
            if designation_resolution["needs_review"]:
                db.execute("""INSERT INTO review_queue_v3 (
                    bill_id, employee_id, source_file, page_no, sr_no, employee_name, employee_id_text,
                    field_name, raw_value, suggested_value, confidence, issue_type, status
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        bill_id, emp_id, fname, emp.page_no, emp.sr_no, emp.raw_name, emp.employee_id,
                        'final_designation_en', emp.raw_designation, designation_resolution.get("suggested_value", ""),
                        designation_resolution["confidence"], 'UNCONFIRMED_DESIGNATION', 'OPEN'
                    )
                )
            
            parsed_gross += emp.gross_amount
            parsed_ded += emp.sum_deduction
            parsed_net += emp.net_total
            
            for comp in emp.components:
                db.execute("""INSERT INTO employee_components_v2 (
                    employee_id, bill_id, component_type, raw_head, normalized_head, amount, source_type, source_page, source_y)
                    VALUES (?,?,?,?,?,?,?,?,?)""",
                    (emp_id, bill_id, comp.component_type, comp.raw_head, comp.normalized_head, comp.amount, "parsed", comp.source_page, comp.source_y)
                )
                
            for issue in emp.issues:
                db.execute("""INSERT INTO validation_issues_v2 (
                    bill_id, employee_id, issue_type, severity, message, raw_value, suggested_value, status)
                    VALUES (?,?,?,?,?,?,?,?)""",
                    (bill_id, emp_id, issue.issue_type, issue.severity, issue.message, issue.raw_value, issue.suggested_value, 'OPEN')
                )
            # Update job status every 10 blocks to show progress
            if i % 10 == 0:
                percent_complete = round(18 + (((i + 1) / total_blocks) * 62), 2) if total_blocks else 80
                db.execute(
                    """UPDATE upload_jobs_v2
                       SET detected_employees = ?,
                           employees_saved = ?,
                           processed_pages = ?,
                           percent_complete = ?,
                           review_issue_count = ?,
                           issue_count = ?,
                           warning_count = ?,
                           current_stage = 'parsing employees',
                           message = ?
                       WHERE source_filename = ? AND status = 'PROCESSING'""",
                    (total_blocks, saved_count, emp.page_no, percent_complete, review_issue_count,
                     review_issue_count + warning_issue_count, warning_issue_count,
                     f'Parsed {saved_count} of {total_blocks} employees', fname)
                )
                db.commit()

        db.execute("""
            UPDATE upload_jobs_v2
            SET current_stage = 'validating totals',
                message = 'Validating bill totals',
                employees_saved = ?,
                processed_pages = ?,
                percent_complete = 88,
                review_issue_count = ?,
                issue_count = ?,
                warning_count = ?
            WHERE source_filename = ? AND status = 'PROCESSING'
        """, (
            saved_count,
            reader.get_page_count(),
            review_issue_count,
            review_issue_count + warning_issue_count,
            warning_issue_count,
            fname,
        ))
        db.commit()
        official_gross = float(meta.get('official_gross', 0) or 0)
        official_ded = float(meta.get('official_deduction', 0) or 0)
        official_net = float(meta.get('official_net', 0) or 0)
        if official_gross > 0 or official_ded > 0 or official_net > 0:
            is_verified = (
                round(parsed_gross) == round(official_gross)
                and round(parsed_ded) == round(official_ded)
                and round(parsed_net) == round(official_net)
            )
            verification_status = 'VERIFIED' if is_verified else 'FAILED_RECONCILIATION'
        else:
            verification_status = 'NEEDS_REVIEW'

        db.execute(
            'UPDATE bills_v2 SET parsed_gross=?, parsed_deduction=?, parsed_net=?, verification_status=? WHERE id=?',
            (parsed_gross, parsed_ded, parsed_net, verification_status, bill_id)
        )
        if verification_status == 'FAILED_RECONCILIATION':
            db.execute(
                """INSERT INTO validation_issues_v2 (
                    bill_id, employee_id, issue_type, severity, message, raw_value, suggested_value, status
                ) VALUES (?,?,?,?,?,?,?,?)""",
                (
                    bill_id,
                    None,
                    'BILL_RECONCILIATION_FAILED',
                    'HIGH',
                    'Bill-level parsed totals do not match official PDF totals.',
                    f'parsed(gross={parsed_gross},ded={parsed_ded},net={parsed_net})',
                    f'official(gross={official_gross},ded={official_ded},net={official_net})',
                    'OPEN'
                )
            )
            review_issue_count += 1

        db.execute("""
            UPDATE upload_jobs_v2
            SET current_stage = 'applying dictionary',
                message = 'Applying dictionary and finalizing audit',
                percent_complete = 95,
                review_issue_count = ?,
                issue_count = ?
            WHERE source_filename = ? AND status = 'PROCESSING'
        """, (review_issue_count, review_issue_count + warning_issue_count, fname))
        
        job_status = 'NEEDS_REVIEW' if review_issue_count else 'COMPLETED'
        db.execute(
            """UPDATE upload_jobs_v2
               SET status = ?,
                   detected_employees = ?,
                   employees_saved = ?,
                   processed_pages = ?,
                   percent_complete = 100,
                   review_issue_count = ?,
                   issue_count = ?,
                   warning_count = ?,
                   current_stage = ?,
                   message = ?,
                   finished_at = CURRENT_TIMESTAMP
               WHERE source_filename = ? AND status = 'PROCESSING'""",
            (
                job_status,
                total_blocks,
                saved_count,
                reader.get_page_count(),
                review_issue_count,
                review_issue_count + warning_issue_count,
                warning_issue_count,
                'completed with warnings' if review_issue_count else 'completed',
                'Completed with review items' if review_issue_count else 'Completed successfully',
                fname,
            )
        )
        db.commit()
        db.close()
        reader.close()
    except Exception as e:
        try:
            db.rollback()
            db.close()
        except Exception:
            pass
        db = sqlite3.connect(app_config['DATABASE'])
        db.execute(
            """UPDATE upload_jobs_v2
               SET status = 'FAILED',
                   current_stage = 'failed',
                   message = ?,
                   error_message = ?,
                   percent_complete = COALESCE(NULLIF(percent_complete, 0), 100),
                   finished_at = CURRENT_TIMESTAMP
               WHERE source_filename = ? AND status IN ('QUEUED', 'PROCESSING')""",
            (str(e), str(e), fname)
        )
        db.commit()
        db.close()
        print(f"Error processing {fname}: {e}")

from flask import Flask, request, jsonify, render_template, send_file, g
from werkzeug.utils import secure_filename
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from parser import parse_inner_sheet, parse_outer_sheet, detect_sheet_type
from schedule_parser import parse_schedule_pdf, get_all_employee_ids, get_schedule_type_name
from ifms_engine.schedule_engine_v3 import (
    SOURCE_TYPE as SCHEDULE_SOURCE_TYPE_V3,
    deactivate_component_mapping,
    ignore_component,
    init_schedule_mapping_schema,
    list_schedule_types as list_schedule_types_v3,
    list_component_mappings,
    query_schedule_entries as query_schedule_entries_v3,
    reactivate_component_mapping,
    reconciliation_rows,
    unignore_component,
    unmapped_components,
    upsert_component_mapping,
)
from ifms_engine.official_schedule_compare import (
    classify_pdf as classify_official_schedule_pdf,
    comparison_rows as official_comparison_rows,
    comparison_summary as official_comparison_summary,
    export_comparison_workbook,
    init_official_schedule_schema,
    official_aggregates,
    run_comparison as run_official_schedule_comparison,
    store_official_upload,
    upload_summary as official_upload_summary,
)
from ifms_engine.monthly_comparison_v3 import (
    delete_run as delete_monthly_comparison_run,
    export_workbook as export_monthly_comparison_workbook,
    init_monthly_comparison_schema,
    list_runs as list_monthly_comparison_runs,
    mark_reviewed as mark_monthly_comparison_reviewed,
    options as monthly_comparison_options,
    results as monthly_comparison_results,
    run_comparison as run_monthly_comparison,
)
from ifms_engine.master_employee_register_v3 import (
    archive_employee as archive_master_employee,
    build_master_register,
    employee_detail as master_employee_detail,
    employees as master_register_employees,
    export_workbook as export_master_register_workbook,
    history as master_register_history,
    init_master_register_schema,
    issues as master_register_issues,
    options as master_register_options,
    resolve_issue as resolve_master_register_issue,
    summary as master_register_summary,
    update_master_register,
)

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = os.path.join(os.path.dirname(__file__), 'uploads')
app.config['OFFICIAL_SCHEDULE_UPLOAD_FOLDER'] = os.path.join(app.config['UPLOAD_FOLDER'], 'official_schedules')
app.config['EXPORT_FOLDER'] = os.path.join(os.path.dirname(__file__), 'exports')
app.config['BACKUP_FOLDER'] = os.path.join(os.path.dirname(__file__), 'backups')
app.config['USE_V2_PARSER'] = True
app.config['DATABASE'] = os.path.join(os.path.dirname(__file__), 'ifms_data.db')
app.config['JSON_AS_ASCII'] = False
app.json.ensure_ascii = False
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['OFFICIAL_SCHEDULE_UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['EXPORT_FOLDER'], exist_ok=True)
os.makedirs(app.config['BACKUP_FOLDER'], exist_ok=True)

# ===================== DATABASE =====================

def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(app.config['DATABASE'])
        g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close_db(exception):
    db = g.pop('db', None)
    if db: db.close()

def init_db():
    db = sqlite3.connect(app.config['DATABASE'])
    db.executescript('''
    CREATE TABLE IF NOT EXISTS uploads (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        filename TEXT, sheet_type TEXT, group_name TEXT,
        month_year TEXT, bill_no TEXT, reference_no TEXT,
        upload_time TEXT, employee_count INTEGER,
        grand_total_gross REAL DEFAULT 0,
        grand_total_deduction REAL DEFAULT 0,
        grand_total_net REAL DEFAULT 0,
        parsed_gross REAL DEFAULT 0,
        parsed_deduction REAL DEFAULT 0,
        parsed_net REAL DEFAULT 0,
        validation_status TEXT DEFAULT ''
    );
    CREATE TABLE IF NOT EXISTS employees (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        upload_id INTEGER, group_name TEXT, month_year TEXT,
        sr_no INTEGER, name TEXT, designation TEXT,
        gpf_pran TEXT, dob TEXT, belt_no TEXT, pan TEXT,
        employee_id TEXT, st_ins_no TEXT, pay_scale TEXT,
        grade_pay TEXT, bank_acc TEXT, aadhar TEXT,
        basic INTEGER DEFAULT 0, da INTEGER DEFAULT 0,
        hra INTEGER DEFAULT 0, cca INTEGER DEFAULT 0,
        hda INTEGER DEFAULT 0, mess_all INTEGER DEFAULT 0,
        rop INTEGER DEFAULT 0, ca INTEGER DEFAULT 0,
        gross_amount REAL DEFAULT 0, sum_deduction REAL DEFAULT 0,
        net_total REAL DEFAULT 0,
        itax INTEGER DEFAULT 0, lic INTEGER DEFAULT 0,
        gpf INTEGER DEFAULT 0, sip INTEGER DEFAULT 0,
        sil INTEGER DEFAULT 0, sii INTEGER DEFAULT 0,
        rghs INTEGER DEFAULT 0, road_ways INTEGER DEFAULT 0,
        advance_salary INTEGER DEFAULT 0, hba_hdfc INTEGER DEFAULT 0,
        all_allowances TEXT, all_deductions TEXT,
        FOREIGN KEY(upload_id) REFERENCES uploads(id)
    );
    CREATE INDEX IF NOT EXISTS idx_emp_group ON employees(group_name);
    CREATE INDEX IF NOT EXISTS idx_emp_name ON employees(name);
    CREATE INDEX IF NOT EXISTS idx_emp_empid ON employees(employee_id);
    CREATE INDEX IF NOT EXISTS idx_emp_month ON employees(month_year);
    
    CREATE TABLE IF NOT EXISTS schedule_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        upload_id INTEGER,
        group_name TEXT,
        month_year TEXT,
        schedule_type TEXT,
        sr_no INTEGER,
        employee_id TEXT,
        name TEXT,
        designation TEXT,
        basic_pay REAL DEFAULT 0,
        pan TEXT DEFAULT '',
        policy_no TEXT DEFAULT '',
        gpf_no TEXT DEFAULT '',
        si_policy_no TEXT DEFAULT '',
        loan_ac_no TEXT DEFAULT '',
        bank_ac_no TEXT DEFAULT '',
        coop_ser_no TEXT DEFAULT '',
        nominee_name TEXT DEFAULT '',
        reason TEXT DEFAULT '',
        income_tax REAL DEFAULT 0,
        sit REAL DEFAULT 0,
        health_edu_cess REAL DEFAULT 0,
        hi_edu_cess REAL DEFAULT 0,
        premium REAL DEFAULT 0,
        gpf_contribution REAL DEFAULT 0,
        gpf_loan REAL DEFAULT 0,
        gpf_2004 REAL DEFAULT 0,
        si_premium REAL DEFAULT 0,
        si_loan REAL DEFAULT 0,
        si_interest REAL DEFAULT 0,
        s_tax_on_si REAL DEFAULT 0,
        rghs REAL DEFAULT 0,
        total REAL DEFAULT 0,
        extra_data TEXT DEFAULT '{}',
        FOREIGN KEY(upload_id) REFERENCES uploads(id)
    );
    CREATE INDEX IF NOT EXISTS idx_sched_type ON schedule_entries(schedule_type);
    CREATE INDEX IF NOT EXISTS idx_sched_empid ON schedule_entries(employee_id);
    CREATE INDEX IF NOT EXISTS idx_sched_group ON schedule_entries(group_name);
    CREATE INDEX IF NOT EXISTS idx_sched_month ON schedule_entries(month_year);
    ''')
    try:
        cols = [row[1] for row in db.execute("PRAGMA table_info(upload_jobs_v2)").fetchall()]
        job_migrations = {
            'error_message': 'TEXT',
            'percent_complete': 'REAL DEFAULT 0',
            'review_issue_count': 'INTEGER DEFAULT 0',
            'current_stage': "TEXT DEFAULT 'queued'",
            'employees_saved': 'INTEGER DEFAULT 0',
            'issue_count': 'INTEGER DEFAULT 0',
            'warning_count': 'INTEGER DEFAULT 0',
        }
        for col, typ in job_migrations.items():
            if col not in cols:
                db.execute(f'ALTER TABLE upload_jobs_v2 ADD COLUMN {col} {typ}')
    except Exception:
        pass
    
    # Migrate: add new columns to existing uploads table if missing
    try:
        cols = [row[1] for row in db.execute("PRAGMA table_info(uploads)").fetchall()]
        migrations = {
            'grand_total_gross': 'REAL DEFAULT 0',
            'grand_total_deduction': 'REAL DEFAULT 0',
            'grand_total_net': 'REAL DEFAULT 0',
            'parsed_gross': 'REAL DEFAULT 0',
            'parsed_deduction': 'REAL DEFAULT 0',
            'parsed_net': 'REAL DEFAULT 0',
            'validation_status': 'TEXT DEFAULT ""',
        }
        for col, typ in migrations.items():
            if col not in cols:
                db.execute(f'ALTER TABLE uploads ADD COLUMN {col} {typ}')
    except Exception:
        pass
    
    # Migrate schedule_entries: add new columns if missing
    try:
        cols = [row[1] for row in db.execute("PRAGMA table_info(schedule_entries)").fetchall()]
        sched_migrations = {
            'bank_ac_no': 'TEXT DEFAULT ""',
            'coop_ser_no': 'TEXT DEFAULT ""',
            'nominee_name': 'TEXT DEFAULT ""',
            'reason': 'TEXT DEFAULT ""',
            'rghs': 'REAL DEFAULT 0',
        }
        for col, typ in sched_migrations.items():
            if col not in cols:
                db.execute(f'ALTER TABLE schedule_entries ADD COLUMN {col} {typ}')
    except Exception:
        pass
    
    db.commit()
    db.close()

init_db()

# V2 schema initialization
from ifms_engine.db import init_v2_schema
init_v2_schema(app.config['DATABASE'])
try:
    _schedule_schema_db = sqlite3.connect(app.config['DATABASE'])
    init_schedule_mapping_schema(_schedule_schema_db)
    init_official_schedule_schema(_schedule_schema_db)
    init_monthly_comparison_schema(_schedule_schema_db)
    init_master_register_schema(_schedule_schema_db)
    _schedule_schema_db.commit()
    _schedule_schema_db.close()
except Exception:
    pass

def load_dictionary_from_db():
    try:
        db = sqlite3.connect(app.config['DATABASE'])
        terms = db.execute('SELECT dictionary_type, raw_value, normalized_value FROM dictionary_terms_v2 WHERE active=1').fetchall()
        from ifms_engine.dictionary import DictionaryManager
        for t in terms:
            if t[0] == 'designation':
                DictionaryManager.DEFAULT_DESIGNATIONS[t[1]] = t[2]
            elif t[0] == 'head':
                DictionaryManager.DEFAULT_COMPONENTS[t[1]] = t[2]
        db.close()
    except Exception:
        pass

load_dictionary_from_db()

DERIVED_SCHEDULE_COMPONENTS = {
    'income_tax': 'income_tax',
    'lic': 'lic',
    'gpf': 'gpf',
    'gpf_2004': 'gpf_2004',
    'si_premium': 'sip',
    'si_loan': 'sil',
    'rghs': 'rghs',
    'cooperative_roadways': 'roadways_deduction',
    'advance_salary': 'advance_salary',
    'ppss_jdr': 'ppss_jdr',
    'rop': 'rop',
}
DERIVED_COMPONENT_SCHEDULES = {v: k for k, v in DERIVED_SCHEDULE_COMPONENTS.items()}

EMPLOYEE_COMPONENT_FIELDS = {
    'basic_pay': 'basic',
    'da': 'da',
    'hra': 'hra',
    'cca': 'cca',
    'hda': 'hda',
    'mess_allowance': 'mess_allowance',
    'kma_allowance': 'kma_allowance',
    'fta': 'fta',
    'mta': 'mta',
    'special_pay': 'special_pay',
    'wash_allowance': 'wash_allowance',
    'rop': 'rop',
    'lic': 'lic',
    'gpf': 'gpf',
    'gpf_2004': 'gpf_2004',
    'ais_gpf': 'ais_gpf',
    'rghs': 'rghs',
    'sip': 'sip',
    'sil': 'sil',
    'sii': 'sii',
    'income_tax': 'income_tax',
    'roadways_deduction': 'roadways_deduction',
    'advance_salary': 'advance_salary',
    'ppss_jdr': 'ppss_jdr',
    'per_ded_1': 'per_ded_1',
    'per_ded_2': 'per_ded_2',
    'eacpenf': 'eacpenf',
    'giais': 'giais',
    'car_rent': 'car_rent',
    'hfrh': 'hfrh',
}

EMPLOYEE_COMPONENT_ALIASES = {
    'itax': 'income_tax',
    'income_tax': 'income_tax',
    'basic': 'basic_pay',
    'basic_pay': 'basic_pay',
    'gpf2004': 'gpf_2004',
    'gpf_2004': 'gpf_2004',
    'roadways': 'roadways_deduction',
    'roadways_deduction': 'roadways_deduction',
}

ADVANCED_FILTER_COLUMNS = {
    'group_name': ('text', 'e.group_name'),
    'source_file': ('text', 'e.source_filename'),
    'bill_no': ('text', 'b.bill_no'),
    'month_year': ('text', 'e.month_year'),
    'employee_name': ('text', 'e.raw_name'),
    'employee_id': ('text', 'e.employee_id'),
    'pan': ('text', 'e.pan_raw'),
    'pay_scale': ('pay_scale', 'e.pay_scale_raw'),
    'raw_designation_en': ('text', "COALESCE(e.raw_designation_en, '')"),
    'final_designation': ('text', "COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation)"),
    'final_designation_en': ('text', "COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation)"),
    'normalized_designation': ('text', "COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation)"),
    'employee_data_status': ('text', "COALESCE(e.employee_data_status, e.row_verification_status)"),
    'gross': ('number', 'e.gross_amount'),
    'gross_amount': ('number', 'e.gross_amount'),
    'deduction': ('number', 'e.sum_deduction'),
    'total_deduction': ('number', 'e.sum_deduction'),
    'net': ('number', 'e.net_total'),
    'net_total': ('number', 'e.net_total'),
    'basic': ('component', 'basic_pay'),
    'basic_pay': ('component', 'basic_pay'),
    'da': ('component', 'da'),
    'hra': ('component', 'hra'),
    'cca': ('component', 'cca'),
    'hda': ('component', 'hda'),
    'mess_allowance': ('component', 'mess_allowance'),
    'kma_allowance': ('component', 'kma_allowance'),
    'fta': ('component', 'fta'),
    'mta': ('component', 'mta'),
    'special_pay': ('component', 'special_pay'),
    'wash_allowance': ('component', 'wash_allowance'),
    'lic': ('component', 'lic'),
    'itax': ('component', 'income_tax'),
    'income_tax': ('component', 'income_tax'),
    'gpf': ('component', 'gpf'),
    'gpf_2004': ('component', 'gpf_2004'),
    'ais_gpf': ('component', 'ais_gpf'),
    'rghs': ('component', 'rghs'),
    'sip': ('component', 'sip'),
    'sil': ('component', 'sil'),
    'sii': ('component', 'sii'),
    'rop': ('component', 'rop'),
    'ppss_jdr': ('component', 'ppss_jdr'),
    'advance_salary': ('component', 'advance_salary'),
    'roadways_deduction': ('component', 'roadways_deduction'),
    'eacpenf': ('component', 'eacpenf'),
    'giais': ('component', 'giais'),
    'car_rent': ('component', 'car_rent'),
    'hfrh': ('component', 'hfrh'),
}

def _normalize_filter_key(value):
    return re.sub(r'[^a-z0-9_]+', '_', (value or '').strip().lower()).strip('_')

def _component_amount_expr(head):
    return f"""COALESCE((
        SELECT SUM(c.amount) FROM employee_components_v2 c
        WHERE c.employee_id = e.id AND c.normalized_head = '{head}'
    ), 0)"""

def _pay_scale_level_expr(col):
    return f"CAST(REPLACE(REPLACE(REPLACE(UPPER(COALESCE({col},'')), 'L', ''), '-FIXED', ''), 'C', '') AS INTEGER)"

def _add_operator_condition(conditions, params, expr, kind, operator, value, value2=None):
    operator = _normalize_filter_key(operator).replace('_', ' ')
    if operator in ('exists',):
        if kind == 'component':
            conditions.append(f"{expr} != 0")
        else:
            conditions.append(f"COALESCE({expr}, '') != ''")
        return
    if operator in ('not exists', 'not_exists'):
        if kind == 'component':
            conditions.append(f"{expr} = 0")
        else:
            conditions.append(f"COALESCE({expr}, '') = ''")
        return
    if value is None or str(value).strip() == '':
        return
    if kind == 'pay_scale':
        text_expr = f"COALESCE({expr}, '')"
        if operator == 'contains':
            conditions.append(f"{text_expr} LIKE ?")
            params.append(f"%{value}%")
            return
        if operator in ('equals', 'equal'):
            conditions.append(f"TRIM(UPPER({text_expr})) = TRIM(UPPER(?))")
            params.append(value)
            return
        if operator in ('not equals', 'not_equals'):
            conditions.append(f"TRIM(UPPER({text_expr})) != TRIM(UPPER(?))")
            params.append(value)
            return
    if kind in ('number', 'component', 'pay_scale'):
        numeric_expr = _pay_scale_level_expr(expr) if kind == 'pay_scale' else f"COALESCE({expr},0)"
        def to_num(v):
            if kind == 'pay_scale':
                m = re.search(r'\d+', str(v))
                return float(m.group(0)) if m else None
            try:
                return float(v)
            except (TypeError, ValueError):
                return None
        n1 = to_num(value)
        n2 = to_num(value2)
        pay_scale_has_level = f"UPPER(COALESCE({expr},'')) GLOB '*[0-9]*'"
        if operator == 'between' and n1 is not None and n2 is not None:
            lo, hi = sorted([n1, n2])
            conditions.append(f"{pay_scale_has_level} AND {numeric_expr} BETWEEN ? AND ?" if kind == 'pay_scale' else f"{numeric_expr} BETWEEN ? AND ?")
            params.extend([lo, hi])
        elif operator in ('equals', 'equal') and n1 is not None:
            conditions.append(f"{numeric_expr} = ?")
            params.append(n1)
        elif operator in ('not equals', 'not_equals') and n1 is not None:
            conditions.append(f"{numeric_expr} != ?")
            params.append(n1)
        elif operator in ('greater than', 'gt') and n1 is not None:
            conditions.append(f"{pay_scale_has_level} AND {numeric_expr} > ?" if kind == 'pay_scale' else f"{numeric_expr} > ?")
            params.append(n1)
        elif operator in ('greater than or equal', 'gte') and n1 is not None:
            conditions.append(f"{pay_scale_has_level} AND {numeric_expr} >= ?" if kind == 'pay_scale' else f"{numeric_expr} >= ?")
            params.append(n1)
        elif operator in ('less than', 'lt') and n1 is not None:
            conditions.append(f"{pay_scale_has_level} AND {numeric_expr} < ?" if kind == 'pay_scale' else f"{numeric_expr} < ?")
            params.append(n1)
        elif operator in ('less than or equal', 'lte') and n1 is not None:
            conditions.append(f"{pay_scale_has_level} AND {numeric_expr} <= ?" if kind == 'pay_scale' else f"{numeric_expr} <= ?")
            params.append(n1)
        return
    text_expr = f"COALESCE({expr}, '')"
    if operator == 'contains':
        conditions.append(f"{text_expr} LIKE ?")
        params.append(f"%{value}%")
    elif operator in ('equals', 'equal'):
        conditions.append(f"TRIM(LOWER({text_expr})) = TRIM(LOWER(?))")
        params.append(value)
    elif operator in ('not equals', 'not_equals'):
        conditions.append(f"TRIM(LOWER({text_expr})) != TRIM(LOWER(?))")
        params.append(value)
    elif operator == 'between' and value2:
        conditions.append(f"{text_expr} BETWEEN ? AND ?")
        params.extend([value, value2])

def add_advanced_filter_conditions(args, conditions, params):
    raw_filters = args.get('advanced_filters', '').strip()
    filters = []
    if raw_filters:
        try:
            parsed = json.loads(raw_filters)
            if isinstance(parsed, list):
                filters.extend(parsed)
        except json.JSONDecodeError:
            pass
    for i in range(1, 11):
        col = args.get(f'filter_col_{i}', '').strip()
        if not col:
            continue
        filters.append({
            'column': col,
            'operator': args.get(f'filter_op_{i}', ''),
            'value': args.get(f'filter_value_{i}', ''),
            'value2': args.get(f'filter_value2_{i}', ''),
        })
    for flt in filters:
        key = _normalize_filter_key(str(flt.get('column', '')))
        operator = flt.get('operator', '')
        value = flt.get('value', '')
        value2 = flt.get('value2', '')
        config = ADVANCED_FILTER_COLUMNS.get(key)
        if not config:
            continue
        kind, expr = config
        if kind == 'component':
            expr = _component_amount_expr(expr)
        _add_operator_condition(conditions, params, expr, kind, operator, value, value2)

FILTER_CHIP_LABELS = {
    'search': 'Search',
    'group': 'Group',
    'group_name': 'Group',
    'month': 'Month/Year',
    'month_year': 'Month/Year',
    'pay_scale': 'Pay Scale',
    'normalized_designation': 'Final Designation',
    'final_designation': 'Final Designation',
    'final_designation_en': 'Final Designation',
    'employee_data_status': 'Employee Data Status',
    'validation_status': 'Employee Data Status',
    'employee_name': 'Employee Name',
    'employee_id': 'Employee ID',
    'pan': 'PAN',
    'source_file': 'Source File',
    'bill_no': 'Bill No',
    'raw_designation': 'Raw Designation',
    'raw_designation_en': 'Raw English Designation',
    'designation_display': 'Designation Display',
    'basic': 'Basic Pay',
    'basic_pay': 'Basic Pay',
    'gross': 'Gross',
    'gross_amount': 'Gross',
    'net': 'Net Total',
    'net_total': 'Net Total',
    'deduction': 'Total Deduction',
    'total_deduction': 'Total Deduction',
    'itax': 'Income Tax',
    'income_tax': 'Income Tax',
    'gpf': 'GPF',
    'gpf_2004': 'GPF 2004',
    'lic': 'LIC',
    'rghs': 'RGHS',
    'sip': 'SIP',
    'sil': 'SIL',
    'rop': 'ROP',
    'component_head': 'Component',
    'show_deleted': 'Show Deleted Rows',
}

def _filter_label(key):
    normalized = _normalize_filter_key(key)
    return FILTER_CHIP_LABELS.get(normalized, str(key or '').replace('_', ' ').strip().title())

def _parse_advanced_filters(args):
    filters = []
    raw_filters = args.get('advanced_filters', '').strip()
    if raw_filters:
        try:
            parsed = json.loads(raw_filters)
            if isinstance(parsed, list):
                filters.extend([f for f in parsed if isinstance(f, dict)])
        except json.JSONDecodeError:
            pass
    for i in range(1, 11):
        col = args.get(f'filter_col_{i}', '').strip()
        if not col:
            continue
        filters.append({
            'column': col,
            'operator': args.get(f'filter_op_{i}', ''),
            'value': args.get(f'filter_value_{i}', ''),
            'value2': args.get(f'filter_value2_{i}', ''),
        })
    return filters

def build_employee_filter_chips(args):
    chips = []

    def add_chip(label, operator, value='', value2=''):
        value = '' if value is None else str(value)
        value2 = '' if value2 is None else str(value2)
        text_parts = [label]
        if operator:
            text_parts.append(operator)
        if value:
            text_parts.append(value)
        if value2:
            text_parts.extend(['to', value2])
        chips.append({
            'label': label,
            'operator': operator,
            'value': value,
            'value2': value2,
            'text': ' '.join(text_parts),
        })

    simple_keys = [
        'search', 'group', 'group_name', 'month', 'month_year', 'pay_scale',
        'normalized_designation', 'final_designation', 'final_designation_en',
        'employee_data_status', 'validation_status', 'employee_name',
        'employee_id', 'pan', 'source_file', 'bill_no',
        'raw_designation', 'raw_designation_en',
    ]
    for key in simple_keys:
        value = args.get(key, '').strip()
        if value:
            add_chip(_filter_label(key), 'is' if key not in ('search', 'employee_name', 'raw_designation', 'raw_designation_en') else 'contains', value)
    if (args.get('show_deleted') or '').strip().lower() in ('1', 'true', 'yes', 'on'):
        add_chip(_filter_label('show_deleted'), 'enabled')
    display_mode = (args.get('designation_display') or '').strip()
    if display_mode and display_mode.lower() not in ('english',):
        add_chip(_filter_label('designation_display'), 'is', display_mode.replace('_', ' '))

    range_keys = [
        ('basic', 'Basic Pay'), ('gross', 'Gross'), ('net', 'Net Total'),
        ('deduction', 'Total Deduction'), ('itax', 'Income Tax'), ('income_tax', 'Income Tax'),
        ('gpf', 'GPF'), ('gpf_2004', 'GPF 2004'), ('lic', 'LIC'), ('rghs', 'RGHS'),
        ('sip', 'SIP'), ('sil', 'SIL'), ('rop', 'ROP'), ('advance_salary', 'Advance Salary'),
    ]
    for key, label in range_keys:
        min_value = args.get(f'{key}_min', '').strip()
        max_value = args.get(f'{key}_max', '').strip()
        if min_value and max_value:
            add_chip(label, 'between', min_value, max_value)
        elif min_value:
            add_chip(label, '>=', min_value)
        elif max_value:
            add_chip(label, '<=', max_value)

    for key, label in [
        ('has_basic', 'Basic Pay'), ('has_lic', 'LIC'), ('has_income_tax', 'Income Tax'),
        ('has_gpf', 'GPF'), ('has_gpf_2004', 'GPF 2004'), ('has_rghs', 'RGHS'),
        ('has_sip', 'SIP'), ('has_sil', 'SIL'), ('has_rop', 'ROP'),
    ]:
        value = args.get(key, '').strip().lower()
        if value:
            add_chip(label, 'exists' if value in ('1', 'true', 'yes') else 'does not exist')

    component_head = args.get('component_head', '').strip()
    component_operator = args.get('component_operator', '').strip()
    if component_head and component_operator:
        add_chip(_filter_label(component_head), component_operator.replace('_', ' '), args.get('component_value', '').strip(), args.get('component_value2', '').strip())

    for flt in _parse_advanced_filters(args):
        key = flt.get('column', '')
        add_chip(
            _filter_label(key),
            str(flt.get('operator', '')).replace('_', ' '),
            flt.get('value', ''),
            flt.get('value2', ''),
        )
    return chips

def is_valid_derived_schedule_component(normalized_head, amount, raw_text_block=''):
    """Keep derived schedules strict while allowing negative ROP when it is a real row component."""
    if amount is None or amount == 0:
        return False
    if normalized_head == 'rop':
        return bool(re.search(r'\bBasic\s+[-\d,]+(?:\.\d+)?\s+(?:IFMS\s+)?ROP\s+-?\d', raw_text_block or '', re.IGNORECASE))
    return True

def build_employee_filter_conditions(args):
    """Build V2 Employee Data filters for API/export consistency."""
    conditions, params = [], []
    show_deleted = (args.get('show_deleted') or '').strip().lower() in ('1', 'true', 'yes', 'on')
    if not show_deleted:
        conditions.append('COALESCE(e.is_deleted, 0) = 0')
    group = (args.get('group_name') or args.get('group') or '').strip()
    month = (args.get('month_year') or args.get('month') or '').strip()
    simple_filters = [
        ('source_file', 'e.source_filename'),
        ('bill_no', 'b.bill_no'),
        ('normalized_designation', "COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation)"),
        ('final_designation_en', "COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation)"),
        ('final_designation', "COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation)"),
        ('employee_id', 'e.employee_id'),
        ('pan', 'e.pan_raw'),
    ]
    if group:
        conditions.append('TRIM(LOWER(e.group_name)) = TRIM(LOWER(?))')
        params.append(group)
    if month:
        conditions.append('e.month_year = ?')
        params.append(month)
    pay_scale = args.get('pay_scale', '').strip()
    if pay_scale:
        conditions.append('TRIM(UPPER(e.pay_scale_raw)) = TRIM(UPPER(?))')
        params.append(pay_scale)
    for arg, col in simple_filters:
        val = args.get(arg, '').strip()
        if val:
            conditions.append(f'{col} = ?')
            params.append(val)
    raw_designation = args.get('raw_designation', '').strip()
    if raw_designation:
        conditions.append('e.raw_designation LIKE ?')
        params.append(f'%{raw_designation}%')
    raw_designation_en = args.get('raw_designation_en', '').strip()
    if raw_designation_en:
        conditions.append("COALESCE(e.raw_designation_en, '') LIKE ?")
        params.append(f'%{raw_designation_en}%')
    employee_name = args.get('employee_name', '').strip()
    if employee_name:
        conditions.append('e.raw_name LIKE ?')
        params.append(f'%{employee_name}%')
    employee_data_status = (args.get('employee_data_status') or args.get('validation_status') or '').strip()
    if employee_data_status:
        status_sql = employee_status_condition(employee_data_status)
        if status_sql:
            conditions.append(status_sql)
    search = args.get('search', '').strip()
    if search:
        conditions.append('(e.raw_name LIKE ? OR e.normalized_name LIKE ? OR e.employee_id LIKE ? OR e.pan_raw LIKE ? OR e.belt_no LIKE ?)')
        s = f'%{search}%'
        params.extend([s, s, s, s, s])

    numeric_filters = [
        ('basic_min', 'basic_pay', '>='), ('basic_max', 'basic_pay', '<='),
        ('gross_min', 'e.gross_amount', '>='), ('gross_max', 'e.gross_amount', '<='),
        ('net_min', 'e.net_total', '>='), ('net_max', 'e.net_total', '<='),
        ('deduction_min', 'e.sum_deduction', '>='), ('deduction_max', 'e.sum_deduction', '<='),
    ]
    for arg, field, op in numeric_filters:
        value = args.get(arg, '').strip()
        if not value:
            continue
        try:
            amount = float(value)
        except ValueError:
            continue
        if field == 'basic_pay':
            conditions.append("""EXISTS (
                SELECT 1 FROM employee_components_v2 c
                WHERE c.employee_id = e.id AND c.normalized_head = 'basic_pay' AND COALESCE(c.amount,0) """ + op + " ?)")
        else:
            conditions.append(f'COALESCE({field},0) {op} ?')
        params.append(amount)

    for head, api_key in EMPLOYEE_COMPONENT_FIELDS.items():
        for suffix, op in [('_min', '>='), ('_max', '<=')]:
            value = args.get(api_key + suffix, '').strip()
            if not value:
                continue
            try:
                amount = float(value)
            except ValueError:
                continue
            conditions.append(f"""EXISTS (
                SELECT 1 FROM employee_components_v2 c
                WHERE c.employee_id = e.id AND c.normalized_head = ? AND COALESCE(c.amount,0) {op} ?
            )""")
            params.extend([head, amount])

    presence_filters = [
        ('has_basic', 'basic_pay'), ('has_lic', 'lic'), ('has_income_tax', 'income_tax'),
        ('has_gpf', 'gpf'), ('has_gpf_2004', 'gpf_2004'), ('has_rghs', 'rghs'),
        ('has_sip', 'sip'), ('has_sil', 'sil'), ('has_sii', 'sii'), ('has_rop', 'rop'),
        ('has_advance_salary', 'advance_salary'), ('has_ppss_jdr', 'ppss_jdr'),
        ('has_roadways_deduction', 'roadways_deduction')
    ]
    for arg, head in presence_filters:
        value = args.get(arg, '').strip().lower()
        if not value:
            continue
        exists_sql = """EXISTS (
            SELECT 1 FROM employee_components_v2 c
            WHERE c.employee_id = e.id AND c.normalized_head = ? AND COALESCE(c.amount,0) != 0
        )"""
        conditions.append(exists_sql if value in ('1', 'true', 'yes') else f'NOT {exists_sql}')
        params.append(head)

    any_component = args.get('has_any_component', '').strip()
    if any_component:
        head = EMPLOYEE_COMPONENT_ALIASES.get(any_component, any_component)
        conditions.append("""EXISTS (
            SELECT 1 FROM employee_components_v2 c
            WHERE c.employee_id = e.id AND c.normalized_head = ? AND COALESCE(c.amount,0) != 0
        )""")
        params.append(head)

    component_head = args.get('component_head', '').strip()
    component_operator = args.get('component_operator', '').strip().lower()
    component_value = args.get('component_value', '').strip()
    if component_head and component_operator:
        head = EMPLOYEE_COMPONENT_ALIASES.get(component_head, component_head)
        if component_operator in ('exists', 'nonzero'):
            conditions.append("""EXISTS (
                SELECT 1 FROM employee_components_v2 c
                WHERE c.employee_id = e.id AND c.normalized_head = ? AND COALESCE(c.amount,0) != 0
            )""")
            params.append(head)
        elif component_operator == 'not_exists':
            conditions.append("""NOT EXISTS (
                SELECT 1 FROM employee_components_v2 c
                WHERE c.employee_id = e.id AND c.normalized_head = ? AND COALESCE(c.amount,0) != 0
            )""")
            params.append(head)
        elif component_operator in ('equals', 'gt', 'gte', 'lt', 'lte', 'between') and component_value:
            op_map = {'equals': '=', 'gt': '>', 'gte': '>=', 'lt': '<', 'lte': '<='}
            try:
                amount = float(component_value)
            except ValueError:
                amount = None
            if amount is not None:
                if component_operator == 'between':
                    try:
                        amount2 = float(args.get('component_value2', '').strip())
                    except ValueError:
                        amount2 = None
                    if amount2 is not None:
                        lo, hi = sorted([amount, amount2])
                        conditions.append("""EXISTS (
                            SELECT 1 FROM employee_components_v2 c
                            WHERE c.employee_id = e.id AND c.normalized_head = ? AND COALESCE(c.amount,0) BETWEEN ? AND ?
                        )""")
                        params.extend([head, lo, hi])
                else:
                    conditions.append(f"""EXISTS (
                        SELECT 1 FROM employee_components_v2 c
                        WHERE c.employee_id = e.id AND c.normalized_head = ? AND COALESCE(c.amount,0) {op_map[component_operator]} ?
                    )""")
                    params.extend([head, amount])
    add_advanced_filter_conditions(args, conditions, params)
    return conditions, params

def employee_quality_summary(db, where='', params=None):
    params = params or []
    from_clause = 'FROM employees_v2 e LEFT JOIN bills_v2 b ON b.id = e.bill_id'
    status_case = employee_status_case_sql()
    status_rows = db.execute(f"""
        SELECT status, COUNT(*) AS c
        FROM (
            SELECT e.id, {status_case} AS status
            {from_clause}
            {where}
        ) q
        GROUP BY status
    """, params).fetchall()
    counts = {
        'VERIFIED': 0,
        'VERIFIED_WITH_WARNINGS': 0,
        'NEEDS_REVIEW': 0,
        'MANUALLY_EDITED': 0,
        'MISSING_CRITICAL_FIELD': 0,
        'COMPONENT_REVIEW_REQUIRED': 0,
    }
    for row in status_rows:
        counts[row['status']] = row['c']
    extra = db.execute(f"""
        SELECT
          COUNT(e.id) AS total_employees,
          COALESCE(SUM(CASE WHEN COALESCE(e.designation_source,'') LIKE 'dictionary_%' THEN 1 ELSE 0 END),0) AS dictionary_applied_rows,
          COALESCE(SUM((
              SELECT COUNT(*) FROM validation_issues_v2 v
              WHERE v.employee_id = e.id AND COALESCE(v.status, 'OPEN') != 'RESOLVED'
          )),0) AS unresolved_issues
        {from_clause}
        {where}
    """, params).fetchone()
    bill_totals = db.execute(f"""
        SELECT
          COUNT(id) AS total_bills,
          COALESCE(SUM(CASE
            WHEN ROUND(COALESCE(scope_gross, 0)) = ROUND(COALESCE(NULLIF(official_gross, 0), parsed_gross, scope_gross, 0))
             AND ROUND(COALESCE(scope_deduction, 0)) = ROUND(COALESCE(NULLIF(official_deduction, 0), parsed_deduction, scope_deduction, 0))
             AND ROUND(COALESCE(scope_net, 0)) = ROUND(COALESCE(NULLIF(official_net, 0), parsed_net, scope_net, 0))
            THEN 1 ELSE 0 END), 0) AS matched_bills
        FROM (
            SELECT b.id, b.parsed_gross, b.parsed_deduction, b.parsed_net,
                   b.official_gross, b.official_deduction, b.official_net,
                   COALESCE(SUM(e.gross_amount), 0) AS scope_gross,
                   COALESCE(SUM(e.sum_deduction), 0) AS scope_deduction,
                   COALESCE(SUM(e.net_total), 0) AS scope_net
            {from_clause}
            {where}
            GROUP BY b.id
        ) scoped_bills
        WHERE id IS NOT NULL
    """, params).fetchone()
    row_totals = db.execute(f"""
        SELECT
          COUNT(e.id) AS total_rows,
          COALESCE(SUM(CASE
            WHEN ABS((COALESCE(e.gross_amount,0) - COALESCE(e.sum_deduction,0)) - COALESCE(e.net_total,0)) <= 1
            THEN 1 ELSE 0 END), 0) AS matched_rows
        {from_clause}
        {where}
    """, params).fetchone()
    component_totals = db.execute(f"""
        SELECT
          COUNT(id) AS total_rows,
          COALESCE(SUM(CASE
            WHEN ABS(COALESCE(allowance_total,0) - COALESCE(gross_amount,0)) <= 1
             AND ABS(COALESCE(deduction_total,0) - COALESCE(sum_deduction,0)) <= 1
            THEN 1 ELSE 0 END), 0) AS matched_rows
        FROM (
          SELECT e.id, e.gross_amount, e.sum_deduction,
                 (
                   SELECT COALESCE(SUM(c.amount),0)
                   FROM employee_components_v2 c
                   WHERE c.employee_id = e.id AND c.component_type = 'allowance'
                 ) AS allowance_total,
                 (
                   SELECT COALESCE(SUM(c.amount),0)
                   FROM employee_components_v2 c
                   WHERE c.employee_id = e.id
                     AND c.component_type = 'deduction'
                     AND COALESCE(c.included_in_total, 1) = 1
                 ) AS deduction_total
          {from_clause}
          {where}
        ) scoped_rows
    """, params).fetchone()
    bill_total_count = bill_totals['total_bills'] or 0
    bill_match_count = bill_totals['matched_bills'] or 0
    row_total_count = row_totals['total_rows'] or 0
    row_match_count = row_totals['matched_rows'] or 0
    component_total_count = component_totals['total_rows'] or 0
    component_match_count = component_totals['matched_rows'] or 0
    return {
        'total_employees': extra['total_employees'] or 0,
        'verified': counts['VERIFIED'],
        'verified_with_warnings': counts['VERIFIED_WITH_WARNINGS'],
        'needs_review': counts['NEEDS_REVIEW'],
        'manually_edited': counts['MANUALLY_EDITED'],
        'missing_critical_field': counts['MISSING_CRITICAL_FIELD'],
        'component_review_required': counts['COMPONENT_REVIEW_REQUIRED'],
        'unresolved_issues': extra['unresolved_issues'] or 0,
        'dictionary_applied_rows': extra['dictionary_applied_rows'] or 0,
        'bill_totals_matched': bill_total_count == 0 or bill_match_count == bill_total_count,
        'bill_totals_matched_text': 'Yes' if bill_total_count == 0 or bill_match_count == bill_total_count else 'No',
        'bill_totals_matched_count': bill_match_count,
        'bill_totals_total': bill_total_count,
        'employee_row_totals_matched': row_match_count,
        'employee_row_totals_total': row_total_count,
        'component_ledger_matched': component_match_count,
        'component_ledger_total': component_total_count,
        'status_counts': counts,
    }


def _round_amount(value):
    return round(float(value or 0))


def _amounts_match(left, right, tolerance=1):
    return abs(_round_amount(left) - _round_amount(right)) <= tolerance


def _df_from_query(db, sql, params=()):
    return pd.DataFrame([dict(r) for r in db.execute(sql, params).fetchall()])


def _write_workbook(path, sheets):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for sheet_name, df in sheets.items():
            if df is None:
                df = pd.DataFrame()
            if not isinstance(df, pd.DataFrame):
                df = pd.DataFrame(df)
            safe_name = sheet_name[:31]
            df.to_excel(writer, index=False, sheet_name=safe_name)
            ws = writer.sheets[safe_name]
            ws.freeze_panes = 'A2'
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F4E78')
            for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 42)
    return path


def _export_metadata_rows():
    return [
        {'key': 'app_name', 'value': APP_NAME},
        {'key': 'developer_credit', 'value': DEVELOPER_CREDIT},
        {'key': 'developer_role', 'value': 'Senior Assistant, Rajasthan Police'},
        {'key': 'developer_email', 'value': DEVELOPER_EMAIL},
        {'key': 'export_timestamp', 'value': datetime.now().isoformat(timespec='seconds')},
        {'key': 'disclaimer', 'value': APP_DISCLAIMER},
    ]


def build_pdf_audit_rows(db, bill_id=None, source_file=None):
    conditions = []
    params = []
    if bill_id is not None:
        conditions.append('b.id = ?')
        params.append(bill_id)
    if source_file:
        conditions.append('b.source_filename = ?')
        params.append(source_file)
    where = ('WHERE ' + ' AND '.join(conditions)) if conditions else ''
    bills = db.execute(f"""
        SELECT b.*
        FROM bills_v2 b
        {where}
        ORDER BY b.id
    """, params).fetchall()
    job_rows = db.execute("""
        SELECT *
        FROM upload_jobs_v2
        ORDER BY id DESC
    """).fetchall()
    latest_job_by_file = {}
    for job in job_rows:
        latest_job_by_file.setdefault(job['source_filename'], job)

    rows = []
    bill_source_files = set()
    for bill in bills:
        bill_source_files.add(bill['source_filename'])
        job = latest_job_by_file.get(bill['source_filename'])
        active = db.execute("""
            SELECT COUNT(*) AS c,
                   COALESCE(SUM(gross_amount),0) AS gross,
                   COALESCE(SUM(sum_deduction),0) AS deduction,
                   COALESCE(SUM(net_total),0) AS net
            FROM employees_v2
            WHERE bill_id = ? AND COALESCE(is_deleted,0) = 0
        """, (bill['id'],)).fetchone()
        all_rows = db.execute("""
            SELECT COUNT(*) AS c
            FROM employees_v2
            WHERE bill_id = ?
        """, (bill['id'],)).fetchone()
        deleted = db.execute("""
            SELECT COUNT(*) AS c,
                   COALESCE(SUM(gross_amount),0) AS gross,
                   COALESCE(SUM(sum_deduction),0) AS deduction,
                   COALESCE(SUM(net_total),0) AS net
            FROM employees_v2
            WHERE bill_id = ? AND COALESCE(is_deleted,0) = 1
        """, (bill['id'],)).fetchone()
        row_totals = db.execute("""
            SELECT COUNT(*) AS total_rows,
                   COALESCE(SUM(CASE
                     WHEN ABS((COALESCE(gross_amount,0) - COALESCE(sum_deduction,0)) - COALESCE(net_total,0)) <= 1
                     THEN 1 ELSE 0 END), 0) AS matched_rows
            FROM employees_v2
            WHERE bill_id = ? AND COALESCE(is_deleted,0) = 0
        """, (bill['id'],)).fetchone()
        component_totals = db.execute("""
            SELECT COUNT(id) AS total_rows,
                   COALESCE(SUM(CASE
                     WHEN ABS(COALESCE(allowance_total,0) - COALESCE(gross_amount,0)) <= 1
                      AND ABS(COALESCE(deduction_total,0) - COALESCE(sum_deduction,0)) <= 1
                     THEN 1 ELSE 0 END), 0) AS matched_rows
            FROM (
              SELECT e.id, e.gross_amount, e.sum_deduction,
                     (SELECT COALESCE(SUM(c.amount),0)
                      FROM employee_components_v2 c
                      WHERE c.employee_id = e.id AND c.component_type = 'allowance') AS allowance_total,
                     (SELECT COALESCE(SUM(c.amount),0)
                      FROM employee_components_v2 c
                      WHERE c.employee_id = e.id AND c.component_type = 'deduction'
                        AND COALESCE(c.included_in_total, 1) = 1) AS deduction_total
              FROM employees_v2 e
              WHERE e.bill_id = ? AND COALESCE(e.is_deleted,0) = 0
            ) scoped_rows
        """, (bill['id'],)).fetchone()
        missing = db.execute("""
            SELECT
              COALESCE(SUM(CASE WHEN COALESCE(employee_id,'') = '' THEN 1 ELSE 0 END),0) AS missing_employee_id_count,
              COALESCE(SUM(CASE WHEN COALESCE(pay_scale_raw,'') = '' THEN 1 ELSE 0 END),0) AS missing_pay_scale_count,
              COALESCE(SUM(CASE WHEN COALESCE(pan_raw,'') = '' THEN 1 ELSE 0 END),0) AS missing_pan_count
            FROM employees_v2
            WHERE bill_id = ? AND COALESCE(is_deleted,0) = 0
        """, (bill['id'],)).fetchone()
        issue_counts = db.execute("""
            SELECT
              COALESCE(SUM(CASE WHEN issue_type = 'UNCONFIRMED_DESIGNATION' AND COALESCE(status,'OPEN') != 'RESOLVED' THEN 1 ELSE 0 END),0) AS unconfirmed_designation_count,
              COALESCE(SUM(CASE WHEN issue_type = 'UNKNOWN_COMPONENT_HEAD' AND COALESCE(status,'OPEN') != 'RESOLVED' THEN 1 ELSE 0 END),0) AS unknown_component_count,
              COALESCE(SUM(CASE WHEN UPPER(COALESCE(severity,'')) IN ('HIGH','CRITICAL') AND COALESCE(status,'OPEN') != 'RESOLVED' THEN 1 ELSE 0 END),0) AS critical_issue_count,
              COALESCE(SUM(CASE WHEN UPPER(COALESCE(severity,'')) NOT IN ('HIGH','CRITICAL') AND COALESCE(status,'OPEN') != 'RESOLVED' THEN 1 ELSE 0 END),0) AS warning_issue_count,
              COALESCE(SUM(CASE WHEN COALESCE(status,'OPEN') != 'RESOLVED' THEN 1 ELSE 0 END),0) AS unresolved_issue_count
            FROM validation_issues_v2
            WHERE bill_id = ?
        """, (bill['id'],)).fetchone()
        designation_needs_review = db.execute("""
            SELECT COUNT(*) AS c
            FROM employees_v2
            WHERE bill_id = ?
              AND COALESCE(is_deleted,0) = 0
              AND COALESCE(designation_needs_review,0) = 1
        """, (bill['id'],)).fetchone()['c']
        manual_rows = db.execute("""
            SELECT COUNT(DISTINCT e.id) AS c
            FROM employees_v2 e
            LEFT JOIN edit_audit_log_v3 a
              ON a.table_name IN ('employees_v2','employee_components_v2')
             AND CAST(a.record_id AS TEXT) = CAST(e.id AS TEXT)
             AND a.action_type IN ('employee_edit','employee_component_edit','employee_mark_verified')
            WHERE e.bill_id = ?
              AND COALESCE(e.is_deleted,0) = 0
              AND (COALESCE(e.employee_data_status,'') = 'MANUALLY_EDITED'
                   OR a.id IS NOT NULL)
        """, (bill['id'],)).fetchone()['c']
        dictionary_applied = db.execute("""
            SELECT COUNT(*) AS c
            FROM employees_v2
            WHERE bill_id = ?
              AND COALESCE(is_deleted,0) = 0
              AND COALESCE(designation_source,'') LIKE 'dictionary_%'
        """, (bill['id'],)).fetchone()['c']
        official_gross = bill['official_gross'] or 0
        official_deduction = bill['official_deduction'] or 0
        official_net = bill['official_net'] or 0
        compare_gross = official_gross if official_gross else bill['parsed_gross']
        compare_deduction = official_deduction if official_deduction else bill['parsed_deduction']
        compare_net = official_net if official_net else bill['parsed_net']
        bill_total_match = (
            _amounts_match(active['gross'], compare_gross)
            and _amounts_match(active['deduction'], compare_deduction)
            and _amounts_match(active['net'], compare_net)
        )
        row_total_match_count = row_totals['matched_rows'] or 0
        row_total_mismatch_count = (row_totals['total_rows'] or 0) - row_total_match_count
        component_match_count = component_totals['matched_rows'] or 0
        component_mismatch_count = (component_totals['total_rows'] or 0) - component_match_count
        unconfirmed_count = (issue_counts['unconfirmed_designation_count'] or 0) + designation_needs_review
        missing_total = (
            (missing['missing_employee_id_count'] or 0)
            + (missing['missing_pay_scale_count'] or 0)
            + (missing['missing_pan_count'] or 0)
        )
        job_failed = bool(job and job['status'] == 'FAILED')
        if job_failed:
            status = 'FAILED'
            recommendation = 'FAILED_DO_NOT_USE'
        elif missing_total:
            status = 'NEEDS_REVIEW'
            recommendation = 'REVIEW_MISSING_FIELDS'
        elif component_mismatch_count or not bill_total_match:
            status = 'NEEDS_REVIEW'
            recommendation = 'REVIEW_COMPONENTS'
        elif unconfirmed_count:
            status = 'NEEDS_REVIEW'
            recommendation = 'REVIEW_DESIGNATIONS'
        elif issue_counts['critical_issue_count']:
            status = 'NEEDS_REVIEW'
            recommendation = 'REVIEW_MISSING_FIELDS'
        elif issue_counts['warning_issue_count'] or manual_rows or deleted['c']:
            status = 'VERIFIED_WITH_WARNINGS'
            recommendation = 'READY_FOR_MANUAL_REVIEW'
        else:
            status = 'VERIFIED'
            recommendation = 'READY_FOR_MANUAL_REVIEW'
        rows.append({
            'bill_id': bill['id'],
            'source_file': bill['source_filename'] or '',
            'bill_no': bill['bill_no'] or '',
            'group_name': bill['group_name'] or '',
            'month_year': bill['month_year'] or '',
            'pages_total': bill['page_count'] or (job['total_pages'] if job else 0) or 0,
            'employees_detected': (job['detected_employees'] if job else 0) or all_rows['c'] or 0,
            'employees_saved': all_rows['c'] or 0,
            'active_employee_count': active['c'] or 0,
            'official_gross': official_gross,
            'official_deduction': official_deduction,
            'official_net': official_net,
            'parsed_gross': bill['parsed_gross'] or 0,
            'parsed_deduction': bill['parsed_deduction'] or 0,
            'parsed_net': bill['parsed_net'] or 0,
            'active_gross': active['gross'] or 0,
            'active_deduction': active['deduction'] or 0,
            'active_net': active['net'] or 0,
            'deleted_gross': deleted['gross'] or 0,
            'deleted_deduction': deleted['deduction'] or 0,
            'deleted_net': deleted['net'] or 0,
            'bill_total_match': bool(bill_total_match),
            'row_total_match_count': row_total_match_count,
            'row_total_mismatch_count': row_total_mismatch_count,
            'component_match_count': component_match_count,
            'component_mismatch_count': component_mismatch_count,
            'missing_employee_id_count': missing['missing_employee_id_count'] or 0,
            'missing_pay_scale_count': missing['missing_pay_scale_count'] or 0,
            'missing_pan_count': missing['missing_pan_count'] or 0,
            'unconfirmed_designation_count': unconfirmed_count,
            'unknown_component_count': issue_counts['unknown_component_count'] or 0,
            'manually_edited_count': manual_rows or 0,
            'deleted_row_count': deleted['c'] or 0,
            'dictionary_applied_count': dictionary_applied or 0,
            'critical_issue_count': issue_counts['critical_issue_count'] or 0,
            'warning_issue_count': issue_counts['warning_issue_count'] or 0,
            'unresolved_issue_count': issue_counts['unresolved_issue_count'] or 0,
            'job_status': job['status'] if job else '',
            'job_stage': job['current_stage'] if job and 'current_stage' in job.keys() else '',
            'error_message': job['error_message'] if job and 'error_message' in job.keys() else '',
            'status': status,
            'recommendation': recommendation,
        })

    if bill_id is None:
        for job in job_rows:
            if job['source_filename'] in bill_source_files:
                continue
            if source_file and job['source_filename'] != source_file:
                continue
            if job['status'] not in ('FAILED', 'QUEUED', 'PROCESSING'):
                continue
            status = 'FAILED' if job['status'] == 'FAILED' else 'NEEDS_REVIEW'
            rows.append({
                'bill_id': None,
                'source_file': job['source_filename'] or '',
                'bill_no': '',
                'group_name': '',
                'month_year': '',
                'pages_total': job['total_pages'] or 0,
                'employees_detected': job['detected_employees'] or 0,
                'employees_saved': job['employees_saved'] if 'employees_saved' in job.keys() else 0,
                'active_employee_count': 0,
                'official_gross': 0,
                'official_deduction': 0,
                'official_net': 0,
                'parsed_gross': 0,
                'parsed_deduction': 0,
                'parsed_net': 0,
                'active_gross': 0,
                'active_deduction': 0,
                'active_net': 0,
                'deleted_gross': 0,
                'deleted_deduction': 0,
                'deleted_net': 0,
                'bill_total_match': False,
                'row_total_match_count': 0,
                'row_total_mismatch_count': 0,
                'component_match_count': 0,
                'component_mismatch_count': 0,
                'missing_employee_id_count': 0,
                'missing_pay_scale_count': 0,
                'missing_pan_count': 0,
                'unconfirmed_designation_count': 0,
                'unknown_component_count': 0,
                'manually_edited_count': 0,
                'deleted_row_count': 0,
                'dictionary_applied_count': 0,
                'critical_issue_count': 1 if status == 'FAILED' else 0,
                'warning_issue_count': 0,
                'unresolved_issue_count': 1 if status == 'FAILED' else 0,
                'job_status': job['status'],
                'job_stage': job['current_stage'] if 'current_stage' in job.keys() else '',
                'error_message': job['error_message'] if 'error_message' in job.keys() else '',
                'status': status,
                'recommendation': 'FAILED_DO_NOT_USE' if status == 'FAILED' else 'REVIEW_COMPONENTS',
            })
    return rows


def build_phase1_health(db):
    audits = build_pdf_audit_rows(db)
    total_pdfs = len(audits)
    total_employees = db.execute("""
        SELECT COUNT(*)
        FROM employees_v2
        WHERE COALESCE(is_deleted,0) = 0
    """).fetchone()[0]
    deleted_rows = db.execute("""
        SELECT COUNT(*)
        FROM employees_v2
        WHERE COALESCE(is_deleted,0) = 1
    """).fetchone()[0]
    unresolved = db.execute("""
        SELECT COUNT(*)
        FROM validation_issues_v2
        WHERE COALESCE(status,'OPEN') != 'RESOLVED'
    """).fetchone()[0]
    high_unresolved = db.execute("""
        SELECT COUNT(*)
        FROM validation_issues_v2
        WHERE COALESCE(status,'OPEN') != 'RESOLVED'
          AND UPPER(COALESCE(severity,'')) IN ('HIGH','CRITICAL')
    """).fetchone()[0]
    failed_pdfs = sum(1 for row in audits if row['status'] == 'FAILED')
    verified = sum(1 for row in audits if row['status'] == 'VERIFIED')
    verified_warnings = sum(1 for row in audits if row['status'] == 'VERIFIED_WITH_WARNINGS')
    needs_review = sum(1 for row in audits if row['status'] == 'NEEDS_REVIEW')
    missing_critical = sum(
        row['missing_employee_id_count'] + row['missing_pay_scale_count'] + row['missing_pan_count']
        for row in audits
    )
    component_review_rows = sum(row['component_mismatch_count'] for row in audits)
    dictionary_rows = sum(row['dictionary_applied_count'] for row in audits)
    all_bill_totals_match = all(row['bill_total_match'] for row in audits) if audits else False
    if not total_pdfs and not total_employees:
        status_badge = 'NO DATA'
    elif failed_pdfs:
        status_badge = 'FAILED PARSING'
    elif all_bill_totals_match and not missing_critical and not high_unresolved:
        status_badge = 'READY FOR MANUAL REVIEW'
    else:
        status_badge = 'NEEDS REVIEW'
    return {
        'status_badge': status_badge,
        'total_pdfs_uploaded': total_pdfs,
        'total_employees_parsed': total_employees,
        'pdfs_verified': verified,
        'pdfs_verified_with_warnings': verified_warnings,
        'pdfs_needing_review': needs_review,
        'failed_pdfs': failed_pdfs,
        'unresolved_issues': unresolved,
        'high_unresolved_issues': high_unresolved,
        'missing_critical_fields': missing_critical,
        'component_review_rows': component_review_rows,
        'dictionary_applied_rows': dictionary_rows,
        'deleted_rows_hidden': deleted_rows,
        'all_bill_totals_match': all_bill_totals_match,
        'last_refresh_time': datetime.now().isoformat(timespec='seconds'),
    }


def employee_status_condition(status):
    status = (status or '').upper()
    allowed = {
        'VERIFIED',
        'VERIFIED_WITH_WARNINGS',
        'NEEDS_REVIEW',
        'MANUALLY_EDITED',
        'MISSING_CRITICAL_FIELD',
        'COMPONENT_REVIEW_REQUIRED',
    }
    if status not in allowed:
        return None
    return f"({employee_status_case_sql()} = '{status}')"

# ===================== ROUTES =====================

@app.after_request
def add_local_no_cache_headers(response):
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.route('/')
@app.route('/analyzer')
def index():
    return app.response_class(
        render_template(
            'index.html',
            build_label=BUILD_LABEL,
            server_started_at=SERVER_STARTED_AT,
        ),
        content_type='text/html; charset=utf-8',
    )


@app.route('/api/version')
def api_version():
    template_path = os.path.join(os.path.dirname(__file__), 'templates', 'index.html')
    try:
        template_mtime = datetime.fromtimestamp(os.path.getmtime(template_path)).isoformat(timespec='seconds')
    except OSError:
        template_mtime = None
    port = request.host.rsplit(':', 1)[-1] if ':' in request.host else os.environ.get('PORT', '5000')
    return jsonify({
        'app_name': APP_NAME,
        'build_label': BUILD_LABEL,
        'started_at': SERVER_STARTED_AT,
        'python_file_path': os.path.abspath(__file__),
        'template_mtime': template_mtime,
        'port': port,
        'database_path': str(Path(app.config['DATABASE']).resolve()),
        'database_size_bytes': Path(app.config['DATABASE']).stat().st_size if Path(app.config['DATABASE']).exists() else 0,
        'backups_folder_path': str(Path(app.config['BACKUP_FOLDER']).resolve()),
        'exports_folder_path': str(Path(app.config['EXPORT_FOLDER']).resolve()),
        'uploads_folder_path': str(Path(app.config['UPLOAD_FOLDER']).resolve()),
        'python_version': sys.version.split()[0],
        'platform': platform.platform(),
        'current_working_directory': os.getcwd(),
    })


def _table_exists(db, table_name):
    row = db.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
        (table_name,),
    ).fetchone()
    return row is not None


def _safe_count(db, table_name, where_sql="1=1"):
    if not _table_exists(db, table_name):
        return 0
    return db.execute(f"SELECT COUNT(*) FROM {table_name} WHERE {where_sql}").fetchone()[0]


def _table_columns(db, table_name):
    if not _table_exists(db, table_name):
        return set()
    return {row[1] for row in db.execute(f"PRAGMA table_info({table_name})").fetchall()}


def _folder_stats(path):
    folder = Path(path)
    exists = folder.exists() and folder.is_dir()
    total_size = 0
    file_count = 0
    old_file_count = 0
    latest_mtime = None
    if exists:
        now_ts = datetime.now().timestamp()
        for item in folder.rglob('*'):
            if not item.is_file():
                continue
            try:
                stat = item.stat()
            except OSError:
                continue
            file_count += 1
            total_size += stat.st_size
            latest_mtime = stat.st_mtime if latest_mtime is None else max(latest_mtime, stat.st_mtime)
            if now_ts - stat.st_mtime > 30 * 24 * 3600:
                old_file_count += 1
    return {
        'path': str(folder.resolve()) if exists else str(folder),
        'exists': exists,
        'file_count': file_count,
        'total_size_bytes': total_size,
        'total_size_mb': round(total_size / (1024 * 1024), 2),
        'old_file_count': old_file_count,
        'latest_file_mtime': datetime.fromtimestamp(latest_mtime).isoformat(timespec='seconds') if latest_mtime else None,
    }


def _route_exists(rule, method="GET"):
    method = method.upper()
    for item in app.url_map.iter_rules():
        if str(item.rule) == rule and method in item.methods:
            return True
    return False


def _db_file_stats():
    path = Path(app.config['DATABASE'])
    return {
        'database_path': str(path.resolve()),
        'database_exists': path.exists(),
        'database_size_bytes': path.stat().st_size if path.exists() else 0,
        'database_size_mb': round((path.stat().st_size if path.exists() else 0) / (1024 * 1024), 2),
    }


def _run_integrity_check(db):
    try:
        row = db.execute("PRAGMA integrity_check").fetchone()
        result = row[0] if row else ''
        return {'ok': result == 'ok', 'result': result}
    except Exception as exc:
        return {'ok': False, 'result': str(exc)}


def _health_check_payload():
    checked_at = datetime.now().isoformat(timespec='seconds')
    warnings = []
    critical_errors = []
    recommended_actions = []
    warning_count = 0
    critical_count = 0
    db_ok = False
    required_tables_ok = {}
    integrity = {'ok': False, 'result': 'not checked'}
    counts = {
        'employee_count': 0,
        'bill_count': 0,
        'dictionary_count': 0,
        'schedule_type_count': 0,
        'official_schedule_upload_count': 0,
        'monthly_comparison_run_count': 0,
        'master_employee_count': 0,
        'unresolved_issue_count': 0,
        'stale_job_count': 0,
    }
    db_stats = _db_file_stats()
    required_tables = [
        'bills_v2',
        'employees_v2',
        'employee_components_v2',
        'validation_issues_v2',
        'dictionary_mappings_v3',
        'schedule_component_mappings_v3',
        'official_schedule_uploads_v3',
        'monthly_comparison_runs_v3',
        'master_employees_v3',
    ]
    try:
        db = get_db()
        db_ok = True
        integrity = _run_integrity_check(db)
        if not integrity['ok']:
            critical_errors.append('SQLite integrity check failed: ' + str(integrity['result']))
        required_tables_ok = {name: _table_exists(db, name) for name in required_tables}
        missing_core = [name for name in required_tables[:5] if not required_tables_ok.get(name)]
        if missing_core:
            critical_errors.append('Missing required core tables: ' + ', '.join(missing_core))
        optional_missing = [name for name in required_tables[5:] if not required_tables_ok.get(name)]
        if optional_missing:
            warnings.append('Optional Phase 2/3 tables not found yet: ' + ', '.join(optional_missing))
        counts.update({
            'employee_count': _safe_count(db, 'employees_v2', 'COALESCE(is_deleted,0) = 0'),
            'bill_count': _safe_count(db, 'bills_v2'),
            'dictionary_count': _safe_count(db, 'dictionary_mappings_v3'),
            'official_schedule_upload_count': _safe_count(db, 'official_schedule_uploads_v3'),
            'monthly_comparison_run_count': _safe_count(db, 'monthly_comparison_runs_v3'),
            'master_employee_count': _safe_count(db, 'master_employees_v3'),
            'unresolved_issue_count': (
                _safe_count(db, 'validation_issues_v2', "COALESCE(status,'OPEN') NOT IN ('RESOLVED','CLOSED')")
                + _safe_count(db, 'review_queue_v3', "COALESCE(status,'OPEN') NOT IN ('RESOLVED','CLOSED')")
                + _safe_count(db, 'master_employee_issues_v3', "COALESCE(status,'OPEN') NOT IN ('RESOLVED','CLOSED')")
            ),
            'stale_job_count': 0,
        })
        job_columns = _table_columns(db, 'upload_jobs_v2')
        if {'status', 'created_at'}.issubset(job_columns):
            counts['stale_job_count'] = _safe_count(db, 'upload_jobs_v2', "status IN ('QUEUED','PROCESSING') AND datetime(created_at, '+2 hours') < CURRENT_TIMESTAMP")
        elif 'created_at' not in job_columns and _table_exists(db, 'upload_jobs_v2'):
            warnings.append('Upload job table has no created_at column; stale job age check skipped.')
        try:
            counts['schedule_type_count'] = len(list_schedule_types_v3(db, {}))
        except Exception as exc:
            warnings.append('Schedule type count failed: ' + str(exc))
    except Exception as exc:
        db_ok = False
        critical_errors.append('Database connection failed: ' + str(exc))

    upload_stats = _folder_stats(app.config['UPLOAD_FOLDER'])
    export_stats = _folder_stats(app.config['EXPORT_FOLDER'])
    backup_stats = _folder_stats(app.config['BACKUP_FOLDER'])
    if not upload_stats['exists']:
        critical_errors.append('Upload folder is missing.')
    if not export_stats['exists']:
        critical_errors.append('Export folder is missing.')
    if not backup_stats['exists']:
        critical_errors.append('Backup folder is missing.')
    if export_stats['total_size_mb'] > 500:
        warnings.append('Exports folder is large. Review old generated files manually.')
    if backup_stats['total_size_mb'] > 1000:
        warnings.append('Backups folder is large. Move older backups to secure storage if needed.')
    if backup_stats['file_count'] == 0:
        warnings.append('No recent backup found.')
        recommended_actions.append('Create a full backup from Backup & Restore before release use.')
    if upload_stats['old_file_count'] > 0:
        warnings.append('Old upload files exist. Review manually before cleanup.')
    if export_stats['file_count'] > 100:
        warnings.append('Export history has many files. Review exports folder manually.')
    if counts['unresolved_issue_count'] > 100:
        warnings.append('Many unresolved review issues exist.')
        recommended_actions.append('Review Exception Center and Master Issues before official use.')
    if counts['stale_job_count'] > 0:
        warnings.append('Stale upload jobs found.')
        recommended_actions.append('Refresh Upload PDFs and review stale jobs.')

    critical_count = len(critical_errors)
    warning_count = len(warnings)
    if critical_count:
        overall_status = 'CRITICAL'
    elif counts['unresolved_issue_count'] > 100 or counts['stale_job_count'] > 0:
        overall_status = 'NEEDS_REVIEW'
    elif warning_count:
        overall_status = 'HEALTHY_WITH_WARNINGS'
    else:
        overall_status = 'HEALTHY'
    return {
        'overall_status': overall_status,
        'checked_at': checked_at,
        'build_label': BUILD_LABEL,
        'database_ok': db_ok and integrity.get('ok', False),
        'database_integrity': integrity,
        'required_tables_ok': required_tables_ok,
        'upload_folder_ok': upload_stats['exists'],
        'export_folder_ok': export_stats['exists'],
        'backup_folder_ok': backup_stats['exists'],
        **counts,
        'backup_count': backup_stats['file_count'],
        'export_file_count': export_stats['file_count'],
        'upload_file_count': upload_stats['file_count'],
        'folder_sizes': {
            'uploads': upload_stats,
            'exports': export_stats,
            'backups': backup_stats,
        },
        **db_stats,
        'backup_folder_path': backup_stats['path'],
        'exports_folder_path': export_stats['path'],
        'uploads_folder_path': upload_stats['path'],
        'warning_count': warning_count,
        'critical_count': critical_count,
        'warnings': warnings,
        'critical_errors': critical_errors,
        'recommended_actions': recommended_actions,
    }


@app.route('/api/v3/smoke_status')
def v3_smoke_status():
    warnings = []
    db = get_db()
    required_tables = [
        'bills_v2',
        'employees_v2',
        'employee_components_v2',
        'validation_issues_v2',
        'dictionary_mappings_v3',
        'review_queue_v3',
        'schedule_component_mappings_v3',
    ]
    table_status = {name: _table_exists(db, name) for name in required_tables}
    missing_tables = [name for name, ok in table_status.items() if not ok]
    if missing_tables:
        warnings.append('Missing tables: ' + ', '.join(missing_tables))
    version_ok = BUILD_LABEL == 'Phase3-Full-Stable'
    if not version_ok:
        warnings.append('Unexpected build label')
    employee_count = _safe_count(db, 'employees_v2')
    bill_count = _safe_count(db, 'bills_v2')
    dictionary_count = _safe_count(db, 'dictionary_mappings_v3')
    unresolved_issues_count = _safe_count(db, 'validation_issues_v2', "COALESCE(status,'OPEN') != 'RESOLVED'")
    stale_review_queue_count = _safe_count(
        db,
        'review_queue_v3',
        "COALESCE(status,'OPEN') NOT IN ('RESOLVED','CLOSED') AND datetime(created_at, '+7 days') < CURRENT_TIMESTAMP",
    )
    try:
        schedule_types_count = len(list_schedule_types_v3(db, {}))
    except Exception as exc:
        schedule_types_count = 0
        warnings.append(f'Schedule type check failed: {exc}')
    return jsonify({
        'version_ok': version_ok,
        'db_ok': not missing_tables,
        'required_tables_ok': table_status,
        'employee_count': employee_count,
        'bill_count': bill_count,
        'dictionary_count': dictionary_count,
        'schedule_types_count': schedule_types_count,
        'unresolved_issues_count': unresolved_issues_count,
        'stale_review_queue_count': stale_review_queue_count,
        'last_start_time': SERVER_STARTED_AT,
        'warnings': warnings,
    })


@app.route('/api/v3/health_check')
def v3_health_check():
    return jsonify(_health_check_payload())


@app.route('/api/v3/ui_action_inventory')
def v3_ui_action_inventory():
    sections = {
        'Upload PDFs': [
            ('upload PDF', '/api/upload', 'POST', 'Main upload route for PDFs.'),
            ('list jobs', '/api/v3/smoke_status', 'GET', 'Smoke status includes queue/readiness signals.'),
            ('clear all data', '/api/delete_all', 'POST', 'Clear parsed data action.'),
        ],
        'Employee Data': [
            ('list employees', '/api/employees', 'GET', 'Employee list API.'),
            ('edit employee', '/api/v3/employees/<int:emp_id>', 'GET'),
            ('save employee', '/api/v3/employees/<int:emp_id>/edit', 'POST'),
            ('mark verified', '/api/v3/employees/<int:emp_id>/edit', 'POST'),
            ('delete row', '/api/v3/employees/<int:emp_id>/delete', 'POST'),
            ('restore row', '/api/v3/employees/<int:emp_id>/restore', 'POST'),
            ('view source', '/api/v3/employees/<int:emp_id>/source', 'GET'),
            ('export filtered', '/api/export', 'GET'),
        ],
        'Dictionary': [
            ('add mapping', '/api/v3/dictionary', 'POST'),
            ('edit mapping', '/api/v3/dictionary/<int:mapping_id>', 'PUT'),
            ('apply mapping', '/api/v3/dictionary/apply', 'POST'),
            ('deactivate mapping', '/api/v3/dictionary/deactivate', 'POST'),
            ('delete all dictionary', '/api/v3/dictionary/reset', 'POST'),
            ('clean dictionary', '/api/v3/dictionary/sanitize', 'POST'),
            ('import JSON', '/api/v3/dictionary/import', 'POST'),
            ('export JSON', '/api/v3/dictionary/export', 'GET'),
        ],
        'Schedules': [
            ('schedule types', '/api/v3/schedules/types', 'GET'),
            ('schedule entries', '/api/v3/schedules', 'GET'),
            ('export schedule', '/api/v3/schedules/export', 'POST'),
            ('reconciliation', '/api/v3/schedules/reconciliation', 'GET'),
            ('unmapped components', '/api/v3/schedules/unmapped_components', 'GET'),
            ('map component', '/api/v3/schedules/component_mappings', 'POST'),
            ('ignore component', '/api/v3/schedules/component_mappings/ignore', 'POST'),
        ],
        'Official Schedule Comparison': [
            ('upload official schedule', '/api/v3/official_schedules/upload', 'POST'),
            ('list uploads', '/api/v3/official_schedules/uploads', 'GET'),
            ('list sections', '/api/v3/official_schedules/sections', 'GET'),
            ('list rows', '/api/v3/official_schedules/rows', 'GET'),
            ('compare', '/api/v3/official_schedules/compare', 'POST'),
            ('export', '/api/v3/official_schedules/export', 'POST'),
            ('delete upload', '/api/v3/official_schedules/delete_upload', 'POST'),
            ('mark reviewed', '/api/v3/official_schedules/mark_reviewed', 'POST'),
        ],
        'Monthly Comparison': [
            ('options', '/api/v3/monthly_comparison/options', 'GET'),
            ('run comparison', '/api/v3/monthly_comparison/run', 'POST'),
            ('list runs', '/api/v3/monthly_comparison/runs', 'GET'),
            ('results', '/api/v3/monthly_comparison/results', 'GET'),
            ('export', '/api/v3/monthly_comparison/export', 'POST'),
            ('delete run', '/api/v3/monthly_comparison/delete_run', 'POST'),
            ('mark reviewed', '/api/v3/monthly_comparison/mark_reviewed', 'POST'),
        ],
        'Master Register': [
            ('build', '/api/v3/master_register/build', 'POST'),
            ('update', '/api/v3/master_register/update', 'POST'),
            ('summary', '/api/v3/master_register/summary', 'GET'),
            ('employees', '/api/v3/master_register/employees', 'GET'),
            ('history', '/api/v3/master_register/history/<int:master_id>', 'GET'),
            ('issues', '/api/v3/master_register/issues', 'GET'),
            ('resolve issue', '/api/v3/master_register/issues/<int:issue_id>/resolve', 'POST'),
            ('archive', '/api/v3/master_register/archive/<int:master_id>', 'POST'),
            ('export', '/api/v3/master_register/export', 'POST'),
        ],
        'Backup & Restore': [
            ('create backup', '/api/v3/backup/create', 'POST'),
            ('list backups', '/api/v3/backup/list', 'GET'),
            ('download backup', '/api/v3/backup/download/<path:file_name>', 'GET'),
            ('restore backup', '/api/v3/backup/restore', 'POST'),
            ('delete backup', '/api/v3/backup/delete', 'POST'),
            ('export dictionary', '/api/v3/backup/export_dictionary', 'GET'),
            ('import dictionary', '/api/v3/backup/import_dictionary', 'POST'),
            ('export schedule mappings', '/api/v3/backup/export_schedule_mappings', 'GET'),
            ('import schedule mappings', '/api/v3/backup/import_schedule_mappings', 'POST'),
        ],
        'Export Center': [
            ('history', '/api/v3/export_center/history', 'GET'),
            ('download export', '/api/v3/export_center/download/<path:file_name>', 'GET'),
            ('delete export', '/api/v3/export_center/delete', 'POST'),
            ('review issues export', '/api/v3/export_center/review_issues', 'POST'),
            ('dictionary Excel export', '/api/v3/export_center/dictionary_excel', 'POST'),
            ('schedule mappings Excel export', '/api/v3/export_center/schedule_mappings_excel', 'POST'),
            ('all schedules export', '/api/v3/export_center/all_schedules', 'POST'),
            ('master register export', '/api/v3/export_center/master_register', 'POST'),
        ],
    }
    payload = []
    for section, actions in sections.items():
        payload.append({
            'section': section,
            'actions': [
                {
                    'action_name': name,
                    'name': name,
                    'endpoint': endpoint,
                    'method': method,
                    'exists': _route_exists(endpoint, method),
                    'endpoint_exists': _route_exists(endpoint, method),
                    'notes': notes[0] if notes else '',
                }
                for item in actions
                for name, endpoint, method, *notes in [item]
            ],
        })
    return jsonify({'sections': payload})


def _backup_folder():
    folder = Path(app.config.get('BACKUP_FOLDER') or os.path.join(os.path.dirname(__file__), 'backups'))
    folder.mkdir(parents=True, exist_ok=True)
    return folder.resolve()


def _audit_backup_action(action_type, note="", record_id=None):
    try:
        db = get_db()
        audit_log_v3(
            db, action_type, "backup_center", record_id, "backup", "", "",
            applied_scope="backup_restore", user_note=note,
        )
        db.commit()
    except Exception:
        try:
            get_db().rollback()
        except Exception:
            pass


def _json_from_request_file(default_name="import.json"):
    if request.files:
        uploaded = next(iter(request.files.values()))
        filename = uploaded.filename or default_name
        if not filename.lower().endswith(".json"):
            raise ValueError("Import file must be .json")
        return json.loads(uploaded.read().decode("utf-8"))
    data = request.get_json(silent=True)
    if isinstance(data, dict) and "payload" in data:
        return data["payload"]
    if data is None:
        raise ValueError("JSON payload is required")
    return data


@app.route('/api/v3/backup/create', methods=['POST'])
def v3_backup_create():
    backup = create_full_backup(
        app.config['DATABASE'],
        _backup_folder(),
        APP_NAME,
        BUILD_LABEL,
    )
    _audit_backup_action("BACKUP_CREATED", backup["file_name"])
    return jsonify({
        'success': True,
        'backup': {
            'file_name': backup['file_name'],
            'size_bytes': backup['size_bytes'],
            'manifest': backup['manifest'],
            'download_url': f"/api/v3/backup/download/{backup['file_name']}",
        },
        'backups': list_backups(_backup_folder()),
    })


@app.route('/api/v3/backup/list', methods=['GET'])
def v3_backup_list():
    return jsonify({'backups': list_backups(_backup_folder())})


@app.route('/api/v3/backup/download/<path:file_name>', methods=['GET'])
def v3_backup_download(file_name):
    try:
        path = safe_named_file(_backup_folder(), file_name, ".zip")
        if not path.exists():
            return jsonify({'error': 'Backup not found'}), 404
        return send_file(io.BytesIO(path.read_bytes()), as_attachment=True, download_name=path.name, mimetype="application/zip")
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/backup/delete', methods=['POST'])
def v3_backup_delete():
    data = request.get_json(silent=True) or {}
    try:
        path = safe_named_file(_backup_folder(), data.get('file_name', ''), ".zip")
        if not path.exists():
            return jsonify({'error': 'Backup not found'}), 404
        path.unlink()
        _audit_backup_action("BACKUP_DELETED", path.name)
        return jsonify({'success': True, 'deleted_file': path.name, 'backups': list_backups(_backup_folder())})
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/backup/restore', methods=['POST'])
def v3_backup_restore():
    confirmation = False
    restore_path = None
    temp_zip = None
    try:
        if request.files:
            confirmation = str(request.form.get('confirmation') or request.form.get('confirm') or '').lower() in {'1', 'true', 'yes', 'on'}
            uploaded = next(iter(request.files.values()))
            if not (uploaded.filename or '').lower().endswith('.zip'):
                return jsonify({'error': 'Restore file must be .zip'}), 400
            temp_zip = tempfile.NamedTemporaryFile(delete=False, suffix='.zip')
            uploaded.save(temp_zip.name)
            temp_zip.close()
            restore_path = Path(temp_zip.name)
            restored_from = uploaded.filename
        else:
            data = request.get_json(silent=True) or {}
            confirmation = bool(data.get('confirmation') or data.get('confirm'))
            restore_path = safe_named_file(_backup_folder(), data.get('file_name', ''), ".zip")
            restored_from = restore_path.name
            if not restore_path.exists():
                return jsonify({'error': 'Backup not found'}), 404
        if not confirmation:
            return jsonify({'error': 'Explicit confirmation is required'}), 400

        old_db = g.pop('db', None)
        if old_db:
            old_db.close()
        summary = restore_full_backup(
            app.config['DATABASE'],
            restore_path,
            _backup_folder(),
            APP_NAME,
            BUILD_LABEL,
        )
        summary['restored_from'] = restored_from
        _audit_backup_action("BACKUP_RESTORED", restored_from)
        return jsonify({'success': True, 'restore_summary': summary, 'backups': list_backups(_backup_folder())})
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400
    finally:
        if temp_zip:
            try:
                os.unlink(temp_zip.name)
            except OSError:
                pass


@app.route('/api/v3/backup/export_dictionary', methods=['GET', 'POST'])
def v3_backup_export_dictionary():
    db = get_db()
    payload = export_dictionary_payload(db, BUILD_LABEL)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    fname = f"dictionary_mappings_{timestamp}.json"
    path = safe_named_file(_backup_folder(), fname, ".json")
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    _audit_backup_action("DICTIONARY_EXPORTED", fname)
    return send_file(path, as_attachment=True, download_name=fname, mimetype="application/json")


@app.route('/api/v3/backup/import_dictionary', methods=['POST'])
def v3_backup_import_dictionary():
    try:
        payload = _json_from_request_file("dictionary_mappings.json")
        db = get_db()
        summary = import_dictionary_payload(db, payload, replace=bool((request.get_json(silent=True) or {}).get('replace')))
        _audit_backup_action("DICTIONARY_IMPORTED", json.dumps(summary))
        db.commit()
        return jsonify({'success': True, **summary})
    except (ValueError, json.JSONDecodeError) as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/backup/export_schedule_mappings', methods=['GET', 'POST'])
def v3_backup_export_schedule_mappings():
    db = get_db()
    init_schedule_mapping_schema(db)
    payload = export_schedule_mapping_payload(db, BUILD_LABEL)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    fname = f"schedule_component_mappings_{timestamp}.json"
    path = safe_named_file(_backup_folder(), fname, ".json")
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    _audit_backup_action("SCHEDULE_MAPPINGS_EXPORTED", fname)
    return send_file(path, as_attachment=True, download_name=fname, mimetype="application/json")


@app.route('/api/v3/backup/import_schedule_mappings', methods=['POST'])
def v3_backup_import_schedule_mappings():
    try:
        payload = _json_from_request_file("schedule_component_mappings.json")
        db = get_db()
        init_schedule_mapping_schema(db)
        summary = import_schedule_mapping_payload(db, payload)
        _audit_backup_action("SCHEDULE_MAPPINGS_IMPORTED", json.dumps(summary))
        db.commit()
        return jsonify({'success': True, **summary})
    except (ValueError, json.JSONDecodeError) as exc:
        return jsonify({'error': str(exc)}), 400


# ===================== EXPORT CENTER APIs =====================

EXPORT_CENTER_EXTENSIONS = {'.xlsx', '.json'}


def _export_folder():
    folder = Path(app.config.get('EXPORT_FOLDER') or os.path.join(os.path.dirname(__file__), 'exports'))
    folder.mkdir(parents=True, exist_ok=True)
    return folder.resolve()


def _safe_export_file(file_name):
    raw = str(file_name or '').strip()
    if not raw:
        raise ValueError('file_name is required')
    if raw != Path(raw).name or os.path.isabs(raw) or '..' in Path(raw).parts:
        raise ValueError('Invalid file name')
    if Path(raw).suffix.lower() not in EXPORT_CENTER_EXTENSIONS:
        raise ValueError('Invalid export file extension')
    path = (_export_folder() / raw).resolve()
    if _export_folder() not in path.parents:
        raise ValueError('Invalid export file path')
    return path


def _export_center_metadata(export_type, filters=None, source_scope='current database'):
    filters = filters or {}
    return [
        {'Field': 'Application', 'Value': APP_NAME},
        {'Field': 'Designed and Developed by', 'Value': 'Neeraj Purohit'},
        {'Field': 'Role', 'Value': 'Senior Assistant, Rajasthan Police'},
        {'Field': 'Email', 'Value': DEVELOPER_EMAIL},
        {'Field': 'Export Type', 'Value': export_type},
        {'Field': 'Export Timestamp', 'Value': datetime.now().isoformat(timespec='seconds')},
        {'Field': 'Active Filters', 'Value': json.dumps(filters, ensure_ascii=False, default=str)},
        {'Field': 'Source Data Scope', 'Value': source_scope},
        {'Field': 'App Build Label', 'Value': BUILD_LABEL},
        {'Field': 'Disclaimer', 'Value': APP_DISCLAIMER},
    ]


def _write_export_center_workbook(file_prefix, export_type, sheets, filters=None, source_scope='current database'):
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    safe_prefix = re.sub(r'[^A-Za-z0-9_.-]+', '_', file_prefix).strip('_') or 'Export'
    fname = f"{safe_prefix}_{timestamp}.xlsx"
    path = _safe_export_file(fname)
    all_sheets = {'Export Metadata': pd.DataFrame(_export_center_metadata(export_type, filters, source_scope))}
    all_sheets.update(sheets)
    _write_workbook(path, all_sheets)
    return path


def _send_export_file(path):
    return send_file(
        io.BytesIO(path.read_bytes()),
        as_attachment=True,
        download_name=path.name,
        mimetype='application/octet-stream',
    )


@app.route('/api/v3/export_center/history', methods=['GET'])
def export_center_history():
    rows = []
    for path in sorted(_export_folder().iterdir(), key=lambda p: p.stat().st_mtime if p.exists() else 0, reverse=True):
        if not path.is_file() or path.suffix.lower() not in EXPORT_CENTER_EXTENSIONS:
            continue
        rows.append({
            'export_type': path.stem.split('_')[0],
            'file_name': path.name,
            'created_at': datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec='seconds'),
            'size_bytes': path.stat().st_size,
            'status': 'ready',
            'download_url': f"/api/v3/export_center/download/{path.name}",
        })
    return jsonify({'exports': rows})


@app.route('/api/v3/export_center/download/<path:file_name>', methods=['GET'])
def export_center_download(file_name):
    try:
        path = _safe_export_file(file_name)
        if not path.exists():
            return jsonify({'error': 'Export not found'}), 404
        return _send_export_file(path)
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/export_center/delete', methods=['POST'])
def export_center_delete():
    data = request.get_json(silent=True) or {}
    try:
        path = _safe_export_file(data.get('file_name', ''))
        if not path.exists():
            return jsonify({'error': 'Export not found'}), 404
        path.unlink()
        return jsonify({'success': True, 'deleted_file': path.name, 'exports': export_center_history().get_json()['exports']})
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/export_center/review_issues', methods=['POST'])
def export_center_review_issues():
    db = get_db()
    filters = request.get_json(silent=True) or {}
    validation_issues = _df_from_query(db, """
        SELECT v.*, e.raw_name AS employee_name, e.employee_id AS employee_code, b.source_filename, b.group_name
        FROM validation_issues_v2 v
        LEFT JOIN employees_v2 e ON e.id = v.employee_id
        LEFT JOIN bills_v2 b ON b.id = v.bill_id
        ORDER BY COALESCE(v.status,'OPEN'), v.id
    """)
    review_queue = _df_from_query(db, """
        SELECT *
        FROM review_queue_v3
        ORDER BY COALESCE(status,'OPEN'), id
    """)
    path = _write_export_center_workbook(
        'Review_Issues_Export',
        'Exception / Review Queue Export',
        {'Validation Issues': validation_issues, 'Review Queue': review_queue},
        filters,
        'validation_issues_v2 and review_queue_v3',
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/dictionary_excel', methods=['POST'])
def export_center_dictionary_excel():
    db = get_db()
    rows = _df_from_query(db, """
        SELECT id, category, raw_text, raw_text_key, COALESCE(final_text, final_text_en, '') AS final_text,
               final_text_en, final_text_hi, match_type, priority, active, created_at, updated_at, created_by, notes
        FROM dictionary_mappings_v3
        ORDER BY category, active DESC, priority DESC, raw_text
    """)
    path = _write_export_center_workbook(
        'Dictionary_Mappings_Export',
        'Dictionary Excel Export',
        {'Dictionary Mappings': rows},
        {},
        'dictionary_mappings_v3',
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/schedule_mappings_excel', methods=['POST'])
def export_center_schedule_mappings_excel():
    db = get_db()
    init_schedule_mapping_schema(db)
    rows = _df_from_query(db, """
        SELECT id, component_head, schedule_type, label, active, ignored, created_at, updated_at, notes
        FROM schedule_component_mappings_v3
        ORDER BY component_head, id
    """)
    path = _write_export_center_workbook(
        'Schedule_Mappings_Export',
        'Schedule Mappings Excel Export',
        {'Schedule Component Mappings': rows},
        {},
        'schedule_component_mappings_v3',
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/all_schedules', methods=['POST'])
def export_center_all_schedules():
    db = get_db()
    filters = request.get_json(silent=True) or {}
    sheets = {}
    summary_rows = []
    for item in list_schedule_types_v3(db, filters):
        schedule_key = item.get('key')
        data = query_schedule_entries_v3(db, {**filters, 'schedule_type': schedule_key}, paginate=False)
        entries = data.get('entries', [])
        summary = data.get('summary', {})
        summary_rows.append({
            'schedule_type': schedule_key,
            'label': item.get('label'),
            'entry_count': summary.get('entry_count', 0),
            'total_amount': summary.get('total_amount', 0),
        })
        sheet_name = str(item.get('label') or schedule_key)[:31]
        sheets[sheet_name] = pd.DataFrame(entries)
    sheets = {'Schedule Summary': pd.DataFrame(summary_rows), **sheets}
    path = _write_export_center_workbook(
        'All_Schedules_Export',
        'All Schedule Types Export',
        sheets,
        filters,
        'derived schedule entries from employee component ledger',
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/schedule_reconciliation', methods=['POST'])
def export_center_schedule_reconciliation():
    db = get_db()
    filters = request.get_json(silent=True) or {}
    path = _write_export_center_workbook(
        'Schedule_Reconciliation_Export',
        'Schedule Reconciliation Export',
        {'Reconciliation': pd.DataFrame(reconciliation_rows(db, filters))},
        filters,
        'derived schedule reconciliation',
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/phase1_health', methods=['POST'])
def export_center_phase1_health():
    db = get_db()
    path = _write_export_center_workbook(
        'Phase1_Health_Export',
        'Phase 1 Health Check Export',
        {'Phase 1 Health': pd.DataFrame([build_phase1_health(db)]), 'PDF Audits': pd.DataFrame(build_pdf_audit_rows(db))},
        {},
        'phase 1 health and per-PDF audit',
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/schedule_qa', methods=['POST'])
def export_center_schedule_qa():
    db = get_db()
    filters = request.get_json(silent=True) or {}
    data = query_schedule_entries_v3(db, filters, paginate=False)
    path = _write_export_center_workbook(
        'Schedule_QA_Export',
        'Schedule QA Export',
        {
            'Schedule Summary': pd.DataFrame([data.get('summary', {})]),
            'Reconciliation': pd.DataFrame(reconciliation_rows(db, filters)),
            'Unmapped Components': pd.DataFrame(unmapped_components(db, filters)),
        },
        filters,
        'schedule QA, reconciliation, and unmapped components',
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/deleted_rows', methods=['POST'])
def export_center_deleted_rows():
    db = get_db()
    rows = _df_from_query(db, _employee_snapshot_sql("WHERE COALESCE(e.is_deleted,0) = 1") + " ORDER BY e.source_filename, e.sr_no")
    path = _write_export_center_workbook(
        'Deleted_Rows_Export',
        'Deleted Rows Export',
        {'Deleted Rows': rows},
        {},
        'soft-deleted employee rows',
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/manual_edits_audit', methods=['POST'])
def export_center_manual_edits_audit():
    db = get_db()
    rows = _df_from_query(db, """
        SELECT *
        FROM edit_audit_log_v3
        WHERE action_type IN ('employee_edit','employee_component_edit','employee_mark_verified','EMPLOYEE_ROW_DELETE','EMPLOYEE_ROW_RESTORE','dictionary_designation_apply')
        ORDER BY id DESC
    """)
    path = _write_export_center_workbook(
        'Manual_Edits_Audit_Export',
        'Manual Edits Audit Export',
        {'Manual Edits Audit': rows},
        {},
        'edit_audit_log_v3 employee edit actions',
    )
    return _send_export_file(path)


def format_currency(value):
    try:
        amount = int(round(float(value or 0)))
    except (TypeError, ValueError):
        amount = 0
    sign = "-" if amount < 0 else ""
    digits = str(abs(amount))
    if len(digits) <= 3:
        grouped = digits
    else:
        last3 = digits[-3:]
        prefix = digits[:-3]
        groups = []
        while len(prefix) > 2:
            groups.insert(0, prefix[-2:])
            prefix = prefix[:-2]
        if prefix:
            groups.insert(0, prefix)
        grouped = ",".join(groups + [last3])
    return f"Rs. {sign}{grouped}"


@app.route('/api/v2/analytics', methods=['GET'])
def get_v2_analytics():
    db = get_db()
    
    totals = db.execute('''
        SELECT COUNT(id), SUM(gross_amount), SUM(sum_deduction), SUM(net_total)
        FROM employees_v2
    ''').fetchone()
    
    desig_counts = db.execute('''
        SELECT normalized_designation, COUNT(id)
        FROM employees_v2
        WHERE normalized_designation != ''
        GROUP BY normalized_designation
        ORDER BY COUNT(id) DESC
    ''').fetchall()
    
    components = db.execute('''
        SELECT normalized_head, SUM(amount)
        FROM employee_components_v2
        GROUP BY normalized_head
        ORDER BY SUM(amount) DESC
    ''').fetchall()
    
    return jsonify({
        'totals': {
            'employees': totals[0] or 0,
            'gross': totals[1] or 0,
            'deduction': totals[2] or 0,
            'net': totals[3] or 0
        },
        'designations': [{'name': d[0], 'count': d[1]} for d in desig_counts],
        'components': [{'name': c[0], 'amount': c[1]} for c in components]
    })

@app.route('/api/v2/issues', methods=['GET'])
def get_v2_issues():
    db = get_db()
    issues = db.execute('''
        SELECT v.id, v.bill_id, v.employee_id, v.issue_type, v.severity, v.message, 
               v.raw_value, v.suggested_value, v.status,
               e.raw_name, COALESCE(e.final_designation_en, e.normalized_designation) AS designation,
               b.source_filename, e.page_no, e.sr_no, e.employee_id AS employee_code
        FROM validation_issues_v2 v
        LEFT JOIN employees_v2 e ON v.employee_id = e.id
        LEFT JOIN bills_v2 b ON v.bill_id = b.id
        WHERE COALESCE(v.status, 'OPEN') = 'OPEN'
        ORDER BY CASE WHEN v.issue_type = 'UNCONFIRMED_DESIGNATION' THEN 0 ELSE 1 END, v.id DESC
    ''').fetchall()
    
    result = []
    for r in issues:
        result.append({
            'id': r[0],
            'bill_id': r[1],
            'employee_id': r[2],
            'issue_type': r[3],
            'severity': r[4],
            'message': r[5],
            'raw_value': r[6],
            'suggested_value': r[7],
            'status': r[8],
            'employee_name': r[9],
            'designation': r[10],
            'filename': r[11],
            'source_file': r[11],
            'page_no': r[12] or 0,
            'sr_no': r[13] or 0,
            'employee_code': r[14] or '',
            'field_name': 'final_designation_en' if r[3] == 'UNCONFIRMED_DESIGNATION' else '',
            'confidence': 0 if r[3] == 'UNCONFIRMED_DESIGNATION' else '',
        })
    return jsonify(result)

@app.route('/api/v3/review_queue', methods=['GET'])
def get_v3_review_queue():
    db = get_db()
    conditions = ["COALESCE(v.status, 'OPEN') != 'HIDDEN'"]
    params = []
    issue_type = request.args.get('issue_type', '').strip()
    status = request.args.get('status', '').strip()
    group = request.args.get('group', '').strip()
    source_file = request.args.get('source_file', '').strip()
    employee_search = request.args.get('employee_search', '').strip()
    if issue_type:
        conditions.append("v.issue_type = ?")
        params.append(issue_type)
    if status:
        conditions.append("COALESCE(v.status, 'OPEN') = ?")
        params.append(status)
    if group:
        conditions.append("e.group_name = ?")
        params.append(group)
    if source_file:
        conditions.append("b.source_filename = ?")
        params.append(source_file)
    if employee_search:
        conditions.append("(e.raw_name LIKE ? OR e.employee_id LIKE ?)")
        s = f"%{employee_search}%"
        params.extend([s, s])
    where = " AND ".join(conditions)
    rows = db.execute(f"""
        SELECT v.id, v.bill_id, v.employee_id, v.issue_type, v.severity, v.message,
               v.raw_value, v.suggested_value, COALESCE(v.status, 'OPEN') AS status,
               v.created_at, v.resolved_at,
               e.source_filename, e.page_no, e.sr_no, e.raw_name AS employee_name,
               e.employee_id AS employee_code, e.group_name,
               COALESCE(e.final_designation_en, e.normalized_designation) AS designation,
               COALESCE(e.raw_designation_en, '') AS raw_designation_en,
               COALESCE(e.raw_designation_key, '') AS raw_designation_key,
               e.raw_text_block
        FROM validation_issues_v2 v
        LEFT JOIN employees_v2 e ON e.id = v.employee_id
        LEFT JOIN bills_v2 b ON b.id = v.bill_id
        WHERE {where}
        ORDER BY CASE WHEN v.issue_type = 'UNCONFIRMED_DESIGNATION' THEN 0 ELSE 1 END,
                 v.severity DESC, v.id DESC
    """, params).fetchall()
    field_by_issue = {
        'UNCONFIRMED_DESIGNATION': 'final_designation_en',
        'MISSING_EMPLOYEE_ID': 'employee_id',
        'MISSING_PAY_SCALE': 'pay_scale',
        'MISSING_PAN': 'pan',
        'SUSPICIOUS_PAN': 'pan',
        'UNKNOWN_ALLOWANCE_HEAD': 'allowance_head',
        'UNKNOWN_DEDUCTION_HEAD': 'deduction_head',
        'UNKNOWN_COMPONENT_HEAD': 'component_head',
        'COMPONENT_EXTRACTION_MISMATCH': 'components',
        'FIELD_LOW_CONFIDENCE': 'field',
        'AMOUNT_NEEDS_REVIEW': 'amount',
        'SUMMARY_CONTAMINATION_DETECTED': 'component_head',
    }
    items = []
    for r in rows:
        item = {
            'id': r['id'],
            'employee_db_id': r['employee_id'],
            'source_file': r['source_filename'] or '',
            'page_no': r['page_no'] or 0,
            'sr_no': r['sr_no'] or 0,
            'employee_name': r['employee_name'] or '',
            'employee_id': r['employee_code'] or '',
            'group_name': r['group_name'] or '',
            'designation': r['designation'] or '',
            'raw_designation_en': r['raw_designation_en'] or extract_english_designation(r['raw_value'] or ''),
            'raw_designation_key': r['raw_designation_key'] or clean_designation_key(r['raw_value'] or ''),
            'field_name': field_by_issue.get(r['issue_type'], ''),
            'raw_value': r['raw_value'] or '',
            'suggested_value': r['suggested_value'] or '',
            'confidence': 0 if r['issue_type'] in ('UNCONFIRMED_DESIGNATION', 'SUMMARY_CONTAMINATION_DETECTED') else '',
            'issue_type': r['issue_type'],
            'severity': r['severity'],
            'message': r['message'] or '',
            'status': r['status'],
            'source_block': r['raw_text_block'] or '',
            'actions': [
                'edit_employee_field',
                'add_dictionary_mapping',
                'mark_resolved',
                'ignore_once',
                'ignore_forever',
                'view_source_block',
            ],
            'action_data': {
                'edit_employee_url': f"/api/v3/employees/{r['employee_id']}/edit" if r['employee_id'] else '',
                'resolve_url': '/api/v2/issues/resolve',
                'dictionary_apply_url': '/api/v3/dictionary/apply',
            }
        }
        items.append(item)
    return jsonify({'issues': items, 'total': len(items)})

@app.route('/api/v2/issues/resolve', methods=['POST'])
@app.route('/api/v3/issues/<int:issue_id>/resolve', methods=['POST'])
def resolve_v2_issue(issue_id=None):
    data = request.json or {}
    issue_id = issue_id or data.get('issue_id')
    new_value = data.get('new_value')
    
    db = get_db()
    issue = db.execute('''
        SELECT v.issue_type, v.employee_id, v.raw_value, e.bill_id, e.source_filename,
               e.page_no, e.sr_no, e.raw_name, e.employee_id AS employee_code,
               COALESCE(e.raw_designation_en, '') AS raw_designation_en,
               COALESCE(e.final_designation_en, e.normalized_designation) AS old_designation
        FROM validation_issues_v2 v
        LEFT JOIN employees_v2 e ON e.id = v.employee_id
        WHERE v.id = ?
    ''', (issue_id,)).fetchone()
    if not issue:
        return jsonify({'error': 'Issue not found'}), 404
        
    issue_type = issue[0]
    emp_id = issue[1]
    action = data.get('action', 'employee')

    if action == 'mark_resolved':
        db.execute("UPDATE validation_issues_v2 SET status = 'RESOLVED', resolved_at = CURRENT_TIMESTAMP WHERE id = ?", (issue_id,))
        if emp_id:
            db.execute("""
                UPDATE review_queue_v3
                SET status = 'RESOLVED', action = 'mark_resolved', resolved_at = CURRENT_TIMESTAMP
                WHERE employee_id = ? AND issue_type = ?
            """, (emp_id, issue_type))
            recalculate_employee_data_status(db, emp_id)
        audit_log_v3(
            db, "review_item_mark_resolved", "validation_issues_v2", issue_id, issue_type,
            issue["raw_value"], new_value or "Resolved", issue["source_filename"], issue["page_no"],
            issue["sr_no"], issue["raw_name"], issue["employee_code"], "issue"
        )
        db.commit()
        return jsonify({'status': 'success'})
    
    if issue_type == 'GROSS_MISMATCH':
        db.execute('UPDATE employees_v2 SET gross_amount = ? WHERE id = ?', (new_value, emp_id))
    elif issue_type == 'DEDUCTION_MISMATCH':
        db.execute('UPDATE employees_v2 SET sum_deduction = ? WHERE id = ?', (new_value, emp_id))
    elif issue_type == 'NET_MISMATCH':
        db.execute('UPDATE employees_v2 SET net_total = ? WHERE id = ?', (new_value, emp_id))
    elif issue_type in ('UNMAPPED_DESIGNATION', 'UNCONFIRMED_DESIGNATION'):
        raw_for_mapping = issue['raw_designation_en'] or extract_english_designation(issue['raw_value'])
        match_type = data.get('match_type') or default_designation_match_type(raw_for_mapping)
        if action == 'global':
            mapping_id = save_dictionary_mapping_v3(db, raw_for_mapping, new_value, "", match_type)
            apply_designation_to_employees(db, raw_for_mapping, new_value, "", match_type, 'global', mapping_id=mapping_id)
        elif action == 'pdf':
            apply_designation_to_employees(db, raw_for_mapping, new_value, "", match_type, 'pdf', issue['bill_id'])
        elif action == 'ignore':
            db.execute("""
                UPDATE review_queue_v3 SET status = 'RESOLVED', action = 'ignore_once', resolved_at = CURRENT_TIMESTAMP
                WHERE employee_id = ? AND issue_type = 'UNCONFIRMED_DESIGNATION'
            """, (emp_id,))
            db.execute("""
                UPDATE employees_v2
                SET designation_needs_review = 0
                WHERE id = ?
            """, (emp_id,))
        else:
            db.execute("""
                UPDATE employees_v2
                SET final_designation = ?, final_designation_en = ?, final_designation_hi = '', normalized_designation = ?,
                    designation_source = 'user_manual', designation_confidence = 1,
                    designation_needs_review = 0, employee_data_status = 'VERIFIED',
                    row_verification_status = 'VERIFIED'
                WHERE id = ?
            """, (new_value, new_value, new_value, emp_id))
            db.execute("""
                UPDATE review_queue_v3 SET status = 'RESOLVED', action = 'employee', resolved_at = CURRENT_TIMESTAMP
                WHERE employee_id = ? AND issue_type = 'UNCONFIRMED_DESIGNATION'
            """, (emp_id,))
        audit_log_v3(
            db, "review_item_resolved", "employees_v2", emp_id, "final_designation_en",
            issue["old_designation"], new_value, issue["source_filename"], issue["page_no"],
            issue["sr_no"], issue["raw_name"], issue["employee_code"], action
        )
        
    db.execute("UPDATE validation_issues_v2 SET status = 'RESOLVED', resolved_at = CURRENT_TIMESTAMP WHERE id = ?", (issue_id,))
    if emp_id:
        recalculate_employee_data_status(db, emp_id)
    db.commit()
    
    return jsonify({'status': 'success'})

@app.route('/api/v2/dictionary', methods=['GET'])
@app.route('/api/v3/dictionary', methods=['GET'])
def get_dictionary():
    db = get_db()
    active_filter = (request.args.get('active') or ('1' if request.path.startswith('/api/v3/') else 'all')).strip().lower()
    conditions = []
    params = []
    if active_filter in ('1', 'true', 'active'):
        conditions.append('active = 1')
    elif active_filter in ('0', 'false', 'inactive'):
        conditions.append('active = 0')
    where = ('WHERE ' + ' AND '.join(conditions)) if conditions else ''
    terms = db.execute('''
        SELECT id, category, raw_text, COALESCE(raw_text_key, '') AS raw_text_key,
               COALESCE(final_text, final_text_en, '') AS final_text,
               final_text_en, final_text_hi, match_type,
               priority, active, created_at, updated_at, notes
        FROM dictionary_mappings_v3
        {where}
        ORDER BY active DESC, category, priority DESC, raw_text
    '''.format(where=where), params).fetchall()
    result = []
    for t in terms:
        item = {
            'id': t['id'],
            'type': t['category'],
            'category': t['category'],
            'raw': t['raw_text'],
            'raw_text': t['raw_text'],
            'raw_text_key': t['raw_text_key'],
            'normalized': t['final_text'],
            'final_text': t['final_text'],
            'final_text_en': t['final_text'],
            'match_type': t['match_type'],
            'priority': t['priority'],
            'active': bool(t['active']),
            'created_at': t['created_at'],
            'updated_at': t['updated_at'],
            'notes': t['notes'] or '',
        }
        if request.path.startswith('/api/v2/'):
            item['final_text_hi'] = ''
        result.append(item)
    return jsonify(result)

@app.route('/api/v2/dictionary', methods=['POST'])
@app.route('/api/v3/dictionary', methods=['POST'])
def add_dictionary_term():
    data = request.json
    dtype = data.get('category') or data.get('type') or 'designation'
    raw = data.get('raw_text') or data.get('raw')
    norm = data.get('final_text') or data.get('final_text_en') or data.get('normalized')
    final_hi = ""
    match_type = data.get('match_type') or default_designation_match_type(raw)
    priority = int(data.get('priority', 100))
    
    db = get_db()
    active = int(data.get('active', 1))
    mapping_id = save_dictionary_mapping_v3(db, raw, norm, final_hi, match_type, priority, dtype, active)
    if dtype == 'designation':
        if active and data.get('apply_existing', True):
            changed = apply_designation_to_employees(db, raw, norm, "", match_type, 'global', mapping_id=mapping_id)
            audit_log_v3(db, "mapping_applied_globally", "dictionary_mappings_v3", mapping_id, "final_designation", "", norm, applied_scope="global", user_note=f"{changed['updated_count']} existing rows updated")
    elif dtype in ('allowance_head', 'deduction_head', 'head'):
        db.execute('INSERT INTO dictionary_terms_v2 (dictionary_type, raw_value, normalized_value) VALUES (?,?,?)', ('head', raw, norm))
        if data.get('apply_existing', True):
            db.execute("UPDATE employee_components_v2 SET normalized_head = ? WHERE raw_head LIKE ?", (norm, f'%{raw}%'))
    elif dtype == 'head':
        pass
    db.commit()
        
    return jsonify({'status': 'success', 'id': mapping_id, **(changed if dtype == 'designation' and active and data.get('apply_existing', True) else {})})

@app.route('/api/v2/dictionary/<int:mapping_id>', methods=['PUT'])
@app.route('/api/v3/dictionary/<int:mapping_id>', methods=['PUT'])
def update_dictionary_term(mapping_id):
    data = request.json or {}
    db = get_db()
    old = db.execute('SELECT * FROM dictionary_mappings_v3 WHERE id = ?', (mapping_id,)).fetchone()
    if not old:
        return jsonify({'error': 'Mapping not found'}), 404
    fields = {
        'category': data.get('category', old['category']),
        'raw_text': data.get('raw_text', old['raw_text']),
        'raw_text_key': '',
        'final_text': data.get('final_text') or data.get('final_text_en') or old['final_text'] or old['final_text_en'],
        'final_text_en': data.get('final_text') or data.get('final_text_en') or old['final_text'] or old['final_text_en'],
        'final_text_hi': '',
        'match_type': data.get('match_type', old['match_type']),
        'priority': int(data.get('priority', old['priority'] or 0)),
        'active': int(data.get('active', old['active'])),
        'notes': data.get('notes', old['notes']),
    }
    if fields['category'] == 'designation':
        fields['raw_text'] = extract_english_designation(fields['raw_text'])
        fields['raw_text_key'] = clean_designation_key(fields['raw_text'])
        fields['final_text'] = extract_english_designation(fields['final_text']) or fields['final_text']
        fields['final_text_en'] = fields['final_text']
    db.execute("""
        UPDATE dictionary_mappings_v3
        SET category=?, raw_text=?, raw_text_key=?, final_text=?, final_text_en=?, final_text_hi=?, match_type=?,
            priority=?, active=?, notes=?, updated_at=CURRENT_TIMESTAMP
        WHERE id=?
    """, (
        fields['category'], fields['raw_text'], fields['raw_text_key'], fields['final_text'],
        fields['final_text_en'], fields['final_text_hi'],
        fields['match_type'], fields['priority'], fields['active'], fields['notes'], mapping_id
    ))
    audit_log_v3(db, "designation_dictionary_edit", "dictionary_mappings_v3", mapping_id, "mapping", dict(old), fields)
    if data.get('apply_existing') and fields['active'] and fields['category'] == 'designation':
        changed = apply_designation_to_employees(db, fields['raw_text'], fields['final_text_en'], "", fields['match_type'], 'global', mapping_id=mapping_id)
        audit_log_v3(db, "mapping_applied_globally", "dictionary_mappings_v3", mapping_id, "final_designation", "", fields['final_text_en'], applied_scope="global", user_note=f"{changed['updated_count']} existing rows updated")
    db.commit()
    return jsonify({'status': 'success', **(changed if data.get('apply_existing') and fields['active'] and fields['category'] == 'designation' else {})})

@app.route('/api/v2/dictionary/<int:mapping_id>', methods=['DELETE'])
@app.route('/api/v3/dictionary/<int:mapping_id>', methods=['DELETE'])
def delete_dictionary_term(mapping_id):
    db = get_db()
    old = db.execute('SELECT * FROM dictionary_mappings_v3 WHERE id = ?', (mapping_id,)).fetchone()
    if not old:
        return jsonify({'error': 'Mapping not found'}), 404
    hard = str(request.args.get('hard', '')).lower() in ('1', 'true', 'yes')
    if hard:
        db.execute('DELETE FROM dictionary_mappings_v3 WHERE id = ?', (mapping_id,))
        action = "designation_dictionary_delete"
    else:
        db.execute('UPDATE dictionary_mappings_v3 SET active = 0, updated_at = CURRENT_TIMESTAMP WHERE id = ?', (mapping_id,))
        action = "designation_dictionary_deactivate"
    audit_log_v3(db, action, "dictionary_mappings_v3", mapping_id, "active", old["active"], 0, applied_scope="global")
    db.commit()
    return jsonify({'status': 'success'})

@app.route('/api/v2/dictionary/<int:mapping_id>/apply', methods=['POST'])
@app.route('/api/v3/dictionary/<int:mapping_id>/apply', methods=['POST'])
def apply_dictionary_term(mapping_id):
    db = get_db()
    row = db.execute('SELECT * FROM dictionary_mappings_v3 WHERE id = ? AND active = 1', (mapping_id,)).fetchone()
    if not row:
        return jsonify({'error': 'Active mapping not found'}), 404
    changed = 0
    if row['category'] == 'designation':
        final_text = row['final_text'] or row['final_text_en']
        changed = apply_designation_to_employees(db, row['raw_text'], final_text, "", row['match_type'], 'global', mapping_id=mapping_id)
        audit_log_v3(db, "mapping_applied_globally", "dictionary_mappings_v3", mapping_id, "final_designation", "", final_text, applied_scope="global", user_note=f"{changed['updated_count']} existing rows updated")
    db.commit()
    return jsonify({'status': 'success', 'updated': changed['updated_count'] if isinstance(changed, dict) else changed, **(changed if isinstance(changed, dict) else {})})

@app.route('/api/v3/dictionary/apply', methods=['POST'])
def apply_dictionary_term_v3():
    data = request.json or {}
    mapping_id = data.get('mapping_id') or data.get('id')
    if not mapping_id:
        return jsonify({'error': 'mapping_id is required'}), 400
    db = get_db()
    row = db.execute('SELECT * FROM dictionary_mappings_v3 WHERE id = ? AND active = 1', (mapping_id,)).fetchone()
    if not row:
        return jsonify({'error': 'Active mapping not found'}), 404
    changed = 0
    if row['category'] == 'designation':
        final_text = row['final_text'] or row['final_text_en']
        changed = apply_designation_to_employees(db, row['raw_text'], final_text, "", row['match_type'], 'global', mapping_id=mapping_id)
        audit_log_v3(db, "mapping_applied_globally", "dictionary_mappings_v3", mapping_id, "final_designation", "", final_text, applied_scope="global", user_note=f"{changed['updated_count']} existing rows updated")
    db.commit()
    return jsonify({'status': 'success', 'updated': changed['updated_count'] if isinstance(changed, dict) else changed, **(changed if isinstance(changed, dict) else {})})

@app.route('/api/v3/dictionary/deactivate', methods=['POST'])
def deactivate_dictionary_v3():
    data = request.json or {}
    mapping_id = data.get('mapping_id') or data.get('id')
    if not mapping_id:
        return jsonify({'error': 'mapping_id is required'}), 400
    db = get_db()
    old = db.execute('SELECT * FROM dictionary_mappings_v3 WHERE id = ?', (mapping_id,)).fetchone()
    if not old:
        return jsonify({'error': 'Mapping not found'}), 404
    db.execute('UPDATE dictionary_mappings_v3 SET active = 0, updated_at = CURRENT_TIMESTAMP WHERE id = ?', (mapping_id,))
    audit_log_v3(db, "designation_dictionary_deactivate", "dictionary_mappings_v3", mapping_id, "active", old["active"], 0, applied_scope="global")
    db.commit()
    return jsonify({'status': 'success'})

@app.route('/api/v2/dictionary/export', methods=['GET'])
@app.route('/api/v3/dictionary/export', methods=['GET'])
def export_dictionary_json():
    db = get_db()
    rows = db.execute('''
        SELECT id, category, raw_text, COALESCE(raw_text_key, '') AS raw_text_key,
               COALESCE(final_text, final_text_en, '') AS final_text,
               match_type, priority, active, created_at, updated_at, notes
        FROM dictionary_mappings_v3
        ORDER BY category, priority DESC, raw_text
    ''').fetchall()
    payload = {'app_name': APP_NAME, 'exported_at': datetime.now().isoformat(timespec='seconds'), 'mappings': [dict(r) for r in rows]}
    return jsonify(payload)

@app.route('/api/v2/dictionary/import', methods=['POST'])
@app.route('/api/v3/dictionary/import', methods=['POST'])
def import_dictionary_json():
    data = request.json or {}
    mappings = data.get('mappings', data if isinstance(data, list) else [])
    db = get_db()
    count = 0
    for item in mappings:
        save_dictionary_mapping_v3(
            db,
            item.get('raw_text') or item.get('raw'),
            item.get('final_text') or item.get('final_text_en') or item.get('normalized'),
            "",
            item.get('match_type', 'exact'),
            item.get('priority', 100),
            item.get('category', 'designation'),
            item.get('active', 1),
            item.get('created_by', 'import'),
            item.get('notes', '')
        )
        count += 1
    db.commit()
    return jsonify({'status': 'success', 'imported': count})

@app.route('/api/v3/dictionary/sanitize', methods=['POST'])
def sanitize_dictionary_v3():
    db = get_db()
    rows = db.execute("""
        SELECT *
        FROM dictionary_mappings_v3
        WHERE category = 'designation'
        ORDER BY id
    """).fetchall()
    scanned_count = len(rows)
    sanitized_count = 0
    hindi_removed_count = 0
    numeric_removed_count = 0
    devanagari_re = re.compile(r"[\u0900-\u097F]")
    long_number_re = re.compile(r"\b\d{7,}\b")

    for row in rows:
        old_raw = row['raw_text'] or ''
        old_final = row['final_text'] or row['final_text_en'] or ''
        raw_clean = extract_english_designation(old_raw)
        raw_key = clean_designation_key(raw_clean)
        final_clean = extract_english_designation(old_final) or str(old_final or '').strip()
        changed = (
            raw_clean != old_raw
            or raw_key != (row['raw_text_key'] or '')
            or final_clean != old_final
            or (row['final_text_hi'] or '') != ''
        )
        if devanagari_re.search(old_raw) or devanagari_re.search(row['final_text_hi'] or ''):
            hindi_removed_count += 1
        if long_number_re.search(old_raw):
            numeric_removed_count += 1
        if changed:
            sanitized_count += 1
        db.execute("""
            UPDATE dictionary_mappings_v3
            SET raw_text = ?, raw_text_key = ?, final_text = ?, final_text_en = ?,
                final_text_hi = '', updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
        """, (raw_clean, raw_key, final_clean, final_clean, row['id']))

    duplicate_groups_found = 0
    duplicates_deactivated = 0
    dup_groups = db.execute("""
        SELECT category, raw_text_key, COALESCE(final_text, final_text_en, '') AS final_text,
               COALESCE(match_type, 'exact') AS match_type, COUNT(*) AS c
        FROM dictionary_mappings_v3
        WHERE active = 1
        GROUP BY category, raw_text_key, final_text, match_type
        HAVING c > 1
    """).fetchall()
    duplicate_groups_found = len(dup_groups)
    for group in dup_groups:
        matches = db.execute("""
            SELECT id
            FROM dictionary_mappings_v3
            WHERE active = 1
              AND category = ?
              AND COALESCE(raw_text_key, '') = ?
              AND COALESCE(final_text, final_text_en, '') = ?
              AND COALESCE(match_type, 'exact') = ?
            ORDER BY COALESCE(priority, 0) DESC,
                     COALESCE(updated_at, created_at, '') DESC,
                     COALESCE(created_at, '') DESC,
                     id DESC
        """, (group['category'], group['raw_text_key'] or '', group['final_text'] or '', group['match_type'] or 'exact')).fetchall()
        keep_id = matches[0]['id']
        deactivate_ids = [m['id'] for m in matches[1:]]
        if deactivate_ids:
            placeholders = ",".join("?" * len(deactivate_ids))
            db.execute(f"""
                UPDATE dictionary_mappings_v3
                SET active = 0, updated_at = CURRENT_TIMESTAMP
                WHERE id IN ({placeholders})
            """, deactivate_ids)
            duplicates_deactivated += len(deactivate_ids)
            audit_log_v3(
                db, "DICTIONARY_DUPLICATE_DEACTIVATED", "dictionary_mappings_v3", keep_id,
                "duplicate_ids", ",".join(map(str, deactivate_ids)), "kept_active",
                applied_scope="dictionary_sanitize"
            )

    audit_log_v3(
        db, "DICTIONARY_SANITIZE", "dictionary_mappings_v3", None,
        "designation_mappings", scanned_count, sanitized_count,
        applied_scope="all_dictionary_mappings",
        user_note=f"{duplicates_deactivated} duplicate active mapping(s) deactivated"
    )
    db.commit()
    return jsonify({
        'status': 'success',
        'scanned_count': scanned_count,
        'sanitized_count': sanitized_count,
        'hindi_removed_count': hindi_removed_count,
        'numeric_removed_count': numeric_removed_count,
        'duplicate_groups_found': duplicate_groups_found,
        'duplicates_deactivated': duplicates_deactivated,
    })

@app.route('/api/v3/dictionary/reset', methods=['POST'])
def reset_dictionary_v3():
    db = get_db()
    deleted_count = db.execute("SELECT COUNT(*) FROM dictionary_mappings_v3").fetchone()[0]
    db.execute("DELETE FROM dictionary_mappings_v3")
    audit_log_v3(
        db, "DICTIONARY_RESET", "dictionary_mappings_v3", None,
        "all_dictionary_mappings", deleted_count, 0,
        applied_scope="all_dictionary_mappings",
        user_note="Deleted all dictionary mappings only"
    )
    db.commit()
    return jsonify({
        'status': 'success',
        'deleted_count': deleted_count,
        'deactivated_count': 0,
    })

def employee_detail_payload(db, emp_id):
    emp = db.execute("""
        SELECT e.*, b.bill_no
        FROM employees_v2 e
        LEFT JOIN bills_v2 b ON b.id = e.bill_id
        WHERE e.id = ?
    """, (emp_id,)).fetchone()
    if not emp:
        return None
    comps = db.execute("""
        SELECT normalized_head, SUM(amount) AS amount
        FROM employee_components_v2
        WHERE employee_id = ?
        GROUP BY normalized_head
    """, (emp_id,)).fetchall()
    comp_map = {c["normalized_head"]: c["amount"] or 0 for c in comps}
    final_en = emp["final_designation"] or emp["final_designation_en"] or emp["normalized_designation"] or "Needs Review"
    parts = extract_designation_parts(emp["raw_designation_full"] or emp["raw_designation"] or "")
    raw_hi = emp["raw_designation_hi"] or parts["raw_designation_hi"] if "raw_designation_hi" in emp.keys() else parts["raw_designation_hi"]
    final_hi_display = (
        emp["final_designation_hi_display"]
        if "final_designation_hi_display" in emp.keys() and emp["final_designation_hi_display"]
        else raw_hi
    )
    final_hi = emp["final_designation_hi"] if "final_designation_hi" in emp.keys() and emp["final_designation_hi"] else ""
    item = {
        "id": emp["id"],
        "source_file": emp["source_filename"] or "",
        "page_no": emp["page_no"] or 0,
        "sr_no": emp["sr_no"] or 0,
        "group_name": emp["group_name"] or "",
        "month_year": emp["month_year"] or "",
        "bill_no": emp["bill_no"] or "",
        "name": emp["raw_name"] or "",
        "employee_name": emp["raw_name"] or "",
        "employee_id": emp["employee_id"] or "",
        "gpf_pran": emp["gpf_pran_raw"] or "",
        "belt_no": emp["belt_no"] or "",
        "pan": emp["pan_raw"] or "",
        "pay_scale": emp["pay_scale_raw"] or "",
        "bank_account": emp["bank_acc"] or "",
        "aadhaar_masked": emp["aadhar_raw"] or "",
        "raw_designation": emp["raw_designation"] or "",
        "raw_designation_full": emp["raw_designation_full"] or emp["raw_designation"] or "",
        "raw_designation_en": emp["raw_designation_en"] or parts["raw_designation_en"],
        "raw_designation_hi": raw_hi,
        "raw_designation_key": emp["raw_designation_key"] or parts["raw_designation_key"],
        "designation": final_en,
        "final_designation": final_en,
        "final_designation_en": final_en,
        "final_designation_hi": final_hi,
        "final_designation_hi_display": final_hi_display,
        "designation_display": f"{final_en} / {final_hi_display}" if final_hi_display else final_en,
        "designation_source": emp["designation_source"] or "",
        "designation_needs_review": bool(emp["designation_needs_review"]),
        "employee_data_status": emp["employee_data_status"] or emp["row_verification_status"] or "",
        "is_deleted": bool(emp["is_deleted"]) if "is_deleted" in emp.keys() else False,
        "deleted_at": emp["deleted_at"] if "deleted_at" in emp.keys() else "",
        "gross_amount": emp["gross_amount"] or 0,
        "sum_deduction": emp["sum_deduction"] or 0,
        "net_total": emp["net_total"] or 0,
    }
    for head, api_key in EMPLOYEE_COMPONENT_FIELDS.items():
        item[api_key] = comp_map.get(head, 0)
    item["basic_pay"] = item.get("basic", 0)
    item["itax"] = item.get("income_tax", 0)
    return item

@app.route('/api/employees/<int:emp_id>', methods=['GET'])
@app.route('/api/v3/employees/<int:emp_id>', methods=['GET'])
def get_employee_detail(emp_id):
    db = get_db()
    item = employee_detail_payload(db, emp_id)
    if not item:
        return jsonify({'error': 'Employee not found'}), 404
    return jsonify(item)


@app.route('/api/v3/employees/<int:emp_id>/source', methods=['GET'])
def get_employee_source(emp_id):
    db = get_db()
    detail = employee_detail_payload(db, emp_id)
    if not detail:
        return jsonify({'error': 'Employee not found'}), 404
    emp = db.execute("""
        SELECT raw_text_block
        FROM employees_v2
        WHERE id = ?
    """, (emp_id,)).fetchone()
    components = db.execute("""
        SELECT id, component_type, raw_head, normalized_head, amount,
               included_in_total, mapping_status, confidence,
               COALESCE(source_type, 'parsed') AS source_type, source_page, source_y
        FROM employee_components_v2
        WHERE employee_id = ?
        ORDER BY component_type, id
    """, (emp_id,)).fetchall()
    issues = db.execute("""
        SELECT id, issue_type, severity, message, raw_value, suggested_value, status,
               created_at, resolved_at
        FROM validation_issues_v2
        WHERE employee_id = ?
        ORDER BY id
    """, (emp_id,)).fetchall()
    audit_rows = db.execute("""
        SELECT *
        FROM edit_audit_log_v3
        WHERE table_name IN ('employees_v2', 'employee_components_v2')
          AND CAST(record_id AS TEXT) = CAST(? AS TEXT)
        ORDER BY id DESC
        LIMIT 200
    """, (emp_id,)).fetchall()
    return jsonify({
        **detail,
        'raw_text_block': emp['raw_text_block'] if emp else '',
        'components': [dict(row) for row in components],
        'validation_issues': [dict(row) for row in issues],
        'audit_log': [dict(row) for row in audit_rows],
    })


@app.route('/api/v3/audit_log', methods=['GET'])
def get_audit_log_v3():
    db = get_db()
    limit = max(1, min(500, int(request.args.get('limit', 100))))
    rows = db.execute("""
        SELECT *
        FROM edit_audit_log_v3
        ORDER BY id DESC
        LIMIT ?
    """, (limit,)).fetchall()
    return jsonify({'audit_log': [dict(r) for r in rows], 'total': len(rows)})

@app.route('/api/employees/<int:emp_id>', methods=['PUT'])
@app.route('/api/v3/employees/<int:emp_id>/edit', methods=['POST', 'PUT'])
def update_employee(emp_id):
    data = request.json or {}
    db = get_db()
    emp = db.execute('SELECT * FROM employees_v2 WHERE id = ?', (emp_id,)).fetchone()
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404

    allowed = {
        'employee_name': 'raw_name',
        'employee_id': 'employee_id',
        'gpf_pran': 'gpf_pran_raw',
        'belt_no': 'belt_no',
        'pan': 'pan_raw',
        'pay_scale': 'pay_scale_raw',
        'bank_account': 'bank_acc',
        'aadhaar_masked': 'aadhar_raw',
        'final_designation_hi_display': 'final_designation_hi_display',
        'employee_data_status': 'employee_data_status',
    }
    updates = []
    params = []
    for api_field, db_field in allowed.items():
        if api_field not in data:
            continue
        old_value = emp[db_field] if db_field in emp.keys() else ''
        new_value = data.get(api_field, '')
        updates.append(f"{db_field} = ?")
        params.append(new_value)
        audit_log_v3(
            db, "employee_edit", "employees_v2", emp_id, db_field, old_value, new_value,
            emp["source_filename"], emp["page_no"], emp["sr_no"], emp["raw_name"], emp["employee_id"],
            data.get("designation_scope", "employee")
        )

    total_fields = {
        'gross': 'gross_amount',
        'gross_amount': 'gross_amount',
        'deduction': 'sum_deduction',
        'sum_deduction': 'sum_deduction',
        'net': 'net_total',
        'net_total': 'net_total',
    }
    for api_field, db_field in total_fields.items():
        if api_field not in data:
            continue
        try:
            new_value = float(data.get(api_field) or 0)
        except (TypeError, ValueError):
            continue
        old_value = emp[db_field]
        updates.append(f"{db_field} = ?")
        params.append(new_value)
        audit_log_v3(db, "employee_edit", "employees_v2", emp_id, db_field, old_value, new_value,
                     emp["source_filename"], emp["page_no"], emp["sr_no"], emp["raw_name"], emp["employee_id"],
                     data.get("designation_scope", "employee"))

    component_api = {api_key: head for head, api_key in EMPLOYEE_COMPONENT_FIELDS.items()}
    component_api.update({
        'Basic': 'basic_pay', 'DA': 'da', 'HRA': 'hra', 'LIC': 'lic',
        'GPF': 'gpf', 'GPF 2004': 'gpf_2004', 'RGHS': 'rghs', 'SIP': 'sip',
        'SIL': 'sil', 'SII': 'sii', 'ITax': 'income_tax', 'ROP': 'rop',
        'Roadways Deduction': 'roadways_deduction', 'Advance Salary': 'advance_salary',
        'PPSS JDR': 'ppss_jdr', 'AIS GPF': 'ais_gpf', 'EACPenF': 'eacpenf',
        'GIAIS': 'giais', 'Car Rent': 'car_rent', 'HFRH': 'hfrh',
        'Mess Allowance': 'mess_allowance', 'KMA Allowance': 'kma_allowance',
        'FTA': 'fta', 'MTA': 'mta', 'Special Allowance': 'special_pay',
        'Special Pay': 'special_pay', 'Wash Allowance': 'wash_allowance',
    })
    deduction_heads = {
        'lic', 'gpf', 'gpf_2004', 'rghs', 'sip', 'sil', 'sii', 'income_tax',
        'roadways_deduction', 'advance_salary', 'ppss_jdr', 'ais_gpf',
        'eacpenf', 'giais', 'car_rent', 'hfrh', 'per_ded_1', 'per_ded_2'
    }
    component_payload = data.get('components') if isinstance(data.get('components'), dict) else data
    for api_field, normalized_head in component_api.items():
        if api_field not in component_payload:
            continue
        try:
            amount = float(component_payload.get(api_field) or 0)
        except (TypeError, ValueError):
            continue
        existing = db.execute("""
            SELECT id, amount FROM employee_components_v2
            WHERE employee_id = ? AND normalized_head = ?
            ORDER BY id
        """, (emp_id, normalized_head)).fetchall()
        if existing:
            old_value = existing[0]["amount"]
            db.execute("""
                UPDATE employee_components_v2
                SET amount = ?,
                    mapping_status = 'user_manual',
                    confidence = 1,
                    source_type = 'manual_edit'
                WHERE employee_id = ? AND normalized_head = ?
            """, (amount, emp_id, normalized_head))
        else:
            old_value = ""
            comp_type = "deduction" if normalized_head in deduction_heads else "allowance"
            db.execute("""
                INSERT INTO employee_components_v2 (
                    employee_id, bill_id, component_type, raw_head, normalized_head, amount,
                    mapping_status, confidence, source_type
                ) VALUES (?,?,?,?,?,?,?,?,?)
            """, (emp_id, emp["bill_id"], comp_type, api_field, normalized_head, amount, "user_manual", 1, "manual_edit"))
        audit_log_v3(db, "employee_component_edit", "employee_components_v2", emp_id, normalized_head,
                     old_value, amount, emp["source_filename"], emp["page_no"], emp["sr_no"],
                     emp["raw_name"], emp["employee_id"], data.get("designation_scope", "employee"))

    if 'final_designation' in data and 'final_designation_en' not in data:
        data['final_designation_en'] = data.get('final_designation')
    final_en_in_payload = 'final_designation_en' in data
    designation_scope = data.get('designation_scope', 'employee')
    designation_changed = final_en_in_payload and data.get('final_designation_en') != (emp['final_designation_en'] or emp['normalized_designation'])
    designation_mapping_requested = final_en_in_payload and designation_scope in ('global', 'pdf')
    if designation_changed or designation_mapping_requested:
        final_en = extract_english_designation(data.get('final_designation_en', '').strip()) or data.get('final_designation_en', '').strip()
        final_hi = ""
        scope = designation_scope
        raw_text = emp['raw_designation_en'] or extract_english_designation(emp['raw_designation_full'] or emp['raw_designation'] or '')
        match_type = data.get('match_type') or default_designation_match_type(raw_text)
        if scope == 'global':
            mapping_id = save_dictionary_mapping_v3(db, raw_text, final_en, final_hi, match_type, data.get('priority', 100))
            changed = apply_designation_to_employees(db, raw_text, final_en, final_hi, match_type, 'global', mapping_id=mapping_id)
            audit_log_v3(db, "mapping_applied_globally", "dictionary_mappings_v3", mapping_id, "final_designation", "", final_en, emp["source_filename"], emp["page_no"], emp["sr_no"], emp["raw_name"], emp["employee_id"], "global", f"{changed['updated_count']} existing rows updated")
        elif scope == 'pdf':
            changed = apply_designation_to_employees(db, raw_text, final_en, final_hi, match_type, 'pdf', emp['bill_id'])
            audit_log_v3(db, "mapping_applied_to_current_pdf", "employees_v2", emp_id, "final_designation", emp['final_designation_en'], final_en, emp["source_filename"], emp["page_no"], emp["sr_no"], emp["raw_name"], emp["employee_id"], "pdf", f"{changed['updated_count']} rows updated in current PDF")
        else:
            updates.extend([
                "final_designation = ?",
                "final_designation_en = ?",
                "final_designation_hi = ''",
                "normalized_designation = ?",
                "designation_source = 'user_manual'",
                "designation_confidence = 1",
                "designation_needs_review = 0",
            ])
            params.extend([final_en, final_en, final_en])
            db.execute("""
                UPDATE review_queue_v3 SET status = 'RESOLVED', action = 'employee', resolved_at = CURRENT_TIMESTAMP
                WHERE employee_id = ? AND issue_type = 'UNCONFIRMED_DESIGNATION'
            """, (emp_id,))
            db.execute("""
                UPDATE validation_issues_v2 SET status = 'RESOLVED', resolved_at = CURRENT_TIMESTAMP
                WHERE employee_id = ? AND issue_type = 'UNCONFIRMED_DESIGNATION'
            """, (emp_id,))
            audit_log_v3(db, "employee_edit", "employees_v2", emp_id, "final_designation_en",
                         emp["final_designation_en"] or emp["normalized_designation"], final_en,
                         emp["source_filename"], emp["page_no"], emp["sr_no"], emp["raw_name"],
                         emp["employee_id"], "employee")

    if updates:
        db.execute(f"UPDATE employees_v2 SET {', '.join(updates)} WHERE id = ?", params + [emp_id])
    recalculate_employee_data_status(db, emp_id)
    if data.get('mark_verified'):
        db.execute("""
            UPDATE employees_v2
            SET employee_data_status = 'VERIFIED',
                row_verification_status = 'VERIFIED',
                designation_needs_review = 0
            WHERE id = ?
        """, (emp_id,))
        db.execute("""
            UPDATE validation_issues_v2
            SET status = 'RESOLVED', resolved_at = CURRENT_TIMESTAMP
            WHERE employee_id = ? AND COALESCE(status, 'OPEN') != 'RESOLVED'
        """, (emp_id,))
        db.execute("""
            UPDATE review_queue_v3
            SET status = 'RESOLVED', action = 'mark_verified', resolved_at = CURRENT_TIMESTAMP
            WHERE employee_id = ? AND COALESCE(status, 'OPEN') != 'RESOLVED'
        """, (emp_id,))
        audit_log_v3(
            db, "employee_mark_verified", "employees_v2", emp_id, "employee_data_status",
            emp["employee_data_status"] or emp["row_verification_status"] or "",
            "VERIFIED", emp["source_filename"], emp["page_no"], emp["sr_no"],
            emp["raw_name"], emp["employee_id"], "employee"
        )
    elif updates or component_payload is not data or any(k in data for k in component_api):
        db.execute("""
            UPDATE employees_v2
            SET employee_data_status = 'MANUALLY_EDITED',
                row_verification_status = 'MANUALLY_EDITED'
            WHERE id = ?
              AND COALESCE(employee_data_status, '') NOT IN ('NEEDS_REVIEW', 'MISSING_CRITICAL_FIELD', 'COMPONENT_REVIEW_REQUIRED')
        """, (emp_id,))
    db.commit()
    return jsonify({
        'status': 'success',
        'employee': employee_detail_payload(db, emp_id),
        'quality': employee_quality_summary(db),
    })


@app.route('/api/v3/employees/<int:emp_id>/delete', methods=['POST'])
@app.route('/api/v3/employees/<int:emp_id>', methods=['DELETE'])
def delete_employee_row(emp_id):
    db = get_db()
    emp = db.execute('SELECT * FROM employees_v2 WHERE id = ?', (emp_id,)).fetchone()
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404
    old_deleted = int(emp['is_deleted'] or 0) if 'is_deleted' in emp.keys() else 0
    db.execute("""
        UPDATE employees_v2
        SET is_deleted = 1,
            deleted_at = COALESCE(deleted_at, CURRENT_TIMESTAMP),
            employee_data_status = 'MANUALLY_EDITED',
            row_verification_status = 'MANUALLY_EDITED'
        WHERE id = ?
    """, (emp_id,))
    audit_log_v3(
        db, "EMPLOYEE_ROW_DELETE", "employees_v2", emp_id, "is_deleted",
        old_deleted, 1, emp["source_filename"], emp["page_no"], emp["sr_no"],
        emp["raw_name"], emp["employee_id"], "employee",
        "Soft deleted from active Employee Data"
    )
    db.commit()
    active_where = " WHERE COALESCE(e.is_deleted, 0) = 0"
    return jsonify({
        'status': 'success',
        'deleted': True,
        'employee': employee_detail_payload(db, emp_id),
        'quality': employee_quality_summary(db, active_where, []),
    })


@app.route('/api/v3/employees/<int:emp_id>/restore', methods=['POST'])
def restore_employee_row(emp_id):
    db = get_db()
    emp = db.execute('SELECT * FROM employees_v2 WHERE id = ?', (emp_id,)).fetchone()
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404
    old_deleted = int(emp['is_deleted'] or 0) if 'is_deleted' in emp.keys() else 0
    db.execute("""
        UPDATE employees_v2
        SET is_deleted = 0,
            deleted_at = NULL
        WHERE id = ?
    """, (emp_id,))
    audit_log_v3(
        db, "EMPLOYEE_ROW_RESTORE", "employees_v2", emp_id, "is_deleted",
        old_deleted, 0, emp["source_filename"], emp["page_no"], emp["sr_no"],
        emp["raw_name"], emp["employee_id"], "employee",
        "Restored to active Employee Data"
    )
    recalculate_employee_data_status(db, emp_id)
    db.commit()
    active_where = " WHERE COALESCE(e.is_deleted, 0) = 0"
    return jsonify({
        'status': 'success',
        'restored': True,
        'employee': employee_detail_payload(db, emp_id),
        'quality': employee_quality_summary(db, active_where, []),
    })

@app.route('/api/upload', methods=['POST'])
def upload_pdf():
    if 'files' not in request.files:
        return jsonify({'error': 'No files uploaded'}), 400
    
    files = request.files.getlist('files')
    results = []
    
    db = get_db()
    for f in files:
        if not f.filename.lower().endswith('.pdf'):
            results.append({'file': f.filename, 'error': 'Not a PDF file'})
            continue
        
        fname = secure_filename(f.filename)
        fpath = os.path.join(app.config['UPLOAD_FOLDER'], fname)
        f.save(fpath)
        
        if app.config.get('USE_V2_PARSER', False):
            # Insert into jobs queue
            db.execute("""
                INSERT INTO upload_jobs_v2 (
                    source_filename, status, started_at, current_stage,
                    message, percent_complete
                ) VALUES (?, 'QUEUED', CURRENT_TIMESTAMP, 'queued', 'Queued for parsing', 0)
            """, (fname,))
            db.commit()
            
            # Submit to thread pool
            executor.submit(process_file_v2, fpath, fname, dict(app.config))
            results.append({'file': fname, 'status': 'queued'})
        else:
            # V2 parser is the default production flow.
            # Legacy/V1 fallback is intentionally disabled and retained only for
            # emergency/manual debugging when explicitly re-enabled with a known
            # stable V1 ingestion implementation.
            results.append({
                'file': fname,
                'status': 'error',
                'error': 'Legacy V1 upload fallback is disabled. Enable USE_V2_PARSER or wire a stable V1 ingestion path.'
            })
    
    return jsonify({'results': results})

@app.route('/api/v2/jobs', methods=['GET'])
def get_v2_jobs():
    db = get_db()
    stale_rows = db.execute("""
        SELECT id
        FROM upload_jobs_v2
        WHERE status IN ('QUEUED', 'PROCESSING')
          AND started_at IS NOT NULL
          AND datetime(started_at, '+30 minutes') < CURRENT_TIMESTAMP
    """).fetchall()
    for row in stale_rows:
        db.execute("""
            UPDATE upload_jobs_v2
            SET status = 'FAILED',
                current_stage = 'failed',
                message = 'Job timed out before completion',
                error_message = 'Job timed out before completion',
                finished_at = CURRENT_TIMESTAMP
            WHERE id = ?
        """, (row['id'],))
    if stale_rows:
        db.commit()
    jobs = db.execute("""
        SELECT id, source_filename, status, processed_pages, total_pages,
               detected_employees, message, error_message, started_at,
               percent_complete, review_issue_count, current_stage,
               employees_saved, issue_count, warning_count, finished_at
        FROM upload_jobs_v2
        ORDER BY id DESC
        LIMIT 50
    """).fetchall()
    status_map = {
        'PROCESSING': 'parsing',
        'COMPLETED': 'verified',
        'NEEDS_REVIEW': 'needs_review',
        'FAILED': 'failed',
        'QUEUED': 'queued',
    }
    return jsonify([{
        'id': j[0], 'filename': j[1], 'status': status_map.get(j[2], (j[2] or '').lower()), 'raw_status': j[2],
        'processed_pages': j[3], 'total_pages': j[4], 'pages_processed': j[3], 'pages_total': j[4],
        'detected_employees': j[5], 'employees_detected': j[5],
        'message': j[6], 'error_message': j[7], 'started_at': j[8],
        'percentage': round(j[9] or (100 if j[2] in ('COMPLETED', 'NEEDS_REVIEW', 'FAILED') else 0), 2),
        'progress_percent': round(j[9] or (100 if j[2] in ('COMPLETED', 'NEEDS_REVIEW', 'FAILED') else 0), 2),
        'review_issue_count': j[10] or 0,
        'current_stage': j[11] or status_map.get(j[2], 'queued'),
        'employees_saved': j[12] or 0,
        'issue_count': j[13] or j[10] or 0,
        'warning_count': j[14] or 0,
        'finished_at': j[15],
    } for j in jobs])

@app.route('/api/employees')
def get_employees():
    db = get_db()
    if app.config.get('USE_V2_PARSER', False):
        conditions, params = build_employee_filter_conditions(request.args)
        where = (' WHERE ' + ' AND '.join(conditions)) if conditions else ''
        sort = request.args.get('sort', 'sr_no')
        dirn = request.args.get('dir', 'asc').lower()
        if dirn not in ('asc', 'desc'):
            dirn = 'asc'
        allowed = {
            'sr_no': 'e.sr_no',
            'group_name': 'e.group_name',
            'raw_name': 'e.raw_name',
            'gross_amount': 'e.gross_amount',
            'sum_deduction': 'e.sum_deduction',
            'net_total': 'e.net_total',
            'pay_scale_raw': 'e.pay_scale_raw',
            'employee_id': 'e.employee_id',
            'pan_raw': 'e.pan_raw',
            'normalized_designation': "COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation)",
            'final_designation_en': "COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation)",
        }
        if sort not in allowed:
            sort = 'sr_no'
        display_mode = (request.args.get('designation_display') or 'english').strip().lower().replace('_', '-')
        page = max(1, int(request.args.get('page', 1)))
        per_page = max(1, min(200, int(request.args.get('per_page', 50))))
        offset = (page - 1) * per_page

        # Employee Data source of truth: employees_v2 identity/totals plus
        # employee_components_v2 canonical normalized_head amounts. Do not read
        # legacy employees or schedule_entries for this route.
        from_clause = 'FROM employees_v2 e LEFT JOIN bills_v2 b ON b.id = e.bill_id'
        total = db.execute(f'SELECT COUNT(*) {from_clause}{where}', params).fetchone()[0]
        rows = db.execute(
            f'''SELECT e.id, e.source_filename, e.page_no, e.sr_no, e.group_name, e.raw_name,
                       e.raw_designation, e.normalized_designation, e.employee_id, e.gpf_pran_raw,
                       e.dob, e.belt_no, e.pan_raw, e.pay_scale_raw, e.bank_acc, e.aadhar_raw,
                       e.gross_amount, e.sum_deduction, e.net_total, e.month_year,
                       COALESCE(e.raw_designation_full, e.raw_designation) AS raw_designation_full,
                       COALESCE(e.raw_designation_en, '') AS raw_designation_en,
                       COALESCE(e.raw_designation_hi, '') AS raw_designation_hi,
                       COALESCE(e.raw_designation_key, '') AS raw_designation_key,
                       COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation) AS final_designation,
                       COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation) AS final_designation_en,
                       COALESCE(e.final_designation_hi, '') AS final_designation_hi,
                       COALESCE(e.final_designation_hi_display, e.final_designation_hi, '') AS final_designation_hi_display,
                       COALESCE(e.designation_source, '') AS designation_source,
                       COALESCE(e.designation_needs_review, 0) AS designation_needs_review,
                       COALESCE(e.is_deleted, 0) AS is_deleted,
                       COALESCE(e.deleted_at, '') AS deleted_at,
                       {employee_status_case_sql()} AS employee_data_status_calc,
                       COALESCE(e.employee_data_status, e.row_verification_status, 'OK') AS employee_data_status,
                       b.bill_no
                {from_clause}{where}
                ORDER BY {allowed[sort]} {dirn}
                LIMIT ? OFFSET ?''',
            params + [per_page, offset]
        ).fetchall()
        summary = db.execute(
            f'''SELECT COUNT(*) as count, COALESCE(SUM(e.gross_amount),0) as total_gross,
                       COALESCE(SUM(e.sum_deduction),0) as total_deduction, COALESCE(SUM(e.net_total),0) as total_net
                {from_clause}{where}''',
            params
        ).fetchone()
        row_ids = [r['id'] for r in rows]
        comp_map = {}
        issue_map = {}
        if row_ids:
            ph = ",".join("?" * len(row_ids))
            comps = db.execute(
                f"SELECT employee_id, normalized_head, amount FROM employee_components_v2 WHERE employee_id IN ({ph})",
                row_ids
            ).fetchall()
            for c in comps:
                comp_map.setdefault(c['employee_id'], {})
                comp_map[c['employee_id']][c['normalized_head']] = comp_map[c['employee_id']].get(c['normalized_head'], 0) + (c['amount'] or 0)
            issue_rows = db.execute(
                f"""SELECT employee_id, issue_type, severity, message
                    FROM validation_issues_v2
                    WHERE employee_id IN ({ph}) AND COALESCE(status, 'OPEN') != 'RESOLVED'
                    ORDER BY severity DESC, issue_type""",
                row_ids
            ).fetchall()
            for issue in issue_rows:
                issue_map.setdefault(issue['employee_id'], []).append(dict(issue))

        def cget(cm, *keys):
            for k in keys:
                if k in cm:
                    return cm[k]
            return 0
        employees = []
        for r in rows:
            cm = comp_map.get(r['id'], {})
            issues = issue_map.get(r['id'], [])
            status = r['employee_data_status_calc'] or 'VERIFIED'
            final_en = r['final_designation'] or r['final_designation_en'] or r['normalized_designation'] or 'Needs Review'
            parts = extract_designation_parts(r['raw_designation_full'] or r['raw_designation'] or '')
            raw_hi = r['raw_designation_hi'] or parts["raw_designation_hi"]
            final_hi = r['final_designation_hi_display'] or raw_hi
            designation_warning = ''
            if display_mode in ('english+hindi', 'english-hindi', 'combined'):
                display_designation = f"{final_en} / {final_hi}" if final_hi else final_en
                display_language = 'english+hindi'
            elif display_mode in ('raw', 'raw-source', 'source'):
                display_designation = r['raw_designation_full'] or r['raw_designation'] or final_en
                display_language = 'raw_source'
            else:
                display_designation = final_en or 'Needs Review'
                display_language = 'english'
            item = {
                'id': r['id'],
                'source_file': r['source_filename'] or '',
                'page_no': r['page_no'] or 0,
                'sr_no': r['sr_no'] or 0,
                'group_name': r['group_name'] or '',
                'month_year': r['month_year'] or '',
                'bill_no': r['bill_no'] or '',
                'name': r['raw_name'] or '',
                'employee_name': r['raw_name'] or '',
                'raw_designation': r['raw_designation'] or '',
                'raw_designation_full': r['raw_designation_full'] or r['raw_designation'] or '',
                'raw_designation_en': r['raw_designation_en'] or parts["raw_designation_en"],
                'raw_designation_hi': raw_hi,
                'raw_designation_key': r['raw_designation_key'] or parts["raw_designation_key"],
                'designation': display_designation,
                'display_designation': display_designation,
                'designation_display_language': display_language,
                'designation_display': display_designation,
                'designation_warning': designation_warning,
                'normalized_designation': final_en,
                'final_designation': final_en,
                'final_designation_en': final_en,
                'final_designation_hi': r['final_designation_hi'] or '',
                'final_designation_hi_display': final_hi,
                'designation_source': r['designation_source'] or '',
                'designation_needs_review': bool(r['designation_needs_review']),
                'is_deleted': bool(r['is_deleted']),
                'deleted_at': r['deleted_at'] or '',
                'employee_id': r['employee_id'] or '',
                'gpf_pran': r['gpf_pran_raw'] or '',
                'dob': r['dob'] or '',
                'belt_no': r['belt_no'] or '',
                'pan': r['pan_raw'] or '',
                'pay_scale': r['pay_scale_raw'] or '',
                'bank_account': r['bank_acc'] or '',
                'aadhaar_masked': r['aadhar_raw'] or '',
                'gross_amount': r['gross_amount'] or 0,
                'sum_deduction': r['sum_deduction'] or 0,
                'net_total': r['net_total'] or 0,
                'printed_gross': r['gross_amount'] or 0,
                'printed_deduction': r['sum_deduction'] or 0,
                'printed_net': r['net_total'] or 0,
                'employee_data_status': status,
                'has_issue': bool(issues),
                'issue_count': len(issues),
                'issue_messages': [i['message'] for i in issues if i.get('message')],
            }
            for head, api_key in EMPLOYEE_COMPONENT_FIELDS.items():
                item[api_key] = cget(cm, head)
            item['basic_pay'] = item['basic']
            item['itax'] = item['income_tax']
            employees.append(item)

        filtered_emp_ids = [r['id'] for r in db.execute(f'SELECT e.id {from_clause}{where}', params).fetchall()]
        filtered_ph = ','.join('?' * len(filtered_emp_ids)) if filtered_emp_ids else ''
        def component_total(heads):
            if not filtered_emp_ids:
                return 0
            placeholders = ','.join('?' * len(heads))
            return db.execute(
                f'''SELECT COALESCE(SUM(amount),0) FROM employee_components_v2
                    WHERE employee_id IN ({filtered_ph}) AND normalized_head IN ({placeholders})''',
                filtered_emp_ids + list(heads)
            ).fetchone()[0]

        quality = employee_quality_summary(db, where, params)
        filter_chips = build_employee_filter_chips(request.args)
        return jsonify({
            'employees': employees,
            'total': total,
            'page': page,
            'per_page': per_page,
            'pages': (total + per_page - 1) // per_page,
            'quality': quality,
            'filters': {
                'chips': filter_chips,
                'count': len(filter_chips),
                'no_result_message': 'No matching records found.' if total == 0 else '',
            },
            'summary': {
                'count': summary['count'] or 0,
                'total_gross': summary['total_gross'] or 0,
                'total_deduction': summary['total_deduction'] or 0,
                'total_net': summary['total_net'] or 0,
                'total_itax': component_total(['income_tax', 'itax']),
                'total_gpf': component_total(['gpf']),
                'avg_basic': 0,
                'avg_net': (summary['total_net'] or 0) / (summary['count'] or 1),
            }
        })
    return jsonify({'employees': [], 'total': 0, 'page': 1, 'per_page': 50, 'pages': 0, 'summary': {'count': 0}})


# ===================== V2 API REPLACEMENTS =====================

@app.route('/api/stats')
def get_stats():
    # Mixed-risk route: this still supports V1+V2 branches for compatibility.
    # Source-of-truth direction is V2-only; keep this dual path until route conversion is completed.
    db = get_db()
    if app.config.get('USE_V2_PARSER', False):
        active_employee_where = 'COALESCE(is_deleted, 0) = 0'
        total_emp = db.execute(f'SELECT COUNT(*) FROM employees_v2 WHERE {active_employee_where}').fetchone()[0]
        total_groups = db.execute(f'SELECT COUNT(DISTINCT group_name) FROM employees_v2 WHERE {active_employee_where}').fetchone()[0]
        total_uploads = db.execute('SELECT COUNT(*) FROM bills_v2').fetchone()[0]
        
        groups = [r[0] for r in db.execute(f'SELECT DISTINCT group_name FROM employees_v2 WHERE {active_employee_where} ORDER BY group_name').fetchall()]
        months = [r[0] for r in db.execute(f'SELECT DISTINCT month_year FROM employees_v2 WHERE {active_employee_where} ORDER BY month_year').fetchall()]
        
        gt_row = db.execute("""SELECT 
            COALESCE(SUM(official_gross), 0) as gt_gross,
            COALESCE(SUM(official_deduction), 0) as gt_ded,
            COALESCE(SUM(official_net), 0) as gt_net,
            COALESCE(SUM(parsed_gross), 0) as p_gross,
            COALESCE(SUM(parsed_deduction), 0) as p_ded,
            COALESCE(SUM(parsed_net), 0) as p_net
        FROM bills_v2""").fetchone()
        active_totals = db.execute(f"""
            SELECT COALESCE(SUM(gross_amount),0) AS gross,
                   COALESCE(SUM(sum_deduction),0) AS deduction,
                   COALESCE(SUM(net_total),0) AS net
            FROM employees_v2
            WHERE {active_employee_where}
        """).fetchone()
        
        gt_gross = active_totals['gross']
        gt_ded = active_totals['deduction']
        gt_net = active_totals['net']
        
        parsed_gross = gt_row['p_gross']
        parsed_ded = gt_row['p_ded']
        parsed_net = gt_row['p_net']
        
        pay_scales = [r[0] for r in db.execute(f'SELECT DISTINCT pay_scale_raw FROM employees_v2 WHERE pay_scale_raw != "" AND {active_employee_where} ORDER BY pay_scale_raw').fetchall()]
    else:
        total_emp = db.execute('SELECT COUNT(*) FROM employees').fetchone()[0]
        total_groups = db.execute('SELECT COUNT(DISTINCT group_name) FROM employees').fetchone()[0]
        total_uploads = db.execute('SELECT COUNT(*) FROM uploads').fetchone()[0]
        groups = [r[0] for r in db.execute('SELECT DISTINCT group_name FROM employees ORDER BY group_name').fetchall()]
        months = [r[0] for r in db.execute('SELECT DISTINCT month_year FROM employees ORDER BY month_year').fetchall()]
        
        gt_row = db.execute("""SELECT 
            COALESCE(SUM(grand_total_gross), 0) as gt_gross,
            COALESCE(SUM(grand_total_deduction), 0) as gt_ded,
            COALESCE(SUM(grand_total_net), 0) as gt_net,
            COALESCE(SUM(parsed_gross), 0) as p_gross,
            COALESCE(SUM(parsed_deduction), 0) as p_ded,
            COALESCE(SUM(parsed_net), 0) as p_net
        FROM uploads WHERE sheet_type = 'inner' """).fetchone()
        
        gt_gross = gt_row['gt_gross'] if gt_row['gt_gross'] > 0 else gt_row['p_gross']
        gt_ded = gt_row['gt_ded'] if gt_row['gt_ded'] > 0 else gt_row['p_ded']
        gt_net = gt_row['gt_net'] if gt_row['gt_net'] > 0 else gt_row['p_net']
        
        parsed_gross = db.execute('SELECT COALESCE(SUM(gross_amount),0) FROM employees').fetchone()[0]
        parsed_net = db.execute('SELECT COALESCE(SUM(net_total),0) FROM employees').fetchone()[0]
        parsed_ded = db.execute('SELECT COALESCE(SUM(sum_deduction),0) FROM employees').fetchone()[0]
        pay_scales = [r[0] for r in db.execute('SELECT DISTINCT pay_scale FROM employees WHERE pay_scale != "" ORDER BY pay_scale').fetchall()]
    
    return jsonify({
        'total_employees': total_emp, 'total_groups': total_groups,
        'total_uploads': total_uploads, 'groups': groups, 'months': months,
        'pay_scales': pay_scales,
        'total_gross': gt_gross, 'total_net': gt_net, 'total_deduction': gt_ded,
        'parsed_gross': parsed_gross, 'parsed_net': parsed_net, 'parsed_deduction': parsed_ded,
    })


@app.route('/api/v3/phase1_health', methods=['GET'])
def phase1_health_api():
    db = get_db()
    return jsonify(build_phase1_health(db))


@app.route('/api/v3/pdf_audits', methods=['GET'])
def pdf_audits_api():
    db = get_db()
    bill_id = request.args.get('bill_id')
    source_file = request.args.get('source_file') or request.args.get('source_filename')
    try:
        bill_id = int(bill_id) if bill_id else None
    except ValueError:
        return jsonify({'error': 'bill_id must be an integer'}), 400
    rows = build_pdf_audit_rows(db, bill_id=bill_id, source_file=source_file)
    return jsonify({'audits': rows, 'pdf_audits': rows, 'total': len(rows), 'health': build_phase1_health(db)})


def _bill_filter_from_request(data):
    bill_id = data.get('bill_id') or request.args.get('bill_id')
    source_file = data.get('source_file') or data.get('source_filename') or request.args.get('source_file')
    if bill_id in ('', None):
        bill_id = None
    if bill_id is not None:
        try:
            bill_id = int(bill_id)
        except (TypeError, ValueError):
            return None, None, 'bill_id must be an integer'
    return bill_id, source_file, None


def _employee_snapshot_sql(where=''):
    status_case = employee_status_case_sql()
    return f"""
        SELECT e.id AS employee_db_id, e.source_filename AS source_file, b.bill_no,
               e.group_name, e.month_year, e.page_no, e.sr_no,
               e.raw_name AS employee_name, e.employee_id, e.gpf_pran_raw AS gpf_pran,
               e.belt_no, e.pan_raw AS pan, e.pay_scale_raw AS pay_scale,
               COALESCE(e.raw_designation_full, e.raw_designation, '') AS raw_designation_full,
               COALESCE(e.raw_designation_en, '') AS raw_designation_en,
               COALESCE(e.raw_designation_hi, '') AS raw_designation_hi,
               COALESCE(e.final_designation, e.final_designation_en, e.normalized_designation, '') AS final_designation,
               COALESCE(e.final_designation_hi_display, '') AS final_designation_hi_display,
               COALESCE(e.designation_source, '') AS designation_source,
               COALESCE(e.is_deleted,0) AS is_deleted,
               e.deleted_at,
               e.gross_amount AS gross,
               e.sum_deduction AS total_deduction,
               e.net_total,
               (COALESCE(e.gross_amount,0) - COALESCE(e.sum_deduction,0)) AS calculated_net,
               CASE WHEN ABS((COALESCE(e.gross_amount,0) - COALESCE(e.sum_deduction,0)) - COALESCE(e.net_total,0)) <= 1
                    THEN 1 ELSE 0 END AS row_total_match,
               (SELECT COALESCE(SUM(c.amount),0) FROM employee_components_v2 c
                WHERE c.employee_id = e.id AND c.component_type = 'allowance') AS allowance_component_total,
               (SELECT COALESCE(SUM(c.amount),0) FROM employee_components_v2 c
                WHERE c.employee_id = e.id AND c.component_type = 'deduction'
                  AND COALESCE(c.included_in_total,1) = 1) AS deduction_component_total,
               {status_case} AS employee_data_status
        FROM employees_v2 e
        LEFT JOIN bills_v2 b ON b.id = e.bill_id
        {where}
    """


@app.route('/api/v3/pdf_audit_export', methods=['POST'])
def pdf_audit_export():
    data = request.json or request.form.to_dict() or {}
    db = get_db()
    bill_id, source_file, error = _bill_filter_from_request(data)
    if error:
        return jsonify({'error': error}), 400
    audits = build_pdf_audit_rows(db, bill_id=bill_id, source_file=source_file)
    if not audits:
        return jsonify({'error': 'No matching PDF audit found'}), 404
    audit = audits[0]
    if audit['bill_id'] is None:
        return jsonify({'error': 'Audit export requires a processed bill'}), 400
    bill_id = audit['bill_id']
    emp_ids = [r['id'] for r in db.execute("SELECT id FROM employees_v2 WHERE bill_id = ?", (bill_id,)).fetchall()]
    emp_ph = ','.join('?' * len(emp_ids)) if emp_ids else ''
    employee_where = "WHERE e.bill_id = ?"
    snapshot = _df_from_query(db, _employee_snapshot_sql(employee_where) + " ORDER BY e.sr_no", (bill_id,))
    needs_review = _df_from_query(db, _employee_snapshot_sql(employee_where + " AND COALESCE(e.is_deleted,0) = 0 AND (COALESCE(e.designation_needs_review,0) = 1 OR COALESCE(e.employee_data_status,'') IN ('NEEDS_REVIEW','MISSING_CRITICAL_FIELD','COMPONENT_REVIEW_REQUIRED'))") + " ORDER BY e.sr_no", (bill_id,))
    missing_critical = _df_from_query(db, _employee_snapshot_sql(employee_where + " AND COALESCE(e.is_deleted,0) = 0 AND (COALESCE(e.employee_id,'') = '' OR COALESCE(e.pay_scale_raw,'') = '' OR COALESCE(e.pan_raw,'') = '')") + " ORDER BY e.sr_no", (bill_id,))
    component_review = _df_from_query(db, _employee_snapshot_sql(employee_where + " AND COALESCE(e.is_deleted,0) = 0 AND (ABS((SELECT COALESCE(SUM(c.amount),0) FROM employee_components_v2 c WHERE c.employee_id = e.id AND c.component_type = 'allowance') - COALESCE(e.gross_amount,0)) > 1 OR ABS((SELECT COALESCE(SUM(c.amount),0) FROM employee_components_v2 c WHERE c.employee_id = e.id AND c.component_type = 'deduction' AND COALESCE(c.included_in_total,1) = 1) - COALESCE(e.sum_deduction,0)) > 1)") + " ORDER BY e.sr_no", (bill_id,))
    row_mismatches = _df_from_query(db, _employee_snapshot_sql(employee_where + " AND COALESCE(e.is_deleted,0) = 0 AND ABS((COALESCE(e.gross_amount,0) - COALESCE(e.sum_deduction,0)) - COALESCE(e.net_total,0)) > 1") + " ORDER BY e.sr_no", (bill_id,))
    unknown_components = _df_from_query(db, """
        SELECT v.*, e.raw_name AS employee_name, e.employee_id AS employee_code
        FROM validation_issues_v2 v
        LEFT JOIN employees_v2 e ON e.id = v.employee_id
        WHERE v.bill_id = ? AND v.issue_type = 'UNKNOWN_COMPONENT_HEAD'
        ORDER BY v.id
    """, (bill_id,))
    dictionary_applied = _df_from_query(db, _employee_snapshot_sql(employee_where + " AND COALESCE(e.is_deleted,0) = 0 AND COALESCE(e.designation_source,'') LIKE 'dictionary_%'") + " ORDER BY e.sr_no", (bill_id,))
    deleted_rows = _df_from_query(db, _employee_snapshot_sql(employee_where + " AND COALESCE(e.is_deleted,0) = 1") + " ORDER BY e.sr_no", (bill_id,))
    audit_log = pd.DataFrame()
    if emp_ids:
        audit_log = _df_from_query(db, f"""
            SELECT *
            FROM edit_audit_log_v3
            WHERE (table_name IN ('employees_v2','employee_components_v2')
                   AND CAST(record_id AS TEXT) IN ({emp_ph}))
               OR source_file = ?
            ORDER BY id DESC
        """, emp_ids + [audit['source_file']])
    sheets = {
        'Audit Summary': pd.DataFrame([audit]),
        'Employee Master Snapshot': snapshot,
        'Needs Review Rows': needs_review,
        'Missing Critical Fields': missing_critical,
        'Component Review Rows': component_review,
        'Row Total Mismatches': row_mismatches,
        'Unknown Components': unknown_components,
        'Dictionary Applied Rows': dictionary_applied,
        'Deleted Rows': deleted_rows,
        'Audit Log': audit_log,
        'Export Metadata': pd.DataFrame(_export_metadata_rows()),
    }
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    safe_source = re.sub(r'[^A-Za-z0-9_.-]+', '_', audit['source_file'] or f"bill_{bill_id}")
    path = os.path.join(app.config['EXPORT_FOLDER'], f"PDF_Audit_{safe_source}_{timestamp}.xlsx")
    _write_workbook(path, sheets)
    return send_file(path, as_attachment=True, download_name=os.path.basename(path))


@app.route('/api/v3/manual_qa_export', methods=['POST'])
def manual_qa_export():
    db = get_db()
    audits = build_pdf_audit_rows(db)
    health = build_phase1_health(db)
    active_where = "WHERE COALESCE(e.is_deleted,0) = 0"
    employee_data = _df_from_query(db, _employee_snapshot_sql(active_where) + " ORDER BY e.source_filename, e.sr_no")
    needs_review = _df_from_query(db, _employee_snapshot_sql(active_where + " AND (COALESCE(e.designation_needs_review,0) = 1 OR COALESCE(e.employee_data_status,'') IN ('NEEDS_REVIEW','MISSING_CRITICAL_FIELD','COMPONENT_REVIEW_REQUIRED'))") + " ORDER BY e.source_filename, e.sr_no")
    missing_critical = _df_from_query(db, _employee_snapshot_sql(active_where + " AND (COALESCE(e.employee_id,'') = '' OR COALESCE(e.pay_scale_raw,'') = '' OR COALESCE(e.pan_raw,'') = '')") + " ORDER BY e.source_filename, e.sr_no")
    component_review = _df_from_query(db, _employee_snapshot_sql(active_where + " AND (ABS((SELECT COALESCE(SUM(c.amount),0) FROM employee_components_v2 c WHERE c.employee_id = e.id AND c.component_type = 'allowance') - COALESCE(e.gross_amount,0)) > 1 OR ABS((SELECT COALESCE(SUM(c.amount),0) FROM employee_components_v2 c WHERE c.employee_id = e.id AND c.component_type = 'deduction' AND COALESCE(c.included_in_total,1) = 1) - COALESCE(e.sum_deduction,0)) > 1)") + " ORDER BY e.source_filename, e.sr_no")
    designation_review = _df_from_query(db, _employee_snapshot_sql(active_where + " AND COALESCE(e.designation_needs_review,0) = 1") + " ORDER BY e.source_filename, e.sr_no")
    dictionary_mappings = _df_from_query(db, """
        SELECT id, category, raw_text, raw_text_key, COALESCE(final_text, final_text_en, '') AS final_text,
               match_type, priority, active, created_at, updated_at, notes
        FROM dictionary_mappings_v3
        ORDER BY category, active DESC, priority DESC, raw_text
    """)
    deleted_rows = _df_from_query(db, _employee_snapshot_sql("WHERE COALESCE(e.is_deleted,0) = 1") + " ORDER BY e.source_filename, e.sr_no")
    manual_edits = _df_from_query(db, """
        SELECT *
        FROM edit_audit_log_v3
        WHERE action_type IN ('employee_edit','employee_component_edit','employee_mark_verified','EMPLOYEE_ROW_DELETE','EMPLOYEE_ROW_RESTORE','dictionary_designation_apply')
        ORDER BY id DESC
    """)
    sheets = {
        'Phase 1 Health Check': pd.DataFrame([health]),
        'PDF Audit Summary': pd.DataFrame(audits),
        'Employee Data': employee_data,
        'Needs Review': needs_review,
        'Missing Critical Fields': missing_critical,
        'Component Review': component_review,
        'Designation Review': designation_review,
        'Dictionary Mappings': dictionary_mappings,
        'Deleted Rows': deleted_rows,
        'Manual Edits Audit': manual_edits,
        'Export Metadata': pd.DataFrame(_export_metadata_rows()),
    }
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    path = os.path.join(app.config['EXPORT_FOLDER'], f"Manual_QA_Export_{timestamp}.xlsx")
    _write_workbook(path, sheets)
    return send_file(path, as_attachment=True, download_name=os.path.basename(path))



@app.route('/api/group_summary')
def group_summary():
    db = get_db()
    if not app.config.get('USE_V2_PARSER', False): return jsonify({'groups': []})
    
    month = request.args.get('month', '')
    params = []
    conditions = ['COALESCE(e.is_deleted, 0) = 0']
    if month:
        conditions.append('e.month_year = ?')
        params = [month]
    where = ' WHERE ' + ' AND '.join(conditions)
    
    query = f"""
    SELECT e.group_name,
        COUNT(e.id) as emp_count,
        SUM(e.gross_amount) as total_gross,
        SUM(e.net_total) as total_net,
        SUM(e.sum_deduction) as total_deduction,
        AVG(e.net_total) as avg_net,
        MAX(b.official_gross) as gt_gross,
        MAX(b.official_deduction) as gt_ded,
        MAX(b.official_net) as gt_net,
        MAX(b.verification_status) as validation,
        (SELECT COALESCE(SUM(amount),0) FROM employee_components_v2 c JOIN employees_v2 e2 ON c.employee_id=e2.id WHERE e2.group_name=e.group_name AND COALESCE(e2.is_deleted,0)=0 AND c.normalized_head='basic_pay') as total_basic,
        (SELECT COALESCE(SUM(amount),0) FROM employee_components_v2 c JOIN employees_v2 e2 ON c.employee_id=e2.id WHERE e2.group_name=e.group_name AND COALESCE(e2.is_deleted,0)=0 AND c.normalized_head='income_tax') as total_itax,
        (SELECT COALESCE(SUM(amount),0) FROM employee_components_v2 c JOIN employees_v2 e2 ON c.employee_id=e2.id WHERE e2.group_name=e.group_name AND COALESCE(e2.is_deleted,0)=0 AND c.normalized_head='gpf') as total_gpf,
        (SELECT COALESCE(SUM(amount),0) FROM employee_components_v2 c JOIN employees_v2 e2 ON c.employee_id=e2.id WHERE e2.group_name=e.group_name AND COALESCE(e2.is_deleted,0)=0 AND c.normalized_head='lic') as total_lic,
        (SELECT COALESCE(SUM(amount),0) FROM employee_components_v2 c JOIN employees_v2 e2 ON c.employee_id=e2.id WHERE e2.group_name=e.group_name AND COALESCE(e2.is_deleted,0)=0 AND c.normalized_head='si_premium') as total_sip,
        (SELECT COALESCE(SUM(amount),0) FROM employee_components_v2 c JOIN employees_v2 e2 ON c.employee_id=e2.id WHERE e2.group_name=e.group_name AND COALESCE(e2.is_deleted,0)=0 AND c.normalized_head='si_loan') as total_sil,
        (SELECT COALESCE(SUM(amount),0) FROM employee_components_v2 c JOIN employees_v2 e2 ON c.employee_id=e2.id WHERE e2.group_name=e.group_name AND COALESCE(e2.is_deleted,0)=0 AND c.normalized_head='rghs') as total_rghs
    FROM employees_v2 e
    LEFT JOIN bills_v2 b ON e.bill_id = b.id
    {where}
    GROUP BY e.group_name
    ORDER BY e.group_name
    """
    rows = db.execute(query, params).fetchall()
    return jsonify({'groups': [dict(r) for r in rows]})



@app.route('/api/validation')
def get_validation():
    """Return validation details for all uploads."""
    # Legacy V1 route: backed by `uploads` table. Keep for compatibility during V2 transition.
    db = get_db()
    rows = db.execute('''SELECT group_name, employee_count,
        grand_total_gross, grand_total_deduction, grand_total_net,
        parsed_gross, parsed_deduction, parsed_net,
        validation_status, upload_time
    FROM uploads ORDER BY upload_time DESC''').fetchall()
    return jsonify({'uploads': [dict(r) for r in rows]})


@app.route('/api/export')
@app.route('/api/export_filtered')
def export_excel():
    db = get_db()
    if not app.config.get('USE_V2_PARSER', False): return jsonify({'error': 'V1 export disabled'}), 400

    conditions, params = build_employee_filter_conditions(request.args)
    where = ' WHERE ' + ' AND '.join(conditions) if conditions else ''
    display_mode = (request.args.get('designation_display') or 'english').strip().lower().replace('_', '-')

    emp_rows = db.execute(f"""
        SELECT e.*, b.bill_no AS bill_no, b.verification_status AS bill_verification_status,
               {employee_status_case_sql()} AS employee_data_status_calc
        FROM employees_v2 e
        LEFT JOIN bills_v2 b ON b.id = e.bill_id
        {where}
        ORDER BY e.source_filename, e.group_name, e.sr_no
    """, params).fetchall()
    if not emp_rows:
        return jsonify({'error': 'No data to export'}), 404

    emp_ids = [r['id'] for r in emp_rows]
    bill_ids = sorted({r['bill_id'] for r in emp_rows if r['bill_id'] is not None})
    emp_ph = ','.join('?' * len(emp_ids))
    bill_ph = ','.join('?' * len(bill_ids)) if bill_ids else 'NULL'

    component_columns = [
        ('basic_pay', 'Basic'), ('da', 'DA'), ('hra', 'HRA'), ('cca', 'CCA'),
        ('hda', 'HDA'), ('mess_allowance', 'Mess Allowance'),
        ('kma_allowance', 'KMA Allowance'), ('fta', 'FTA'), ('mta', 'MTA'),
        ('special_pay', 'Special Pay'), ('wash_allowance', 'Wash Allowance'),
        ('rop', 'ROP'), ('lic', 'LIC'), ('gpf', 'GPF'), ('gpf_2004', 'GPF 2004'),
        ('rghs', 'RGHS'), ('sip', 'SIP'), ('sil', 'SIL'), ('sii', 'SII'),
        ('income_tax', 'Income Tax'), ('roadways_deduction', 'Roadways Deduction'),
        ('advance_salary', 'Advance Salary'), ('ppss_jdr', 'PPSS JDR'),
        ('car_rent', 'Car Rent'), ('hfrh', 'HFRH'), ('eacpenf', 'EACPenF'),
        ('giais', 'GIAIS'), ('ais_gpf', 'AIS GPF'),
    ]
    component_label_by_head = dict(component_columns)

    comp_rows = db.execute(f"""
        SELECT c.*, e.source_filename, e.page_no, e.group_name, e.sr_no,
               e.raw_name, e.employee_id AS employee_code, e.normalized_designation,
               e.gpf_pran_raw
        FROM employee_components_v2 c
        JOIN employees_v2 e ON e.id = c.employee_id
        WHERE c.employee_id IN ({emp_ph})
        ORDER BY e.source_filename, e.sr_no, c.component_type, c.id
    """, emp_ids).fetchall()

    comp_map = {}
    dynamic_heads = []
    seen_dynamic_heads = set()
    for comp in comp_rows:
        head = comp['normalized_head'] or ''
        comp_map.setdefault(comp['employee_id'], {})
        comp_map[comp['employee_id']][head] = comp_map[comp['employee_id']].get(head, 0) + (comp['amount'] or 0)
        if head and head not in component_label_by_head and head not in seen_dynamic_heads:
            seen_dynamic_heads.add(head)
            dynamic_heads.append(head)

    issue_details = {}
    issue_rows = db.execute(f"""
        SELECT employee_id, issue_type, severity, message
        FROM validation_issues_v2
        WHERE employee_id IN ({emp_ph}) AND COALESCE(status, 'OPEN') != 'RESOLVED'
        ORDER BY employee_id, severity DESC, issue_type
    """, emp_ids).fetchall()
    for row in issue_rows:
        issue_details.setdefault(row['employee_id'], []).append(dict(row))

    bill_register = db.execute(f"""
        SELECT b.source_filename AS source_file, b.group_name, b.bill_no, b.bill_date,
               b.reference_no, b.month_year, b.office_id, b.ddo_name, b.ddo_code,
               b.demand_no, b.object_head, b.budget_head_full AS budget_head,
               b.sf_amount, b.ca_amount, b.na_amount, b.tan_no, b.page_count,
               COUNT(e.id) AS parsed_employee_count,
               b.parsed_gross, b.parsed_deduction, b.parsed_net,
               b.official_gross, b.official_deduction, b.official_net,
               b.verification_status AS reconciliation_status,
               CASE
                 WHEN COUNT(v.id) > 0 THEN 'HAS_OPEN_ISSUES'
                 WHEN b.verification_status = 'VERIFIED' THEN 'VERIFIED'
                 ELSE COALESCE(b.verification_status, '')
               END AS validation_status
        FROM bills_v2 b
        LEFT JOIN employees_v2 e ON e.bill_id = b.id AND e.id IN ({emp_ph})
        LEFT JOIN validation_issues_v2 v ON v.bill_id = b.id AND COALESCE(v.status, 'OPEN') != 'RESOLVED'
        WHERE b.id IN ({bill_ph})
        GROUP BY b.id
        ORDER BY CAST(REPLACE(b.source_filename, '.pdf', '') AS INTEGER), b.source_filename
    """, emp_ids + bill_ids).fetchall() if bill_ids else []

    employee_master = []
    for emp in emp_rows:
        issues = issue_details.get(emp['id'], [])
        status = emp['employee_data_status_calc'] or 'VERIFIED'
        final_designation = emp['final_designation'] or emp['final_designation_en'] or emp['normalized_designation'] or ''
        designation_parts = extract_designation_parts(emp['raw_designation_full'] or emp['raw_designation'] or '')
        raw_designation_en = emp['raw_designation_en'] or designation_parts['raw_designation_en']
        raw_designation_hi = emp['raw_designation_hi'] or designation_parts['raw_designation_hi']
        final_designation_hi_display = emp['final_designation_hi_display'] or emp['final_designation_hi'] or raw_designation_hi
        if display_mode in ('english+hindi', 'english-hindi', 'combined'):
            designation_display = f"{final_designation} / {final_designation_hi_display}" if final_designation_hi_display else final_designation
        elif display_mode in ('raw', 'raw-source', 'source'):
            designation_display = emp['raw_designation_full'] or emp['raw_designation'] or final_designation
        else:
            designation_display = final_designation
        row = {
            'source_file': emp['source_filename'] or '',
            'page_no': emp['page_no'] or 0,
            'group_name': emp['group_name'] or '',
            'bill_no': emp['bill_no'] or '',
            'month_year': emp['month_year'] or '',
            'sr_no': emp['sr_no'] or 0,
            'employee_name': emp['raw_name'] or emp['normalized_name'] or '',
            'employee_id': emp['employee_id'] or '',
            'gpf_pran': emp['gpf_pran_raw'] or '',
            'dob': emp['dob'] or '',
            'belt_no': emp['belt_no'] or '',
            'pan': emp['pan_raw'] or '',
            'pay_scale': emp['pay_scale_raw'] or '',
            'raw_designation': emp['raw_designation'] or '',
            'raw_designation_full': emp['raw_designation_full'] or emp['raw_designation'] or '',
            'raw_designation_en': raw_designation_en,
            'raw_designation_hi': raw_designation_hi,
            'final_designation': final_designation,
            'final_designation_en': final_designation,
            'final_designation_hi': emp['final_designation_hi'] or '',
            'final_designation_hi_display': final_designation_hi_display,
            'designation_display': designation_display,
            'designation_source': emp['designation_source'] or '',
            'designation_needs_review': bool(emp['designation_needs_review']),
            'normalized_designation': emp['normalized_designation'] or '',
            'bank_account': emp['bank_acc'] or '',
            'aadhaar_masked': emp['aadhar_raw'] or '',
            'printed_gross': emp['gross_amount'] or 0,
            'printed_deduction': emp['sum_deduction'] or 0,
            'printed_net': emp['net_total'] or 0,
            'validation_status': status,
            'employee_data_status': status,
        }
        cmap = comp_map.get(emp['id'], {})
        for head, label in component_columns:
            row[label] = cmap.get(head, '')
        row['Gross'] = emp['gross_amount'] or 0
        row['Total Deduction'] = emp['sum_deduction'] or 0
        row['Net Total'] = emp['net_total'] or 0
        row['ITax'] = cmap.get('income_tax', '')
        for head in dynamic_heads:
            row[f"Component: {head}"] = cmap.get(head, '')
        employee_master.append(row)

    allowance_ledger = []
    deduction_ledger = []
    for comp in comp_rows:
        base = {
            'source_file': comp['source_filename'] or '',
            'page_no': comp['source_page'] or comp['page_no'] or 0,
            'group_name': comp['group_name'] or '',
            'sr_no': comp['sr_no'] or 0,
            'employee_name': comp['raw_name'] or '',
            'employee_id': comp['employee_code'] or '',
            'raw_head': comp['raw_head'] or '',
            'normalized_head': comp['normalized_head'] or '',
            'amount': comp['amount'] or 0,
            'confidence': comp['confidence'] if comp['confidence'] is not None else '',
            'source_type': 'inner_sheet_component',
        }
        if comp['component_type'] == 'deduction':
            drow = dict(base)
            drow['included_in_total'] = comp['included_in_total']
            deduction_ledger.append(drow)
        else:
            allowance_ledger.append(base)

    schedule_detail_rows = []
    for comp in comp_rows:
        schedule_type = DERIVED_COMPONENT_SCHEDULES.get(comp['normalized_head'] or '')
        if not schedule_type:
            continue
        if not is_valid_derived_schedule_component(comp['normalized_head'], comp['amount'], db.execute(
            'SELECT raw_text_block FROM employees_v2 WHERE id = ?', (comp['employee_id'],)
        ).fetchone()['raw_text_block']):
            continue
        schedule_detail_rows.append({
            'schedule_type': schedule_type,
            'normalized_head': comp['normalized_head'] or '',
            'group_name': comp['group_name'] or '',
            'sr_no': comp['sr_no'] or 0,
            'employee_name': comp['raw_name'] or '',
            'employee_id': comp['employee_code'] or '',
            'designation': comp['normalized_designation'] or '',
            'basic_pay': comp_map.get(comp['employee_id'], {}).get('basic_pay', ''),
            'gpf_pran': comp['gpf_pran_raw'] or '',
            'component_amount': comp['amount'] or 0,
            'source_file': comp['source_filename'] or '',
            'source_page': comp['source_page'] or comp['page_no'] or 0,
            'source_type': 'derived_from_inner_sheet',
        })
    schedule_summary_map = {}
    for row in schedule_detail_rows:
        key = (row['schedule_type'], row['normalized_head'])
        schedule_summary_map.setdefault(key, {'employee_ids': set(), 'total_amount': 0})
        schedule_summary_map[key]['employee_ids'].add(row['employee_id'])
        schedule_summary_map[key]['total_amount'] += row['component_amount']
    schedule_summary_rows = [
        {
            'schedule_type': key[0],
            'normalized_head': key[1],
            'employee_count': len(data['employee_ids']),
            'total_amount': data['total_amount'],
            'source_type': 'derived_from_inner_sheet',
        }
        for key, data in sorted(schedule_summary_map.items())
    ]

    group_summary = db.execute(f"""
        SELECT e.group_name, COUNT(e.id) AS employee_count,
               COALESCE(SUM(e.gross_amount),0) AS gross,
               COALESCE(SUM(e.sum_deduction),0) AS deduction,
               COALESCE(SUM(e.net_total),0) AS net,
               COALESCE(SUM(DISTINCT b.official_gross),0) AS official_gross,
               COALESCE(SUM(DISTINCT b.official_deduction),0) AS official_deduction,
               COALESCE(SUM(DISTINCT b.official_net),0) AS official_net,
               CASE
                 WHEN ROUND(COALESCE(SUM(e.gross_amount),0)) = ROUND(COALESCE(SUM(DISTINCT b.official_gross),0))
                  AND ROUND(COALESCE(SUM(e.sum_deduction),0)) = ROUND(COALESCE(SUM(DISTINCT b.official_deduction),0))
                  AND ROUND(COALESCE(SUM(e.net_total),0)) = ROUND(COALESCE(SUM(DISTINCT b.official_net),0))
                 THEN 'VERIFIED' ELSE 'FAILED_RECONCILIATION'
               END AS reconciliation_status
        FROM employees_v2 e
        LEFT JOIN bills_v2 b ON b.id = e.bill_id
        WHERE e.id IN ({emp_ph})
        GROUP BY e.group_name
        ORDER BY e.group_name
    """, emp_ids).fetchall()

    designation_summary = db.execute(f"""
        SELECT COALESCE(NULLIF(e.normalized_designation,''), 'Unmapped') AS normalized_designation,
               COUNT(DISTINCT e.raw_designation) AS raw_designation_count,
               COUNT(e.id) AS employee_count,
               COALESCE(SUM(e.gross_amount),0) AS gross,
               COALESCE(SUM(e.sum_deduction),0) AS deduction,
               COALESCE(SUM(e.net_total),0) AS net
        FROM employees_v2 e
        WHERE e.id IN ({emp_ph})
        GROUP BY normalized_designation
        ORDER BY normalized_designation
    """, emp_ids).fetchall()

    pay_scale_summary = db.execute(f"""
        SELECT COALESCE(NULLIF(e.pay_scale_raw,''), 'Unmapped') AS pay_scale,
               COUNT(e.id) AS employee_count,
               COALESCE(SUM(e.gross_amount),0) AS gross,
               COALESCE(SUM(e.sum_deduction),0) AS deduction,
               COALESCE(SUM(e.net_total),0) AS net
        FROM employees_v2 e
        WHERE e.id IN ({emp_ph})
        GROUP BY pay_scale
        ORDER BY pay_scale
    """, emp_ids).fetchall()

    budget_head_summary = db.execute(f"""
        SELECT COALESCE(NULLIF(b.budget_head_full,''), 'Unmapped') AS budget_head,
               COUNT(DISTINCT b.group_name) AS group_count,
               COUNT(e.id) AS employee_count,
               COALESCE(SUM(e.gross_amount),0) AS gross,
               COALESCE(SUM(e.sum_deduction),0) AS deduction,
               COALESCE(SUM(e.net_total),0) AS net
        FROM employees_v2 e
        JOIN bills_v2 b ON b.id = e.bill_id
        WHERE e.id IN ({emp_ph})
        GROUP BY budget_head
        ORDER BY budget_head
    """, emp_ids).fetchall()

    validation_report = db.execute(f"""
        SELECT v.issue_type, v.severity, b.source_filename AS source_file, b.group_name,
               e.page_no, e.sr_no, e.raw_name AS employee_name, e.employee_id,
               v.message, v.status, v.created_at, v.resolved_at
        FROM validation_issues_v2 v
        LEFT JOIN bills_v2 b ON b.id = v.bill_id
        LEFT JOIN employees_v2 e ON e.id = v.employee_id
        WHERE (v.employee_id IN ({emp_ph}) OR v.bill_id IN ({bill_ph}))
        ORDER BY v.severity DESC, b.source_filename, e.sr_no, v.id
    """, emp_ids + bill_ids).fetchall()

    dictionary_review = []
    dict_terms = db.execute("""
        SELECT raw_value AS raw_text, dictionary_type AS category, '' AS suggested_mapping,
               normalized_value, '' AS first_seen_group, '' AS first_seen_employee,
               CASE WHEN active = 1 THEN 'ACTIVE' ELSE 'INACTIVE' END AS status
        FROM dictionary_terms_v2
        ORDER BY dictionary_type, raw_value
    """).fetchall()
    dictionary_review.extend([dict(r) for r in dict_terms])
    unknown_rows = db.execute(f"""
        SELECT v.raw_value AS raw_text, v.issue_type AS category, v.suggested_value AS suggested_mapping,
               v.suggested_value AS normalized_value, b.group_name AS first_seen_group,
               e.raw_name AS first_seen_employee, v.status
        FROM validation_issues_v2 v
        LEFT JOIN bills_v2 b ON b.id = v.bill_id
        LEFT JOIN employees_v2 e ON e.id = v.employee_id
        WHERE v.issue_type IN ('UNKNOWN_COMPONENT_HEAD','UNMAPPED_DESIGNATION')
          AND (v.employee_id IN ({emp_ph}) OR v.bill_id IN ({bill_ph}))
        ORDER BY v.issue_type, v.raw_value
    """, emp_ids + bill_ids).fetchall()
    dictionary_review.extend([dict(r) for r in unknown_rows])

    all_verified = all((r['reconciliation_status'] == 'VERIFIED') for r in bill_register)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    fname = f"IFMS_V2_Verified_Export_{timestamp}.xlsx" if all_verified else f"IFMS_V2_DRAFT_NOT_VERIFIED_Export_{timestamp}.xlsx"
    fpath = os.path.join(app.config['EXPORT_FOLDER'], fname)

    metadata = [
        {'key': 'app_name', 'value': APP_NAME},
        {'key': 'developer_credit', 'value': DEVELOPER_CREDIT},
        {'key': 'developer_email', 'value': DEVELOPER_EMAIL},
        {'key': 'disclaimer', 'value': APP_DISCLAIMER},
        {'key': 'export_timestamp', 'value': datetime.now().isoformat(timespec='seconds')},
        {'key': 'database', 'value': app.config['DATABASE']},
        {'key': 'bill_count', 'value': len(bill_register)},
        {'key': 'employee_count', 'value': len(employee_master)},
        {'key': 'all_bills_verified', 'value': all_verified},
        {'key': 'active_filters', 'value': request.query_string.decode('utf-8')},
    ]
    applied_filters = [{'filter': k, 'value': v} for k, values in request.args.lists() for v in values]
    if not applied_filters:
        applied_filters = [{'filter': 'none', 'value': ''}]
    quality_summary = employee_quality_summary(db, where, params)
    quality_rows = [{'metric': k, 'value': v} for k, v in quality_summary.items() if k != 'status_counts']

    sheets = {
        'Export Metadata': pd.DataFrame(metadata),
        'Filtered Employee Data': pd.DataFrame(employee_master),
        'Applied Filters': pd.DataFrame(applied_filters),
        'Data Quality Summary': pd.DataFrame(quality_rows),
        'Bill Register': pd.DataFrame([dict(r) for r in bill_register]),
        'Employee Master': pd.DataFrame(employee_master),
        'Allowance Ledger': pd.DataFrame(allowance_ledger),
        'Deduction Ledger': pd.DataFrame(deduction_ledger),
        'Derived Schedule Summary': pd.DataFrame(schedule_summary_rows),
        'Derived Schedule Details': pd.DataFrame(schedule_detail_rows),
        'Group Summary': pd.DataFrame([dict(r) for r in group_summary]),
        'Designation Summary': pd.DataFrame([dict(r) for r in designation_summary]),
        'Pay Scale Summary': pd.DataFrame([dict(r) for r in pay_scale_summary]),
        'Budget Head Summary': pd.DataFrame([dict(r) for r in budget_head_summary]),
        'Validation Report': pd.DataFrame([dict(r) for r in validation_report]),
        'Unknown Dictionary Review': pd.DataFrame(dictionary_review),
    }
    expected_columns = {
        'Filtered Employee Data': ['source_file','group_name','bill_no','month_year','page_no','sr_no','employee_name','employee_id','pan','pay_scale','raw_designation_full','raw_designation_en','raw_designation_hi','final_designation','final_designation_hi_display','designation_display','designation_source','employee_data_status','Basic','DA','HRA','CCA','HDA','Mess Allowance','FTA','ROP','Gross','Total Deduction','Net Total','LIC','GPF','GPF 2004','AIS GPF','RGHS','SIP','SIL','SII','ITax','Income Tax','Roadways Deduction','Advance Salary','PPSS JDR','EACPenF','GIAIS','Car Rent','HFRH'],
        'Applied Filters': ['filter','value'],
        'Data Quality Summary': ['metric','value'],
        'Bill Register': ['source_file','group_name','bill_no','bill_date','reference_no','month_year','office_id','ddo_name','ddo_code','demand_no','object_head','budget_head','sf_amount','ca_amount','na_amount','tan_no','page_count','parsed_employee_count','parsed_gross','parsed_deduction','parsed_net','official_gross','official_deduction','official_net','reconciliation_status','validation_status'],
        'Employee Master': ['source_file','page_no','group_name','bill_no','month_year','sr_no','employee_name','employee_id','gpf_pran','dob','belt_no','pan','pay_scale','raw_designation','raw_designation_full','raw_designation_en','raw_designation_hi','final_designation','final_designation_en','final_designation_hi','final_designation_hi_display','designation_display','designation_source','designation_needs_review','normalized_designation','bank_account','aadhaar_masked','printed_gross','printed_deduction','printed_net','validation_status','employee_data_status'] + [label for _, label in component_columns] + [f"Component: {head}" for head in dynamic_heads],
        'Allowance Ledger': ['source_file','page_no','group_name','sr_no','employee_name','employee_id','raw_head','normalized_head','amount','confidence','source_type'],
        'Deduction Ledger': ['source_file','page_no','group_name','sr_no','employee_name','employee_id','raw_head','normalized_head','amount','included_in_total','source_type','confidence'],
        'Derived Schedule Summary': ['schedule_type','normalized_head','employee_count','total_amount','source_type'],
        'Derived Schedule Details': ['schedule_type','normalized_head','group_name','sr_no','employee_name','employee_id','designation','basic_pay','gpf_pran','component_amount','source_file','source_page','source_type'],
        'Group Summary': ['group_name','employee_count','gross','deduction','net','official_gross','official_deduction','official_net','reconciliation_status'],
        'Designation Summary': ['normalized_designation','raw_designation_count','employee_count','gross','deduction','net'],
        'Pay Scale Summary': ['pay_scale','employee_count','gross','deduction','net'],
        'Budget Head Summary': ['budget_head','group_count','employee_count','gross','deduction','net'],
        'Validation Report': ['issue_type','severity','source_file','group_name','page_no','sr_no','employee_name','employee_id','message','status','created_at','resolved_at'],
        'Unknown Dictionary Review': ['raw_text','category','suggested_mapping','normalized_value','first_seen_group','first_seen_employee','status'],
    }
    for sheet_name, columns in expected_columns.items():
        if sheets[sheet_name].empty:
            sheets[sheet_name] = pd.DataFrame(columns=columns)
        else:
            for col in columns:
                if col not in sheets[sheet_name].columns:
                    sheets[sheet_name][col] = ''
            sheets[sheet_name] = sheets[sheet_name][columns]

    with pd.ExcelWriter(fpath, engine='openpyxl') as writer:
        for sheet_name, df in sheets.items():
            safe_name = sheet_name[:31]
            df.to_excel(writer, index=False, sheet_name=safe_name)
            ws = writer.sheets[safe_name]
            ws.freeze_panes = 'A2'
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F4E78')
            amount_headers = {'amount','component_amount','total_amount','gross','deduction','net','Gross','Total Deduction','Net Total','parsed_gross','parsed_deduction','parsed_net','official_gross','official_deduction','official_net','printed_gross','printed_deduction','printed_net','sf_amount','ca_amount','na_amount','basic_pay'}
            for idx, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
                header = str(col[0].value or '')
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 32)
                if header in amount_headers or header in [label for _, label in component_columns] or header.startswith('Component: '):
                    for cell in col[1:]:
                        cell.number_format = '#,##,##0.00'
            if safe_name in ('Bill Register', 'Group Summary'):
                status_cols = [cell.column for cell in ws[1] if str(cell.value or '').endswith('status')]
                for row in ws.iter_rows(min_row=2):
                    for col_idx in status_cols:
                        value = str(row[col_idx - 1].value or '').upper()
                        if 'VERIFIED' in value:
                            row[col_idx - 1].fill = PatternFill('solid', fgColor='C6EFCE')
                        elif 'FAILED' in value or 'ISSUES' in value:
                            row[col_idx - 1].fill = PatternFill('solid', fgColor='FFC7CE')
            if safe_name == 'Validation Report':
                severity_col = next((cell.column for cell in ws[1] if cell.value == 'severity'), None)
                if severity_col:
                    for row in ws.iter_rows(min_row=2):
                        severity = str(row[severity_col - 1].value or '').upper()
                        if severity in ('CRITICAL', 'HIGH'):
                            row[severity_col - 1].fill = PatternFill('solid', fgColor='FFC7CE')
                        elif severity in ('MEDIUM', 'LOW'):
                            row[severity_col - 1].fill = PatternFill('solid', fgColor='FFEB9C')
            if safe_name in ('Employee Master', 'Filtered Employee Data'):
                status_col = next((cell.column for cell in ws[1] if cell.value in ('employee_data_status', 'validation_status')), None)
                if status_col:
                    for row in ws.iter_rows(min_row=2):
                        status = str(row[status_col - 1].value or '').upper()
                        if status in ('NEEDS_REVIEW', 'MISSING_CRITICAL_FIELD', 'COMPONENT_REVIEW_REQUIRED'):
                            for cell in row:
                                cell.fill = PatternFill('solid', fgColor='FFC7CE')
                        elif status == 'VERIFIED_WITH_WARNINGS':
                            for cell in row:
                                cell.fill = PatternFill('solid', fgColor='FFEB9C')
                        elif status == 'MANUALLY_EDITED':
                            for cell in row:
                                cell.fill = PatternFill('solid', fgColor='D9EAD3')

    return send_file(fpath, as_attachment=True, download_name=fname)



# ===================== OFFICIAL SCHEDULE COMPARISON APIs =====================

def _official_schedule_upload_folder():
    folder = Path(app.config.get('OFFICIAL_SCHEDULE_UPLOAD_FOLDER') or os.path.join(app.config['UPLOAD_FOLDER'], 'official_schedules'))
    folder.mkdir(parents=True, exist_ok=True)
    return folder.resolve()


def _safe_official_upload_path(file_name):
    raw = str(file_name or '').strip()
    if not raw:
        raise ValueError('file_name is required')
    if raw != Path(raw).name or os.path.isabs(raw) or '..' in Path(raw).parts:
        raise ValueError('Invalid official schedule file name')
    if Path(raw).suffix.lower() != '.pdf':
        raise ValueError('Only PDF files are supported')
    path = (_official_schedule_upload_folder() / raw).resolve()
    if _official_schedule_upload_folder() not in path.parents:
        raise ValueError('Invalid official schedule file path')
    return path


def _official_args():
    data = dict(request.args)
    if request.is_json:
        data.update(request.get_json(silent=True) or {})
    return data


@app.route('/api/v3/official_schedules/upload', methods=['POST'])
def v3_official_schedule_upload():
    if 'file' not in request.files:
        return jsonify({'error': 'file is required'}), 400
    file = request.files['file']
    original_name = file.filename or ''
    if not original_name.lower().endswith('.pdf'):
        return jsonify({'error': 'Only official schedule PDF files are supported'}), 400
    safe_base = secure_filename(original_name) or 'official_schedule.pdf'
    source_file = f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{safe_base}"
    try:
        path = _safe_official_upload_path(source_file)
        file.save(path)
        classification = classify_official_schedule_pdf(path)
        if classification.get('classification') != 'official_schedule':
            path.unlink(missing_ok=True)
            return jsonify({'error': 'Uploaded PDF is not an official schedule PDF', 'classification': classification}), 400
        db = get_db()
        summary = store_official_upload(db, path, original_name, source_file)
        db.commit()
        return jsonify({'success': True, **summary})
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400
    except Exception as exc:
        return jsonify({'error': str(exc)}), 500


@app.route('/api/v3/official_schedules/uploads', methods=['GET'])
def v3_official_schedule_uploads():
    db = get_db()
    init_official_schedule_schema(db)
    rows = [dict(row) for row in db.execute("""
        SELECT *
        FROM official_schedule_uploads_v3
        ORDER BY id DESC
    """).fetchall()]
    return jsonify({'uploads': rows})


@app.route('/api/v3/official_schedules/sections', methods=['GET'])
def v3_official_schedule_sections():
    db = get_db()
    init_official_schedule_schema(db)
    args = _official_args()
    where, params = [], []
    if args.get('upload_id'):
        where.append('upload_id = ?')
        params.append(args.get('upload_id'))
    if args.get('schedule_type'):
        where.append('schedule_type = ?')
        params.append(args.get('schedule_type'))
    sql = "SELECT * FROM official_schedule_sections_v3"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY upload_id DESC, start_page, id"
    return jsonify({'sections': [dict(row) for row in db.execute(sql, params).fetchall()]})


@app.route('/api/v3/official_schedules/rows', methods=['GET'])
def v3_official_schedule_rows():
    db = get_db()
    init_official_schedule_schema(db)
    args = _official_args()
    where, params = [], []
    for key in ('upload_id', 'section_id', 'schedule_type', 'row_status'):
        if args.get(key):
            where.append(f"{key} = ?")
            params.append(args.get(key))
    search = str(args.get('search') or '').strip()
    if search:
        where.append('(employee_id LIKE ? OR employee_name LIKE ? OR pan LIKE ? OR gpf_no LIKE ? OR policy_no LIKE ?)')
        params.extend([f'%{search}%'] * 5)
    sql = "SELECT * FROM official_schedule_rows_v3"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY schedule_type, sr_no, id"
    return jsonify({'rows': [dict(row) for row in db.execute(sql, params).fetchall()]})


@app.route('/api/v3/official_schedules/compare', methods=['POST'])
def v3_official_schedule_compare():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        upload_id = int(data.get('upload_id') or 0)
    except (TypeError, ValueError):
        return jsonify({'error': 'upload_id must be numeric'}), 400
    if not upload_id:
        return jsonify({'error': 'upload_id is required'}), 400
    summary = run_official_schedule_comparison(db, upload_id, data)
    db.commit()
    rows = official_comparison_rows(db, {
        'upload_id': upload_id,
        'section_id': data.get('section_id') or '',
        'schedule_type': data.get('schedule_type') or '',
    })
    return jsonify({'success': True, 'summary': summary, 'comparison': rows})


@app.route('/api/v3/official_schedules/comparison', methods=['GET'])
def v3_official_schedule_comparison():
    db = get_db()
    init_official_schedule_schema(db)
    args = _official_args()
    rows = official_comparison_rows(db, args)
    summary = {}
    if args.get('upload_id'):
        try:
            upload_id = int(args.get('upload_id'))
            section_id = int(args.get('section_id')) if args.get('section_id') else None
        except (TypeError, ValueError):
            return jsonify({'error': 'upload_id and section_id must be numeric'}), 400
        summary = official_comparison_summary(
            db,
            upload_id,
            section_id,
            args.get('schedule_type') or None,
        )
    return jsonify({'summary': summary, 'comparison': rows})


@app.route('/api/v3/official_schedules/export', methods=['POST'])
def v3_official_schedule_export():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        upload_id = int(data.get('upload_id') or 0)
    except (TypeError, ValueError):
        return jsonify({'error': 'upload_id must be numeric'}), 400
    if not upload_id:
        latest = db.execute("SELECT id FROM official_schedule_uploads_v3 ORDER BY id DESC LIMIT 1").fetchone()
        upload_id = latest['id'] if latest else 0
    if not upload_id:
        return jsonify({'error': 'No official schedule upload available for export'}), 404
    path = export_comparison_workbook(
        db,
        _export_folder(),
        upload_id,
        data,
        {
            'app_name': APP_NAME,
            'developer_email': DEVELOPER_EMAIL,
            'build_label': BUILD_LABEL,
            'disclaimer': APP_DISCLAIMER,
        },
    )
    return _send_export_file(path)


@app.route('/api/v3/official_schedules/delete_upload', methods=['POST'])
def v3_official_schedule_delete_upload():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        upload_id = int(data.get('upload_id') or 0)
    except (TypeError, ValueError):
        return jsonify({'error': 'upload_id must be numeric'}), 400
    if not upload_id:
        return jsonify({'error': 'upload_id is required'}), 400
    row = db.execute("SELECT source_file FROM official_schedule_uploads_v3 WHERE id = ?", (upload_id,)).fetchone()
    if not row:
        return jsonify({'error': 'Official schedule upload not found'}), 404
    try:
        path = _safe_official_upload_path(row['source_file'])
        if path.exists():
            path.unlink()
    except ValueError:
        return jsonify({'error': 'Stored official schedule file path is invalid'}), 400
    db.execute("DELETE FROM official_schedule_comparison_v3 WHERE upload_id = ?", (upload_id,))
    db.execute("DELETE FROM official_schedule_rows_v3 WHERE upload_id = ?", (upload_id,))
    db.execute("DELETE FROM official_schedule_sections_v3 WHERE upload_id = ?", (upload_id,))
    db.execute("DELETE FROM official_schedule_uploads_v3 WHERE id = ?", (upload_id,))
    db.commit()
    return jsonify({'success': True, 'deleted_upload_id': upload_id})


@app.route('/api/v3/official_schedules/mark_reviewed', methods=['POST'])
def v3_official_schedule_mark_reviewed():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        comparison_id = int(data.get('comparison_id') or data.get('id') or 0)
    except (TypeError, ValueError):
        return jsonify({'error': 'comparison_id must be numeric'}), 400
    if not comparison_id:
        return jsonify({'error': 'comparison_id is required'}), 400
    db.execute("""
        UPDATE official_schedule_comparison_v3
        SET reviewed_status = 'REVIEWED', reviewed_at = CURRENT_TIMESTAMP
        WHERE id = ?
    """, (comparison_id,))
    db.commit()
    return jsonify({'success': True, 'comparison_id': comparison_id})


@app.route('/api/v3/export_center/official_schedule_comparison', methods=['POST'])
def export_center_official_schedule_comparison():
    return v3_official_schedule_export()


# ===================== MONTHLY COMPARISON APIs =====================

@app.route('/api/v3/monthly_comparison/options', methods=['GET'])
def v3_monthly_comparison_options():
    db = get_db()
    init_monthly_comparison_schema(db)
    return jsonify(monthly_comparison_options(db))


@app.route('/api/v3/monthly_comparison/run', methods=['POST'])
def v3_monthly_comparison_run():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        summary = run_monthly_comparison(db, data)
        db.commit()
        rows = monthly_comparison_results(db, {'run_id': summary['run_id']})
        return jsonify({'success': True, 'summary': summary, 'results': rows})
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/monthly_comparison/runs', methods=['GET'])
def v3_monthly_comparison_runs():
    db = get_db()
    return jsonify({'runs': list_monthly_comparison_runs(db)})


@app.route('/api/v3/monthly_comparison/results', methods=['GET'])
def v3_monthly_comparison_results():
    db = get_db()
    return jsonify({'results': monthly_comparison_results(db, dict(request.args))})


@app.route('/api/v3/monthly_comparison/export', methods=['POST'])
def v3_monthly_comparison_export():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        run_id = int(data.get('run_id') or 0)
    except (TypeError, ValueError):
        return jsonify({'error': 'run_id must be numeric'}), 400
    if not run_id:
        latest = db.execute("SELECT id FROM monthly_comparison_runs_v3 ORDER BY id DESC LIMIT 1").fetchone()
        run_id = latest['id'] if latest else 0
    if not run_id:
        return jsonify({'error': 'No monthly comparison run available for export'}), 404
    try:
        path = export_monthly_comparison_workbook(
            db,
            _export_folder(),
            run_id,
            {
                'app_name': APP_NAME,
                'developer_email': DEVELOPER_EMAIL,
                'build_label': BUILD_LABEL,
                'disclaimer': APP_DISCLAIMER,
            },
        )
        return _send_export_file(path)
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 404


@app.route('/api/v3/monthly_comparison/delete_run', methods=['POST'])
def v3_monthly_comparison_delete_run():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        run_id = int(data.get('run_id') or 0)
    except (TypeError, ValueError):
        return jsonify({'error': 'run_id must be numeric'}), 400
    if not run_id:
        return jsonify({'error': 'run_id is required'}), 400
    delete_monthly_comparison_run(db, run_id)
    db.commit()
    return jsonify({'success': True, 'deleted_run_id': run_id})


@app.route('/api/v3/monthly_comparison/mark_reviewed', methods=['POST'])
def v3_monthly_comparison_mark_reviewed():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        result_id = int(data.get('result_id') or data.get('id') or 0)
    except (TypeError, ValueError):
        return jsonify({'error': 'result_id must be numeric'}), 400
    if not result_id:
        return jsonify({'error': 'result_id is required'}), 400
    updated = mark_monthly_comparison_reviewed(db, result_id)
    db.commit()
    return jsonify({'success': True, 'updated_count': updated, 'result_id': result_id})


@app.route('/api/v3/export_center/monthly_comparison', methods=['POST'])
def export_center_monthly_comparison():
    return v3_monthly_comparison_export()


# ===================== MASTER EMPLOYEE REGISTER APIs =====================

@app.route('/api/v3/master_register/build', methods=['POST'])
def v3_master_register_build():
    db = get_db()
    summary = build_master_register(db)
    db.commit()
    return jsonify({'success': True, 'summary': summary})


@app.route('/api/v3/master_register/update', methods=['POST'])
def v3_master_register_update():
    db = get_db()
    summary = update_master_register(db)
    db.commit()
    return jsonify({'success': True, 'summary': summary})


@app.route('/api/v3/master_register/summary', methods=['GET'])
def v3_master_register_summary():
    db = get_db()
    return jsonify({'summary': master_register_summary(db), 'options': master_register_options(db)})


@app.route('/api/v3/master_register/employees', methods=['GET'])
def v3_master_register_employees():
    db = get_db()
    return jsonify({'employees': master_register_employees(db, dict(request.args))})


@app.route('/api/v3/master_register/employees/<int:master_id>', methods=['GET'])
def v3_master_register_employee_detail(master_id):
    db = get_db()
    row = master_employee_detail(db, master_id)
    if not row:
        return jsonify({'error': 'Master employee not found'}), 404
    return jsonify({'employee': row})


@app.route('/api/v3/master_register/history/<int:master_id>', methods=['GET'])
def v3_master_register_history(master_id):
    db = get_db()
    return jsonify({'history': master_register_history(db, master_id)})


@app.route('/api/v3/master_register/issues', methods=['GET'])
def v3_master_register_issues():
    db = get_db()
    return jsonify({'issues': master_register_issues(db, dict(request.args))})


@app.route('/api/v3/master_register/issues/<int:issue_id>/resolve', methods=['POST'])
def v3_master_register_resolve_issue(issue_id):
    db = get_db()
    updated = resolve_master_register_issue(db, issue_id)
    db.commit()
    return jsonify({'success': True, 'updated_count': updated, 'issue_id': issue_id})


@app.route('/api/v3/master_register/archive/<int:master_id>', methods=['POST'])
def v3_master_register_archive(master_id):
    db = get_db()
    updated = archive_master_employee(db, master_id)
    db.commit()
    return jsonify({'success': True, 'updated_count': updated, 'master_employee_id': master_id})


@app.route('/api/v3/master_register/export', methods=['POST'])
def v3_master_register_export():
    db = get_db()
    path = export_master_register_workbook(
        db,
        _export_folder(),
        {
            'app_name': APP_NAME,
            'developer_email': DEVELOPER_EMAIL,
            'build_label': BUILD_LABEL,
            'disclaimer': APP_DISCLAIMER,
        },
    )
    return _send_export_file(path)


@app.route('/api/v3/export_center/master_register', methods=['POST'])
def export_center_master_register():
    return v3_master_register_export()


# ===================== SCHEDULE APIs =====================

def _schedule_v3_args():
    data = dict(request.args)
    if request.is_json:
        data.update(request.get_json(silent=True) or {})
    return data


@app.route('/api/v3/schedules/types', methods=['GET'])
def v3_schedule_types():
    db = get_db()
    return jsonify({'schedule_types': list_schedule_types_v3(db, dict(request.args))})


@app.route('/api/v3/schedules', methods=['GET'])
def v3_schedules():
    db = get_db()
    payload = query_schedule_entries_v3(db, _schedule_v3_args(), paginate=True)
    return jsonify(payload)


@app.route('/api/v3/schedules/unmapped_components', methods=['GET'])
def v3_schedules_unmapped_components():
    db = get_db()
    return jsonify({'unmapped_components': unmapped_components(db, dict(request.args))})


@app.route('/api/v3/schedules/reconciliation', methods=['GET'])
def v3_schedules_reconciliation():
    db = get_db()
    rows = reconciliation_rows(db, dict(request.args))
    return jsonify({'reconciliation': rows})


@app.route('/api/v3/schedules/component_mappings', methods=['GET'])
def v3_schedule_component_mappings():
    db = get_db()
    return jsonify(list_component_mappings(db))


@app.route('/api/v3/schedules/component_mappings', methods=['POST'])
def v3_schedule_component_mapping_add():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        mapping = upsert_component_mapping(
            db,
            data.get('component_head', ''),
            data.get('schedule_type', 'other_components'),
            data.get('label', ''),
            data.get('notes', ''),
            bool(data.get('active', True)),
        )
        audit_log_v3(
            db, 'schedule_component_mapping_add', 'schedule_component_mappings_v3',
            mapping.get('id'), 'component_head', '', mapping.get('component_head', ''),
            user_note=f"Mapped to {mapping.get('schedule_type')}"
        )
        db.commit()
        return jsonify({'mapping': mapping})
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/schedules/component_mappings/deactivate', methods=['POST'])
def v3_schedule_component_mapping_deactivate():
    db = get_db()
    data = request.get_json(silent=True) or {}
    component_head = data.get('component_head', '')
    mapping_id = data.get('id') or data.get('mapping_id')
    count = deactivate_component_mapping(db, component_head, int(mapping_id) if mapping_id else None)
    audit_log_v3(
        db, 'schedule_component_mapping_deactivate', 'schedule_component_mappings_v3',
        mapping_id, 'component_head', component_head, '',
        user_note=f"Deactivated {count} mapping(s)"
    )
    db.commit()
    return jsonify({'deactivated_count': count})


@app.route('/api/v3/schedules/component_mappings/reactivate', methods=['POST'])
def v3_schedule_component_mapping_reactivate():
    db = get_db()
    data = request.get_json(silent=True) or {}
    try:
        mapping = reactivate_component_mapping(db, int(data.get('id') or data.get('mapping_id')))
        audit_log_v3(
            db, 'schedule_component_mapping_reactivate', 'schedule_component_mappings_v3',
            mapping.get('id'), 'component_head', '', mapping.get('component_head', ''),
            user_note=f"Reactivated mapping to {mapping.get('schedule_type')}"
        )
        db.commit()
        return jsonify({'mapping': mapping})
    except (TypeError, ValueError) as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/schedules/component_mappings/ignore', methods=['POST'])
def v3_schedule_component_mapping_ignore():
    db = get_db()
    data = request.get_json(silent=True) or {}
    component_head = data.get('component_head', '')
    try:
        mapping = ignore_component(db, component_head, data.get('notes', ''))
        audit_log_v3(
            db, 'schedule_component_mapping_ignore', 'schedule_component_mappings_v3',
            mapping.get('id'), 'component_head', '', component_head,
            user_note='Ignored from unmapped component review'
        )
        db.commit()
        return jsonify({'mapping': mapping})
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400


@app.route('/api/v3/schedules/component_mappings/unignore', methods=['POST'])
def v3_schedule_component_mapping_unignore():
    db = get_db()
    data = request.get_json(silent=True) or {}
    component_head = data.get('component_head', '')
    mapping_id = data.get('id') or data.get('mapping_id')
    count = unignore_component(db, component_head, int(mapping_id) if mapping_id else None)
    audit_log_v3(
        db, 'schedule_component_mapping_unignore', 'schedule_component_mappings_v3',
        mapping_id, 'component_head', component_head, '',
        user_note=f"Unignored {count} mapping(s)"
    )
    db.commit()
    return jsonify({'unignored_count': count})


@app.route('/api/v3/schedules/unmapped_components/export', methods=['POST'])
def v3_schedules_unmapped_components_export():
    db = get_db()
    filters = _schedule_v3_args()
    unmapped = unmapped_components(db, filters)
    heads = [row.get('component_head') for row in unmapped if row.get('component_head')]

    sample_rows = []
    if heads:
        ph = ','.join('?' for _ in heads)
        sample_rows = [dict(row) for row in db.execute(f"""
            SELECT
              c.normalized_head AS component_head,
              c.raw_head,
              c.amount,
              e.raw_name AS employee_name,
              e.employee_id,
              e.group_name,
              e.month_year,
              e.source_filename AS source_file,
              e.sr_no,
              COALESCE(e.is_deleted,0) AS is_deleted
            FROM employee_components_v2 c
            JOIN employees_v2 e ON e.id = c.employee_id
            WHERE c.normalized_head IN ({ph})
              AND c.amount IS NOT NULL
              AND c.amount != 0
              AND COALESCE(e.is_deleted,0) = 0
            ORDER BY c.normalized_head, e.source_filename, e.sr_no
            LIMIT 500
        """, heads).fetchall()]

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    fname = f"IFMS_V3_Unmapped_Components_Review_{timestamp}.xlsx"
    fpath = os.path.join(app.config['EXPORT_FOLDER'], fname)
    metadata_rows = [
        {'Field': 'Application', 'Value': APP_NAME},
        {'Field': 'Designed and Developed by', 'Value': 'Neeraj Purohit'},
        {'Field': 'Role', 'Value': 'Senior Assistant, Rajasthan Police'},
        {'Field': 'Email', 'Value': DEVELOPER_EMAIL},
        {'Field': 'Export Timestamp', 'Value': datetime.now().isoformat(timespec='seconds')},
        {'Field': 'Disclaimer', 'Value': APP_DISCLAIMER},
    ]
    suggested_rows = [
        {
            'component_head': row.get('component_head', ''),
            'display_label': row.get('display_label', ''),
            'suggested_schedule_type': row.get('suggested_schedule_type', 'other_components'),
            'total_amount': row.get('total_amount', 0),
            'sample_employees': ', '.join(row.get('sample_employees') or []),
        }
        for row in unmapped
    ]
    notes_rows = [
        {'Note': 'Review each unmapped non-zero component before using schedule totals.'},
        {'Note': 'Ignored components remain preserved in Employee Data and employee_components_v2.'},
        {'Note': 'Mapping a component changes derived schedule grouping only; it does not modify parsed PDF data.'},
    ]
    with pd.ExcelWriter(fpath, engine='openpyxl') as writer:
        sheets = {
            'Export Metadata': pd.DataFrame(metadata_rows),
            'Unmapped Components': pd.DataFrame(unmapped),
            'Sample Employees': pd.DataFrame(sample_rows),
            'Suggested Mappings': pd.DataFrame(suggested_rows),
            'Notes': pd.DataFrame(notes_rows),
        }
        for sheet_name, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            ws = writer.sheets[sheet_name]
            ws.freeze_panes = 'A2'
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F4E78')
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 42)
            for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
                header = str(col[0].value or '')
                if header in {'amount', 'total_amount', 'positive_total', 'negative_total'}:
                    for cell in col[1:]:
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##,##0.00'

    return send_file(fpath, as_attachment=True, download_name=fname)


@app.route('/api/v3/schedules/export', methods=['POST'])
def v3_schedules_export():
    db = get_db()
    filters = _schedule_v3_args()
    payload = query_schedule_entries_v3(db, filters, paginate=False)
    entries = payload['entries']
    summary = payload['summary']
    reconciliation = reconciliation_rows(db, filters)
    unmapped = unmapped_components(db, filters)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    fname = f"IFMS_V3_Derived_Schedule_Export_{timestamp}.xlsx"
    fpath = os.path.join(app.config['EXPORT_FOLDER'], fname)

    metadata_rows = [
        {'Field': 'Application', 'Value': APP_NAME},
        {'Field': 'Designed and Developed by', 'Value': 'Neeraj Purohit'},
        {'Field': 'Role', 'Value': 'Senior Assistant, Rajasthan Police'},
        {'Field': 'Email', 'Value': DEVELOPER_EMAIL},
        {'Field': 'Export Timestamp', 'Value': datetime.now().isoformat(timespec='seconds')},
        {'Field': 'source_type', 'Value': SCHEDULE_SOURCE_TYPE_V3},
        {'Field': 'Disclaimer', 'Value': APP_DISCLAIMER},
    ]
    entry_columns = [
        'employee_sr_no', 'group_name', 'employee_name', 'employee_id', 'pan', 'pay_scale',
        'final_designation', 'basic_pay', 'schedule_label', 'component_head', 'amount',
        'component_source_type', 'employee_data_status', 'source_type', 'source_file',
        'bill_no', 'page_no', 'row_audit_status'
    ]
    entry_rows = [{col: row.get(col, '') for col in entry_columns} for row in entries]
    summary_rows = [
        {'Metric': 'schedule_type', 'Value': summary.get('schedule_type', '')},
        {'Metric': 'entry_count', 'Value': summary.get('entry_count', 0)},
        {'Metric': 'unique_employees', 'Value': summary.get('unique_employees', 0)},
        {'Metric': 'total_amount', 'Value': summary.get('total_amount', 0)},
        {'Metric': 'total_positive', 'Value': summary.get('total_positive', 0)},
        {'Metric': 'total_negative', 'Value': summary.get('total_negative', 0)},
        {'Metric': 'groups_count', 'Value': summary.get('groups_count', 0)},
        {'Metric': 'files_count', 'Value': summary.get('files_count', 0)},
        {'Metric': 'source_type', 'Value': summary.get('source_type', SCHEDULE_SOURCE_TYPE_V3)},
    ] + [{'Metric': key, 'Value': value} for key, value in (summary.get('audit') or {}).items()]
    filter_rows = [{'Filter': key, 'Value': value} for key, value in sorted(filters.items()) if value not in ('', None)]
    notes_rows = [
        {'Note': 'Schedule rows are derived from verified Employee Data and the employee component ledger.'},
        {'Note': 'Deleted employees are excluded unless show_deleted=1 is selected.'},
        {'Note': 'Rows with zero or missing component amounts are excluded.'},
        {'Note': 'Human review of the data is compulsory before using the data.'},
    ]

    with pd.ExcelWriter(fpath, engine='openpyxl') as writer:
        sheets = {
            'Export Metadata': pd.DataFrame(metadata_rows),
            'Schedule Entries': pd.DataFrame(entry_rows, columns=entry_columns),
            'Schedule Summary': pd.DataFrame(summary_rows),
            'Applied Filters': pd.DataFrame(filter_rows or [{'Filter': 'None', 'Value': ''}]),
            'Reconciliation': pd.DataFrame(reconciliation),
            'Unmapped Components': pd.DataFrame(unmapped),
            'Data Quality Notes': pd.DataFrame(notes_rows),
        }
        for sheet_name, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            ws = writer.sheets[sheet_name]
            ws.freeze_panes = 'A2'
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F4E78')
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 36)
            amount_cols = {'amount', 'basic_pay', 'Value'}
            for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
                header = str(col[0].value or '')
                if header in amount_cols:
                    for cell in col[1:]:
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##,##0.00'

    return send_file(fpath, as_attachment=True, download_name=fname)


@app.route('/api/schedule_stats')
def schedule_stats():
    """Get schedule data statistics."""
    # Legacy V1 route: currently reads `schedule_entries`.
    db = get_db()
    total = db.execute('SELECT COUNT(*) FROM schedule_entries').fetchone()[0]
    types = [r[0] for r in db.execute('SELECT DISTINCT schedule_type FROM schedule_entries ORDER BY schedule_type').fetchall()]
    if total == 0 and app.config.get('USE_V2_PARSER', False):
        heads = list(DERIVED_COMPONENT_SCHEDULES.keys())
        ph = ','.join('?' * len(heads))
        component_rows = db.execute("""
            SELECT
              c.normalized_head,
              c.amount,
              e.employee_id,
              e.raw_text_block
            FROM employee_components_v2 c
            JOIN employees_v2 e ON e.id = c.employee_id
            WHERE c.normalized_head IN (""" + ph + """)
              AND c.amount IS NOT NULL
              AND c.amount != 0
            ORDER BY c.normalized_head
        """, heads).fetchall()
        by_head = {}
        employees_by_head = {}
        for r in component_rows:
            if not is_valid_derived_schedule_component(r['normalized_head'], r['amount'], r['raw_text_block']):
                continue
            by_head.setdefault(r['normalized_head'], {'entry_count': 0, 'total_amount': 0})
            employees_by_head.setdefault(r['normalized_head'], set())
            by_head[r['normalized_head']]['entry_count'] += 1
            by_head[r['normalized_head']]['total_amount'] += r['amount'] or 0
            if r['employee_id']:
                employees_by_head[r['normalized_head']].add(r['employee_id'])
        type_summary = [
            {
                'schedule_type': DERIVED_COMPONENT_SCHEDULES.get(head, ''),
                'entry_count': data['entry_count'],
                'total_amount': data['total_amount'],
                'unique_employees': len(employees_by_head.get(head, set())),
            }
            for head, data in sorted(by_head.items())
        ]
        types = [r['schedule_type'] for r in type_summary]
        total = sum(r['entry_count'] for r in type_summary)
        return jsonify({
            'total_entries': total,
            'schedule_types': types,
            'type_summary': type_summary
        })
    
    type_summary = []
    for st in types:
        row = db.execute('''SELECT schedule_type, COUNT(*) as entry_count,
            SUM(total) as total_amount,
            COUNT(DISTINCT employee_id) as unique_employees
        FROM schedule_entries WHERE schedule_type = ?
        GROUP BY schedule_type''', (st,)).fetchone()
        if row:
            type_summary.append(dict(row))
    
    return jsonify({
        'total_entries': total,
        'schedule_types': types,
        'type_summary': type_summary
    })

@app.route('/api/schedule_entries')
def get_schedule_entries():
    """Get schedule entries with filters."""
    # Legacy V1 route: currently reads `schedule_entries`.
    db = get_db()
    
    conditions, params = [], []
    
    stype = request.args.get('schedule_type', '')
    if stype:
        conditions.append('schedule_type = ?')
        params.append(stype)
    
    month = request.args.get('month', '')
    if month:
        conditions.append('month_year = ?')
        params.append(month)
    
    search = request.args.get('search', '')
    if search:
        conditions.append('(name LIKE ? OR employee_id LIKE ? OR pan LIKE ? OR policy_no LIKE ? OR gpf_no LIKE ?)')
        s = f'%{search}%'
        params.extend([s, s, s, s, s])
    
    where = ' WHERE ' + ' AND '.join(conditions) if conditions else ''
    sched_total = db.execute('SELECT COUNT(*) FROM schedule_entries').fetchone()[0]
    if sched_total == 0 and app.config.get('USE_V2_PARSER', False):
        stype = request.args.get('schedule_type', '')
        search = request.args.get('search', '')
        comp_head = DERIVED_SCHEDULE_COMPONENTS.get(stype) if stype else None
        if stype and comp_head is None:
            page_num = int(request.args.get('page', 1))
            per_page = int(request.args.get('per_page', 50))
            return jsonify({
                'entries': [],
                'total': 0,
                'page': page_num,
                'per_page': per_page,
                'pages': 0,
                'summary': {
                    'count': 0,
                    'grand_total': 0,
                    'total_itax': 0,
                    'total_premium': 0,
                    'total_gpf': 0,
                    'total_gpf2004': 0,
                    'total_si_premium': 0,
                    'total_si_loan': 0,
                    'unique_employees': 0,
                }
            })
        params = []
        q = """
            SELECT
              e.sr_no, e.employee_id, e.raw_name as employee_name, e.normalized_designation as designation,
              e.gpf_pran_raw, e.source_filename, e.raw_text_block, c.raw_head, c.normalized_head,
              c.amount as component_amount, c.source_page,
              e.month_year, e.group_name, e.pan_raw as pan, e.id as employee_row_id,
              COALESCE((
                SELECT c2.amount FROM employee_components_v2 c2
                WHERE c2.employee_id = e.id AND c2.normalized_head = 'basic_pay'
                ORDER BY c2.id LIMIT 1
              ), 0) as basic_pay
            FROM employee_components_v2 c
            JOIN employees_v2 e ON e.id = c.employee_id
            WHERE c.amount IS NOT NULL
              AND c.amount != 0
        """
        if comp_head:
            q += " AND c.normalized_head = ?"
            params.append(comp_head)
        else:
            heads = list(DERIVED_COMPONENT_SCHEDULES.keys())
            q += " AND c.normalized_head IN (" + ",".join("?" * len(heads)) + ")"
            params.extend(heads)
        if search:
            q += " AND (e.raw_name LIKE ? OR e.employee_id LIKE ? OR e.pan_raw LIKE ?)"
            s = f"%{search}%"
            params.extend([s, s, s])
        q += " ORDER BY c.normalized_head, e.sr_no"
        rows = db.execute(q, params).fetchall()
        entries = []
        for r in rows:
            if not is_valid_derived_schedule_component(r['normalized_head'], r['component_amount'], r['raw_text_block']):
                continue
            schedule_type = DERIVED_COMPONENT_SCHEDULES.get(r['normalized_head'], '')
            entries.append({
                'schedule_type': schedule_type,
                'normalized_head': r['normalized_head'],
                'sr_no': r['sr_no'] or 0,
                'employee_id': r['employee_id'] or '',
                'employee_name': r['employee_name'] or '',
                'name': r['employee_name'] or '',
                'designation': r['designation'] or '',
                'basic_pay': r['basic_pay'] or 0,
                'gpf_pran': r['gpf_pran_raw'] or '',
                'component_head': r['raw_head'] or '',
                'component_amount': r['component_amount'] or 0,
                'source_file': r['source_filename'] or '',
                'source_page': r['source_page'] or 0,
                'pan': r['pan'] or '',
                'policy_no': '',
                'gpf_no': '',
                'si_policy_no': '',
                'loan_ac_no': '',
                'bank_ac_no': '',
                'coop_ser_no': '',
                'nominee_name': '',
                'reason': '',
                'income_tax': r['component_amount'] if schedule_type == 'income_tax' else 0,
                'sit': 0,
                'health_edu_cess': 0,
                'hi_edu_cess': 0,
                'premium': r['component_amount'] if schedule_type == 'lic' else 0,
                'gpf_contribution': r['component_amount'] if schedule_type == 'gpf' else 0,
                'gpf_loan': 0,
                'gpf_2004': r['component_amount'] if schedule_type == 'gpf_2004' else 0,
                'si_premium': r['component_amount'] if schedule_type == 'si_premium' else 0,
                'si_loan': r['component_amount'] if schedule_type == 'si_loan' else 0,
                'si_interest': 0,
                's_tax_on_si': 0,
                'rghs': r['component_amount'] if schedule_type == 'rghs' else 0,
                'total': r['component_amount'] or 0,
                'month_year': r['month_year'] or '',
                'group_name': r['group_name'] or '',
                'source_type': 'derived_from_inner_sheet'
            })
        count = len(entries)
        page_num = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 50))
        start = (page_num - 1) * per_page
        paged = entries[start:start + per_page]
        summary = {
            'count': count,
            'grand_total': sum(e['total'] for e in entries),
            'total_itax': sum(e['income_tax'] for e in entries),
            'total_premium': sum(e['premium'] for e in entries),
            'total_gpf': sum(e['gpf_contribution'] for e in entries),
            'total_gpf2004': sum(e['gpf_2004'] for e in entries),
            'total_si_premium': sum(e['si_premium'] for e in entries),
            'total_si_loan': sum(e['si_loan'] for e in entries),
            'unique_employees': len({e['employee_id'] for e in entries if e['employee_id']}),
        }
        return jsonify({
            'entries': paged,
            'total': count,
            'page': page_num,
            'per_page': per_page,
            'pages': (count + per_page - 1) // per_page,
            'summary': summary
        })
    
    # Sorting
    sort_by = request.args.get('sort', 'sr_no')
    sort_dir = request.args.get('dir', 'asc')
    allowed = ['sr_no','name','employee_id','basic_pay','total','income_tax','premium',
               'gpf_contribution','gpf_2004','si_premium','si_loan','schedule_type']
    if sort_by not in allowed:
        sort_by = 'sr_no'
    order = f' ORDER BY {sort_by} {sort_dir}'
    
    # Pagination
    page_num = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 50))
    offset = (page_num - 1) * per_page
    
    count = db.execute(f'SELECT COUNT(*) FROM schedule_entries{where}', params).fetchone()[0]
    
    rows = db.execute(f'SELECT * FROM schedule_entries{where}{order} LIMIT ? OFFSET ?',
                      params + [per_page, offset]).fetchall()
    
    entries = [dict(r) for r in rows]
    
    # Summary
    summary = db.execute(f'''SELECT COUNT(*) as count,
        COALESCE(SUM(total),0) as grand_total,
        COALESCE(SUM(income_tax),0) as total_itax,
        COALESCE(SUM(premium),0) as total_premium,
        COALESCE(SUM(gpf_contribution),0) as total_gpf,
        COALESCE(SUM(gpf_2004),0) as total_gpf2004,
        COALESCE(SUM(si_premium),0) as total_si_premium,
        COALESCE(SUM(si_loan),0) as total_si_loan,
        COUNT(DISTINCT employee_id) as unique_employees
    FROM schedule_entries{where}''', params).fetchone()
    
    return jsonify({
        'entries': entries,
        'total': count,
        'page': page_num,
        'per_page': per_page,
        'pages': (count + per_page - 1) // per_page,
        'summary': dict(summary)
    })

@app.route('/api/schedule_export')
def schedule_export():
    """Export schedule data to Excel."""
    # Legacy V1 route: currently exports from `schedule_entries`.
    db = get_db()
    
    conditions, params = [], []
    stype = request.args.get('schedule_type', '')
    if stype:
        conditions.append('schedule_type = ?')
        params.append(stype)
    
    search = request.args.get('search', '')
    if search:
        conditions.append('(name LIKE ? OR employee_id LIKE ?)')
        s = f'%{search}%'
        params.extend([s, s])
    
    where = ' WHERE ' + ' AND '.join(conditions) if conditions else ''
    
    rows = db.execute(f'''SELECT schedule_type, sr_no, employee_id, name, designation,
        basic_pay, pan, policy_no, gpf_no, si_policy_no, loan_ac_no,
        bank_ac_no, coop_ser_no, nominee_name, reason,
        income_tax, sit, health_edu_cess, hi_edu_cess,
        premium, gpf_contribution, gpf_loan, gpf_2004,
        si_premium, si_loan, si_interest, s_tax_on_si,
        rghs, total, month_year, group_name
    FROM schedule_entries{where} ORDER BY schedule_type, sr_no''', params).fetchall()

    if not rows and app.config.get('USE_V2_PARSER', False):
        comp_head = DERIVED_SCHEDULE_COMPONENTS.get(stype) if stype else None
        if stype and comp_head is None:
            return jsonify({'error': 'No schedule data to export'}), 404

        v2_params = []
        v2_query = """
            SELECT
              e.sr_no, e.employee_id, e.raw_name as name, e.normalized_designation as designation,
              e.gpf_pran_raw, e.raw_text_block, c.raw_head, c.normalized_head, c.amount as component_amount,
              e.month_year, e.group_name, e.pan_raw as pan,
              COALESCE((
                SELECT c2.amount FROM employee_components_v2 c2
                WHERE c2.employee_id = e.id AND c2.normalized_head = 'basic_pay'
                ORDER BY c2.id LIMIT 1
              ), 0) as basic_pay
            FROM employee_components_v2 c
            JOIN employees_v2 e ON e.id = c.employee_id
            WHERE c.amount IS NOT NULL
              AND c.amount != 0
        """
        if comp_head:
            v2_query += " AND c.normalized_head = ?"
            v2_params.append(comp_head)
        else:
            heads = list(DERIVED_COMPONENT_SCHEDULES.keys())
            v2_query += " AND c.normalized_head IN (" + ",".join("?" * len(heads)) + ")"
            v2_params.extend(heads)
        if search:
            v2_query += " AND (e.raw_name LIKE ? OR e.employee_id LIKE ?)"
            s = f'%{search}%'
            v2_params.extend([s, s])
        v2_query += " ORDER BY c.normalized_head, e.sr_no"

        rows = []
        for r in db.execute(v2_query, v2_params).fetchall():
            if not is_valid_derived_schedule_component(r['normalized_head'], r['component_amount'], r['raw_text_block']):
                continue
            schedule_type = DERIVED_COMPONENT_SCHEDULES.get(r['normalized_head'], '')
            amount = r['component_amount'] or 0
            rows.append({
                'schedule_type': schedule_type,
                'sr_no': r['sr_no'] or 0,
                'employee_id': r['employee_id'] or '',
                'name': r['name'] or '',
                'designation': r['designation'] or '',
                'basic_pay': r['basic_pay'] or 0,
                'pan': r['pan'] or '',
                'policy_no': '',
                'gpf_no': r['gpf_pran_raw'] or '',
                'si_policy_no': '',
                'loan_ac_no': '',
                'bank_ac_no': '',
                'coop_ser_no': '',
                'nominee_name': '',
                'reason': '',
                'income_tax': amount if schedule_type == 'income_tax' else 0,
                'sit': 0,
                'health_edu_cess': 0,
                'hi_edu_cess': 0,
                'premium': amount if schedule_type == 'lic' else 0,
                'gpf_contribution': amount if schedule_type == 'gpf' else 0,
                'gpf_loan': 0,
                'gpf_2004': amount if schedule_type == 'gpf_2004' else 0,
                'si_premium': amount if schedule_type == 'si_premium' else 0,
                'si_loan': amount if schedule_type == 'si_loan' else 0,
                'si_interest': 0,
                's_tax_on_si': 0,
                'rghs': amount if schedule_type == 'rghs' else 0,
                'total': amount,
                'month_year': r['month_year'] or '',
                'group_name': r['group_name'] or '',
            })
    
    if not rows:
        return jsonify({'error': 'No schedule data to export'}), 404
    
    df = pd.DataFrame([dict(r) for r in rows])
    
    col_map = {
        'schedule_type': 'Schedule Type', 'sr_no': 'Sr.No', 'employee_id': 'Employee ID',
        'name': 'Name', 'designation': 'Designation', 'basic_pay': 'Basic Pay',
        'pan': 'PAN', 'policy_no': 'Policy No', 'gpf_no': 'GPF No',
        'si_policy_no': 'SI Policy No', 'loan_ac_no': 'Loan A/C No',
        'bank_ac_no': 'Bank A/C No', 'coop_ser_no': 'Coop Ser. No',
        'nominee_name': 'Nominee Name', 'reason': 'Reason',
        'income_tax': 'Income Tax', 'sit': 'SIT',
        'health_edu_cess': 'Health Edu Cess', 'hi_edu_cess': 'Hi Edu Cess',
        'premium': 'LIC Premium', 'gpf_contribution': 'GPF Contribution',
        'gpf_loan': 'GPF Loan', 'gpf_2004': 'GPF 2004',
        'si_premium': 'SI Premium', 'si_loan': 'SI Loan',
        'si_interest': 'SI Interest', 's_tax_on_si': 'S.Tax on SI',
        'rghs': 'RGHS', 'total': 'Total', 'month_year': 'Month/Year',
        'group_name': 'Group'
    }
    df.rename(columns=col_map, inplace=True)
    
    # Remove columns that are all zero/empty
    for col in df.columns:
        if df[col].dtype in ['float64', 'int64'] and (df[col] == 0).all():
            df.drop(col, axis=1, inplace=True)
        elif df[col].dtype == 'object' and (df[col] == '').all():
            df.drop(col, axis=1, inplace=True)
    
    fname = f"IFMS_Schedule_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    fpath = os.path.join(app.config['EXPORT_FOLDER'], fname)
    
    with pd.ExcelWriter(fpath, engine='openpyxl') as writer:
        # If specific type, single sheet; else separate sheets per type
        if stype:
            df.to_excel(writer, index=False, sheet_name=stype.replace('_', ' ').title())
        else:
            types_in_data = df['Schedule Type'].unique() if 'Schedule Type' in df.columns else [stype]
            for t in types_in_data:
                sheet_df = df[df['Schedule Type'] == t].drop('Schedule Type', axis=1, errors='ignore')
                sheet_name = t.replace('_', ' ').title()[:31]
                sheet_df.to_excel(writer, index=False, sheet_name=sheet_name)
        
        for ws in writer.sheets.values():
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)
    
    return send_file(fpath, as_attachment=True, download_name=fname)

# ===================== DELETE APIs =====================


@app.route('/api/delete_all', methods=['POST'])
def delete_all():
    db = get_db()
    try:
        counts = {
            'v1_employees': db.execute('SELECT COUNT(*) FROM employees').fetchone()[0],
            'v1_uploads': db.execute('SELECT COUNT(*) FROM uploads').fetchone()[0],
            'v1_schedule_entries': db.execute('SELECT COUNT(*) FROM schedule_entries').fetchone()[0],
            'v2_bills': db.execute('SELECT COUNT(*) FROM bills_v2').fetchone()[0],
            'v2_employees': db.execute('SELECT COUNT(*) FROM employees_v2').fetchone()[0],
            'v2_components': db.execute('SELECT COUNT(*) FROM employee_components_v2').fetchone()[0],
            'v2_validation_issues': db.execute('SELECT COUNT(*) FROM validation_issues_v2').fetchone()[0],
            'v3_review_queue': db.execute('SELECT COUNT(*) FROM review_queue_v3').fetchone()[0],
            'v2_upload_jobs': db.execute('SELECT COUNT(*) FROM upload_jobs_v2').fetchone()[0],
        }

        db.execute('BEGIN')
        # V1 tables
        db.execute('DELETE FROM employees')
        db.execute('DELETE FROM schedule_entries')
        db.execute('DELETE FROM uploads')
        # V2 tables (children first)
        db.execute('DELETE FROM employee_components_v2')
        db.execute('DELETE FROM validation_issues_v2')
        db.execute('DELETE FROM review_queue_v3')
        db.execute('DELETE FROM employees_v2')
        db.execute('DELETE FROM upload_jobs_v2')
        db.execute('DELETE FROM bills_v2')
        db.commit()
        return jsonify({'status': 'All data deleted', 'deleted_counts': counts})
    except Exception as e:
        db.rollback()
        return jsonify({'error': str(e)}), 500



@app.route('/api/delete_group', methods=['POST'])
def delete_group():
    data = request.json
    group = data.get('group_name', '')
    if not group:
        return jsonify({'error': 'No group specified'}), 400
    db = get_db()
    try:
        bill_ids = [r[0] for r in db.execute('SELECT id FROM bills_v2 WHERE group_name = ?', (group,)).fetchall()]
        emp_ids = [r[0] for r in db.execute('SELECT id FROM employees_v2 WHERE group_name = ?', (group,)).fetchall()]
        filenames = [r[0] for r in db.execute('SELECT source_filename FROM bills_v2 WHERE group_name = ?', (group,)).fetchall()]

        counts = {
            'v1_employees': db.execute('SELECT COUNT(*) FROM employees WHERE group_name = ?', (group,)).fetchone()[0],
            'v1_uploads': db.execute('SELECT COUNT(*) FROM uploads WHERE group_name = ?', (group,)).fetchone()[0],
            'v1_schedule_entries': db.execute('SELECT COUNT(*) FROM schedule_entries WHERE group_name = ?', (group,)).fetchone()[0],
            'v2_bills': len(bill_ids),
            'v2_employees': len(emp_ids),
            'v2_components': 0,
            'v2_validation_issues': 0,
            'v2_upload_jobs': 0,
        }

        if emp_ids:
            ph_emp = ','.join('?' * len(emp_ids))
            counts['v2_components'] = db.execute(
                f'SELECT COUNT(*) FROM employee_components_v2 WHERE employee_id IN ({ph_emp})',
                emp_ids
            ).fetchone()[0]

        if emp_ids and bill_ids:
            ph_emp = ','.join('?' * len(emp_ids))
            ph_bill = ','.join('?' * len(bill_ids))
            counts['v2_validation_issues'] = db.execute(
                f'SELECT COUNT(DISTINCT id) FROM validation_issues_v2 WHERE employee_id IN ({ph_emp}) OR bill_id IN ({ph_bill})',
                emp_ids + bill_ids
            ).fetchone()[0]
        elif emp_ids:
            ph_emp = ','.join('?' * len(emp_ids))
            counts['v2_validation_issues'] = db.execute(
                f'SELECT COUNT(*) FROM validation_issues_v2 WHERE employee_id IN ({ph_emp})',
                emp_ids
            ).fetchone()[0]
        elif bill_ids:
            ph_bill = ','.join('?' * len(bill_ids))
            counts['v2_validation_issues'] = db.execute(
                f'SELECT COUNT(*) FROM validation_issues_v2 WHERE bill_id IN ({ph_bill})',
                bill_ids
            ).fetchone()[0]

        if filenames:
            ph_file = ','.join('?' * len(filenames))
            counts['v2_upload_jobs'] = db.execute(
                f'SELECT COUNT(*) FROM upload_jobs_v2 WHERE source_filename IN ({ph_file})',
                filenames
            ).fetchone()[0]

        db.execute('BEGIN')
        # V1 tables
        db.execute('DELETE FROM employees WHERE group_name = ?', (group,))
        db.execute('DELETE FROM schedule_entries WHERE group_name = ?', (group,))
        db.execute('DELETE FROM uploads WHERE group_name = ?', (group,))

        # V2 child rows first
        if emp_ids:
            ph_emp = ','.join('?' * len(emp_ids))
            db.execute(f'DELETE FROM employee_components_v2 WHERE employee_id IN ({ph_emp})', emp_ids)
            db.execute(f'DELETE FROM validation_issues_v2 WHERE employee_id IN ({ph_emp})', emp_ids)
        if bill_ids:
            ph_bill = ','.join('?' * len(bill_ids))
            db.execute(f'DELETE FROM validation_issues_v2 WHERE bill_id IN ({ph_bill})', bill_ids)
            db.execute(f'DELETE FROM employees_v2 WHERE bill_id IN ({ph_bill})', bill_ids)
            db.execute(f'DELETE FROM bills_v2 WHERE id IN ({ph_bill})', bill_ids)
        db.execute('DELETE FROM employees_v2 WHERE group_name = ?', (group,))
        if filenames:
            ph_file = ','.join('?' * len(filenames))
            db.execute(f'DELETE FROM upload_jobs_v2 WHERE source_filename IN ({ph_file})', filenames)

        db.commit()
        return jsonify({'status': f'Group "{group}" deleted', 'deleted_counts': counts})
    except Exception as e:
        db.rollback()
        return jsonify({'error': str(e)}), 500

from ifms_suite.web import register_suite
register_suite(app)


@app.route('/groups')
def group_analyzer():
    return send_file(os.path.join(os.path.dirname(__file__), 'IFMS Group Management - Analyzer.html'))

@app.route('/xlsx.full.min.js')
def group_excel_library():
    return send_file(os.path.join(os.path.dirname(__file__), 'static', 'xlsx.full.min.js'))

if __name__ == '__main__':
    selected_port = int(os.environ.get('PORT', '5000'))
    print("\n" + "="*60)
    print("  IFMS 3.0 Salary Analytics Tool v7")
    print(f"  Open in browser: http://127.0.0.1:{selected_port}")
    print("="*60 + "\n")
    app.run(host='0.0.0.0', port=selected_port, debug=False)
